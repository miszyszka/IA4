"""
IA 4 — testy zespołu agentów 1.41 na danych syntetycznych — część `python -m ia4.ai.selftest`.
Wersja projektu: 1.55 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Badacz: cechy bez zaglądania w przyszłość, symulacja hipotezy (SL przed TP, luka, limit, koszty), sejf AI,
zasiana przewaga odnaleziona i przechodzi sito, szum nie przechodzi, cykl z atrapą modelu (błąd → ponowienie,
powtórka pominięta). Krytyk: granice poziomu nauki, werdykty sprawdzane przez kod, wnioski. Inwestor w trybie nauki:
model w każdej rundzie, kandydaci z hipotez „próba”, automat bierności, wymagane wejście na poziomie 3.
Dyspozytor: jeden agent naraz, pierwszeństwo Inwestor › Krytyk › Badacz, dziennik pracy, status.json, branch `ai`.
"""

from __future__ import annotations

import copy
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

from . import config as aicfg
from . import market, research


def _check(name, ok, detail=""):
    from .selftest import check
    return check(name, ok, detail)


def planted_fx(sym: str, start="2024-10-01", end="2026-10-09T19:00:00Z", seed=5, edge=True) -> pd.DataFrame:
    """Świece 1h 24/7: szum, a świeca o 10:00 UTC (gdy `edge`) zawsze rośnie o 0,6 zmienności."""
    rng = np.random.default_rng(seed + sum(map(ord, sym)))
    t0 = int(datetime.fromisoformat(start).replace(tzinfo=timezone.utc).timestamp())
    t1 = int(datetime.fromisoformat(end.replace("Z", "+00:00")).timestamp())
    t = np.arange(t0, t1 + 1, 3600)
    from ..fx.synth import SPEC
    p = SPEC.get(sym, (1.2, 0))[0]
    vol = 0.0012
    o = np.empty(t.size); c = np.empty(t.size); h = np.empty(t.size); lo = np.empty(t.size)
    for i in range(t.size):
        o[i] = p
        if edge and (t[i] % 86400) // 3600 == 10:
            c[i] = p * (1 + 0.6 * vol)
            h[i], lo[i] = c[i] * (1 + 0.05 * vol), o[i] * (1 - 0.05 * vol)
        else:
            c[i] = p * (1 + rng.normal(0, vol))
            h[i] = max(o[i], c[i]) * (1 + abs(rng.normal(0, vol / 2)))
            lo[i] = min(o[i], c[i]) * (1 - abs(rng.normal(0, vol / 2)))
        p = c[i]
    return pd.DataFrame({"time": t, "o": o, "h": h, "l": lo, "c": c, "filled": False})


EDGE_HYP = {"market": "fx", "direction": "long", "instruments": "all",
            "conditions": [{"feature": "hour", "op": ">", "value": 8.5}, {"feature": "hour", "op": "<", "value": 9.5}],
            "sl_atr": 0.9, "tp_atr": 0.3, "max_bars": 2, "idea": "świeca 10:00 UTC rośnie (zasiane)"}
NOISE_HYP = {"market": "fx", "direction": "long", "instruments": "all",
             "conditions": [{"feature": "rsi14", "op": "<", "value": 40}], "sl_atr": 1.5, "tp_atr": 1.5, "max_bars": 12,
             "idea": "odbicie po spadku (na szumie — bez przewagi)"}


class FakeLLM:
    """Atrapa warstwy modelu: odpowiedź ze skryptu (agent, brief, nr wywołania)."""

    def __init__(self, script, done=None):
        self.script, self.done, self.n, self.calls, self.sec_total, self.last, self.last_by = script, done, 0, 0, 0.0, {}, {}

    def ask(self, agent, system, user, schema, check, temperature):
        meta = {"attempts": 0, "sec": 0.0, "errors": [], "promptTokens": len(user) // 4, "evalTokens": 100}
        msgs = []
        for a in range(3):
            self.n += 1
            self.calls += 1
            self.sec_total += 1.5
            meta["attempts"] = a + 1
            meta["sec"] += 1.5
            obj = self.script(agent, user, self.n, msgs)
            err = list(check(obj)) if isinstance(obj, dict) else ["brak JSON"]
            meta["errors"] = err
            self.last = self.last_by[agent] = {"agent": agent, "sec": 1.5, "ok": not err}
            if not err:
                if self.done:
                    self.done(agent)
                return obj, meta
            msgs.append(err)
        if self.done:
            self.done(agent)
        return (obj if isinstance(obj, dict) else None), meta


def _cfg(tmp: Path) -> dict:
    cfg = copy.deepcopy(aicfg.DEFAULTS)
    cfg["xlsx_path"] = str(tmp / "AI-INVESTOR.xlsx")
    return cfg


# ---------------------------------------------------------------------------
def test_research(tmp: Path) -> bool:
    ok = True
    cfg = _cfg(tmp)
    # cechy: wartość na świecy t z serii uciętej w t = wartość z pełnej serii
    df = planted_fx("EURUSD", edge=False)
    s = market.fx_series("EURUSD", df)
    F, _ = research.features(s.o, s.h, s.l, s.c, s.to, s.kind)
    t = 3000
    G, _ = research.features(s.o[:t + 1], s.h[:t + 1], s.l[:t + 1], s.c[:t + 1], s.to[:t + 1], s.kind)
    same = all((np.isnan(F[f][t]) and np.isnan(G[f][t])) or abs(F[f][t] - G[f][t]) < 1e-9 for f in F)
    ok &= _check("cechy Badacza bez zaglądania w przyszłość (seria ucięta w t = pełna do t)", same)
    # symulacja: SL przed TP, luka, limit, koszty
    o = np.array([100, 100, 100, 100, 100, 100, 100, 100.]); h = np.array([101, 101, 103, 101, 101, 101, 101, 101.])
    l = np.array([99, 99, 97, 99, 99, 99, 99, 99.]); c = np.array([100, 100, 100, 100, 100, 100, 100, 100.])
    atr = np.full(8, 2.0)
    m = np.zeros(8, bool); m[1] = True
    E, X, R, W = research._sim(o, h, l, c, atr, m, 1, 1.0, 1.0, 5, 0.0, 0)
    ok &= _check("symulacja hipotezy: SL i TP w tej samej świecy → SL (pesymistycznie)", len(R) == 1 and W[0] == 1 and abs(R[0] + 2.0) < 1e-9,
                 f"wynik {R[0]:+.2f}%" if len(R) else "brak")
    o2 = o.copy(); o2[3] = 95.0
    h2 = h.copy(); h2[2] = 101; l2 = l.copy(); l2[2] = 99
    E, X, R, W = research._sim(o2, h2, l2, c, atr, m, 1, 1.0, 3.0, 5, 0.0, 0)
    ok &= _check("luka na otwarciu przez SL → wyjście po cenie otwarcia (strata większa niż SL)", len(R) == 1 and abs(R[0] + 5.0) < 1e-9,
                 f"{R[0]:+.2f}%" if len(R) else "brak")
    E, X, R, W = research._sim(o, h2, l2, c, atr, m, 1, 5.0, 5.0, 3, 0.1, 0)
    ok &= _check("limit świec → wyjście po close; koszt odjęty od wyniku", len(R) == 1 and W[0] == 2 and X[0] - E[0] + 1 == 3
                 and abs(R[0] + 0.1) < 1e-9, f"{R[0]:+.3f}%" if len(R) else "brak")
    # pracownia: sejf FX odcięty, zasiana przewaga przechodzi sito, szum — nie
    fx = {s: planted_fx(s) for s in ("EURUSD", "GBPUSD", "USDJPY")}
    pr = market.Prices(fx=fx)
    lab = research.Lab(pr, [], cfg)
    a, _ = lab.span["fx"]
    cut = research.ts(cfg["fx_research_from"])
    ok &= _check("sejf AI: świece FX sprzed 2025-04-01 odcięte przed liczeniem cech", all(d["to"][0] >= cut for d in lab.M["fx"].values())
                 and a >= cut, datetime.fromtimestamp(a, timezone.utc).strftime("%Y-%m-%d"))
    r = lab.evaluate(EDGE_HYP)
    why = research.sieve(r, cfg)
    ok &= _check("zasiana przewaga: hipoteza odnaleziona, lepsza niż wejście na ślepo, przechodzi sito kodu",
                 not why and r["edge"] > 0.03, research.res_line(r) + (f" · {why}" if why else ""))
    lab0 = research.Lab(market.Prices(fx={s: planted_fx(s, edge=False) for s in fx}), [], cfg)
    rn = lab0.evaluate(NOISE_HYP)
    ok &= _check("szum: hipoteza bez przewagi odrzucona przez sito kodu", research.sieve(rn, cfg) != [], research.res_line(rn))
    return ok


# ---------------------------------------------------------------------------
def test_team(tmp: Path) -> bool:
    from .agents import AgentWorker
    from .critic import Critic
    from .hub import Hub
    from .investor import Investor
    from .research import Researcher
    from .store import AiStore, restore
    from . import status
    from .selftest2 import frames_stocks, strategies
    ok = True
    cfg = _cfg(tmp)
    cfg["passive_rounds"] = 2
    fx = {s: planted_fx(s) for s in ("EURUSD", "GBPUSD", "USDJPY")}
    fr = frames_stocks(n_inst=6, sessions=200)
    traded = [s for s in fr if s not in ("SPY", "QQQ")]
    hub = Hub(tmp / "team.sqlite")
    clock = {"t": float(fx["EURUSD"]["time"].iloc[-1]) + 3600 + 90}

    class Data:
        def __init__(self):
            self.traded, self.context = traded, ["SPY", "QQQ"]
            self.cut = None

        def get(self, need):
            f = fx if self.cut is None else {s: d[d["time"] <= self.cut] for s, d in fx.items()}
            return market.Prices(stocks=fr, fx=f), None, None

    data = Data()
    plan = {"researcher": [], "critic": [], "investor": []}

    def script(agent, user, n, retries):
        q = plan[agent]
        return q[0](user, retries) if q else None

    llm = FakeLLM(script, done=lambda agent: plan[agent].pop(0) if plan[agent] else None)
    inv = Investor(cfg, hub, llm, clock=lambda: clock["t"], log=lambda *a: None)
    inv.set_playbook({"builtAt": 0, "source": {}, "strategies": strategies()}, None)
    crit = Critic(cfg, hub, llm, inv, clock=lambda: clock["t"], log=lambda *a: None)
    res = Researcher(cfg, hub, llm, data, clock=lambda: clock["t"], log=lambda *a: None)
    store = AiStore(push=False, root=tmp / "ai-repo")
    w = AgentWorker(inv, data, store=store, clock=lambda: clock["t"], log=lambda *a: None, critic=crit, researcher=res)
    w.research_after = 0
    w._round = lambda job: (setattr(data, "cut", job["key"]), _fx_round(inv, data, job))[1]

    # Badacz: 1. odpowiedź z błędem (nieznana cecha) → ponowienie; 2. dwie poprawne (przewaga + szum)
    bad = dict(EDGE_HYP, conditions=[{"feature": "magia", "op": ">", "value": 1}])
    noise = dict(NOISE_HYP, direction="short", conditions=[{"feature": "rsi14", "op": ">", "value": 60}],
                 idea="short po wzroście (tu: wbrew zasianemu wzrostowi)")
    plan["researcher"].append(lambda u, r: {"hypotheses": [bad], "note": "x"} if not r else
                              {"hypotheses": [EDGE_HYP, noise], "note": "pierwsze pomysły"})
    kind = w.step()
    hs = {h["idea"]: h for h in hub.hypotheses(10)}
    he, hn = hs.get(EDGE_HYP["idea"]), hs.get(noise["idea"])
    ok &= _check("Badacz: błędna cecha → ponowienie z listą błędów; hipotezy policzone, przewaga „do oceny”, szum odrzucony przez kod",
                 kind == "research" and bool(he) and he["status"] == "do oceny" and bool(hn) and hn["status"] == "odrzucona przez kod",
                 f"{he and he['status']} / {hn and hn['status']}")
    plan["researcher"].append(lambda u, r: {"hypotheses": [EDGE_HYP], "note": "powtórka"})
    w.research_after = 0
    w.step()
    ok &= _check("Badacz: ta sama hipoteza drugi raz nie jest liczona (powtórka)", len(hub.hypotheses(10)) == 2)

    # Krytyk: chce poziom 3 (kod: najwyżej +1), próba dla przewagi, zatwierdzenie odrzucone (n < 20), wniosek
    hid = he["id"]
    lid = hub.lessons_active()[0]["id"]
    plan["critic"].append(lambda u, r: {"explore": 3, "explore_reason": "Inwestor bierny", "investor_notes": ["wchodź małym ryzykiem"],
                                        "verdicts": [{"id": hid, "verdict": "próba", "note": "warto sprawdzić na żywo"},
                                                     {"id": hn["id"], "verdict": "zatwierdzona", "note": "x"}],
                                        "researcher_notes": ["badaj pory dnia"], "lessons_add": ["pora dnia ma znaczenie"],
                                        "lessons_remove": [lid], "summary": "pierwszy przegląd"})
    for _ in range(3):
        hub.hyp_add(dict(NOISE_HYP, sig=f"x{_}", status="do oceny", result={"n": 0}, at=clock["t"]))
    t_before = crit.due()
    w.research_after = 1e18
    kind = w.step()
    rv = hub.reviews(1)[0]
    ex = inv.explore()
    ok &= _check("Krytyk: pora na przegląd (≥ 3 do oceny); poziom nauki 1 → 2 (najwyżej +1), próba nadana, zatwierdzenie bez 20 transakcji odrzucone",
                 t_before and kind == "critic" and ex["base"] == 2 and hub.hyp_get(hid)["status"] == "próba"
                 and any(not v["applied"] and v["verdict"] == "zatwierdzona" for v in rv["verdicts"]),
                 f"poziom {ex['base']}, {hub.hyp_get(hid)['status']}")
    ok &= _check("Krytyk: wniosek dodany i usunięty, uwagi dla Inwestora i Badacza w przeglądzie",
                 any(x["text"] == "pora dnia ma znaczenie" for x in hub.lessons_active()) and lid not in {x["id"] for x in hub.lessons_active()}
                 and rv["investor_notes"] and rv["researcher_notes"])

    # Inwestor w trybie nauki: runda FX o 9:00 UTC → kandydat z hipotezy w próbie; uwagi Krytyka w briefie
    t9 = int(fx["EURUSD"]["time"].iloc[-1])
    t9 -= ((t9 % 86400) // 3600 - 9) * 3600
    seen = {}

    def inv_reply(u, r):
        seen["brief"] = u
        C = re.findall(r"^(C\d+) ", u, flags=re.M)
        return {"candidates": [{"ref": x, "action": "enter", "confidence": 55, "note": "test hipotezy"} for x in C],
                "positions": [{"ref": x, "action": "hold", "new_sl": None, "note": "t"} for x in re.findall(r"^(P\d+) ", u, flags=re.M)],
                "own_ideas": [], "round_note": "runda testowa"}
    plan["investor"].append(inv_reply)
    clock["t"] = t9 + 3600 + 90
    w.submit("fx", t9, "FX 9–10", t9 + 3600)
    w.step()
    rec = hub.rounds(1)[0]
    hc = [c for c in rec["candidates"] if c["basis"] == hid]
    pos = [p for p in inv.main.pos if p.basis == hid]
    ok &= _check("Inwestor: hipoteza w próbie daje kandydatów na żywo (mnożnik 0,25), uwagi Krytyka i poziom nauki w briefie",
                 len(hc) == 3 and pos and all(p.mult == 0.25 for p in pos) and "UWAGI KRYTYKA:" in seen["brief"]
                 and "wchodź małym ryzykiem" in seen["brief"] and "POZIOM NAUKI 2" in seen["brief"], f"kandydatów {len(hc)}, pozycji {len(pos)}")
    # bierność: rundy bez wejść → automat podnosi poziom do 3, a na poziomie 3 brak wejścia = błąd i ponowienie
    hub.set_meta("explore", {"level": 0, "by": "test", "reason": "test"})
    for i in range(5):
        t = t9 + 3600 * (3 + i)
        clock["t"] = t + 3600 + 90
        plan["investor"].append(lambda u, r: {"candidates": [{"ref": x, "action": "skip", "confidence": 50, "note": "nie"}
                                                             for x in re.findall(r"^(C\d+) ", u, flags=re.M)],
                                              "positions": [{"ref": x, "action": "hold", "new_sl": None, "note": "t"}
                                                            for x in re.findall(r"^(P\d+) ", u, flags=re.M)],
                                              "own_ideas": [], "round_note": "pomijam"})
        w.submit("fx", t, f"FX {i}", t + 3600)
        w.step()
    ex = inv.explore()
    rr = hub.rounds(5)
    ok &= _check("automat bierności: rundy bez wejścia podnoszą poziom nauki do 3 mimo ustawienia 0",
                 ex["level"] == 3 and ex["auto"] and ex["passive"] >= 4, f"poziom {ex['level']}, bez wejścia {ex['passive']}")
    ok &= _check("model wołany w każdej rundzie (także bez kandydatów); poziom 3: brak wejścia → ponowienie z błędem",
                 all(r.get("llm") for r in rr) and any((r.get("llm") or {}).get("attempts", 1) > 1 for r in rr[:2]),
                 ", ".join(str((r.get('llm') or {}).get('attempts')) for r in rr))

    # dyspozytor: Inwestor przed Krytykiem i Badaczem; dziennik pracy; status.json
    for _ in range(3):
        hub.hyp_add(dict(NOISE_HYP, sig=f"y{_}", status="do oceny", result={"n": 0}, at=clock["t"]))
    w.research_after = 0
    w.submit("fx", t9 + 3600 * 9, "FX x", t9 + 3600 * 10)
    plan["investor"].append(inv_reply)
    first = w.step()
    runs = hub.agent_runs(100)
    ag = {r["agent"] for r in runs}
    ok &= _check("dyspozytor: nowa świeca ma pierwszeństwo przed Krytykiem i Badaczem; każde zadanie w dzienniku pracy (czas, czas modelu)",
                 first == "investor" and {"investor", "critic", "researcher"} <= ag and all(r["sec"] is not None and r["llm_sec"] >= 0 for r in runs),
                 f"pierwsze: {first}, zadań {len(runs)}")

    class Sup:
        version, step, started_at, ollama, fx_ratings, fx_note = "t", "AI-3", clock["t"] - 7200, {}, None, None
        stocks = fx = None
        agent = w

        def reads_today(self):
            return 0
    Sup.hub = hub
    inv._persist(market.Prices(fx=fx), inv._risk_state(market.Prices(fx=fx), clock["t"]), None, [])
    doc = status.document(Sup(), clock["t"])
    txt = json.dumps(doc, allow_nan=False)
    wk = doc.get("work") or {}
    ok &= _check("status.json: praca agentów (statystyka 24 h na agenta, ostatnie zadania), hipotezy, Krytyk — bez NaN",
                 set(wk.get("stats", {})) >= {"investor", "critic", "researcher"} and wk["runs"] and doc["research"]["total"] >= 5
                 and doc["critic"]["explore"]["level"] == 3 and [a["id"] for a in doc["agents"]] == ["investor", "critic", "researcher"],
                 f"{len(txt) // 1024} KB")
    # branch ai: praca, hipotezy, przeglądy; odtworzenie huba
    store.save_work(hub, inv)
    d = store.repo.dir
    rows = lambda f: [json.loads(x) for x in (d / f).read_text().splitlines() if x.strip()]
    ok &= _check("branch ai: runs.jsonl, hypotheses.jsonl, reviews.jsonl (bez powtórek przy kolejnym zapisie)",
                 len(rows("runs.jsonl")) == len(hub.agent_runs(1000)) and len(rows("reviews.jsonl")) == len(hub.reviews(100))
                 and {h["id"] for h in rows("hypotheses.jsonl")} == {h["id"] for h in hub.hypotheses(100)}
                 and not store.save_work(hub, inv))
    store.save(inv, None, [], {}, market.Prices(fx=fx))
    h2 = Hub(tmp / "restored.sqlite")
    restore(h2, d, log=lambda *a: None)
    ok &= _check("odtworzenie huba: hipotezy ze statusami, przeglądy, poziom nauki",
                 {(h["id"], h["status"]) for h in h2.hypotheses(100)} == {(h["id"], h["status"]) for h in hub.hypotheses(100)}
                 and len(h2.reviews(100)) == len(hub.reviews(100)) and (h2.meta("explore") or {}).get("level") == 0)
    from . import excel
    from openpyxl import load_workbook
    excel.write(inv, inv.public.get("bench") or {}, market.Prices(fx=fx))
    wb = load_workbook(cfg["xlsx_path"], read_only=True)
    ok &= _check("AI-INVESTOR.xlsx: arkusze HIPOTEZY i PRZEGLĄDY wypełnione",
                 wb["HIPOTEZY"].max_row - 1 == len(hub.hypotheses(1000)) and wb["PRZEGLĄDY"].max_row - 1 == len(hub.reviews(100)))
    return ok


def _fx_round(inv, data, job):
    prices, _, _ = data.get(False)
    rec = inv.run(job["kind"], job["key"], job["label"], job["close_t"], prices, None, None)
    return {"ok": True, "summary": rec["id"]}
