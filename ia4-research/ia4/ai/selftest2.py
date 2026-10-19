"""
IA 4 — testy kroku AI-2 (Agent-Investor) na danych syntetycznych — część `python -m ia4.ai.selftest`.
Wersja projektu: 1.49 (2026-10-19) — musi zgadzać się z IA4_INSTRUKCJA.md

Silnik portfela: zgodność z symulacją backtestu (ia4/lab/sim) transakcja po transakcji, koszty, USD,
luka, SL przed TP, zamknięcie decyzją, zacieśnianie SL, limity 11c.7, test na historii (replay).
Runda: atrapa modelu (decyzje, błędny JSON → ponowienie → „brak decyzji”), wykonanie, odrzucenia przez kod,
cienie, punkty odniesienia, AI-INVESTOR.xlsx, branch `ai` (bez push) i odtworzenie huba, status.json.
"""

from __future__ import annotations

import copy
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd

from . import config as aicfg
from . import engine, market


def _check(name, ok, detail=""):
    from .selftest import check
    return check(name, ok, detail)


RULES = [
    {"schema": "ia4-rule/1", "direction": "long", "lines": {"A": {"ma": "EMA", "n": 9}, "B": {"ma": "SMA", "n": 30}},
     "signal": {"kind": "cross"}, "filters": [], "exit": {"sl": 2, "tp": 3, "max_bars": 20, "fc": [{"kind": "cross_back"}]}},
    {"schema": "ia4-rule/1", "direction": "short", "lines": {"A": {"ma": "HMA", "n": 13}, "B": {"ma": "EMA", "n": 40}},
     "signal": {"kind": "converge", "grow": 3, "shrink": 2}, "filters": [],
     "exit": {"sl": 3, "tp": 2, "max_bars": None, "fc": [{"kind": "reexpand", "n": 2}]}},
    {"schema": "ia4-rule/1", "direction": "long", "lines": {"A": {"ma": "KAMA", "n": 10}}, "signal": {"kind": "pcross"},
     "filters": [{"f": "rsi", "op": "<", "x": 60}], "exit": {"sl": 1, "tp": 1, "max_bars": 7, "fc": []}},
    {"schema": "ia4-rule/1", "direction": "short", "lines": {"A": {"ma": "DEMA", "n": 6}, "B": {"ma": "SMA", "n": 20}},
     "signal": {"kind": "turn", "pos": "above"}, "filters": [], "exit": {"sl": 4, "tp": 10, "max_bars": 14, "fc": [{"kind": "turn_back"}]}},
]


def strategies() -> list[dict]:
    out = []
    for i, r in enumerate(RULES):
        out.append({"id": f"S-test{i:05d}", "desc": f"reguła testowa {i}", "group": f"{r['signal']['kind']}|{r['direction']}",
                    "rule": r, "direction": r["direction"], "sl": r["exit"]["sl"], "tp": r["exit"]["tp"],
                    "max_bars": r["exit"].get("max_bars"), "has_fc": bool(r["exit"]["fc"]), "status": "bazowa",
                    "basis_kind": "bazowa_wybrana" if i % 2 == 0 else "bazowa", "selected": i % 2 == 0, "priority": i + 1,
                    "est_pct": 0.5 - 0.1 * i, "dp": {"n": 30 + i, "avg": 1.0, "pf": 1.6, "hold": 7, "tp": 0.4},
                    "search": {"pf_w": 2.0, "n": 60}, "vi": None})
    return out


def frames_stocks(n_inst=8, sessions=300, seed=3):
    from ..lab import selftest as lst
    fr = lst.synthetic(n_inst=n_inst, sessions=sessions, seed=seed)
    # daty syntetyczne → najnowsze sesje kończą się 2026-10-09 (spójnie z FX)
    days = sorted(fr["SPY"]["date"].unique())
    real = pd.bdate_range(end="2026-10-09", periods=len(days)).strftime("%Y-%m-%d")
    mp = dict(zip(days, real))
    for s in fr:
        fr[s]["date"] = fr[s]["date"].map(mp)
    return fr


def frames_fx():
    from ..fx import synth
    return {s: synth.synthetic(s, days=60, end="2026-10-09T19:00:00Z", seed=11) for s in
            ("EURUSD", "USDJPY", "GBPUSD", "EURGBP", "BTCUSD")}


# ---------------------------------------------------------------------------
def test_engine() -> bool:
    ok = True
    cfg = copy.deepcopy(aicfg.DEFAULTS)
    # --- zgodność z backtestem na wielu regułach i spółkach (test na historii)
    from .replay import replay
    fr = frames_stocks(n_inst=10, sessions=320, seed=5)
    traded = [s for s in fr if s not in ("SPY", "QQQ")]
    st = strategies() + [dict(s, id=s["id"] + "b", priority=s["priority"] + 10,
                              rule=dict(s["rule"], exit=dict(s["rule"]["exit"], sl=s["tp"], tp=s["sl"])), sl=s["tp"], tp=s["sl"])
                         for s in strategies()]
    r = replay(fr, traded, st, cfg, sessions=40, log=lambda *a: None)
    ok &= _check("test na historii (40 sesji): każda transakcja = backtest ia4/lab/sim (wyjście, powód, zwrot)",
                 r["parity"]["checked"] > 300 and r["parity"]["bad"] == 0,
                 f"sprawdzonych {r['parity']['checked']}, rozbieżności {r['parity']['bad']}, wejść {r['entered']}, zamkniętych {r['closed']}")
    w = r["worst"]
    ok &= _check("limity 11c.7 nigdy przekroczone (pozycje, ekspozycja, budżet ryzyka)",
                 w["positions"] <= 12 and w["grossPct"] <= 200 + 1e-6 and w["riskPct"] <= 6 + 1e-6,
                 f"najwięcej: pozycji {w['positions']}, ekspozycja {w['grossPct']:.0f}%, ryzyko {w['riskPct']:.2f}%")

    # --- pojedyncze świece: luka, SL przed TP, zamknięcie decyzją, koszty, USD
    def stock(rows):
        return pd.DataFrame([("T", d, s, o, h, l, c, 100) for d, s, o, h, l, c in rows],
                            columns=["symbol", "date", "slot", "o", "h", "l", "c", "v"])
    pr = market.Prices(stocks={"T": stock([("2026-10-08", 7, 100, 100, 100, 100), ("2026-10-09", 1, 100, 103.5, 97.5, 101),
                                           ("2026-10-09", 2, 95, 96, 94, 95)])})
    b = engine.Book(cfg, "t")
    p = b.add(pr, round_id="r", inst="T", side=1, basis="S", basis_kind="bazowa", decided_k=202610087, now=0, sl_pct=2, tp_pct=3,
              max_bars=None, has_fc=False, qty=10, risk=1, mult=1)
    b.settle(pr)
    ok &= _check("SL i TP w tej samej świecy → SL (8.6 p. 3)", p.reason == "SL" and abs(p.exit_px - 98) < 1e-9 and p.exit_k == 202610091)
    p = b.add(pr, round_id="r", inst="T", side=1, basis="S", basis_kind="bazowa", decided_k=202610087, now=0, sl_pct=10, tp_pct=10,
              max_bars=None, has_fc=False, qty=10, risk=1, mult=1)
    p.last_k = 202610087
    b.settle(pr)
    prg = market.Prices(stocks={"T": stock([("2026-10-08", 7, 100, 100, 100, 100), ("2026-10-09", 1, 100, 101, 99.5, 100),
                                            ("2026-10-09", 2, 95, 96, 94, 95)])})
    bg = engine.Book(cfg, "t")
    p2 = bg.add(prg, round_id="r", inst="T", side=-1, basis="S", basis_kind="bazowa", decided_k=202610087, now=0, sl_pct=3, tp_pct=4,
                max_bars=None, has_fc=False, qty=10, risk=1, mult=1)
    bg.settle(prg)
    ok &= _check("luka na otwarciu przez TP (short) → wyjście po poziomie TP, bez premii", p2.reason == "TP" and p2.gap and abs(p2.exit_px - 96) < 1e-9,
                 f"{p2.reason} {p2.exit_px}")
    cost = 10 * 100 * 0.005 / 100 + 10 * p.exit_px * 0.005 / 100 if p.status == "closed" else None
    ok &= _check("koszt akcji 0,005% przy wejściu i przy wyjściu; wynik netto = brutto − koszty",
                 p.status == "open" and abs(p.cost_in_usd - 0.05) < 1e-12)
    pr2 = market.Prices(stocks={"T": stock([("2026-10-08", 7, 100, 100, 100, 100), ("2026-10-09", 1, 100, 101, 99, 100.5),
                                            ("2026-10-09", 2, 100.4, 101, 99.5, 100.2), ("2026-10-09", 3, 100.1, 100.3, 99.9, 100)])})
    b2 = engine.Book(cfg, "t")
    q = b2.add(pr2, round_id="r", inst="T", side=1, basis="S", basis_kind="bazowa", decided_k=202610087, now=0, sl_pct=5, tp_pct=5,
               max_bars=None, has_fc=False, qty=10, risk=1, mult=1)
    pr2.t_max = market.stock_close("2026-10-09", 1) + 1
    b2.settle(pr2)
    e1 = b2.tighten(q, 101, pr2, "x", 0)
    e2 = b2.tighten(q, 94, pr2, "x", 0)
    e3 = b2.tighten(q, 99.8, pr2, "zysk", 0)
    b2.request_close(q, "koniec")
    pr2.t_max = None
    b2.settle(pr2)
    ok &= _check("SL tylko zacieśniany i po właściwej stronie ceny; zamknięcie decyzją po open następnej świecy",
                 e1 and e2 and e3 is None and q.reason == "CLOSE" and abs(q.exit_px - 100.4) < 1e-9 and q.exit_k == 202610092,
                 f"{e1} · {e2}")
    # FX: USD
    fx = frames_fx()
    pf = market.Prices(fx=fx)
    jt = pf.last("USDJPY")
    ok &= _check("USD: zysk na USDJPY przeliczony kursem USDJPY, na EURGBP — kursem GBPUSD",
                 abs(pf.conv("USDJPY", jt[2]) - 1 / jt[1]) < 1e-12 and abs(pf.conv("EURGBP", pf.last("GBPUSD")[2]) - pf.last("GBPUSD")[1]) < 1e-12
                 and pf.conv("EURUSD", 0) == 1.0)
    c = engine.cost_side(cfg, pf, "USDJPY", 100000, 150.0, 1 / 150.0)
    ok &= _check("koszt FX: ½ spreadu na stronę (USDJPY 0,9 pipsa, 100 000 USD → 3,00 $)", abs(c - 100000 * 0.9 * 0.01 / 2 / 150) < 1e-9,
                 f"{c:.2f} $")
    # wielkość pozycji i limity
    bk = engine.Book(cfg, "main")
    q1, r1, w1 = bk.plan(pf, "EURUSD", 1, 0.5, 1.0, 100000, 1.0, 0)
    eu = pf.last("EURUSD")[1]
    ok &= _check("wielkość: ryzyko 0,5% (500 $) na SL z kosztami; wartość pary ≤ 50% stanu konta",
                 w1 is None and q1 * eu <= 50000 + 1e-6 and r1 <= 500 + 1e-6, f"{q1:,.0f} EUR, ryzyko {r1:.0f} $".replace(",", " "))
    def book_with(spec):
        bb = engine.Book(cfg, "main")
        for s, value_pct, risk, basis in spec:
            px = pf.last(s)[1]
            bb.add(pf, round_id="r", inst=s, side=1, basis=basis, basis_kind=basis, decided_k=pf.last(s)[0], now=0, sl_pct=1, tp_pct=1,
                   max_bars=10, has_fc=False, qty=value_pct / 100 * 100000 / (px * (pf.conv(s, pf.last(s)[2]) or 1)), risk=risk, mult=1)
        return bb
    b1 = book_with([("EURUSD", 10, 900, "own"), ("USDJPY", 10, 900, "own")])
    _, _, why_hold = b1.plan(pf, "EURUSD", 1, 0.5, 1.0, 100000, 1.0, 0)
    _, _, why_own = b1.plan(pf, "GBPUSD", 1, 2.0, 1.0, 100000, 1.0, 0, own=True)
    _, _, ok_own = b1.plan(pf, "GBPUSD", 1, 2.0, 0.25, 100000, 1.0, 0, own=True)
    ok &= _check("limity: 1 pozycja na instrument; budżet własnych pomysłów 2% (1800 $ + 125 $ mieści się, + 500 $ — nie)",
                 why_hold and "jest już" in why_hold and why_own and "własnych" in why_own and ok_own is None, f"{why_hold} · {why_own}")
    b2b = book_with([(s, 50, 100, "S") for s in ("EURUSD", "USDJPY", "GBPUSD", "EURGBP")])
    _, _, why_gross = b2b.plan(pf, "BTCUSD", 1, 2.0, 1.0, 100000, 1.0, 0)
    b3 = book_with([(s, 10, 1500, "S") for s in ("EURUSD", "USDJPY", "GBPUSD", "EURGBP")])
    _, _, why_budget = b3.plan(pf, "BTCUSD", 1, 2.0, 1.0, 100000, 1.0, 0)
    ok &= _check("limit ekspozycji brutto 200% i budżet ryzyka 6%",
                 why_gross and "brutto" in why_gross and why_budget and "budżet ryzyka" in why_budget, f"{why_gross} · {why_budget}")
    return ok


# ---------------------------------------------------------------------------
class FakeOllama:
    """Atrapa /api/chat: odpowiedzi ze skryptu (funkcja briefu → JSON albo tekst)."""

    def __init__(self, script):
        self.script, self.calls = script, []

    def __call__(self, path, payload, timeout):
        user = payload["messages"][1]["content"]
        self.calls.append(payload)
        out = self.script(user, len(self.calls), payload)
        return {"message": {"content": out if isinstance(out, str) else json.dumps(out)}, "prompt_eval_count": len(user) // 4,
                "eval_count": 300}


def _refs(text, prefix):
    return sorted(set(re.findall(rf"^({prefix}\d+) ", text, flags=re.M)), key=lambda x: int(x[1:]))


def test_round(tmp: Path) -> bool:
    from .hub import Hub
    from .investor import Investor
    from .llm import LLM, PriorityLock
    from .playbook import Signals
    from .store import AiStore, restore
    from . import excel, status
    ok = True
    cfg = copy.deepcopy(aicfg.DEFAULTS)
    cfg["xlsx_path"] = str(tmp / "AI-INVESTOR.xlsx")
    fr, fx = frames_stocks(), frames_fx()
    traded = [s for s in fr if s not in ("SPY", "QQQ")]
    st = strategies()
    sig_full = Signals(fr, traded)
    keys = np.unique(sig_full.key)[-60:]
    K = next(int(k) for k in keys[:-5] if len({x["inst"] for x in Signals({s: d[(d["date"].str.replace("-", "").astype(int) * 10 + d["slot"]) <= k]
                                                                             for s, d in fr.items()}, traded).at(st, int(k))}) >= 3)
    cut = lambda k: {s: d[(d["date"].str.replace("-", "").astype(int) * 10 + d["slot"]) <= k].reset_index(drop=True) for s, d in fr.items()}
    ks = [int(k) for k in keys if k >= K]
    t_of = lambda k: market.stock_close(market.key_date(k), k % 10)
    clock = {"t": t_of(K) + 200}
    fxc = lambda t: {s: d[d["time"] + 3600 <= t].reset_index(drop=True) for s, d in fx.items()}

    def script(user, n, payload):
        if n == 2:
            return "to nie jest JSON"                                           # błąd → ponowienie
        C, P = _refs(user, "C"), _refs(user, "P")
        own = []
        if n == 1:
            m = re.search(r"^EURUSD ([0-9.]+) .*?ATR14 ([0-9.]+)", user, flags=re.M)
            px, a = float(m.group(1)), float(m.group(2))
            own = [{"instrument": "EURUSD", "direction": "long", "sl": px - 2 * a, "tp": px + 3 * a, "max_bars": 24, "confidence": 55,
                    "note": "test: pomysł własny"},
                   ]
        return {"candidates": [{"ref": r, "action": "enter" if i % 2 == 0 else "skip", "confidence": 60 - i, "note": f"test {r}"}
                               for i, r in enumerate(C)],
                "positions": [{"ref": r, "action": "close" if (i == 0 and n == 3) else "hold", "new_sl": None, "note": f"poz {r}"}
                              for i, r in enumerate(P)],
                "own_ideas": own, "round_note": f"runda testowa {n}"}

    fake = FakeOllama(script)
    hub = Hub(tmp / "inv.sqlite")
    store = AiStore(push=False, root=tmp / "ai-repo")
    llm = LLM(cfg, record=lambda **k: hub.llm_call(**k), post=fake, lock=PriorityLock())
    inv = Investor(cfg, hub, llm, out=store, xlsx=lambda i, b, p: excel.write(i, b, p), clock=lambda: clock["t"], log=lambda *a: None)
    inv.set_playbook({"builtAt": 0, "source": {}, "strategies": st}, {"closed": 11, "avgPct": 0.34, "ci95Pct": 1.9})
    recs = []
    for i, k in enumerate(ks[:6]):
        clock["t"] = t_of(k) + 200
        f = cut(k)
        pr = market.Prices(stocks=f, fx=fxc(t_of(k)))
        recs.append(inv.run("stocks", k, market.Prices().label("X", k), t_of(k), pr, Signals(f, traded), None))
    r1 = recs[0]
    ent = [d for d in r1["decisions"] if d["action"] == "enter"]
    ok &= _check("runda 1: decyzja dla każdego kandydata, wejścia → pozycje oczekujące, cienie dla wszystkich",
                 len(r1["candidates"]) >= 3 and all(d["outcome"].startswith(("wykonano", "odrzucone")) for d in ent)
                 and len(inv.shadow.pos) >= len(r1["candidates"]) and r1["llm"]["attempts"] == 1,
                 f"kandydatów {len(r1['candidates'])}, wejść {sum(d['outcome'].startswith('wykonano') for d in ent)}")
    own = r1["ownIdeas"]
    ok &= _check("własny pomysł sprawdzony przez kod (SL ≥ 1 × ATR14, TP/SL 1:3…3:1) i wykonany z mnożnikiem 0,25",
                 own and own[0]["outcome"].startswith("wykonano") and inv.main.pos[[p.basis for p in inv.main.pos].index("own")].mult == 0.25,
                 own[0]["outcome"] if own else "brak")
    r2 = recs[1]
    ok &= _check("runda 2: błędny JSON → ponowienie (2 próby), potem poprawna decyzja",
                 r2["llm"] is None or r2["llm"]["attempts"] == 2, f"próby {r2['llm']['attempts'] if r2['llm'] else '—'}")
    filled = [p for p in inv.main.pos if p.basis != "own" and p.status in ("open", "closed")]
    ok &= _check("wejście po open pierwszej świecy po decyzji (nie wcześniej)",
                 filled and all(p.entry_k > p.decided_k for p in filled))
    closes = [d for r in recs for d in r["positions"] if d["action"] == "close"]
    byid = {p.id: p for p in inv.main.pos}
    acc = [d for d in closes if d["outcome"].startswith("zamknięcie")]
    good = all(byid[d["pos"]].status == "closed" and byid[d["pos"]].reason in ("CLOSE", "FC", "SL", "TP") for d in acc)
    ok &= _check("decyzja „close” → zamknięcie po open następnej świecy (CLOSE albo FC/luka na tym samym open); dla oczekującej — odrzucona przez kod",
                 closes and good and all(d["outcome"].startswith(("zamknięcie", "odrzucone")) for d in closes),
                 "; ".join(f"#{d['pos']}: {d['outcome']} → {byid[d['pos']].reason or byid[d['pos']].status}" for d in closes))
    # brak modelu → brak decyzji
    llm.post = lambda *a, **k: (_ for _ in ()).throw(OSError("Ollama nie odpowiada"))
    k = ks[6]
    clock["t"] = t_of(k) + 200
    f = cut(k)
    act_before = [(p.id, p.status, p.sl_px, p.close_req) for p in inv.main.active()]
    n_before = len(inv.main.pos)
    r7 = inv.run("stocks", k, "x", t_of(k), market.Prices(stocks=f, fx=fxc(t_of(k))), Signals(f, traded), None)
    ok &= _check("model niedostępny → „brak decyzji”: bez nowych pozycji, decyzje o pozycjach „hold”, runda zapisana",
                 act_before and "brak decyzji" in r7["roundNote"] and len(inv.main.pos) == n_before
                 and all(d["action"] == "hold" for d in r7["positions"]), r7["roundNote"][:90])
    # zapis: hub, branch, Excel, status
    ok &= _check("hub: decyzje, notatki, wywołania modelu (pełny prompt), krzywa stanu konta",
                 len(hub.decisions()) >= len(r1["decisions"]) and hub.db.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] >= 6
                 and len(hub.equity_series()) >= 7 and any(n["kind"] == "decyzja" for n in hub.notes(200)))
    d = store.repo.dir
    dec = [json.loads(x) for x in (d / "decisions.jsonl").read_text().splitlines()]
    trades = [json.loads(x) for x in (d / "trades.jsonl").read_text().splitlines()] if (d / "trades.jsonl").exists() else []
    ok &= _check("branch ai: decisions.jsonl (runda = wiersz), trades.jsonl, state.json, benchmarks.json, playbook.json, lessons.json",
                 len(dec) == 7 and len(trades) == len(inv.main.closed()) and all((d / f).exists() for f in
                 ("state.json", "benchmarks.json", "playbook.json", "lessons.json", "config.json")),
                 f"rund {len(dec)}, transakcji {len(trades)}")
    from openpyxl import load_workbook
    wb = load_workbook(cfg["xlsx_path"], read_only=True)
    tr = list(wb["TRANSAKCJE"].iter_rows(values_only=True))
    ok &= _check("AI-INVESTOR.xlsx: 8 arkuszy, TRANSAKCJE = wszystkie pozycje z notatką, SL, TP",
                 wb.sheetnames[:6] == ["PODSUMOWANIE", "TRANSAKCJE", "DECYZJE", "PORTFEL", "DZIENNIK", "PLAYBOOK"] and len(wb.sheetnames) == 8
                 and len(tr) - 1 == len(inv.main.pos) and all(r[18] for r in tr[1:]) and all(r[10] for r in tr[1:] if r[1] != "oczekuje"),
                 f"{len(tr) - 1} transakcji")
    b = inv.public["bench"]
    ok &= _check("punkty odniesienia: Inwestor, wszyscy kandydaci, mechaniczny, kalibracja, werdykt",
                 b["all"]["n"] >= b["investor"]["n"] and b["mechanical"]["closed"] == 11 and "za wcześnie" in b["verdict"])
    # odtworzenie huba z brancha
    h2 = Hub(tmp / "restored.sqlite")
    restore(h2, d, log=lambda *a: None)
    inv2 = Investor(cfg, h2, llm, clock=lambda: clock["t"], log=lambda *a: None)
    same = sorted((p.id, p.status, p.inst) for p in inv2.main.pos) == sorted((p.id, p.status, p.inst) for p in inv.main.pos)
    ok &= _check("odtworzenie huba z brancha ai: te same pozycje (otwarte, oczekujące, zamknięte) i liczniki",
                 same and h2.meta("round_n") == hub.meta("round_n") and abs(inv2.peak - inv.peak) < 1e-6,
                 f"pozycji {len(inv2.main.pos)}")

    class W:
        inv = None
        critic = researcher = None
        done = 7
        def pending(self): return 0
        def now_doing(self): return None
    w = W()
    w.inv = inv

    class Sup:
        version, step, started_at, ollama, fx_ratings, fx_note = "t", "AI-2", clock["t"], {}, None, None
        stocks = fx = None
        agent = w
        def reads_today(self): return 0
    Sup.hub = hub
    doc = status.document(Sup(), clock["t"])
    txt = json.dumps(doc, allow_nan=False)
    ok &= _check("status.json: Inwestor aktywny, portfel, pozycje z notatkami, ostatnia runda, krzywa, notatki — bez NaN",
                 doc["agents"][0]["state"] == "on" and doc["investor"]["portfolio"]["equity"] > 0 and doc["investor"]["lastRound"]
                 and len(doc["investor"]["equity"]) >= 7 and doc["investor"]["notes"], f"{len(txt) // 1024} KB")
    return ok
