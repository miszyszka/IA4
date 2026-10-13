"""
IA 4 — AI-INVESTOR.xlsx: dziennik Agent-Investora w Excelu (instrukcja, sekcja 11c.11).
Wersja projektu: 1.38 (2026-10-13) — musi zgadzać się z IA4_INSTRUKCJA.md

Przepisywany w całości po każdej rundzie (zapis atomowy: plik tymczasowy → zamiana; plik otwarty
w Excelu/Numbers pokazuje stan z chwili otwarcia). Ścieżka: ai/config.json → xlsx_path (domyślnie
~/Desktop/IA4/AI-INVESTOR.xlsx; w folderze Dysku Google ten sam plik widać w Arkuszach Google).
Arkusze: PODSUMOWANIE, TRANSAKCJE, DECYZJE, PORTFEL, DZIENNIK, PLAYBOOK, HIPOTEZY (AI-3), PRZEGLĄDY (AI-4).
Pełna historia jest też na branchu `ai`; DECYZJE, PORTFEL i DZIENNIK — ostatnie 5000 wierszy.
"""

from __future__ import annotations

import os
import tempfile
from datetime import datetime
from pathlib import Path

from .config import xlsx_path
from .market import PL

MAX_ROWS = 5000
HEAD_FILL = "1F2A37"
HEAD_FONT = "FFFFFF"


def _dt(t):
    if t is None:
        return None
    return datetime.fromtimestamp(float(t), PL).replace(tzinfo=None)


def _sheet(wb, title, headers, rows, widths, wrap=(), fmts=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color=HEAD_FONT)
        c.fill = PatternFill("solid", fgColor=HEAD_FILL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    fmts = fmts or {}
    for col, fmt in fmts.items():
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = fmt
    for col in wrap:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def write(inv, bench: dict, prices, path: Path | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    path = path or xlsx_path(inv.cfg)
    pub = inv.public or {}
    pf = pub.get("portfolio") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "PODSUMOWANIE"
    ws["A1"] = "IA 4 — AI Investor (Agent-Investor)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"stan z {datetime.now(PL):%Y-%m-%d %H:%M} PL · wszystko na papierze · instrukcja, sekcja 11c"
    rows = [("Kapitał startowy $", inv.cfg["capital"]), ("Od", _dt(inv.started_at)), ("Stan konta $", pf.get("equity")),
            ("Wynik zrealizowany $", pf.get("realized")), ("Wynik niezrealizowany $", pf.get("unrealized")),
            ("Szczyt stanu konta $", pf.get("peak")), ("Obsunięcie od szczytu %", pf.get("ddPct")), ("Wynik doby UTC %", pf.get("dayPct")),
            ("Ekspozycja brutto %", pf.get("grossPct")), ("Ryzyko otwarte %", pf.get("riskPct")),
            ("Pozycje otwarte / oczekujące", f"{pf.get('open', 0)} / {pf.get('pending', 0)}"), ("Zamkniętych transakcji", pf.get("closed")),
            ("Koszty transakcyjne $", pf.get("costs")), ("risk_dial", pf.get("dial")),
            ("Blokady", "; ".join(pf.get("blocks") or []) or "brak")]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
        if isinstance(v, float):
            ws.cell(i, 2).number_format = "#,##0.00"
        if isinstance(v, datetime):
            ws.cell(i, 2).number_format = "yyyy-mm-dd hh:mm"
    r0 = 4 + len(rows) + 1
    ws.cell(r0, 1, "Punkty odniesienia (11c.10)").font = Font(bold=True)
    b = bench or {}
    bl = [("Inwestor — transakcje z playbooka", b.get("investor")), ("Inwestor — własne pomysły", b.get("own")),
          ("Wszyscy kandydaci (cienie, bez limitów)", b.get("all"))]
    ws.cell(r0 + 1, 1, "")
    for j, h in enumerate(("transakcji", "śr. wynik % netto", "± 95%", "zyskownych"), start=2):
        ws.cell(r0 + 1, j, h).font = Font(bold=True)
    for i, (k, s) in enumerate(bl, start=r0 + 2):
        s = s or {}
        ws.cell(i, 1, k)
        for j, key in enumerate(("n", "avg", "ci95", "win"), start=2):
            ws.cell(i, j, s.get(key))
            ws.cell(i, j).number_format = "0.00" if key != "win" else "0%"
    i += 1
    rnd = b.get("random")
    ws.cell(i, 1, "Losowy wybór (200 losowań)")
    ws.cell(i, 2, f"percentyl Inwestora {rnd['pctl']:.0f} · losowo śr. {rnd['randomAvg']:+.2f}% (95%: {rnd['lo']:+.2f}…{rnd['hi']:+.2f})"
            if rnd else "za mało rund z zamkniętymi cieniami")
    i += 1
    mech = b.get("mechanical") or {}
    ws.cell(i, 1, "Mechaniczny inwestor (Apps Script)")
    ws.cell(i, 2, f"{mech.get('closed')} zamkniętych · śr. {mech.get('avgPct'):+.2f}%" if mech.get("avgPct") is not None else "brak danych")
    i += 1
    spy = b.get("spy")
    ws.cell(i, 1, "SPY kup i trzymaj od startu")
    ws.cell(i, 2, f"{spy['pct']:+.2f}%" if spy else "—")
    i += 1
    ws.cell(i, 1, "Werdykt")
    ws.cell(i, 2, b.get("verdict"))
    i += 2
    ws.cell(i, 1, "Kalibracja pewności").font = Font(bold=True)
    for c in b.get("calibration") or []:
        i += 1
        ws.cell(i, 1, f"pewność {c['from']}–{c['to']}")
        ws.cell(i, 2, f"{c['n']} transakcji · deklarowana {c['conf']:.0%} · zyskownych {c['win']:.0%}")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 70
    for col in "CDE":
        ws.column_dimensions[col].width = 16

    # TRANSAKCJE
    allp = inv.main.active() + sorted(inv.main.closed(), key=lambda p: -(p.exit_t or 0))
    st = {"pending": "oczekuje", "open": "otwarta", "closed": "zamknięta"}
    rows = []
    for p in allp:
        m = inv.main.mark(p, prices) if p.status == "open" else {}
        rows.append((p.id, st[p.status], _dt(p.decided_at), _dt(p.entry_t), p.inst, p.kind, p.direction, p.basis, p.confidence,
                     p.entry_px, p.sl_px, p.tp_px, p.sl_pct, p.tp_pct, p.qty, p.value_usd, p.risk_usd, p.max_bars, p.note_in,
                     _dt(p.exit_t), p.exit_px if p.status == "closed" else m.get("px"), p.reason or ("teraz" if p.status == "open" else None),
                     p.pnl_net_usd if p.status == "closed" else m.get("pnl"), p.ret_net_pct if p.status == "closed" else m.get("pct"),
                     p.r_mult if p.status == "closed" else m.get("R"), p.cost_in_usd + p.cost_out_usd, p.bars,
                     (p.note_out + (" · SL: " + "; ".join(f"{x['from']:.6g}→{x['to']:.6g} ({x['note'][:80]})" for x in p.sl_moves)
                                    if p.sl_moves else "")).strip(" ·")))
    _sheet(wb, "TRANSAKCJE",
           ["Nr", "Status", "Decyzja (PL)", "Wejście (PL)", "Instrument", "Rynek", "Kierunek", "Podstawa", "Pewność", "Cena wejścia",
            "SL", "TP", "SL %", "TP %", "Wielkość", "Wartość $", "Ryzyko $", "Limit świec", "Notatka — dlaczego wchodzę",
            "Wyjście (PL)", "Cena wyjścia / teraz", "Powód wyjścia", "Wynik $ netto", "Wynik % netto", "R", "Koszty $", "Świec",
            "Notatka — wyjście / zmiany SL"], rows,
           [6, 10, 16, 16, 10, 8, 8, 14, 8, 12, 12, 12, 7, 7, 12, 12, 10, 8, 60, 16, 12, 10, 12, 10, 7, 9, 7, 50], wrap=(19, 28),
           fmts={3: "yyyy-mm-dd hh:mm", 4: "yyyy-mm-dd hh:mm", 13: "0.00", 14: "0.00", 15: "#,##0.####", 16: "#,##0.00", 17: "#,##0.00",
                 20: "yyyy-mm-dd hh:mm", 23: "#,##0.00", 24: "0.00", 25: "0.00", 26: "#,##0.00"})

    # DECYZJE
    dec = inv.hub.decisions(MAX_ROWS)
    _sheet(wb, "DECYZJE", ["Runda", "Czas (PL)", "Ref", "Instrument", "Kierunek", "Podstawa", "Decyzja", "Pewność", "Notatka", "Skutek", "Pozycja"],
           [(d["round_id"], _dt(d["at"]), d["ref"], d["inst"], "long" if (d["side"] or 0) > 0 else "short", d["basis"], d["action"],
             d["confidence"], d["note"], d["outcome"], d["pos_id"]) for d in dec],
           [9, 16, 6, 10, 8, 14, 11, 8, 70, 45, 8], wrap=(9, 10), fmts={2: "yyyy-mm-dd hh:mm"})

    # PORTFEL
    eqs = inv.hub.equity_series(MAX_ROWS)[::-1]
    _sheet(wb, "PORTFEL", ["Czas (PL)", "Runda", "Stan konta $", "Zrealizowany $", "Niezrealizowany $", "Ekspozycja %", "Ryzyko %",
                           "Pozycji", "risk_dial", "Obsunięcie %"],
           [(_dt(e["at"]), e["round_id"], e["equity"], e["realized"], e["unrealized"], 100 * e["gross"] / e["equity"] if e["equity"] else None,
             100 * e["risk"] / e["equity"] if e["equity"] else None, e["positions"], e["dial"], e["dd"]) for e in eqs],
           [16, 9, 14, 14, 14, 12, 10, 8, 9, 12],
           fmts={1: "yyyy-mm-dd hh:mm", 3: "#,##0.00", 4: "#,##0.00", 5: "#,##0.00", 6: "0.0", 7: "0.00", 10: "0.00"})

    # DZIENNIK
    notes = inv.hub.notes(MAX_ROWS)
    _sheet(wb, "DZIENNIK", ["Czas (PL)", "Agent", "Rodzaj", "Treść"],
           [(_dt(n["at"]), n["agent"], n["kind"], n["text"]) for n in notes], [16, 10, 10, 120], wrap=(4,), fmts={1: "yyyy-mm-dd hh:mm"})

    # PLAYBOOK
    from .store import playbook_view
    pv = playbook_view(inv)["strategies"]
    _sheet(wb, "PLAYBOOK", ["Pierw.", "ID", "Status", "Mnożnik ryzyka", "Grupa", "Opis reguły", "SL %", "TP %", "Limit świec",
                            "Szacunek przewagi %", "doubleProof: trans.", "doubleProof: śr. %", "doubleProof: PF",
                            "Mechaniczny na żywo: trans.", "Mechaniczny: śr. %", "AI: trans.", "AI: śr. %", "AI: $"],
           [(s["priority"], s["id"], s["status"], s["mult"], s["group"], s["desc"], s["sl"], s["tp"], s["max_bars"], s["est_pct"],
             (s["dp"] or {}).get("n"), (s["dp"] or {}).get("avg"), (s["dp"] or {}).get("pf"), (s["vi"] or {}).get("n"),
             (s["vi"] or {}).get("avg"), s["ai"]["n"], s["ai"]["avg"], s["ai"]["usd"]) for s in pv],
           [7, 14, 10, 9, 14, 70, 6, 6, 8, 10, 10, 10, 10, 10, 10, 8, 9, 10], wrap=(6,),
           fmts={10: "0.00", 12: "0.00", 13: "0.00", 15: "0.00", 17: "0.00", 18: "#,##0.00"})
    _sheet(wb, "HIPOTEZY", ["Hipotezy Strategy-Makera — od kroku AI-3"], [], [60])
    _sheet(wb, "PRZEGLĄDY", ["Przeglądy Strategy-Managera — od kroku AI-4"], [], [60])

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".AI-INVESTOR-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return path
