"""
IA 4 — strategie jednokierunkowe z SL i TP, które towarzyszą wybiciom (instrukcja, sekcja 7a).
Wersja projektu: 1.59 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Zadanie człowieka (16.10.2026): skoro siłę ruchu da się przewidzieć, a kierunku nie — szukać okoliczności
(średnie kroczące mówiące o trendzie, konwergencja wyznaczająca poziomy, wybicie w górę → long, w dół → short
albo odwrotnie), które dają zysk w JEDNĄ stronę, zawsze z SL i TP.

Siatka (każda strategia = setup × filtr × kierunek × wyjście):
  • 20 par setupów (wersja „w górę” i lustrzana „w dół”, czyli 40): wybicie szczytu / dołka 10, 20, 55, 252 sesji;
    wybicie z ciasnego przedziału (10, 20 sesji) i po konsolidacji; wybicie ze ściśnięcia Bollingera i ze ściśnięcia
    średnich 10/20/50 (konwergencja); wybicie dnia wewnętrznego i NR7; duży dzień (≥ 2 R i ≥ 3 R) z zamknięciem przy
    końcu zakresu; luka ≥ 1 R utrzymana do close; przecięcia SMA20/50, SMA50/200 i ceny z SMA200; korekta w trendzie
    (RSI(2) < 10 przy SMA50 > SMA200); 3 wzrosty z rzędu i szczyt 20 sesji; duży dzień po konsolidacji;
  • 13 filtrów (bez filtra, cena wobec SMA200, SMA50 wobec SMA200, nachylenie SMA200, SPY wobec SMA200,
    zmienność bezwzględna R ≥ 2% / 3%, siła względem SPY z 20 sesji);
  • kierunek long albo short — każdy setup w obie strony (za wybiciem i przeciw niemu);
  • wyjście: SL 0,5 / 1 / 1,5 / 2 R od wejścia, TP 1 / 2 / 3 × SL, najwyżej 1 / 3 / 5 / 10 sesji (potem po close).
  Razem 40 × 13 × 2 × 48 = 49 920 strategii. R = średni zakres dnia z 100 sesji (normalna dynamika spółki).

Rozliczenie (moduł `swing`): wejście po open sesji po sygnale, pesymistycznie (zasada 5), spread 0,3% raz na
transakcję, jedna pozycja na spółkę naraz (sygnały w czasie pozycji pomijane).
Ocena (jak w badaniach One-Day): 53 spółki, 20 okresów kalendarza — najwyżej 3 stratne i najwyżej 4 puste,
≥ 30 transakcji, suma netto > 0, część A (< 2013) i B (≥ 2013) dodatnie, lepiej niż wejście na ślepo (ten sam
kierunek i wyjście każdego dnia); potwierdzenie ✓ na 20 spółkach doubleProof (t ≥ próg Bonferroniego dla liczby
reguł w arkuszu, wynik i przewaga nad ślepym wejściem > 0); placebo — każdy warunek przesunięty w czasie o losową
liczbę ≥ 252 sesji (ta sama częstość, zero związku z dniem): ile reguł przechodzi z czystego przypadku.

    python -m ia4.swtrend              # baza dobowa od 2000 (najpierw python -m ia4.daily); kilka minut
    python -m ia4.swtrend --selftest   # „WYNIK: OK”
Wynik: ~/Desktop/IA4/SWING-TREND.xlsx (poza gitem).
"""

from __future__ import annotations

import argparse
import math
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np

from . import __version__, config
from . import odbreak as ob
from . import odsearch as osr
from . import swing as sw

SL_R = (0.5, 1.0, 1.5, 2.0)
TP_X = (1, 2, 3)
HOLDS = (1, 3, 5, 10)
N_PERIODS, MAX_LOSING, MAX_EMPTY = 20, 3, 4
MIN_TRADES = 30
WARM = 272                     # sesji rozgrzewki (SMA200, percentyl Bollingera z 252 sesji)
PLACEBO_SEED = 2026
TOP = 200
PER_SETUP = 3                  # w arkuszu najwyżej 3 reguły z tego samego setupu i kierunku
PF_SHOW = 3


# ---------------------------------------------------------------------------
#  Cechy i warunki (na ciągu pełnych sesji jednej spółki; wartość w q — znana po close q)
# ---------------------------------------------------------------------------
PAIRS = [
    ("hi10", "close nad szczytem 10 sesji", "close pod dołkiem 10 sesji"),
    ("hi20", "close nad szczytem 20 sesji", "close pod dołkiem 20 sesji"),
    ("hi55", "close nad szczytem 55 sesji", "close pod dołkiem 55 sesji"),
    ("hi252", "close nad szczytem roku (252 sesje)", "close pod dołkiem roku"),
    ("box10", "wybicie w górę z ciasnego przedziału 10 sesji (< 0,8 R·√10)", "wybicie w dół z ciasnego przedziału 10 sesji"),
    ("box20", "wybicie w górę z ciasnego przedziału 20 sesji", "wybicie w dół z ciasnego przedziału 20 sesji"),
    ("cr10", "szczyt 10 sesji po konsolidacji (średni zakres 10 sesji < 0,7 R)", "dołek 10 sesji po konsolidacji"),
    ("bbsq", "ściśnięcie Bollingera (najwęższe 20% roku) i close nad górną wstęgą", "ściśnięcie Bollingera i close pod dolną wstęgą"),
    ("masq", "konwergencja SMA10/20/50 (< 0,5 R) i close przebija wszystkie w górę", "konwergencja SMA10/20/50 i close przebija w dół"),
    ("in", "dzień wewnętrzny wczoraj, dziś close nad jego szczytem", "dzień wewnętrzny wczoraj, dziś close pod jego dołkiem"),
    ("nr7", "NR7 wczoraj, dziś close nad jego szczytem", "NR7 wczoraj, dziś close pod jego dołkiem"),
    ("big2", "duży dzień wzrostowy: zakres ≥ 2 R, close w górnych 30%", "duży dzień spadkowy: zakres ≥ 2 R, close w dolnych 30%"),
    ("big3", "bardzo duży dzień wzrostowy: zakres ≥ 3 R, close w górnych 30%", "bardzo duży dzień spadkowy: ≥ 3 R, close w dolnych 30%"),
    ("gap", "luka w górę ≥ 1 R utrzymana (close > open)", "luka w dół ≥ 1 R utrzymana (close < open)"),
    ("x2050", "SMA20 przecina SMA50 w górę", "SMA20 przecina SMA50 w dół"),
    ("x50200", "SMA50 przecina SMA200 w górę (złoty krzyż)", "SMA50 przecina SMA200 w dół (krzyż śmierci)"),
    ("x200", "close przecina SMA200 w górę", "close przecina SMA200 w dół"),
    ("dip", "korekta w trendzie wzrostowym: SMA50 > SMA200, close > SMA200, RSI(2) < 10",
     "odbicie w trendzie spadkowym: SMA50 < SMA200, close < SMA200, RSI(2) > 90"),
    ("mom3", "3 wzrostowe zamknięcia z rzędu i szczyt 20 sesji", "3 spadkowe zamknięcia z rzędu i dołek 20 sesji"),
    ("bigcon", "duży dzień wzrostowy (≥ 2 R) po konsolidacji 10 sesji", "duży dzień spadkowy (≥ 2 R) po konsolidacji 10 sesji"),
]
FILTERS = [
    ("—", "bez filtra"), ("c>200", "close > SMA200"), ("c<200", "close < SMA200"),
    ("50>200", "SMA50 > SMA200"), ("50<200", "SMA50 < SMA200"),
    ("200↑", "SMA200 rośnie (wyżej niż 20 sesji temu)"), ("200↓", "SMA200 spada"),
    ("SPY↑", "SPY > SMA200 (hossa)"), ("SPY↓", "SPY < SMA200 (bessa)"),
    ("R≥2", "zmienność R ≥ 2% (spread relatywnie tańszy)"), ("R≥3", "zmienność R ≥ 3%"),
    ("RS+", "silniejsza od SPY w 20 sesji"), ("RS−", "słabsza od SPY w 20 sesji"),
]


def spy_features(D: sw.Data) -> dict:
    T = len(D.cal)
    out = {"c": np.full(T, np.nan), "up200": np.zeros(T, bool)}
    s = D.by_sym.get("SPY")
    if s is None:
        return out
    c = np.full(T, np.nan)
    c[s.t] = s.c
    for t in range(1, T):                           # dni bez SPY — ostatnie znane
        if np.isnan(c[t]):
            c[t] = c[t - 1]
    m200 = np.full(T, np.nan)
    m200[s.t] = sw.sma(s.c, 200)
    for t in range(1, T):
        if np.isnan(m200[t]):
            m200[t] = m200[t - 1]
    out["c"], out["up200"] = c, np.nan_to_num(c > m200, nan=False)
    return out


def features(s: sw.Stock, spy: dict) -> tuple[dict, dict, np.ndarray, np.ndarray]:
    """→ (setupy {kod+U/D: bool[Q]}, filtry {kod: bool[Q]}, R[Q] w %, rozgrzewka bool[Q])."""
    o, h, l, c = s.o, s.h, s.l, s.c
    Q = len(c)
    st = ob.state_features(o, h, l, c)
    R = st["R"]
    rng = (h - l) / o * 100.0
    pr = sw.prior
    ma = {n: sw.sma(c, n) for n in (10, 20, 50, 200)}
    hi = {n: pr(sw.rolling(h, n, np.max)) for n in (10, 20, 55, 252)}
    lo = {n: pr(sw.rolling(l, n, np.min)) for n in (10, 20, 55, 252)}
    sd20 = sw.rolling(c, 20, np.std)
    up_b, dn_b = ma[20] + 2 * sd20, ma[20] - 2 * sd20
    bbw = 4 * sd20 / ma[20]
    pct = np.full(Q, np.nan)
    if Q >= 272:
        from numpy.lib.stride_tricks import sliding_window_view
        w = sliding_window_view(np.nan_to_num(bbw, nan=np.inf), 252)
        pct[251:] = (w <= w[:, -1:]).mean(axis=1)
    rel = rng / pr(R)
    clv = np.where(h > l, (c - l) / np.where(h > l, h - l, 1.0), 0.5)
    gap = (o - pr(c)) / pr(c) * 100.0 / pr(R)
    mx3 = np.maximum(np.maximum(ma[10], ma[20]), ma[50])
    mn3 = np.minimum(np.minimum(ma[10], ma[20]), ma[50])
    spread3 = (mx3 - mn3) / c * 100.0
    r2 = sw.rsi(c, 2)
    upc, dnc = sw.streak(c > pr(c)), sw.streak(c < pr(c))
    cr10p, b10p, b20p = pr(st["cr10"]), pr(st["box10"]), pr(st["box20"])

    with np.errstate(invalid="ignore"):
        X = {
            "hi10": (c > hi[10], c < lo[10]), "hi20": (c > hi[20], c < lo[20]), "hi55": (c > hi[55], c < lo[55]),
            "hi252": (c > hi[252], c < lo[252]),
            "box10": ((b10p < 0.8) & (c > hi[10]), (b10p < 0.8) & (c < lo[10])),
            "box20": ((b20p < 0.8) & (c > hi[20]), (b20p < 0.8) & (c < lo[20])),
            "cr10": ((cr10p < 0.7) & (c > hi[10]), (cr10p < 0.7) & (c < lo[10])),
            "bbsq": ((pr(pct) <= 0.2) & (c > up_b), (pr(pct) <= 0.2) & (c < dn_b)),
            "masq": ((pr(spread3) < 0.5 * pr(R)) & (c > mx3) & (pr(c) <= pr(mx3)),
                     (pr(spread3) < 0.5 * pr(R)) & (c < mn3) & (pr(c) >= pr(mn3))),
            "in": ((pr(st["inside"]) > 0) & (c > pr(h)), (pr(st["inside"]) > 0) & (c < pr(l))),
            "nr7": ((pr(st["nr7"]) > 0) & (c > pr(h)), (pr(st["nr7"]) > 0) & (c < pr(l))),
            "big2": ((rel >= 2) & (c > o) & (clv >= 0.7), (rel >= 2) & (c < o) & (clv <= 0.3)),
            "big3": ((rel >= 3) & (c > o) & (clv >= 0.7), (rel >= 3) & (c < o) & (clv <= 0.3)),
            "gap": ((gap >= 1) & (c > o), (gap <= -1) & (c < o)),
            "x2050": ((ma[20] > ma[50]) & (pr(ma[20]) <= pr(ma[50])), (ma[20] < ma[50]) & (pr(ma[20]) >= pr(ma[50]))),
            "x50200": ((ma[50] > ma[200]) & (pr(ma[50]) <= pr(ma[200])), (ma[50] < ma[200]) & (pr(ma[50]) >= pr(ma[200]))),
            "x200": ((c > ma[200]) & (pr(c) <= pr(ma[200])), (c < ma[200]) & (pr(c) >= pr(ma[200]))),
            "dip": ((ma[50] > ma[200]) & (c > ma[200]) & (r2 < 10), (ma[50] < ma[200]) & (c < ma[200]) & (r2 > 90)),
            "mom3": ((upc >= 3) & (c > hi[20]), (dnc >= 3) & (c < lo[20])),
            "bigcon": ((cr10p < 0.7) & (rel >= 2) & (c > o), (cr10p < 0.7) & (rel >= 2) & (c < o)),
        }
        sp = spy["c"][s.t]
        rs = (c / pr(c, 20) - 1.0) - (sp / pr(sp, 20) - 1.0)
        F = {
            "—": np.ones(Q, bool), "c>200": c > ma[200], "c<200": c < ma[200],
            "50>200": ma[50] > ma[200], "50<200": ma[50] < ma[200],
            "200↑": ma[200] > pr(ma[200], 20), "200↓": ma[200] < pr(ma[200], 20),
            "SPY↑": spy["up200"][s.t], "SPY↓": ~spy["up200"][s.t],
            "R≥2": R >= 2.0, "R≥3": R >= 3.0, "RS+": rs > 0, "RS−": rs < 0,
        }
    setups = {}
    for code, (u, d) in X.items():
        setups[code + "↑"] = np.nan_to_num(u, nan=False).astype(bool)
        setups[code + "↓"] = np.nan_to_num(d, nan=False).astype(bool)
    F = {k: np.nan_to_num(v, nan=False).astype(bool) for k, v in F.items()}
    warm = np.zeros(Q, bool)
    warm[WARM:] = True
    warm &= np.isfinite(R) & (R > 0) & np.isfinite(ma[200])
    return setups, F, R, warm


def setup_desc(code: str) -> str:
    base, arrow = code[:-1], code[-1]
    for c, u, d in PAIRS:
        if c == base:
            return u if arrow == "↑" else d
    return code


# ---------------------------------------------------------------------------
#  Przygotowanie: warunki na kalendarzu w układzie SESJI WEJŚCIA (sygnał po close q → wejście q+1)
# ---------------------------------------------------------------------------
@dataclass
class Prep:
    D: sw.Data
    S: int
    T: int
    setups: dict          # kod → bool (S, T)
    filters: dict         # kod → bool (S, T)
    base: np.ndarray      # bool (S, T): wejście możliwe (rozgrzewka sygnału)
    R: list               # na spółkę: R[q] (sygnał) w %
    feats: list           # na spółkę: (setups, F) w układzie sesji sygnału
    rows: np.ndarray
    dprows: np.ndarray


def prepare(D: sw.Data) -> Prep:
    spy = spy_features(D)
    S, T = len(D.stocks), len(D.cal)
    setups: dict = {}
    filters: dict = {}
    base = np.zeros((S, T), bool)
    Rs, feats = [], []
    for i, s in enumerate(D.stocks):
        if s.group == "context":
            Rs.append(None)
            feats.append(None)
            continue
        su, F, R, warm = features(s, spy)
        Rs.append(R)
        feats.append((su, F, warm))
        te = s.t[1:]
        for k, v in su.items():
            setups.setdefault(k, np.zeros((S, T), bool))[i, te] = v[:-1]
        for k, v in F.items():
            filters.setdefault(k, np.zeros((S, T), bool))[i, te] = v[:-1]
        base[i, te] = warm[:-1]
    rows = np.array(D.rows("main", "proof"), dtype=int)
    dprows = np.array(D.rows("doubleProof"), dtype=int)
    return Prep(D, S, T, setups, filters, base, Rs, feats, rows, dprows)


def results(pp: Prep, d: int, a: float, b: int, H: int, cost: float):
    """Wyniki netto (S×T, sesja wejścia) i dzień wyjścia (S×T, indeks kalendarza) — niezależnie od warunku."""
    NET = np.full((pp.S, pp.T), np.nan, dtype=np.float32)
    EX = np.full((pp.S, pp.T), -1, dtype=np.int32)
    for i, s in enumerate(pp.D.stocks):
        R = pp.R[i]
        if R is None:
            continue
        slp = np.empty(len(s.c))
        slp[1:] = a * R[:-1]
        slp[0] = np.nan
        mv, off, _ = sw.sltp(s.o, s.h, s.l, s.c, d, slp, slp * b, H)
        ok = np.isfinite(mv)
        ok[0] = False
        q = np.flatnonzero(ok)
        NET[i, s.t[q]] = mv[q] - cost
        EX[i, s.t[q]] = s.t[q + off[q]]
    return NET, EX


def deoverlap(M: np.ndarray, EX: np.ndarray) -> np.ndarray:
    """M (F, S, T) — sygnały wejścia; jedna pozycja na spółkę: wejście dopiero po dniu wyjścia poprzedniej."""
    F, S, T = M.shape
    busy = np.full((F, S), -1, dtype=np.int32)
    TK = np.zeros_like(M)
    for t in np.flatnonzero(M.any(axis=(0, 1))):
        m = M[:, :, t] & (busy < t)
        if m.any():
            TK[:, :, t] = m
            busy = np.where(m, EX[None, :, t], busy)
    return TK


# ---------------------------------------------------------------------------
#  Ocena
# ---------------------------------------------------------------------------
@dataclass
class Win:
    g0: int
    per: np.ndarray       # okres każdego dnia kalendarza (−1 przed g0)
    split: int            # pierwszy dzień części B
    T: int


def window(pp: Prep, rows) -> Win:
    share = pp.base[rows].mean(axis=0)
    g0 = int(np.argmax(share >= 0.5)) if (share >= 0.5).any() else pp.T
    per = np.full(pp.T, -1, dtype=np.int32)
    for p, idx in enumerate(np.array_split(np.arange(g0, pp.T), N_PERIODS)):
        per[idx] = p
    split = next((t for t, d in enumerate(pp.D.cal) if int(d[:4]) >= sw.SPLIT_YEAR), pp.T)
    return Win(g0, per, max(split, g0), pp.T)


@dataclass
class St:
    n: int
    total: float
    avg: float
    win: float
    t: float
    losing: int
    empty: int
    a: float
    b: float
    excess: float
    per_sum: np.ndarray
    per_n: np.ndarray
    blind: float
    te: float = 0.0       # t-stat przewagi nad ślepym wejściem (dzienne sumy wyniku − ślepy wynik okresu)

    @property
    def ok(self) -> bool:
        return (self.n >= MIN_TRADES and self.total > 0 and self.losing <= MAX_LOSING and self.empty <= MAX_EMPTY
                and self.a > 0 and self.b > 0 and self.excess > 0)

    @property
    def why_not(self) -> str:
        r = []
        if self.n < MIN_TRADES:
            r.append(f"transakcji {self.n} < {MIN_TRADES}")
        if self.total <= 0:
            r.append("suma ≤ 0")
        if self.losing > MAX_LOSING:
            r.append(f"okresy stratne {self.losing} > {MAX_LOSING}")
        if self.empty > MAX_EMPTY:
            r.append(f"okresy puste {self.empty} > {MAX_EMPTY}")
        if self.a <= 0 or self.b <= 0:
            r.append("część A albo B ≤ 0")
        if self.excess <= 0:
            r.append("nie lepiej niż na ślepo")
        return ", ".join(r)


def blind(NET: np.ndarray, base: np.ndarray, rows, w: Win) -> np.ndarray:
    """Wejście na ślepo każdego dnia (ten sam kierunek i wyjście): śr. wynik netto w każdym okresie."""
    m = base[rows] & np.isfinite(NET[rows])
    x = np.where(m, NET[rows], 0.0)
    p = w.per
    s = np.bincount(p[p >= 0], weights=x[:, p >= 0].sum(axis=0), minlength=N_PERIODS)
    n = np.bincount(p[p >= 0], weights=m[:, p >= 0].sum(axis=0), minlength=N_PERIODS)
    return np.where(n > 0, s / np.maximum(n, 1), 0.0)


def stats(TK: np.ndarray, NET: np.ndarray, rows, w: Win, bl: np.ndarray) -> St | None:
    """TK (S, T) — transakcje; ocena na wierszach rows w oknie w."""
    sub = TK[rows]
    ii, tt = np.nonzero(sub[:, w.g0:])
    n = len(tt)
    if n == 0:
        return None
    tt = tt + w.g0
    x = NET[rows][ii, tt].astype(float)
    p = w.per[tt]
    per_sum = np.bincount(p, weights=x, minlength=N_PERIODS)
    per_n = np.bincount(p, minlength=N_PERIODS)
    daily = np.bincount(tt - w.g0, weights=x, minlength=w.T - w.g0)
    sd = daily.std(ddof=1)
    t = float(daily.mean() / sd * math.sqrt(len(daily))) if sd > 0 else 0.0
    a = float(x[tt < w.split].sum())
    b = float(x[tt >= w.split].sum())
    xe = x - bl[p]
    de = np.bincount(tt - w.g0, weights=xe, minlength=w.T - w.g0)
    sde = de.std(ddof=1)
    te = float(de.mean() / sde * math.sqrt(len(de))) if sde > 0 else 0.0
    return St(n, float(x.sum()), float(x.mean()), float((x > 0).mean() * 100), t,
              int(((per_sum < 0) & (per_n > 0)).sum()), int((per_n == 0).sum()), a, b,
              float(xe.mean()), per_sum, per_n, float((bl[p]).mean()), te)


def shift_rows(M: np.ndarray, offs: np.ndarray) -> np.ndarray:
    """Placebo: każda spółka (oś 1) przesunięta cyklicznie w czasie o swoje przesunięcie."""
    out = np.empty_like(M)
    T = M.shape[-1]
    for i, k in enumerate(offs):
        out[:, i] = np.roll(M[:, i], int(k), axis=-1)
    return out


@dataclass(frozen=True)
class Strat:
    setup: str
    filt: str
    d: int
    a: float
    b: int
    H: int

    @property
    def code(self) -> str:
        return f"{'L' if self.d > 0 else 'S'}|{self.setup}|{self.filt}|SL{self.a:g}R|TP{self.b}×|{self.H}s"

    @property
    def desc(self) -> str:
        f = dict(FILTERS)[self.filt]
        return (f"{'LONG' if self.d > 0 else 'SHORT'} po: {setup_desc(self.setup)}"
                + ("" if self.filt == "—" else f"; filtr: {f}")
                + f"; SL {self.a:g} R, TP {self.b} × SL, najwyżej {self.H} {'sesja' if self.H == 1 else 'sesji'}")


def configs():
    for d in (1, -1):
        for a in SL_R:
            for b in TP_X:
                for H in HOLDS:
                    yield d, a, b, H


def search(pp: Prep, cost: float, placebo: bool = True, log=print) -> dict:
    w = window(pp, pp.rows)
    wd = window(pp, pp.dprows) if len(pp.dprows) else None
    rng = np.random.default_rng(PLACEBO_SEED)
    offs = rng.integers(252, max(253, pp.T - 252), size=pp.S)
    fk = [k for k, _ in FILTERS]
    FM = np.stack([pp.filters[k] for k in fk])                     # (F, S, T)
    items, plc = [], {"tested": 0, "ok": 0, "conf_cand": []}
    tested = 0
    t0 = time.time()
    cfgs = list(configs())
    for ci, (d, a, b, H) in enumerate(cfgs):
        NET, EX = results(pp, d, a, b, H, cost)
        fin = np.isfinite(NET) & pp.base
        bl = blind(NET, pp.base, pp.rows, w)
        bld = blind(NET, pp.base, pp.dprows, wd) if wd else None
        for su in pp.setups:
            M = pp.setups[su][None] & FM & fin[None]
            TK = M if H == 1 else deoverlap(M, EX)
            for fi, f in enumerate(fk):
                tested += 1
                st = stats(TK[fi], NET, pp.rows, w, bl)
                if st is None or st.n < MIN_TRADES:
                    continue
                dp = stats(TK[fi], NET, pp.dprows, wd, bld) if wd else None
                items.append({"s": Strat(su, f, d, a, b, H), "st": st, "dp": dp})
            if placebo:
                Mp = shift_rows(pp.setups[su][None] & FM, offs) & fin[None]
                TKp = Mp if H == 1 else deoverlap(Mp, EX)
                for fi, f in enumerate(fk):
                    plc["tested"] += 1
                    st = stats(TKp[fi], NET, pp.rows, w, bl)
                    if st is not None and st.ok:
                        plc["ok"] += 1
                        dp = stats(TKp[fi], NET, pp.dprows, wd, bld) if wd else None
                        plc["conf_cand"].append(dp)
        if log and (ci + 1) % 12 == 0:
            log(f"  … {ci + 1}/{len(cfgs)} wariantów wyjścia, {tested} strategii, {sum(it['st'].ok for it in items)} "
                f"spełnia warunek ({time.time() - t0:.0f} s)")
    ok = [it for it in items if it["st"].ok]
    # Najlepsze wyjście każdego setupu × filtra × kierunku, potem najwyżej PER_SETUP z setupu × kierunku
    best: dict = {}
    for it in sorted(ok, key=lambda it: -it["st"].te):
        k = (it["s"].setup, it["s"].filt, it["s"].d)
        best.setdefault(k, it)
    kept, cnt = [], {}
    for it in sorted(best.values(), key=lambda it: -it["st"].te):
        k = (it["s"].setup, it["s"].d)
        if cnt.get(k, 0) >= PER_SETUP:
            continue
        cnt[k] = cnt.get(k, 0) + 1
        kept.append(it)
        if len(kept) >= TOP:
            break
    tcrit = osr.t_crit(len(kept))
    for it in kept:
        it["conf"] = confirmed(it["dp"], tcrit)
    plc["conf"] = sum(confirmed(dp, tcrit) for dp in plc["conf_cand"])
    # Gdy nic nie spełnia warunku — najlepsze z całej siatki (z powodem)
    near = []
    if not kept:
        seen = set()
        for it in sorted(items, key=lambda it: (-(it["st"].total > 0 and it["st"].excess > 0), it["st"].losing, -it["st"].te)):
            k = (it["s"].setup, it["s"].filt, it["s"].d)
            if k in seen:
                continue
            seen.add(k)
            it["conf"] = confirmed(it["dp"], tcrit)
            near.append(it)
            if len(near) >= 50:
                break
    def fallback(d):
        """Najlepsza strategia kierunku z całej siatki (gdy żadna nie spełnia warunku) — z uwagą w arkuszu."""
        cand = [it for it in items if it["s"].d == d]
        if not cand:
            return None
        it = max(cand, key=lambda it: (it["st"].total > 0 and it["st"].excess > 0, -it["st"].losing, it["st"].te))
        it.setdefault("conf", confirmed(it["dp"], tcrit))
        return it

    pick = []
    conf = next((it for it in kept if it["conf"]), None)
    if conf is not None:
        pick.append(("✓", conf))
    for lab, d in (("LONG", 1), ("SHORT", -1)):
        it = next((it for it in (kept or near) if it["s"].d == d), None) or fallback(d)
        if it is not None and all(it is not x for _, x in pick):
            pick.append((lab, it))
    return {"items": items, "ok": ok, "kept": kept, "near": near, "tcrit": tcrit, "tested": tested, "placebo": plc,
            "w": w, "wd": wd, "pick": pick, "secs": time.time() - t0}


def confirmed(dp: St | None, tc: float) -> bool:
    return bool(dp is not None and dp.n >= 10 and dp.total > 0 and dp.excess > 0 and dp.te >= tc
                and dp.losing <= 4)


# ---------------------------------------------------------------------------
#  Transakcje jednej strategii (do portfela)
# ---------------------------------------------------------------------------
def trades_of(pp: Prep, s: Strat, rows, cost: float, g0: int = 0) -> list[dict]:
    """Transakcje strategii (jak w ocenie: jedna pozycja na spółkę, wejścia od dnia g0 kalendarza)."""
    out = []
    for i in rows:
        st = pp.D.stocks[i]
        su, F, warm = pp.feats[i]
        R = pp.R[i]
        sig = su[s.setup] & F[s.filt] & warm
        slp = np.empty(len(st.c))
        slp[1:] = s.a * R[:-1]
        slp[0] = np.nan
        mv, off, why = sw.sltp(st.o, st.h, st.l, st.c, s.d, slp, slp * s.b, s.H)
        busy = -1
        for q in np.flatnonzero(sig[:-1]) + 1:
            if q <= busy or not np.isfinite(mv[q]) or st.t[q] < g0:
                continue
            x = q + int(off[q])
            pout = st.o[q] * (1 + s.d * mv[q] / 100.0)
            out.append({"si": int(i), "qe": int(q), "qx": x, "d": s.d, "pin": float(st.o[q]), "pout": float(pout),
                        "why": int(why[q])})
            busy = x
    return out


# ---------------------------------------------------------------------------
#  Uruchomienie i Excel
# ---------------------------------------------------------------------------
def run(no_sync: bool = False, out: Path | None = None, since: str | None = sw.DAILY_SINCE, cost: float = sw.COST_PCT,
        placebo: bool = True, data: sw.Data | None = None) -> Path:
    t0 = time.time()
    n_all = len(PAIRS) * 2 * len(FILTERS) * 2 * len(SL_R) * len(TP_X) * len(HOLDS)
    print(f"IA 4 — strategie jednokierunkowe z SL i TP: {n_all} strategii, spread {cost:g}% raz na transakcję, "
          f"{N_PERIODS} okresów (≤ {MAX_LOSING} stratne)\n")
    D = data or sw.load(no_sync, since)
    print(f"  {D.note}\n  dane {D.cal[0]} – {D.cal[-1]}, {len(D.stocks)} instrumentów")
    pp = prepare(D)
    res = search(pp, cost, placebo)
    plc = res["placebo"]
    print(f"\n  spełnia warunek: {len(res['ok'])} z {res['tested']} (w arkuszu {len(res['kept'])} po wyborze najlepszego wyjścia), "
          f"✓ doubleProof: {sum(it['conf'] for it in res['kept'])} (próg t ≥ {res['tcrit']:.2f})")
    if placebo:
        print(f"  placebo (warunki przesunięte w czasie): spełnia {plc['ok']} z {plc['tested']}, ✓ {plc['conf']}")
    for it in (res["kept"] or res["near"])[:8]:
        st, dp = it["st"], it["dp"]
        print(f"   {it['s'].code:44s} n {st.n:5d}  śr. {st.avg:+.3f}%  na ślepo {st.blind:+.3f}%  t przewagi {st.te:5.2f}  "
              f"stratne {st.losing}  dP: n {dp.n if dp else 0} śr. {dp.avg if dp else 0:+.3f}% t przewagi {dp.te if dp else 0:.2f} "
              f"{'✓' if it['conf'] else '—'}" + ("" if st.ok else f"  [{st.why_not}]"))
    pfs = []
    for lab, it in res["pick"]:
        tr = trades_of(pp, it["s"], pp.rows, cost, res["w"].g0)
        p1 = sw.portfolio([dict(t) for t in tr], D, lev=1.0, cost=cost, keep_rows=False)
        p5 = sw.portfolio([dict(t) for t in tr], D, lev=5.0, cost=cost)
        pfs.append((lab, it, p1, p5))
        print(f"  portfel {lab} {it['s'].code}: 1:1 {p1['ret']:+.0f}% (obsunięcie {p1['mdd']:.0f}%), "
              f"1:5 {p5['ret']:+.0f}% (obsunięcie {p5['mdd']:.0f}%{', UPADŁOŚĆ' if p5['ruined'] else ''})")
    path = out or (config.RESEARCH_DIR.parent / "SWING-TREND.xlsx")
    write_xlsx(path, pp, res, pfs, cost, placebo)
    print(f"\nGotowe w {time.time() - t0:.0f} s → {path}")
    return path


def write_xlsx(path: Path, pp: Prep, res: dict, pfs: list, cost: float, placebo: bool) -> Path:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    S = sw.styles()
    D = pp.D
    wb = Workbook()
    ws = wb.active
    ws.title = "STRATEGIE"
    rows = res["kept"] or res["near"]
    ws["A1"] = ("Strategie spełniające warunek (najlepsze wyjście każdego setupu × filtra × kierunku)" if res["kept"] else
                "Żadna strategia nie spełnia warunku — najlepsze 50 z całej siatki, z powodem")
    ws["A1"].font = S.bold
    fixed = ["Nr", "Kod", "Opis", "Spełnia warunek", "Dlaczego nie", "Transakcji", "Zyskownych %", "Śr. netto %",
             "Śr. brutto %", "Na ślepo: śr. netto %", "Przewaga nad ślepym %", "Suma netto %", "t-stat wyniku", "t-stat przewagi", "Okresy stratne",
             "Okresy puste", "Część A suma %", "Część B suma %", "dP transakcji", "dP śr. netto %", "dP przewaga %",
             "dP t-stat przewagi", f"dP ✓ (t przewagi ≥ {res['tcrit']:.2f})"]
    sw.header(ws, 2, fixed + [f"O{p + 1}" for p in range(N_PERIODS)], [5, 38, 80] + [10] * (len(fixed) - 3) + [7] * N_PERIODS, S)
    for k, it in enumerate(rows, start=1):
        st, dp = it["st"], it["dp"]
        ws.append([k, it["s"].code, it["s"].desc, "TAK" if st.ok else "—", st.why_not, st.n, round(st.win, 1),
                   round(st.avg, 4), round(st.avg + cost, 4), round(st.blind, 4), round(st.excess, 4), round(st.total, 1),
                   round(st.t, 2), round(st.te, 2), st.losing, st.empty, round(st.a, 1), round(st.b, 1),
                   dp.n if dp else None, round(dp.avg, 4) if dp else None, round(dp.excess, 4) if dp else None,
                   round(dp.te, 2) if dp else None, "✓" if it["conf"] else "—"] + [round(float(v), 1) for v in st.per_sum])
        r = ws.max_row
        ws.cell(r, 4).fill = S.good if st.ok else S.bad
        for p in range(N_PERIODS):
            c = ws.cell(r, len(fixed) + 1 + p)
            c.fill = S.sub if st.per_n[p] == 0 else (S.good if st.per_sum[p] >= 0 else S.bad)
        if it["conf"]:
            ws.cell(r, len(fixed)).fill = S.good
    ws.freeze_panes = "D3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(fixed))}{ws.max_row}"

    # Przegląd: najlepsze wyniki każdego setupu (bez względu na warunek) — żeby było widać całą siatkę
    sv = wb.create_sheet("SETUPY")
    sv["A1"] = ("Każdy setup × kierunek: najlepsza strategia z całej siatki (najpierw zarabiające netto i lepsze niż wejście na ślepo, "
                "potem największy t przewagi) — także te, które nie spełniają warunku")
    sv["A1"].font = S.bold
    cols = ["Setup", "Kierunek", "Opis setupu", "Strategii z ≥ 30 transakcjami", "Spełnia warunek", "Najlepsza: kod",
            "Transakcji", "Śr. netto %", "Na ślepo %", "t-stat przewagi", "Okresy stratne", "Dlaczego nie", "dP śr. netto %",
            "dP t-stat przewagi"]
    sw.header(sv, 2, cols, [10, 9, 70, 12, 10, 44, 10, 10, 10, 9, 9, 40, 10, 9], S)
    groups: dict = {}
    for it in res["items"]:
        groups.setdefault((it["s"].setup, it["s"].d), []).append(it)
    def rank(x):          # najpierw zarabiające i lepsze niż na ślepo, potem t przewagi
        return (x["st"].total > 0 and x["st"].excess > 0, x["st"].te)

    for (su, d), lst in sorted(groups.items(), key=lambda kv: tuple(-v for v in rank(max(kv[1], key=rank)))):
        b = max(lst, key=rank)
        st, dp = b["st"], b["dp"]
        sv.append([su, "long" if d > 0 else "short", setup_desc(su), len(lst), sum(x["st"].ok for x in lst), b["s"].code,
                   st.n, round(st.avg, 4), round(st.blind, 4), round(st.te, 2), st.losing, st.why_not,
                   round(dp.avg, 4) if dp else None, round(dp.te, 2) if dp else None])
        sv.cell(sv.max_row, 8).fill = S.good if st.avg > 0 else S.bad
    sv.freeze_panes = "C3"

    for lab, it, p1, p5 in pfs:
        note = "" if it["st"].ok else f" — NIE SPEŁNIA WARUNKU ({it['st'].why_not}); najlepsza w tym kierunku"
        sw.portfolio_sheet(wb, f"PORTFEL {lab}", p5,
                           f"Portfel 10 000 $ — {it['s'].code}{' (✓ doubleProof)' if it.get('conf') else ''}{note} — "
                           f"53 spółki, dźwignia 1:5, depozyt {sw.STAKE_PCT:g}% "
                           f"(dla porównania 1:1: wynik {p1['ret']:+.1f}%, obsunięcie {p1['mdd']:.1f}%)", S)

    op = wb.create_sheet("OPIS")
    plc = res["placebo"]
    w = res["w"]
    lines = [
        ("IA 4 — strategie jednokierunkowe z SL i TP (towarzyszące wybiciom)", ""),
        ("Wersja", __version__), ("Wygenerowano", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Dane", f"{D.cal[0]} – {D.cal[-1]}; okno ocen od {D.cal[w.g0]} (rozgrzewka {WARM} sesji); {D.note}"),
        ("Siatka", f"{len(PAIRS) * 2} setupów × {len(FILTERS)} filtrów × 2 kierunki × SL {', '.join(f'{x:g}' for x in SL_R)} R × "
                   f"TP {', '.join(map(str, TP_X))} × SL × najwyżej {', '.join(map(str, HOLDS))} sesji = {res['tested']} strategii"),
        ("R", "średni zakres dnia (H − L)/open ze 100 sesji przed wejściem — normalna dynamika spółki"),
        ("Rozliczenie", "wejście po open sesji po sygnale; pesymistycznie: SL i TP w tej samej świecy → SL, luka przez SL → open, "
                        "luka przez TP → TP; po limicie sesji wyjście po close; jedna pozycja na spółkę naraz"),
        ("Spread", f"{cost:g}% wartości pozycji, raz na transakcję (decyzja człowieka z 16.10.2026)"),
        ("Warunek", f"53 spółki: ≥ {MIN_TRADES} transakcji, suma netto > 0, {N_PERIODS} okresów — najwyżej {MAX_LOSING} stratne i "
                    f"{MAX_EMPTY} puste, część A (< {sw.SPLIT_YEAR}) i B dodatnie, lepiej niż wejście na ślepo "
                    "(ten sam kierunek i wyjście każdego dnia)"),
        ("Potwierdzenie ✓", f"20 spółek doubleProof: ≥ 10 transakcji, suma i przewaga nad ślepym > 0, ≤ 4 stratne okresy, "
                            f"t przewagi nad ślepym ≥ {res['tcrit']:.2f} (Bonferroni dla {len(res['kept'])} reguł w arkuszu)"),
        ("Wynik", f"spełnia warunek {len(res['ok'])} z {res['tested']}; w arkuszu {len(res['kept'])}; "
                  f"✓ {sum(it['conf'] for it in res['kept'])}"),
        ("Placebo", (f"te same warunki przesunięte w czasie o ≥ 252 sesje: spełnia {plc['ok']} z {plc['tested']}, ✓ {plc['conf']} — "
                     "tyle daje sam przypadek; wynik prawdziwy jest wart uwagi dopiero, gdy wyraźnie go przewyższa")
         if placebo else "wyłączone (--no-placebo)"),
        ("Portfele", f"najlepsza z ✓ (jeśli jest), najlepszy long i najlepszy short (gdy żaden short nie spełnia warunku — "
                     f"najlepszy z całej siatki, z uwagą): 10 000 $, depozyt {sw.STAKE_PCT:g}% stanu, 1:5 (obok 1:1); "
                     "bez kosztu nocnego (swap CFD)"),
        ("Arkusze", "STRATEGIE — reguły; SETUPY — najlepsza strategia każdego setupu × kierunku z całej siatki; PORTFEL … — przebieg"),
    ]
    for k, (a, b) in enumerate(lines, start=1):
        op.cell(k, 1, a)
        op.cell(k, 2, b)
    op["A1"].font = S.bold
    op.column_dimensions["A"].width = 18
    op.column_dimensions["B"].width = 150
    wb.move_sheet("OPIS", offset=-len(wb.sheetnames) + 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
#  Selftest
# ---------------------------------------------------------------------------
def _synth(n: int, Q: int, seed: int, edge: float = 0.0) -> sw.Data:
    """Losowy spacer (bez przewagi). edge > 0: po close nad szczytem 20 sesji kolejne 3 sesje rosną średnio o edge %."""
    rng = np.random.default_rng(seed)
    days = np.array([str(np.datetime64("2008-01-02") + np.timedelta64(int(k * 7 // 5), "D")) for k in range(Q)])
    stocks = []
    groups = ["main"] * 3 + ["proof"] * (n - 6) + ["doubleProof"] * 3
    for i in range(n):
        p, o, h, l, c = 100.0, [], [], [], []
        boost = 0
        hist_h: list[float] = []
        for q in range(Q):
            op = p * (1 + rng.normal(0, 0.003))
            mv = rng.normal(0, 1.5) + (edge if boost > 0 else 0.0)
            cl = op * (1 + mv / 100.0)
            hi = max(op, cl) * (1 + abs(rng.normal(0, 0.5)) / 100)
            lo = min(op, cl) * (1 - abs(rng.normal(0, 0.5)) / 100)
            boost = max(0, boost - 1)
            if edge and len(hist_h) >= 20 and cl > max(hist_h[-20:]):
                boost = 3
            hist_h.append(hi)
            o.append(op), h.append(hi), l.append(lo), c.append(cl)
            p = cl
        stocks.append(sw.Stock(f"S{i}", groups[i], days, *(np.array(x) for x in (o, h, l, c)), np.ones(Q)))
    stocks.append(sw.Stock("SPY", "context", days, *(np.array(x) for x in (o, h, l, c)), np.ones(Q)))
    return sw.finish(stocks, "syntetyczne")


def selftest() -> None:
    fails: list[str] = []

    def check(name, got, exp):
        ok = got == exp
        print(f"  {'OK  ' if ok else 'BŁĄD'} {name}: {got}" + ("" if ok else f" (oczekiwano {exp})"))
        if not ok:
            fails.append(name)

    print(f"IA 4 {__version__} — selftest ia4.swtrend\n")
    sw.selftest_core(check)
    # Jedna pozycja na spółkę: dwa sygnały dzień po dniu, trzymanie 3 sesje → druga pominięta
    M = np.zeros((1, 1, 10), bool)
    M[0, 0, [2, 3, 6]] = True
    EX = np.full((1, 10), -1, dtype=np.int32)
    EX[0, 2], EX[0, 3], EX[0, 6] = 4, 5, 8
    check("jedna pozycja naraz", np.flatnonzero(deoverlap(M, EX)[0, 0]).tolist(), [2, 6])
    # Brak przewagi: (prawie) nic nie spełnia warunku; z przewagą: long po szczycie 20 sesji spełnia i ma ✓
    global SL_R, TP_X, HOLDS, PAIRS, FILTERS
    keep = SL_R, TP_X, HOLDS, PAIRS, FILTERS
    SL_R, TP_X, HOLDS = (1.0,), (2,), (3,)
    PAIRS = [p for p in keep[3] if p[0] in ("hi20", "in", "x200")]
    FILTERS = keep[4][:3]
    try:
        D0 = _synth(14, 1400, 1)
        r0 = search(prepare(D0), 0.3, placebo=False, log=None)
        check("bez przewagi: brak ✓", sum(it["conf"] for it in r0["kept"]), 0)
        D1 = _synth(14, 1400, 2, edge=1.2)
        r1 = search(prepare(D1), 0.3, placebo=True, log=None)
        hit = [it for it in r1["kept"] if it["s"].setup == "hi20↑" and it["s"].d > 0]
        check("z przewagą: long po szczycie 20 sesji spełnia warunek", bool(hit), True)
        check("z przewagą: ✓ na doubleProof", bool(hit and hit[0]["conf"]), True)
        check("placebo: brak ✓", r1["placebo"]["conf"], 0)
        if hit:
            p1 = prepare(D1)
            tr = trades_of(p1, hit[0]["s"], p1.rows, 0.3, r1["w"].g0)
            net = np.array([t["d"] * (t["pout"] / t["pin"] - 1) * 100 - 0.3 for t in tr])
            check("transakcje portfela = transakcje oceny", (len(tr), round(float(net.sum()), 4)),
                  (hit[0]["st"].n, round(hit[0]["st"].total, 4)))
        import tempfile
        with tempfile.TemporaryDirectory() as td:
            p = run(no_sync=True, out=Path(td) / "t.xlsx", data=D1, placebo=False)
            from openpyxl import load_workbook
            wb = load_workbook(p)
            check("arkusze", wb.sheetnames[:3], ["OPIS", "STRATEGIE", "SETUPY"])
    finally:
        SL_R, TP_X, HOLDS, PAIRS, FILTERS = keep
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    if fails:
        sys.exit(1)


def main() -> None:
    ap = argparse.ArgumentParser(description="Strategie jednokierunkowe z SL i TP towarzyszące wybiciom (instrukcja 7a)")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--no-sync", action="store_true", help="bez synchronizacji bazy dobowej")
    ap.add_argument("--no-placebo", action="store_true", help="bez placebo (dwa razy szybciej)")
    ap.add_argument("--since", default=sw.DAILY_SINCE, help="początek danych (domyślnie 2000-01-01)")
    ap.add_argument("--cost", type=float, default=sw.COST_PCT, help="spread %% na transakcję (raz; domyślnie 0,3)")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    if a.selftest:
        selftest()
        return
    run(a.no_sync, a.out, a.since, a.cost, not a.no_placebo)


if __name__ == "__main__":
    main()
