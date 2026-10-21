"""
IA 4 — One-Day: siła ruchu po konsolidacji i strategie long + short naraz (instrukcja, sekcja 7).
Wersja projektu: 1.56 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Hipoteza człowieka (15.10.2026): kierunku nie da się przewidzieć, ale może da się przewidzieć SIŁĘ ruchu —
wybicie po konsolidacji; każda spółka ma inną dynamikę.

1. ANALIZA: dla każdego stanu przed open dnia t (konsolidacje różnej długości i ciasności, dni wewnętrzne,
   NR4/NR7, najciaśniejsze 10% / 20% w roku danej spółki, duże ruchy wczoraj …) — jak duży jest ruch dnia t
   w jednostkach NORMALNEJ dynamiki tej spółki (R = średni zakres dnia z 100 sesji). Mierniki: zakres dnia,
   największe oddalenie od open (to łapie long + short), |close − open|, oddalenie w 3 sesje; porównanie
   z wszystkimi dniami (wzrost = średnia stanu / średnia wszystkich dni).
2. STRATEGIE: stan × SL (0,25 / 0,5 / 0,75 / 1 R) × TP = 2 albo 3 × SL × sposób wejścia:
   • „long + short” (decyzja człowieka): na open jednocześnie long i short, każdy z SL i TP; spread płaci
     każda noga;
   • „wybicie” (porównawczo): zlecenia buy-stop i sell-stop SL od open (jedno anuluje drugie), SL z powrotem
     na open, TP = k × SL od wejścia; spread tylko przy wejściu.
   Zawsze w ciągu jednego dnia: niezamknięte nogi — po close. Kolejność ruchów w świecy dobowej nieznana →
   zawsze wariant pesymistyczny (zasada 5): gdy dzień sięgnął obu poziomów, liczymy najgorszą kolejność.
3. OCENA: 20 okresów (najwyżej 3 stratne), t-stat, część A (do 2012) i B (od 2013), doubleProof (20 spółek),
   placebo (stan przesunięty w czasie o losową liczbę ≥ 252 sesji — ta sama częstość, zero związku z dniem).

    python -m ia4.odbreak --cost 0.6        # baza dobowa od 2000 (najpierw python -m ia4.daily); kilka minut
    python -m ia4.odbreak --selftest        # „WYNIK: OK”
Wynik: ~/Desktop/IA4/ONE-DAY-WYBICIA.xlsx (poza gitem).
"""

from __future__ import annotations

import argparse
import math
import os
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from numpy.lib.stride_tricks import sliding_window_view

from . import __version__, config
from . import odsearch as osr
from . import oneday as od

COST_PCT = osr.COST_PCT          # spread na transakcję (wejście + wyjście) — każda noga osobno; --cost
R_LONG = 100                     # R = średni zakres dnia z 100 sesji (normalna dynamika spółki)
SL_MULTS = (0.25, 0.5, 0.75, 1.0)
TP_MULTS = (2, 3)
N_PERIODS, MAX_LOSING = 20, 3
SPLIT_YEAR = 2013                # część A: < 2013, część B: ≥ 2013
DP_ALPHA = 0.05
PF_CAPITAL = 10_000.0
PLACEBO_SEED = 2026


# ---------------------------------------------------------------------------
#  Cechy (stan przed open dnia t) i wyniki dnia t
# ---------------------------------------------------------------------------
def _sma(x, n):
    return osr._sma(x, n)


def _roll(x, n, fn):
    out = np.full(len(x), np.nan)
    if len(x) >= n:
        out[n - 1:] = fn(sliding_window_view(x, n), axis=1)
    return out


def state_features(o, h, l, c) -> dict:
    """Na końcu sesji q (znane przed open sesji q+1)."""
    Q = len(c)
    rng = (h - l) / o * 100.0
    R = _sma(rng, R_LONG)
    f = {"R": R, "rng1": rng / R}
    for n in (3, 5, 10, 20):
        f[f"cr{n}"] = _sma(rng, n) / R
    for n in (5, 10, 20):
        box = (_roll(h, n, np.max) - _roll(l, n, np.min)) / c * 100.0
        f[f"box{n}"] = box / (R * math.sqrt(n))
    cr5 = f["cr5"]
    pct = np.full(Q, np.nan)
    if Q >= 252 + R_LONG:
        w = sliding_window_view(np.nan_to_num(cr5, nan=np.inf), 252)
        pct[251:] = (w <= w[:, -1:]).mean(axis=1)
        pct[: R_LONG + 251] = np.nan
    f["pct5"] = pct
    inside = np.zeros(Q)
    inside[1:] = (h[1:] < h[:-1]) & (l[1:] > l[:-1])
    f["inside"] = inside
    f["inside_n"] = osr._streak(inside.astype(bool))
    for n in (4, 7):
        v = np.full(Q, np.nan)
        if Q > n:
            pm = _roll(rng, n - 1, np.min)
            v[n - 1:] = (rng[n - 1:] < pm[n - 2:-1]).astype(float)
        f[f"nr{n}"] = v
    return f


def outcomes(o, h, l, c) -> dict:
    """Wyniki sesji q (ruch dnia) — tylko do analizy i rozliczenia, nigdy do decyzji."""
    Q = len(c)
    U = (h - o) / o * 100.0
    D = (o - l) / o * 100.0
    out = {"U": U, "D": D, "rngT": (h - l) / o * 100.0, "oc": (c - o) / o * 100.0}
    mx3 = np.full(Q, np.nan)
    if Q >= 3:
        hh = _roll(h, 3, np.max)[2:]
        ll = _roll(l, 3, np.min)[2:]
        mx3[:-2] = np.maximum(hh - o[:-2], o[:-2] - ll) / o[:-2] * 100.0
    out["mx3"] = mx3
    return out


def prepare(P: osr.Panel) -> None:
    S, T = P.o.shape
    P.bk = {}
    for i in range(S):
        vi = np.flatnonzero(P.valid[i])
        if len(vi) < 5:
            continue
        o, h, l, c = P.o[i, vi], P.h[i, vi], P.l[i, vi], P.c[i, vi]
        st = state_features(o, h, l, c)
        keep = vi + 1 < T
        src, tgt = np.flatnonzero(keep), vi[keep] + 1
        for k, v in st.items():
            P.bk.setdefault(k, np.full((S, T), np.nan))[i, tgt] = v[src]
        for k, v in outcomes(o, h, l, c).items():
            P.bk.setdefault("o_" + k, np.full((S, T), np.nan))[i, vi] = v
    P.bk_ok = P.base & np.isfinite(P.bk["R"]) & (P.bk["R"] > 0)


# ---------------------------------------------------------------------------
#  Stany
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class St:
    code: str
    desc: str


def states() -> list:
    def g(k, x):
        return lambda P: np.nan_to_num(P.bk[k], nan=-np.inf) >= x

    def lt(k, x):
        return lambda P: np.nan_to_num(P.bk[k], nan=np.inf) < x

    def both(a, b):
        return lambda P: a(P) & b(P)

    out = [(St("all", "każdy dzień (punkt odniesienia)"), lambda P: np.ones_like(P.base))]
    for n in (3, 5, 10, 20):
        for x in (0.5, 0.7, 0.85):
            out.append((St(f"cr{n}<{x:g}", f"konsolidacja {n} sesji: średni zakres < {x:g} R"), lt(f"cr{n}", x)))
    for n in (5, 10, 20):
        for x in (0.6, 0.8, 1.0):
            out.append((St(f"box{n}<{x:g}", f"ciasny przedział {n} sesji: (max H − min L) < {x:g} R·√{n}"), lt(f"box{n}", x)))
    out.append((St("pct5<0.1", "najciaśniejsze 10% w roku tej spółki (średni zakres 5 sesji)"), lt("pct5", 0.1)))
    out.append((St("pct5<0.2", "najciaśniejsze 20% w roku tej spółki"), lt("pct5", 0.2)))
    out.append((St("inside", "wczoraj dzień wewnętrzny (H < poprzedni H, L > poprzedni L)"), g("inside", 1)))
    out.append((St("inside2", "≥ 2 dni wewnętrzne z rzędu"), g("inside_n", 2)))
    out.append((St("nr4", "wczoraj najwęższy zakres z 4 sesji"), g("nr4", 1)))
    out.append((St("nr7", "wczoraj najwęższy zakres z 7 sesji"), g("nr7", 1)))
    out.append((St("nr7+in", "NR7 i dzień wewnętrzny"), both(g("nr7", 1), g("inside", 1))))
    out.append((St("nr7+cr10<0.7", "NR7 po konsolidacji 10 sesji < 0,7 R"), both(g("nr7", 1), lt("cr10", 0.7))))
    out.append((St("pct5<0.1+box10<0.8", "najciaśniejsze 10% w roku i ciasny przedział 10 sesji"),
                both(lt("pct5", 0.1), lt("box10", 0.8))))
    for x in (1.5, 2.0):
        out.append((St(f"cr5>{x:g}", f"rozruszany rynek: średni zakres 5 sesji > {x:g} R"), g("cr5", x)))
    for x in (2.0, 3.0):
        out.append((St(f"rng1>{x:g}", f"wczoraj duży zakres: > {x:g} R"), g("rng1", x)))
    return out


# ---------------------------------------------------------------------------
#  1. Analiza
# ---------------------------------------------------------------------------
METRICS = [("o_rngT", "zakres dnia"), ("mx", "największe oddalenie od open"), ("aoc", "|close − open|"),
           ("o_mx3", "oddalenie w 3 sesje")]


def analysis(P: osr.Panel, rows: np.ndarray, g0: int) -> dict:
    R = P.bk["R"]
    base = P.bk_ok[rows][:, g0:]
    m_ = {"o_rngT": P.bk["o_rngT"] / R, "mx": np.maximum(P.bk["o_U"], P.bk["o_D"]) / R,
          "aoc": np.abs(P.bk["o_oc"]) / R, "o_mx3": P.bk["o_mx3"] / R}
    vals = {k: v[rows][:, g0:] for k, v in m_.items()}
    out = []
    for stt, fn in states():
        m = fn(P)[rows][:, g0:] & base
        row = {"code": stt.code, "desc": stt.desc, "n": int(m.sum())}
        for k, _ in METRICS:
            v = vals[k][m & np.isfinite(vals[k])]
            row[k] = (float(v.mean()), float(np.median(v))) if len(v) else (None, None)
        mx = vals["mx"][m]
        for x in (0.5, 1.0, 1.5, 2.0):
            row[f"p{x:g}"] = float((mx >= x).mean()) if len(mx) else None
        both = (P.bk["o_U"] / R >= 0.5) & (P.bk["o_D"] / R >= 0.5)
        bb = both[rows][:, g0:][m]
        row["whip"] = float(bb.mean()) if len(bb) else None
        out.append(row)
    b = out[0]
    for r in out:
        r["lift"] = {k: (r[k][0] / b[k][0] if r[k][0] and b[k][0] else None) for k, _ in METRICS}
    # spółki: typowa dynamika i wzrost dla wybranych stanów
    per = []
    pick = ["cr5<0.7", "nr7", "pct5<0.1", "cr5>1.5"]
    fns = {s.code: fn for s, fn in states()}
    for j, i in enumerate(rows):
        bi = P.bk_ok[i, g0:]
        mx = m_["mx"][i, g0:]
        bmean = float(np.nanmean(np.where(bi, mx, np.nan))) if bi.any() else None
        rec = {"sym": P.syms[i], "R": float(np.nanmedian(np.where(bi, R[i, g0:], np.nan))) if bi.any() else None,
               "base": bmean}
        for c in pick:
            mm = fns[c](P)[i, g0:] & bi
            rec[c] = (float(np.nanmean(mx[mm])) / bmean if mm.sum() >= 10 and bmean else None, int(mm.sum()))
        per.append(rec)
    return {"rows": out, "per": per, "pick": pick}


# ---------------------------------------------------------------------------
#  2. Rozliczenie transakcji (pesymistycznie, w % stawki jednej nogi)
# ---------------------------------------------------------------------------
def straddle(U, D, oc, s, k, cost):
    """Long + short na open, każda noga SL = s, TP = k·s (%), zamknięcie na close; spread od każdej nogi."""
    u, d = U >= s, D >= s
    up_tp, dn_tp = U >= k * s, D >= k * s
    res = np.where(u & d, -2 * s,                                        # oba poziomy SL — najgorsza kolejność
          np.where(u, np.where(up_tp, k * s, oc) - s,                   # short zatrzymany, long: TP albo close
          np.where(d, np.where(dn_tp, k * s, -oc) - s, 0.0)))           # long zatrzymany, short: TP albo close
    return res - 2 * cost


def breakout(U, D, oc, s, k, cost):
    """Buy-stop i sell-stop s od open (jedno anuluje drugie), SL na open (s od wejścia), TP k·s od wejścia."""
    u, d = U >= s, D >= s
    long_res = np.where(D > 0, -s, np.where(U >= s + k * s, k * s, oc - s))     # po wejściu cena wróciła pod open?
    short_res = np.where(U > 0, -s, np.where(D >= s + k * s, k * s, -oc - s))
    res = np.where(u & d, -s, np.where(u, long_res, np.where(d, short_res, 0.0)))
    return np.where(u | d, res - cost, 0.0), (u | d)


@dataclass(frozen=True)
class Strat:
    state: St
    sl: float
    k: int
    mode: str      # "ls" — long + short naraz; "bo" — wybicie (zlecenia oczekujące)

    @property
    def code(self) -> str:
        return f"{'LS' if self.mode == 'ls' else 'BO'}|{self.state.code}|SL{self.sl:g}R|TP{self.k}x"

    @property
    def desc(self) -> str:
        how = ("long + short na open" if self.mode == "ls" else "wybicie: buy-stop / sell-stop")
        return f"{how} · {self.state.desc} · SL {self.sl:g} R, TP {self.k} × SL"


def results(P: osr.Panel, sl: float, k: int, mode: str, cost: float):
    """Macierz wyników (S×T, % stawki jednej nogi, netto) i maska transakcji — niezależnie od stanu."""
    R = np.nan_to_num(P.bk["R"], nan=0.0)
    s = sl * R
    U, D, oc = (np.nan_to_num(P.bk[k_], nan=0.0) for k_ in ("o_U", "o_D", "o_oc"))
    if mode == "ls":
        x = straddle(U, D, oc, s, k, cost)
        traded = P.bk_ok.copy()
    else:
        x, trig = breakout(U, D, oc, s, k, cost)
        traded = P.bk_ok & trig
    return np.where(traded, x, 0.0), traded


# ---------------------------------------------------------------------------
#  3. Ocena
# ---------------------------------------------------------------------------
def window(P, rows):
    g0 = osr.periods(P, rows)[0]
    idx = np.array_split(np.arange(g0, len(P.dates)), N_PERIODS)
    return g0, [(int(a[0]), int(a[-1]) + 1) for a in idx if len(a)]


def ev(x, m, win, cost):
    old = osr.MAX_LOSING, osr.N_PERIODS
    osr.MAX_LOSING, osr.N_PERIODS = MAX_LOSING, N_PERIODS
    try:
        st = osr.evaluate_x(np.where(m, x, 0.0), m, win, [0.0] * N_PERIODS, cost)
        ok = st.ok
    finally:
        osr.MAX_LOSING, osr.N_PERIODS = old
    return st, ok


def ab(P, x, m, g0):
    ys = np.array([int(d[:4]) for d in P.dates])
    out = []
    for sel in (ys < SPLIT_YEAR, ys >= SPLIT_YEAR):
        sel = sel & (np.arange(len(ys)) >= g0)
        mm = m[:, sel]
        n = int(mm.sum())
        out.append((n, float(np.where(mm, x[:, sel], 0).sum() / n) if n else 0.0))
    return out


def shifted(mask: np.ndarray, rng) -> np.ndarray:
    """Placebo: stan każdej spółki przesunięty w czasie o losową liczbę ≥ 252 sesji (cyklicznie)."""
    out = np.zeros_like(mask)
    T = mask.shape[1]
    for i in range(mask.shape[0]):
        out[i] = np.roll(mask[i], int(rng.integers(252, max(253, T - 252))))
    return out


def search(P: osr.Panel, cost: float, placebo: bool = True) -> dict:
    rows = np.array([i for i, g in enumerate(P.groups) if g in ("main", "proof")])
    dprows = np.array([i for i, g in enumerate(P.groups) if g == "doubleProof"], dtype=int)
    win = window(P, rows)
    win_dp = window(P, dprows) if len(dprows) else None
    g0 = win[0]
    sts = states()
    smask = {s.code: fn(P) & P.bk_ok for s, fn in sts}
    rng = np.random.default_rng(PLACEBO_SEED)
    pmask = {c: shifted(m, rng) & P.bk_ok for c, m in smask.items()} if placebo else {}
    items, plc = [], []
    for mode in ("ls", "bo"):
        for sl in SL_MULTS:
            for k in TP_MULTS:
                X, traded = results(P, sl, k, mode, cost)
                for stt, _ in sts:
                    m = smask[stt.code] & traded
                    if m[rows].sum() < 30:
                        continue
                    st, ok = ev(X[rows], m[rows], win, cost)
                    dp = ev(X[dprows], m[dprows], win_dp, cost)[0] if win_dp else None
                    items.append({"s": Strat(stt, sl, k, mode), "st": st, "ok": ok, "dp": dp,
                                  "ab": ab(P, X[rows], m[rows], g0), "X": None})
                    if placebo:
                        mp = pmask[stt.code] & traded
                        stp, okp = ev(X[rows], mp[rows], win, cost)
                        plc.append({"s": Strat(stt, sl, k, mode), "st": stp, "ok": okp})
    tcrit = osr.t_crit(len(items))
    for it in items:
        dp = it["dp"]
        it["conf"] = bool(dp is not None and dp.n > 0 and dp.total > 0 and dp.tstat >= tcrit and dp.losing <= 4)
        it["stable"] = all(n > 0 and a > 0 for n, a in it["ab"])
    items.sort(key=lambda it: (not it["ok"], it["st"].losing + it["st"].empty, -it["st"].tstat))
    return {"items": items, "placebo": plc, "tcrit": tcrit, "win": win, "rows": rows, "dprows": dprows}


def pf_sim(P, X, m, rows, g0, capital=PF_CAPITAL):
    """Portfel 10 000 $: na każdą transakcję 5% stanu portfela na nogę; suma stawek dnia ≤ stan portfela."""
    eq = peak = capital
    mdd = 0.0
    out, trades, wins, n_scaled = [], 0, 0, 0
    tag = {"doubleProof": " · dP"}
    for t in range(g0, X.shape[1]):
        idx = [i for i in rows if m[i, t]]
        if not idx:
            continue
        e0 = eq
        stake0 = e0 * 0.05
        scale = min(1.0, e0 / (stake0 * len(idx)))
        n_scaled += scale < 1
        for i in idx:
            stake = stake0 * scale
            pnl = stake * X[i, t] / 100
            eq += pnl
            trades += 1
            wins += pnl > 0
            out.append({"day": P.dates[t], "sym": P.syms[i] + tag.get(P.groups[i], ""), "candle": None, "k": 1,
                        "stake": stake, "pnl": pnl, "trade_pct": X[i, t], "ret": (eq / capital - 1) * 100, "eq": eq,
                        "scale": scale})
        peak = max(peak, eq)
        mdd = max(mdd, (peak - eq) / peak * 100)
    return {"rows": out, "capital": capital, "end": eq, "ret": (eq / capital - 1) * 100, "mdd": mdd, "trades": trades,
            "wins": wins, "ruined": False, "days": len({r["day"] for r in out}), "scaled": n_scaled, "max_k": 1,
            "first": out[0]["day"] if out else None, "last": out[-1]["day"] if out else None}


# ---------------------------------------------------------------------------
#  Excel
# ---------------------------------------------------------------------------
def write_xlsx(path: Path, P, an: dict, res: dict, cost: float, note: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="1F2A37")
    good, bad, sub = (PatternFill("solid", fgColor=c) for c in ("DCFCE7", "FEE2E2", "E5E7EB"))
    white, bold = Font(bold=True, color="FFFFFF"), Font(bold=True)

    def header(ws, row, cols, widths):
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill = white, head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for j, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[row].height = 45

    wb = Workbook()
    ws = wb.active
    ws.title = "ANALIZA"
    ws["A1"] = ("Siła ruchu dnia po danym stanie — w jednostkach R (normalna dynamika spółki: średni zakres dnia z 100 sesji); "
                "„wzrost” = średnia stanu / średnia wszystkich dni (> 1 = ruch większy niż zwykle)")
    cols = ["Stan", "Opis", "Dni (53 spółki)"]
    for _, nm in METRICS:
        cols += [f"{nm}: średnia R", f"{nm}: wzrost"]
    cols += ["P(oddalenie ≥ 0,5 R)", "P(≥ 1 R)", "P(≥ 1,5 R)", "P(≥ 2 R)", "P(obie strony ≥ 0,5 R — piłowanie)"]
    header(ws, 2, cols, [18, 60, 10] + [11] * (2 * len(METRICS) + 5))
    for r in an["rows"]:
        row = [r["code"], r["desc"], r["n"]]
        for k, _ in METRICS:
            row += [round(r[k][0], 3) if r[k][0] is not None else None,
                    round(r["lift"][k], 3) if r["lift"][k] is not None else None]
        row += [round(r[f"p{x:g}"], 3) if r[f"p{x:g}"] is not None else None for x in (0.5, 1.0, 1.5, 2.0)]
        row += [round(r["whip"], 3) if r["whip"] is not None else None]
        ws.append(row)
        for j in range(len(METRICS)):
            c = ws.cell(ws.max_row, 5 + 2 * j)
            if isinstance(c.value, float):
                c.fill = good if c.value >= 1.1 else (bad if c.value <= 0.9 else sub)
    ws.freeze_panes = "C3"

    sp = wb.create_sheet("SPÓŁKI")
    pick = an["pick"]
    header(sp, 1, ["Spółka", "R — typowy zakres dnia %", "Oddalenie od open: średnia R (wszystkie dni)"]
           + [f"wzrost po: {c}" for c in pick] + [f"dni: {c}" for c in pick], [10, 12, 14] + [12] * (2 * len(pick)))
    for rec in an["per"]:
        sp.append([rec["sym"], round(rec["R"], 3) if rec["R"] else None, round(rec["base"], 3) if rec["base"] else None]
                  + [round(rec[c][0], 3) if rec[c][0] is not None else None for c in pick] + [rec[c][1] for c in pick])
    sp.freeze_panes = "B2"

    st_ = wb.create_sheet("STRATEGIE")
    win = res["win"]
    fixed = ["Nr", "Kod", "Opis", "Spełnia 20 okresów", "Transakcji", "Zyskownych %", "Śr. netto %", "Śr. brutto %",
             "Suma netto %", "t-stat", "Okresy stratne", "Okresy puste", "Część A: śr. netto %", "Część B: śr. netto %",
             "Stabilna (A i B > 0)", "dP transakcji", "dP śr. netto %", "dP t-stat", f"dP ✓ (t ≥ {res['tcrit']:.2f})"]
    header(st_, 1, fixed + [f"O{p + 1}" for p in range(len(win[1]))], [5, 34, 70] + [10] * (len(fixed) - 3) + [7] * 20)
    for n, it in enumerate(res["items"], start=1):
        s, st, dp = it["s"], it["st"], it["dp"]
        row = [n, s.code, s.desc, "TAK" if it["ok"] else "—", st.n, round(st.wins / st.n * 100, 1) if st.n else None,
               round(st.avg, 4), round(st.gross, 4), round(st.total, 1), round(st.tstat, 2), st.losing, st.empty,
               round(it["ab"][0][1], 4), round(it["ab"][1][1], 4), "TAK" if it["stable"] else "—",
               dp.n if dp else None, round(dp.avg, 4) if dp else None, round(dp.tstat, 2) if dp else None,
               "✓" if it["conf"] else "—"]
        row += [round(v, 1) for v in st.per_sum]
        st_.append(row)
        r = st_.max_row
        for p in range(len(st.per_sum)):
            c = st_.cell(r, len(fixed) + 1 + p)
            c.fill = sub if st.per_n[p] < 3 else (good if st.per_sum[p] >= 0 else bad)
        st_.cell(r, 4).fill = good if it["ok"] else bad
    st_.freeze_panes = "D2"
    st_.auto_filter.ref = f"A1:{get_column_letter(len(fixed))}{st_.max_row}"

    if res["items"]:
        best_ls = next((it for it in res["items"] if it["s"].mode == "ls"), None)
        for name, it in (("LS", best_ls), ("BO", next((it for it in res["items"] if it["s"].mode == "bo"), None))):
            if it is None:
                continue
            s = it["s"]
            X, traded = results(P, s.sl, s.k, s.mode, cost)
            m = next(fn for st0, fn in states() if st0.code == s.state.code)(P) & traded & P.bk_ok
            allrows = np.concatenate([res["rows"], res["dprows"]]).astype(int)
            pf = pf_sim(P, X, m, allrows, res["win"][0])
            od.write_portfolio(wb, SimpleNamespace(id=name, desc=f"{s.code} — {s.desc} · spread {cost:g}% na nogę · 5% portfela "
                                                                 "na transakcję (wynik long + short razem)"
                                                                 + ("" if it["ok"] else " · UWAGA: nie spełnia warunku 20 okresów")),
                               pf, len(allrows))
            it["pf"] = pf

    op = wb.create_sheet("OPIS")
    op.column_dimensions["A"].width = 28
    op.column_dimensions["B"].width = 120
    pl = res["placebo"]
    meta = [
        ("Wersja", f"IA 4 {__version__} · {datetime.now().strftime('%Y-%m-%d %H:%M')} · {note}"),
        ("Dane", f"{P.dates[0]} – {P.dates[-1]}; ocena od {P.dates[res['win'][0]]}; 53 spółki (doubleProof osobno)"),
        ("R", "normalna dynamika spółki: średni zakres dnia (H − L)/open z 100 poprzednich sesji — każda spółka ma swój R."),
        ("Long + short (LS)", "na open long i short; każda noga SL = s, TP = k·s (s = SL × R); niezamknięte — na close; "
                              f"spread {cost:g}% od każdej nogi (2 × na dzień). Dzień sięgnął obu SL → najgorsza kolejność: −2s."),
        ("Wybicie (BO)", "buy-stop i sell-stop s od open, jedno anuluje drugie; SL z powrotem na open, TP k·s od wejścia; "
                         f"spread {cost:g}% tylko po wejściu. Cena po wejściu mogła wrócić za open → liczymy SL (pesymistycznie)."),
        ("Dlaczego BO obok LS", "long + short naraz to w praktyce wejście przy wybiciu z dodatkową pewną stratą s na drugiej "
                                "nodze i podwójnym spreadem — BO pokazuje tę samą ideę bez tej straty."),
        ("Ocena", f"{N_PERIODS} okresów, najwyżej {MAX_LOSING} stratne; część A < {SPLIT_YEAR}, B ≥ {SPLIT_YEAR}; doubleProof: "
                  f"suma > 0, t ≥ {res['tcrit']:.2f} (Bonferroni dla {len(res['items'])} strategii), ≤ 4 stratne."),
        ("Placebo", f"stan przesunięty w czasie (≥ 252 sesji): spełnia warunek {sum(p['ok'] for p in pl)} z {len(pl)} "
                    f"(prawdziwe: {sum(it['ok'] for it in res['items'])} z {len(res['items'])})." if pl else "—"),
        ("Uwaga", "w świecy dobowej nie widać kolejności ruchów — wyniki LS są zaniżone w dniach „piły”; świece 1h "
                  "(od 2024) pozwolą to doprecyzować."),
    ]
    for a, b in meta:
        op.append([a, b])
        op.cell(op.max_row, 1).font = bold
        op.cell(op.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".ONE-DAY-WYBICIA-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


def run(cost: float, since: str, no_sync: bool, placebo: bool, out: Path | None) -> Path:
    t0 = time.time()
    print(f"IA 4 — One-Day: siła ruchu po konsolidacji · spread {cost:g}% na nogę · od {since}\n")
    P, note = osr.load_panel("daily", no_sync, since)
    prepare(P)
    rows = np.array([i for i, g in enumerate(P.groups) if g in ("main", "proof")])
    g0 = osr.periods(P, rows)[0]
    print(f"  {note}\n  dane {P.dates[0]} – {P.dates[-1]}, {len(P.syms)} spółek\n")
    an = analysis(P, rows, g0)
    print("  ANALIZA — największe oddalenie od open (R) i wzrost wobec wszystkich dni:")
    for r in sorted(an["rows"], key=lambda r: -(r["lift"]["mx"] or 0))[:8] + [an["rows"][0]]:
        print(f"    {r['code']:<20} dni {r['n']:>7}  śr. {r['mx'][0]:.3f} R  wzrost ×{r['lift']['mx']:.2f}  "
              f"P(≥1R) {r['p1']:.2f}")
    low = next(r for r in an["rows"] if r["code"] == "cr5<0.7")
    print(f"    (po konsolidacji cr5<0.7: wzrost ×{low['lift']['mx']:.2f} — < 1 znaczy, że po ciszy przychodzi cisza)")
    res = search(P, cost, placebo)
    n_ok = sum(it["ok"] for it in res["items"])
    n_conf = sum(it["ok"] and it["conf"] and it["stable"] for it in res["items"])
    print(f"\n  STRATEGIE: {len(res['items'])} · spełnia 20 okresów: {n_ok} · do tego stabilne A/B i ✓ doubleProof: {n_conf}")
    if placebo:
        print(f"  placebo (stan przesunięty w czasie): spełnia {sum(p['ok'] for p in res['placebo'])} z {len(res['placebo'])}")
    for it in res["items"][:8]:
        s, st = it["s"], it["st"]
        print(f"    {s.code:<34} trans. {st.n:>6}  netto {st.avg:+.3f}% (brutto {st.gross:+.3f}%)  stratne {st.losing}  "
              f"t {st.tstat:+.2f}  A/B {'tak' if it['stable'] else 'nie'}  dP {'✓' if it['conf'] else '—'}")
    path = write_xlsx(out or config.RESEARCH_DIR.parent / "ONE-DAY-WYBICIA.xlsx", P, an, res, cost, note)
    print(f"\n  {time.time() - t0:.0f} s · plik: {path}")
    return path


# ---------------------------------------------------------------------------
#  Selftest
# ---------------------------------------------------------------------------
def _synth(n_stocks, n_days, seed, squeeze=False):
    """Świece dobowe od 2005 r. squeeze=True: po 5 cichych dniach (zakres ×0,3) dzień o zakresie ×4 (kierunek losowy)."""
    import pandas as pd
    rng = np.random.default_rng(seed)
    cal = pd.bdate_range("2005-01-03", periods=n_days).strftime("%Y-%m-%d").tolist()
    days = {}
    for s in range(n_stocks):
        vol = rng.uniform(0.8, 2.5)                      # różna dynamika spółek
        p, o_, h_, l_, c_ = 100.0, [], [], [], []
        quiet = 0
        for t in range(n_days):
            mult = 1.0
            if squeeze:
                if quiet == 0 and rng.random() < 0.04:
                    quiet = 6
                if quiet > 1:
                    mult = 0.3
                elif quiet == 1:
                    mult = 4.0
                quiet = max(0, quiet - 1)
            sd = vol * mult
            op = p
            cl = op * (1 + rng.normal(0, sd) / 100)
            e1, e2 = abs(rng.normal(0, vol * 0.4)) / 100, abs(rng.normal(0, vol * 0.4)) / 100   # knoty z normalnej dynamiki
            o_.append(op), c_.append(cl), h_.append(max(op, cl) * (1 + e1)), l_.append(min(op, cl) * (1 - e2))
            p = cl
        days[f"S{s}"] = od.Days(cal, o_, h_, l_, c_, [True] * n_days, [], [], [], [])
    return days, cal


def selftest() -> None:
    fails = []

    def check(name, ok, info=""):
        print(f"  {'OK ' if ok else 'BŁĄD'} {name}{': ' + str(info) if info != '' else ''}")
        if not ok:
            fails.append(name)

    print("IA 4 — One-Day: siła ruchu po konsolidacji — selftest\n")
    # rozliczenie long + short (s = 1, k = 2, spread 0,1 na nogę)
    U = np.array([0.5, 1.5, 2.5, 0.2, 1.2, 1.0])
    D = np.array([0.5, 0.2, 0.3, 2.1, 1.1, 0.0])
    oc = np.array([0.1, 0.8, 2.0, -1.5, 0.0, 0.7])
    r = straddle(U, D, oc, 1.0, 2, 0.1)
    exp = np.array([0.0, 0.8 - 1, 2.0 - 1, 2.0 - 1, -2.0, 0.7 - 1]) - 0.2
    check("long + short: cisza / wybicie bez TP / TP / TP w dół / piła / wybicie do close", np.allclose(r, exp), r.round(3).tolist())
    b, trig = breakout(U, D, oc, 1.0, 2, 0.1)
    expb = np.array([0.0, -1.0, -1.0, -1.0, -1.0, 0.7 - 1]) - np.array([0, .1, .1, .1, .1, .1])
    check("wybicie: bez wejścia / powrót pod open → SL / piła → SL / wejście bez powrotu do close",
          np.allclose(b, expb) and trig.tolist() == [False, True, True, True, True, True], b.round(3).tolist())
    U2, D2, oc2 = np.array([3.5]), np.array([0.0]), np.array([3.4])
    b2, _ = breakout(U2, D2, oc2, 1.0, 2, 0.1)
    check("wybicie: TP 2 × SL od wejścia", np.allclose(b2, [2.0 - 0.1]), b2.tolist())
    check("long + short gorszy od wybicia o ≥ s − spread w dniu wybicia",
          bool(straddle(U2, D2, oc2, 1.0, 2, 0.1)[0] <= b2[0] - 1.0 + 1e-12))

    # cechy: brak zaglądania w przyszłość
    days, cal = _synth(12, 900, seed=3)
    grp = {s: ("main" if i < 8 else "doubleProof") for i, s in enumerate(days)}
    P = osr.build(days, grp, None, calendar=cal, source="daily")
    prepare(P)
    t = 500
    before = {k: v[:, t].copy() for k, v in P.bk.items() if not k.startswith("o_")}
    days2, _ = _synth(12, 900, seed=3)
    for dd in days2.values():
        dd.h[t] *= 1.2
        dd.l[t] *= 0.8
    P2 = osr.build(days2, grp, None, calendar=cal, source="daily")
    prepare(P2)
    same = all(np.array_equal(np.nan_to_num(v, nan=-9), np.nan_to_num(P2.bk[k][:, t], nan=-9)) for k, v in before.items())
    check("stan dnia t nie widzi świecy t", same)
    check("wynik dnia t widzi świecę t", not np.array_equal(np.nan_to_num(P.bk["o_rngT"][:, t]), np.nan_to_num(P2.bk["o_rngT"][:, t])))
    i, q = 0, 300
    vi = np.flatnonzero(P.valid[i])
    rng_ = (P.h[i, vi] - P.l[i, vi]) / P.o[i, vi] * 100
    check("R = średni zakres 100 sesji", abs(P.bk["R"][i, vi[q] + 1] - rng_[q - 99:q + 1].mean()) < 1e-9)

    # analiza i strategie: szum vs zaszyte wybicie po ciszy
    days, cal = _synth(30, 2600, seed=7, squeeze=True)
    grp = {s: ("main" if i < 20 else "doubleProof") for i, s in enumerate(days)}
    P = osr.build(days, grp, None, calendar=cal, source="daily")
    prepare(P)
    rows = np.array([k for k, g in enumerate(P.groups) if g == "main"])
    g0 = osr.periods(P, rows)[0]
    an = analysis(P, rows, g0)
    sq = next(r for r in an["rows"] if r["code"] == "cr5<0.5")
    check("analiza: po ciszy zaszyte wybicie widoczne (wzrost > 1,5)", sq["lift"]["mx"] and sq["lift"]["mx"] > 1.5,
          round(sq["lift"]["mx"] or 0, 2))
    t0 = time.time()
    res = search(P, 0.1)
    top = res["items"][0]
    check("strategie: najlepsza korzysta z ciszy i spełnia warunek", top["ok"] and top["s"].state.code not in ("all",)
          and top["s"].state.code.startswith(("cr", "box", "pct", "nr", "inside")), top["s"].code)
    check("strategie: potwierdzona na doubleProof i stabilna", top["conf"] and top["stable"])
    pl_ok = sum(p["ok"] and p["s"].state.code != "all" for p in res["placebo"])
    real_ok = sum(it["ok"] and it["s"].state.code != "all" for it in res["items"])
    check("placebo: wyraźnie mniej niż prawdziwe", pl_ok * 3 <= real_ok, f"placebo {pl_ok}, prawdziwe {real_ok}")
    check("czas", time.time() - t0 < 300, f"{time.time() - t0:.0f} s")
    days_n, cal_n = _synth(30, 2600, seed=8)
    Pn = osr.build(days_n, grp, None, calendar=cal_n, source="daily")
    prepare(Pn)
    rn = search(Pn, 0.3)
    good = [it for it in rn["items"] if it["ok"] and it["conf"] and it["stable"]]
    check("szum (spread 0,3): żadna strategia ze ✓ i stabilna", not good, [it["s"].code for it in good][:3])
    f = Path(tempfile.mkdtemp()) / "w.xlsx"
    write_xlsx(f, P, an, res, 0.1, "selftest")
    from openpyxl import load_workbook
    wb = load_workbook(f, read_only=True)
    check("Excel: arkusze", wb.sheetnames == ["ANALIZA", "SPÓŁKI", "STRATEGIE", "PORTFEL LS", "PORTFEL BO", "OPIS"], wb.sheetnames)
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    sys.exit(0 if not fails else 1)


def main() -> None:
    ap = argparse.ArgumentParser(description="IA 4 — One-Day: siła ruchu po konsolidacji, long + short naraz")
    ap.add_argument("--cost", type=float, default=COST_PCT, help=f"spread na transakcję (na nogę) w %% (domyślnie {COST_PCT:g})")
    ap.add_argument("--since", default=osr.DAILY_SINCE)
    ap.add_argument("--no-sync", action="store_true")
    ap.add_argument("--no-placebo", action="store_true")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        selftest()
    run(max(0.0, a.cost), a.since, a.no_sync, not a.no_placebo, a.out)


if __name__ == "__main__":
    main()
