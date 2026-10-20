"""
IA 4 — portfel strategii serii świec spadkowych: wejście long po serii, trzymanie kilka sesji (instrukcja, sekcja 7a).
Wersja projektu: 1.53 (2026-10-20) — musi zgadzać się z IA4_INSTRUKCJA.md

Zadania człowieka (17.10.2026), liczone jednorazowo:
  • 1.46 — po DOKŁADNIE 7 świecach spadkowych (close < open) long po open 8. sesji, wyjście po close 3. sesji
    pozycji; spread 0,3% raz, dźwignia 1:5, 73 spółki (53 + doubleProof). To ustawienie domyślne.
  • 1.47 — warunki z analizy zdarzeń i limit pozycji. PROPOZYCJA (`--propozycja`): co najmniej 5 świec spadkowych
    z rzędu, seria spadła ≥ 10% (open 1. świecy → close), rynek słaby (SPY w czasie serii ≤ −3% albo SPY pod
    SMA200); wejście w pierwszym dniu serii, w którym wszystko jest spełnione (jedno wejście na serię), long po
    open następnej sesji, 5 sesji, bez SL; najwyżej 5 pozycji naraz (gdy sygnałów więcej — najpierw najgłębszy
    spadek serii).
Portfel 10 000 $ jak w module `swing` (depozyt 10% stanu na pozycję, jedna pozycja na spółkę, wycena dzienna,
≤ 10% kapitału = upadłość; bez kosztu nocnego).

Arkusze: PODSUMOWANIE (wynik, CAGR, obsunięcia, wyniki wg SMA200 / SMA50, spadku w R, grupy, dekady), ANALIZA
(cechy serii przed wejściem: 5 przedziałów z progami z lat 2000–2012 i wynik w latach < 2000, 2000–2012, 2013–;
porównanie warunków — na WSZYSTKICH seriach, także odrzuconych, z wynikiem hipotetycznym), KAPITAŁ (dwa wykresy),
OBSUNIĘCIA, ROK PO ROKU, TRANSAKCJE, ZDARZENIA (każda seria: średnie, spadek % i w R, rynek, warunki — przyjęta /
odrzucona i dlaczego, wynik transakcji albo hipotetyczny, czy weszła do portfela).

    python -m ia4.swportfel                      # 7 spadków, 3 sesje, 1:5 (najpierw python -m ia4.daily)
    python -m ia4.swportfel --propozycja         # ≥ 5 spadków, spadek ≥ 10%, słaby rynek, 5 sesji, max 5 pozycji
    python -m ia4.swportfel --min-n 5 --min-drop 12 --rynek --hold 5 --max-poz 3   # własne warunki
    python -m ia4.swportfel --selftest           # „WYNIK: OK”
Wynik: ~/Desktop/IA4/SWING-PORTFEL-7.xlsx, SWING-PORTFEL-PROPOZYCJA.xlsx albo SWING-PORTFEL-<warunki>.xlsx (poza gitem).
"""

from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path

import numpy as np

from . import __version__, config
from . import swing as sw

N_DEF, HOLD_DEF, LEV_DEF = 7, 3, 5.0
R_DAYS = 100
MAS = (50, 100, 200)
DD_MIN = 2.0                 # obsunięcia od 2% w arkuszu OBSUNIĘCIA
SPY_DROP = 3.0               # słaby rynek: SPY w czasie serii spadł o ≥ 3% …
SPY_MA = 200                 # … albo SPY pod SMA200
SPLIT = ("2000", "2013")
SENS_LEV = (1.0, 2.0, 3.0, 5.0)  # PODSUMOWANIE: wrażliwość wyniku na dźwignię …
SENS_POS = (3, 5, 10, None)      # … i limit pozycji naraz     # ANALIZA: < 2000 · 2000–2012 (progi przedziałów) · ≥ 2013


@dataclass(frozen=True)
class Rule:
    n: int = N_DEF               # seria DOKŁADNIE n świec (gdy min_n = None)
    min_n: int | None = None     # seria co najmniej min_n świec — wejście w pierwszym dniu spełnienia warunków
    min_drop: float = 0.0        # spadek serii co najmniej tyle % (0 — bez warunku)
    market: bool = False         # słaby rynek: SPY w serii ≤ −SPY_DROP% albo SPY < SMA200
    hold: int = HOLD_DEF
    d: int = 1
    lev: float = LEV_DEF
    max_pos: int | None = None

    def short(self) -> str:
        s = (f"≥ {self.min_n}" if self.min_n else f"{self.n}") + " świec spadkowych"
        if self.min_drop:
            s += f", spadek ≥ {self.min_drop:g}%"
        if self.market:
            s += ", słaby rynek"
        s += f", {'long' if self.d > 0 else 'short'} {self.hold} {'sesja' if self.hold == 1 else 'sesje' if self.hold < 5 else 'sesji'}"
        s += f", 1:{self.lev:g}" + (f", max {self.max_pos} pozycji" if self.max_pos else "")
        return s

    def tag(self) -> str:
        if self == Rule():
            return str(N_DEF)
        if self == PROPOSAL:
            return "PROPOZYCJA"
        t = [f"{self.min_n}plus" if self.min_n else f"{self.n}"]
        if self.min_drop:
            t.append(f"spadek{self.min_drop:g}")
        if self.market:
            t.append("rynek")
        t.append(f"{self.hold}s")
        if self.d < 0:
            t.append("short")
        if self.lev != LEV_DEF:
            t.append(f"lev{self.lev:g}")
        if self.max_pos:
            t.append(f"max{self.max_pos}")
        return "-".join(t)

    def base(self) -> "Rule":
        """Ta sama definicja serii i wyjścia, bez warunków — do arkusza ANALIZA."""
        return replace(self, min_drop=0.0, market=False)


PROPOSAL = Rule(min_n=5, min_drop=10.0, market=True, hold=5, max_pos=5)


def market(D: sw.Data) -> dict:
    """SPY na kalendarzu (ostatnia znana wartość) i SMA200 — do warunku słabego rynku."""
    T = len(D.cal)
    c = np.full(T, np.nan)
    m = np.full(T, np.nan)
    s = D.by_sym.get("SPY")
    if s is not None:
        c[s.t] = s.c
        m[s.t] = sw.sma(s.c, SPY_MA)
        for t in range(1, T):
            if np.isnan(c[t]):
                c[t], m[t] = c[t - 1], m[t - 1]
    return {"c": c, "ma": m}


def _runs(st: np.ndarray):
    """Serie świec spadkowych: (pierwsza, ostatnia) — z ciągu „ile z rzędu”."""
    ends = np.flatnonzero((st > 0) & (np.append(st[1:], 0) == 0))
    return [(int(q - st[q] + 1), int(q)) for q in ends]


def events(D: sw.Data, rows: list[int], rule: Rule, cost: float, mk: dict | None = None) -> list[dict]:
    """Każda seria, która osiągnęła wymaganą długość: dzień sygnału, cechy (stan po close tego dnia), warunki
    (przyjęta / odrzucona i dlaczego), wynik transakcji (dla odrzuconej — hipotetyczny, z pierwszego możliwego dnia)."""
    mk = mk or market(D)
    need = rule.min_n or rule.n
    out = []
    for si in rows:
        s = D.stocks[si]
        o, h, l, c = s.o, s.h, s.l, s.c
        Q = len(c)
        st = sw.streak(c < o)
        rng = (h - l) / o * 100.0
        ma = {n: sw.sma(c, n) for n in MAS}
        r14 = sw.rsi(c, 14)
        busy = -1
        for q0, qe in _runs(st):
            if qe - q0 + 1 < need:
                continue
            cands = range(q0 + need - 1, qe + 1) if rule.min_n else [q0 + need - 1]
            R = float(rng[max(0, q0 - R_DAYS):q0].mean()) if q0 >= 20 else float("nan")
            t0 = s.t[q0] - 1 if s.t[q0] > 0 else 0

            def feats(q):
                t = s.t[q]
                drop = (c[q] / o[q0] - 1.0) * 100.0
                spy7 = (mk["c"][t] / mk["c"][t0] - 1.0) * 100.0 if np.isfinite(mk["c"][t]) and np.isfinite(mk["c"][t0]) else None
                spy_dn = bool(mk["c"][t] < mk["ma"][t]) if np.isfinite(mk["ma"][t]) else None
                why = []
                if rule.min_drop and not drop <= -rule.min_drop:
                    why.append(f"spadek {drop:.1f}% < {rule.min_drop:g}%")
                if rule.market and not ((spy7 is not None and spy7 <= -SPY_DROP) or spy_dn):
                    why.append("rynek nie jest słaby")
                return drop, spy7, spy_dn, why

            pick, f = None, None
            for q in cands:
                f = feats(q)
                if not f[3]:
                    pick = q
                    break
            q = pick if pick is not None else cands[0]
            if pick is None:
                f = feats(q)
            drop, spy7, spy_dn, why = f
            ev = {"si": si, "sym": s.sym, "group": s.group, "first": str(s.dates[q0]), "last": str(s.dates[q]),
                  "k": int(q - q0 + 1), "len": int(qe - q0 + 1), "c": float(c[q]), "drop": drop, "R": R,
                  "dropR": drop / R if R and R > 0 else None, "spy7": spy7, "spy_dn": spy_dn,
                  "rsi14": float(r14[q]) if np.isfinite(r14[q]) else None,
                  "ma": {n: (float(ma[n][q]) if np.isfinite(ma[n][q]) else None) for n in MAS},
                  "ma200_up": (bool(ma[200][q] > ma[200][q - 20]) if q >= 20 and np.isfinite(ma[200][q - 20]) else None),
                  "ok": pick is not None, "why": "; ".join(why), "trade": None, "pf": None}
            e, x = q + 1, q + rule.hold
            if x >= Q:
                ev["status"] = "brak sesji wyjścia (koniec danych)" if e < Q else "brak sesji wejścia (koniec danych)"
            else:
                pin, pout = float(o[e]), float(c[x])
                move = rule.d * (pout / pin - 1.0) * 100.0
                worst = ((float(l[e:x + 1].min()) / pin - 1.0) * 100.0 if rule.d > 0
                         else (1.0 - float(h[e:x + 1].max()) / pin) * 100.0)
                ev.update({"in": str(s.dates[e]), "out": str(s.dates[x]), "pin": pin, "pout": pout, "move": move,
                           "net": move - cost, "worst": worst})
                if not ev["ok"]:
                    ev["status"] = "odrzucona: " + ev["why"]
                elif e <= busy:
                    ev["status"] = "spółka już w pozycji"
                else:
                    ev["status"] = "transakcja"
                    ev["trade"] = {"si": si, "qe": int(e), "qx": int(x), "d": rule.d, "pin": pin, "pout": pout,
                                   "why": 3, "prio": drop}
                    busy = x
            out.append(ev)
    out.sort(key=lambda e: (e["last"], e["sym"]))
    return out


def drawdowns(curve: list[tuple[str, float]], min_pct: float = DD_MIN) -> list[dict]:
    """Obsunięcia z wyceny dziennej: szczyt → dno → odrobienie szczytu (albo trwa)."""
    out = []
    peak_v, peak_d = None, None
    cur = None
    for d, v in curve:
        if peak_v is None or v >= peak_v:
            if cur is not None:
                cur["rec"] = d
                cur["days"] = (np.datetime64(d) - np.datetime64(cur["start"])).astype(int)
                if cur["depth"] >= min_pct:
                    out.append(cur)
                cur = None
            peak_v, peak_d = v, d
            continue
        dd = (peak_v - v) / peak_v * 100.0 if peak_v > 0 else 0.0
        if cur is None:
            cur = {"start": peak_d, "peak": peak_v, "trough": d, "low": v, "depth": dd, "rec": None, "days": None}
        elif dd > cur["depth"]:
            cur.update({"trough": d, "low": v, "depth": dd})
    if cur is not None and cur["depth"] >= min_pct:
        cur["days"] = (np.datetime64(curve[-1][0]) - np.datetime64(cur["start"])).astype(int)
        out.append(cur)
    return out


def groupstats(evs: list[dict], key, labels: list[str], lev: float) -> list[tuple]:
    out = []
    for lab in labels:
        x = np.array([e["net"] for e in evs if e.get("net") is not None and key(e) == lab])
        if len(x):
            out.append((lab, len(x), float((x > 0).mean() * 100), float(x.mean()), float(np.median(x)), float(x.mean() * lev)))
        else:
            out.append((lab, 0, None, None, None, None))
    return out


# ---------------------------------------------------------------------------
#  ANALIZA — cechy serii a wynik (wszystkie serie, wynik hipotetyczny dla odrzuconych)
# ---------------------------------------------------------------------------
FEATS = [
    ("drop", "Spadek serii %"), ("dropR", "Spadek serii w R"), ("spy7", "SPY w czasie serii %"),
    ("R", "R spółki % (typowy zakres dnia)"), ("rsi14", "RSI(14)"), ("d50", "Cena wobec SMA50 %"),
    ("d200", "Cena wobec SMA200 %"),
]


def _fval(e, k):
    if k in ("d50", "d200"):
        m = e["ma"][int(k[1:])]
        v = (e["c"] / m - 1.0) * 100.0 if m else None
    else:
        v = e.get(k)
    return v if v is not None and np.isfinite(v) else None


def _per(e) -> int:
    return 0 if e["last"] < SPLIT[0] else (1 if e["last"] < SPLIT[1] else 2)


PER_NAMES = (f"przed {SPLIT[0]}", f"{SPLIT[0]}–{int(SPLIT[1]) - 1}", f"od {SPLIT[1]}")


def analysis(evs: list[dict]) -> dict:
    ev = [e for e in evs if e.get("net") is not None]
    out = {"feats": [], "n": len(ev)}
    train = [e for e in ev if _per(e) == 1]
    for k, name in FEATS:
        xs = np.array([_fval(e, k) for e in train if _fval(e, k) is not None], dtype=float)
        if len(xs) < 25:
            continue
        qs = np.quantile(xs, [0.2, 0.4, 0.6, 0.8])
        cells = [[[] for _ in range(3)] for _ in range(5)]
        for e in ev:
            v = _fval(e, k)
            if v is None:
                continue
            cells[int(np.digitize(v, qs))][_per(e)].append(e["net"])
        out["feats"].append((name, qs, cells))
    cats = [("Wszystkie serie", lambda e: True),
            ("SPY pod SMA200", lambda e: e["spy_dn"] is True), ("SPY nad SMA200", lambda e: e["spy_dn"] is False),
            (f"SPY w czasie serii ≤ −{SPY_DROP:g}%", lambda e: e["spy7"] is not None and e["spy7"] <= -SPY_DROP),
            ("Słaby rynek (SPY ≤ −3% albo pod SMA200)", lambda e: (e["spy7"] is not None and e["spy7"] <= -SPY_DROP) or e["spy_dn"] is True),
            ("Spadek serii ≥ 10%", lambda e: e["drop"] <= -10),
            ("Spadek serii ≥ 12%", lambda e: e["drop"] <= -12),
            ("Spadek ≥ 10% i słaby rynek", lambda e: e["drop"] <= -10 and ((e["spy7"] is not None and e["spy7"] <= -SPY_DROP) or e["spy_dn"] is True)),
            ("RSI(14) < 30", lambda e: e["rsi14"] is not None and e["rsi14"] < 30)]
    out["cats"] = []
    for name, fn in cats:
        sel = [e for e in ev if fn(e)]
        per = [[e["net"] for e in sel if _per(e) == p] for p in range(3)]
        dp = [e["net"] for e in sel if e["group"] == "doubleProof" and _per(e) > 0]
        out["cats"].append((name, per, dp))
    return out


def _cell(x):
    if not len(x):
        return (0, None, None)
    a = np.array(x)
    return (len(a), float(a.mean()), float((a > 0).mean() * 100))


def analysis_sheet(wb, an: dict, rule: Rule, S):
    ws = wb.create_sheet("ANALIZA")
    ws["A1"] = (f"Cechy serii a wynik — wszystkie serie {'≥ ' + str(rule.min_n) if rule.min_n else rule.n} świec spadkowych "
                f"(bez warunków; wejście w pierwszym możliwym dniu, {rule.hold} sesje, wynik netto po spreadzie). "
                f"Progi przedziałów ustalone na latach {PER_NAMES[1]}; pozostałe okresy to sprawdzian.")
    ws["A1"].font = S.bold
    r = 3
    ws.cell(r, 1, "Warunki").font = S.bold
    sw.header(ws, r + 1, ["Warunek"] + [f"{p}: serii" for p in PER_NAMES] + [f"{p}: śr. netto %" for p in PER_NAMES]
              + [f"{p}: zyskownych %" for p in PER_NAMES] + ["doubleProof od 2000: serii", "doubleProof: śr. netto %"],
              [44] + [11] * 11, S, 45)
    r += 2
    for name, per, dp in an["cats"]:
        c = [_cell(x) for x in per]
        d = _cell(dp)
        vals = [name] + [x[0] for x in c] + [None if x[1] is None else round(x[1], 3) for x in c] \
            + [None if x[2] is None else round(x[2], 1) for x in c] + [d[0], None if d[1] is None else round(d[1], 3)]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(r, j, v)
            if j in (5, 6, 7, 12) and isinstance(v, float):
                cell.fill = S.good if v > 0 else S.bad
        r += 1
    r += 2
    ws.cell(r, 1, "Cechy — 5 przedziałów (P1 = najniższe wartości); w komórce średni wynik netto % (liczba serii)").font = S.bold
    r += 1
    for name, qs, cells in an["feats"]:
        lims = ["≤ %.2f" % qs[0], "%.2f – %.2f" % (qs[0], qs[1]), "%.2f – %.2f" % (qs[1], qs[2]),
                "%.2f – %.2f" % (qs[2], qs[3]), "> %.2f" % qs[3]]
        sw.header(ws, r, [name] + [f"P{i + 1}: {lims[i]}" for i in range(5)], [44] + [16] * 5, S, 30)
        r += 1
        for p in range(3):
            ws.cell(r, 1, PER_NAMES[p])
            for i in range(5):
                n, m, _ = _cell(cells[i][p])
                cell = ws.cell(r, 2 + i, None if m is None else f"{m:+.2f} ({n})")
                if m is not None:
                    cell.fill = sw.heat(m, 2.0)
            r += 1
        r += 1
    return ws


def run(no_sync: bool = False, out: Path | None = None, since: str | None = sw.DAILY_SINCE, cost: float = sw.COST_PCT,
        rule: Rule = Rule(), data: sw.Data | None = None) -> Path:
    t0 = time.time()
    print(f"IA 4 — portfel: {rule.short()}; spread {cost:g}% raz na transakcję\n")
    # pełna historia — średnie, R i ANALIZA także sprzed `since`; portfel i ZDARZENIA od `since`
    D = data or sw.load(no_sync, None, groups=("main", "proof", "doubleProof", "context"))
    rows = D.rows("main", "proof", "doubleProof")
    mk = market(D)
    print(f"  {D.note}\n  portfel od {since or D.cal[0]} do {D.cal[-1]}, {len(rows)} spółek")
    allev = events(D, rows, rule, cost, mk)
    evs = [e for e in allev if not since or e["last"] >= since]
    an = analysis(events(D, rows, rule.base(), cost, mk) if (rule.min_drop or rule.market) else allev)
    trades = [e["trade"] for e in evs if e["trade"] is not None]
    pf = sw.portfolio([dict(t) for t in trades], D, lev=rule.lev, cost=cost, max_pos=rule.max_pos)
    p1 = sw.portfolio([dict(t) for t in trades], D, lev=1.0, cost=cost, keep_rows=False, max_pos=rule.max_pos)
    inpf = {(r["sym"], r["in"]): r for r in pf["rows"]}
    skip = {(D.stocks[si].sym, D.cal[t]): why for si, t, why in pf["skipped_list"]}
    for e in evs:
        e["pf"] = inpf.get((e["sym"], e.get("in")))
        if e["trade"] is not None and e["pf"] is None:
            e["status"] = "poza portfelem — " + skip.get((e["sym"], e.get("in")), "upadłość")
    dds = drawdowns(pf["curve"])
    sens = {}                                   # wrażliwość: dźwignia × limit pozycji (te same transakcje)
    for lv in SENS_LEV:
        for mp in SENS_POS:
            q = sw.portfolio([dict(t) for t in trades], D, lev=lv, cost=cost, keep_rows=False, max_pos=mp)
            sens[(lv, mp)] = (q["ret"], q["mdd"], sw.cagr(q, D), q["ruined"])
    acc = [e for e in evs if e["ok"]]
    net = np.array([e["net"] for e in acc if e.get("net") is not None])
    print(f"  serii: {len(evs)}, spełnia warunki: {len(acc)}, w portfelu: {pf['trades']}"
          + (f" (limit {rule.max_pos} pozycji pominął {pf['skipped_limit']})" if rule.max_pos else ""))
    if len(net):
        print(f"  przyjęte serie: zyskownych {(net > 0).mean() * 100:.0f}%, śr. netto {net.mean():+.3f}% "
              f"(× {rule.lev:g} = {net.mean() * rule.lev:+.2f}% depozytu)")
    print(f"  portfel 1:{rule.lev:g}: {pf['ret']:+.1f}% (stan {pf['end']:,.0f} $), największe obsunięcie {pf['mdd']:.1f}%"
          f"{', UPADŁOŚĆ' if pf['ruined'] else ''}; dla porównania 1:1 {p1['ret']:+.1f}% (obsunięcie {p1['mdd']:.1f}%)")
    for x in sorted(dds, key=lambda x: -x["depth"])[:3]:
        print(f"   obsunięcie {x['depth']:.1f}%: szczyt {x['start']} → dno {x['trough']} → "
              + (f"odrobione {x['rec']}" if x["rec"] else "jeszcze nie odrobione"))
    path = out or (config.RESEARCH_DIR.parent / f"SWING-PORTFEL-{rule.tag()}.xlsx")
    write_xlsx(path, D, evs, pf, p1, dds, an, rule, cost, sens)
    print(f"\nGotowe w {time.time() - t0:.0f} s → {path}")
    return path


def write_xlsx(path: Path, D, evs, pf, p1, dds, an, rule: Rule, cost, sens: dict | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter
    S = sw.styles()
    lev = rule.lev
    wb = Workbook()
    ws = wb.active
    ws.title = "PODSUMOWANIE"
    tr = [e for e in evs if e["pf"] is not None]
    acc = [e for e in evs if e["ok"] and e.get("net") is not None]
    net = np.array([e["net"] for e in acc]) if acc else np.zeros(0)
    cg = sw.cagr(pf, D)
    need = f"co najmniej {rule.min_n}" if rule.min_n else f"dokładnie {rule.n}"
    lines = [
        (f"IA 4 — portfel: {rule.short()}", ""),
        ("Wersja", __version__), ("Wygenerowano", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Dane", f"zdarzenia i portfel od {evs[0]['last'] if evs else '—'} do {D.cal[-1]} (średnie, R i ANALIZA z pełnej "
                 f"historii bazy dobowej, od {D.cal[0]}); {D.note}"),
        ("Spółki", f"{len(D.rows('main', 'proof'))} handlowanych + {len(D.rows('doubleProof'))} doubleProof"),
        ("Seria", f"{need} świec spadkowych z rzędu (close < open)" + (
            "; wejście w pierwszym dniu serii, w którym spełnione są wszystkie warunki (jedno wejście na serię)"
            if rule.min_n else "; dłuższa seria nie daje drugiego wejścia")),
        ("Warunki", "; ".join(x for x in (
            f"spadek serii ≥ {rule.min_drop:g}% (open 1. świecy → close)" if rule.min_drop else "",
            f"słaby rynek: SPY w czasie serii ≤ −{SPY_DROP:g}% albo SPY pod SMA{SPY_MA}" if rule.market else "") if x) or "brak"),
        ("Transakcja", f"{'long' if rule.d > 0 else 'short'} po open następnej sesji, wyjście po close {rule.hold}. sesji pozycji; bez SL i TP"),
        ("Portfel", f"10 000 $, dźwignia 1:{lev:g}, depozyt {sw.STAKE_PCT:g}% stanu portfela na pozycję (wartość pozycji "
                    f"{sw.STAKE_PCT * lev:g}% stanu); jedna pozycja na spółkę"
                    + (f"; najwyżej {rule.max_pos} pozycji naraz — przy nadmiarze sygnałów najpierw najgłębszy spadek serii"
                       if rule.max_pos else "; brak wolnego depozytu → stawki dnia mniejsze")),
        ("Spread", f"{cost:g}% wartości pozycji raz na transakcję; bez kosztu nocnego (swap CFD) i dywidend"),
        ("Uwaga", "73 spółki to dzisiejsze duże spółki — te, które przetrwały; spółki, które po takim spadku upadły, "
                  "nie są w bazie, więc kupowanie spadków wygląda tu lepiej niż w rzeczywistości."),
        ("", ""),
        ("Serii", len(evs)), ("Spełnia warunki", len(acc)), ("Transakcji w portfelu", pf["trades"]),
        ("Pominięte przez limit pozycji", pf.get("skipped_limit", 0)),
        ("Zyskownych % (serie spełniające warunki)", round(float((net > 0).mean() * 100), 1) if len(net) else None),
        ("Śr. wynik transakcji netto %", round(float(net.mean()), 3) if len(net) else None),
        (f"Śr. wynik na depozycie % (× {lev:g})", round(float(net.mean() * lev), 3) if len(net) else None),
        ("Mediana wyniku netto %", round(float(np.median(net)), 3) if len(net) else None),
        ("Stan końcowy $", round(pf["end"], 2)), ("Wynik portfela %", round(pf["ret"], 1)),
        ("CAGR %", round(cg, 2) if cg is not None else None),
        ("Największe obsunięcie % (wycena dzienna)", round(pf["mdd"], 1)), ("Upadłość", "TAK" if pf["ruined"] else "nie"),
        ("Najwięcej pozycji naraz", pf["max_open"]), ("Dni z pomniejszoną stawką", pf["scaled"]),
        ("Dla porównania 1:1", f"{p1['ret']:+.1f}% (obsunięcie {p1['mdd']:.1f}%)"),
    ]
    for k, (a, b) in enumerate(lines, start=1):
        ws.cell(k, 1, a)
        ws.cell(k, 2, b)
    ws["A1"].font = S.bold
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    for col in "CDEFG":
        ws.column_dimensions[col].width = 14

    r = len(lines) + 2
    ws.cell(r, 1, "Najgłębsze obsunięcia").font = S.bold
    sw.header(ws, r + 1, ["Szczyt (od)", "Dno", "Obsunięcie %", "Odrobione", "Dni od szczytu do odrobienia"], [44, 16, 14, 14, 14], S, 30)
    r += 2
    for x in sorted(dds, key=lambda x: -x["depth"])[:10]:
        ws.append([x["start"], x["trough"], round(x["depth"], 1), x["rec"] or "jeszcze nie", int(x["days"]) if x["days"] is not None else None])
        r += 1

    if sens:
        r += 2
        ws.cell(r, 1, "Wrażliwość: te same transakcje przy innej dźwigni i limicie pozycji — wynik % / obsunięcie % / CAGR %").font = S.bold
        sw.header(ws, r + 1, ["Dźwignia"] + [f"max {mp} pozycji" if mp else "bez limitu" for mp in SENS_POS],
                  [44, 16, 14, 14, 14], S, 30)
        r += 2
        for lv in SENS_LEV:
            ws.cell(r, 1, f"1:{lv:g}")
            for j, mp in enumerate(SENS_POS, start=2):
                ret, dd, cg_, ru = sens[(lv, mp)]
                c = ws.cell(r, j, "upadłość" if ru else f"{ret:+.0f}% / {dd:.0f}% / {cg_ or 0:.1f}%")
                c.fill = S.bad if ru or dd >= 50 else (S.good if ret > 0 else S.sub)
            r += 1

    def table(title, rowsx, r):
        r += 2
        ws.cell(r, 1, title).font = S.bold
        sw.header(ws, r + 1, ["Grupa", "Transakcji", "Zyskownych %", "Śr. netto %", "Mediana netto %", f"Śr. × {lev:g} %"],
                  [44, 16, 14, 14, 14, 14], S, 30)
        r += 2
        for lab, n, w, a, m, al in rowsx:
            ws.cell(r, 1, lab)
            ws.cell(r, 2, n)
            for j, v in enumerate((w, a, m, al), start=3):
                c = ws.cell(r, j, None if v is None else round(v, 3 if j > 3 else 1))
                if j == 4 and v is not None:
                    c.fill = S.good if v > 0 else S.bad
            r += 1
        return r

    def pos(e, n):
        m = e["ma"][n]
        return "brak średniej (za mało sesji)" if m is None else ("cena nad SMA%d" % n if e["c"] > m else "cena pod SMA%d" % n)

    r = table("Wynik wg położenia wobec SMA200 (transakcje portfela)", groupstats(tr, lambda e: pos(e, 200), ["cena nad SMA200", "cena pod SMA200",
              "brak średniej (za mało sesji)"], lev), r)
    r = table("Wynik wg nachylenia SMA200 (wobec 20 sesji wcześniej)", groupstats(
        tr, lambda e: "SMA200 rośnie" if e["ma200_up"] else ("SMA200 spada" if e["ma200_up"] is False else "brak"),
        ["SMA200 rośnie", "SMA200 spada", "brak"], lev), r)
    r = table("Wynik wg położenia wobec SMA50", groupstats(tr, lambda e: pos(e, 50), ["cena nad SMA50", "cena pod SMA50"], lev), r)

    def drop_lab(e):
        x = e["dropR"]
        if x is None:
            return "brak R"
        return ("spadek serii > 8 R" if x < -8 else "spadek serii 5–8 R" if x < -5 else
                "spadek serii 3–5 R" if x < -3 else "spadek serii < 3 R")
    r = table("Wynik wg głębokości spadku serii (w R spółki)", groupstats(
        tr, drop_lab, ["spadek serii > 8 R", "spadek serii 5–8 R", "spadek serii 3–5 R", "spadek serii < 3 R"], lev), r)
    gname = {"main": "główne (3)", "proof": "kontrolne (50)", "doubleProof": "doubleProof (20)"}
    r = table("Wynik wg grupy spółek", groupstats(tr, lambda e: gname[e["group"]], list(gname.values()), lev), r)
    dec = sorted({e["in"][:3] + "0–" + e["in"][:3] + "9" for e in tr})
    table("Wynik wg dekady", groupstats(tr, lambda e: e["in"][:3] + "0–" + e["in"][:3] + "9", dec, lev), r)

    analysis_sheet(wb, an, rule, S)

    # KAPITAŁ + wykresy
    kp = wb.create_sheet("KAPITAŁ")
    sw.header(kp, 1, ["Dzień", f"Stan portfela $ (1:{lev:g})", "Szczyt $", "Obsunięcie %"], [12, 16, 14, 12], S, 30)
    peak = 0.0
    for dday, v in pf["curve"]:
        peak = max(peak, v)
        kp.append([dday, round(v, 2), round(peak, 2), round(-(peak - v) / peak * 100.0 if peak > 0 else 0.0, 2)])
    n = len(pf["curve"])
    if n > 1:
        ch = LineChart()
        ch.title = f"Stan portfela 10 000 $ — {rule.short()} (skala logarytmiczna)"
        ch.y_axis.title = "$"
        ch.y_axis.scaling.logBase = 10
        ch.add_data(Reference(kp, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(kp, min_col=1, min_row=2, max_row=n + 1))
        ch.x_axis.tickLblSkip = max(1, n // 12)
        ch.x_axis.tickMarkSkip = max(1, n // 12)
        ch.legend = None
        ch.height, ch.width = 11, 30
        ch.series[0].graphicalProperties.line.width = 12000
        kp.add_chart(ch, "F2")
        c2 = LineChart()
        c2.title = "Obsunięcie od szczytu %"
        c2.add_data(Reference(kp, min_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
        c2.set_categories(Reference(kp, min_col=1, min_row=2, max_row=n + 1))
        c2.x_axis.tickLblSkip = max(1, n // 12)
        c2.x_axis.tickMarkSkip = max(1, n // 12)
        c2.legend = None
        c2.x_axis.tickLblPos = "low"
        c2.height, c2.width = 8, 30
        c2.series[0].graphicalProperties.line.solidFill = "C0392B"
        c2.series[0].graphicalProperties.line.width = 9000
        kp.add_chart(c2, "F25")
    kp.freeze_panes = "A2"

    ob = wb.create_sheet("OBSUNIĘCIA")
    ob["A1"] = f"Każde obsunięcie portfela ≥ {DD_MIN:g}% (wycena po close): od szczytu do dna i do odrobienia szczytu"
    ob["A1"].font = S.bold
    sw.header(ob, 2, ["Szczyt (od)", "Stan w szczycie $", "Dno", "Stan na dnie $", "Obsunięcie %", "Odrobione",
                      "Dni od szczytu do odrobienia"], [12, 14, 12, 14, 12, 14, 14], S, 30)
    for x in dds:
        ob.append([x["start"], round(x["peak"], 2), x["trough"], round(x["low"], 2), round(x["depth"], 1),
                   x["rec"] or "jeszcze nie", int(x["days"]) if x["days"] is not None else None])
        ob.cell(ob.max_row, 5).fill = S.bad if x["depth"] >= 20 else S.sub
    ob.freeze_panes = "A3"

    yr = wb.create_sheet("ROK PO ROKU")
    sw.header(yr, 1, ["Rok", "Wynik roku %", "Transakcji", "Śr. netto %", "Zyskownych %"], [8, 12, 11, 11, 11], S, 30)
    byy: dict = {}
    for e in tr:
        byy.setdefault(e["in"][:4], []).append(e["net"])
    for y, v in pf["years"]:
        x = np.array(byy.get(y, []))
        yr.append([y, round(v, 2), len(x), round(float(x.mean()), 3) if len(x) else None,
                   round(float((x > 0).mean() * 100), 1) if len(x) else None])
        yr.cell(yr.max_row, 2).fill = S.good if v >= 0 else S.bad

    sw.portfolio_sheet(wb, "TRANSAKCJE", pf, f"Przebieg portfela 10 000 $ — {rule.short()}, depozyt {sw.STAKE_PCT:g}% na pozycję", S)

    zd = wb.create_sheet("ZDARZENIA")
    zd["A1"] = ("Każda seria świec spadkowych (close < open), która osiągnęła wymaganą długość. Spadek serii = open 1. świecy → "
                f"close świecy sygnału; R = średni zakres dnia (H − L)/open ze {R_DAYS} sesji przed serią; średnie z zamknięć, "
                "stan po close świecy sygnału. Dla serii odrzuconych wynik jest hipotetyczny (gdyby weszła).")
    zd["A1"].font = S.bold
    cols = ["Nr", "Spółka", "Grupa", "1. świeca serii", "Świeca sygnału", "Świec do sygnału", "Seria łącznie świec",
            "Close sygnału", "Spadek serii %", "R spółki %", "Spadek w R", "SPY w czasie serii %", "SPY pod SMA200",
            "RSI(14)", "SMA50", "Cena wobec SMA50 %", "SMA100", "Cena wobec SMA100 %", "SMA200", "Cena wobec SMA200 %",
            "SMA200 rośnie (20 sesji)", "Warunki", "Wejście (open)", "Wyjście (close)", "Cena wejścia", "Cena wyjścia",
            "Najgłębiej w trakcie %", "Wynik % (przed spreadem)", "Wynik netto %", f"Wynik na depozycie % (× {lev:g})",
            "W portfelu", "Zysk / strata $ w portfelu", "Uwagi"]
    sw.header(zd, 2, cols, [6, 9, 11, 11, 11, 8, 8, 10, 10, 9, 9, 10, 9, 8, 10, 10, 10, 10, 10, 10, 10, 11, 11, 11,
                            10, 10, 10, 10, 10, 11, 9, 11, 30], S)
    gshort = {"main": "główna", "proof": "kontrolna", "doubleProof": "doubleProof"}
    rr = lambda v, p=2: None if v is None or (isinstance(v, float) and np.isnan(v)) else round(v, p)
    for k, e in enumerate(evs, start=1):
        ma = e["ma"]
        dist = {m: ((e["c"] / ma[m] - 1.0) * 100.0 if ma[m] else None) for m in MAS}
        p = e.get("pf")
        note = e["status"] if e["status"] != "transakcja" else (
            f"stawka pomniejszona × {p['scale']:.2f}" if p and p.get("scale", 1.0) < 0.999 else "")
        zd.append([k, e["sym"], gshort[e["group"]], e["first"], e["last"], e["k"], e["len"], rr(e["c"], 4), rr(e["drop"]),
                   rr(e["R"]), rr(e["dropR"]), rr(e["spy7"]), None if e["spy_dn"] is None else ("tak" if e["spy_dn"] else "nie"),
                   rr(e["rsi14"], 1), rr(ma[50], 4), rr(dist[50]), rr(ma[100], 4), rr(dist[100]), rr(ma[200], 4), rr(dist[200]),
                   None if e["ma200_up"] is None else ("tak" if e["ma200_up"] else "nie"),
                   "spełnione" if e["ok"] else "odrzucona",
                   e.get("in"), e.get("out"), rr(e.get("pin"), 4), rr(e.get("pout"), 4), rr(e.get("worst")),
                   rr(e.get("move"), 3), rr(e.get("net"), 3), rr(e["net"] * lev if e.get("net") is not None else None, 2),
                   "tak" if p else "nie", rr(p["pnl"]) if p else None, note])
        rw = zd.max_row
        if e.get("net") is not None:
            zd.cell(rw, 29).fill = (S.good if e["net"] > 0 else S.bad) if e["ok"] else S.sub
        zd.cell(rw, 22).fill = S.good if e["ok"] else S.sub
        for j, m in ((16, 50), (18, 100), (20, 200)):
            if dist[m] is not None:
                zd.cell(rw, j).fill = S.good if dist[m] > 0 else S.bad
    zd.freeze_panes = "C3"
    zd.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{zd.max_row}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
def selftest() -> None:
    fails: list[str] = []

    def check(name, got, exp):
        ok = got == exp
        print(f"  {'OK  ' if ok else 'BŁĄD'} {name}: {got}" + ("" if ok else f" (oczekiwano {exp})"))
        if not ok:
            fails.append(name)

    print(f"IA 4 {__version__} — selftest ia4.swportfel\n")
    sw.selftest_core(check)
    # A: 30 sesji płasko, 7 spadków po 1%, 4 wzrosty;  B: 30 płasko, 9 spadków po 2%, 6 wzrostów;  SPY: płasko
    def mk(down, step, up, sym):
        o, c, p = [100.0] * 30, [100.2] * 30, 100.0
        for _ in range(down):
            o.append(p), c.append(p * (1 - step))
            p *= 1 - step
        for _ in range(up):
            o.append(p), c.append(p * 1.01)
            p *= 1.01
        o, c = np.array(o), np.array(c)
        h, l = np.maximum(o, c) * 1.004, np.minimum(o, c) * 0.996
        dts = np.array([str(np.datetime64("2010-01-04") + np.timedelta64(k, "D")) for k in range(len(c))])
        return sw.Stock(sym, "main", dts, o, h, l, c, np.ones(len(c)))
    A = mk(7, 0.01, 4, "A")
    D = sw.finish([A])
    ev = events(D, [0], Rule(), 0.3)
    check("jedna seria 7 spadków", len(ev), 1)
    e = ev[0]
    dts, o, c = A.dates, A.o, A.c
    check("sygnał po 7. świecy, wejście 8., wyjście po close 10.", (e["last"], e["in"], e["out"]), (dts[36], dts[37], dts[39]))
    check("spadek serii %", round(e["drop"], 4), round((0.99 ** 7 - 1) * 100, 4))
    check("wynik netto = ruch − spread", round(e["net"], 6), round((c[39] / o[37] - 1) * 100 - 0.3, 6))
    pf = sw.portfolio([dict(e["trade"])], D, lev=5.0)
    check("portfel 1:5, depozyt 10%", round(pf["ret"], 6), round(e["net"] * 5 * 0.10, 6))
    dd = drawdowns([("2020-01-01", 100.0), ("2020-01-02", 90.0), ("2020-01-03", 80.0), ("2020-01-04", 101.0), ("2020-01-05", 99.0)], 0.5)
    check("obsunięcia: głębokość, dno, odrobienie", [(round(x["depth"], 4), x["trough"], x["rec"]) for x in dd],
          [(20.0, "2020-01-03", "2020-01-04"), (round(2 / 101 * 100, 4), "2020-01-05", None)])
    # Warunki: ≥ 5 świec, spadek ≥ 10%: seria B (2% dziennie) spełnia na 6. świecy (−11,4%), A (−6,8%) odrzucona
    B = mk(9, 0.02, 6, "B")
    spy = sw.Stock("SPY", "context", A.dates, np.full(41, 100.0), np.full(41, 100.5), np.full(41, 99.5), np.full(41, 100.0), np.ones(41))
    D2 = sw.finish([A, B, spy])
    r = Rule(min_n=5, min_drop=10.0, hold=5)
    ev2 = {e["sym"]: e for e in events(D2, [0, 1], r, 0.3)}
    check("seria A odrzucona (spadek < 10%)", (ev2["A"]["ok"], ev2["A"]["trade"] is None), (False, True))
    check("seria B: wejście po 6. świecy (pierwszy dzień ≥ 10%)", (ev2["B"]["ok"], ev2["B"]["k"]), (True, 6))
    ev3 = {e["sym"]: e for e in events(D2, [0, 1], replace(r, market=True), 0.3)}
    check("słaby rynek: SPY płasko i bez SMA200 → odrzucona", ev3["B"]["ok"], False)
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        pth = run(no_sync=True, out=Path(td) / "t.xlsx", data=D2, rule=r)
        from openpyxl import load_workbook
        wb = load_workbook(pth)
        check("arkusze", wb.sheetnames, ["PODSUMOWANIE", "ANALIZA", "KAPITAŁ", "OBSUNIĘCIA", "ROK PO ROKU", "TRANSAKCJE", "ZDARZENIA"])
    check("nazwy plików", (Rule().tag(), PROPOSAL.tag(), r.tag()), ("7", "PROPOZYCJA", "5plus-spadek10-5s"))
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    if fails:
        sys.exit(1)


def main() -> None:
    ap = argparse.ArgumentParser(description="Portfel strategii serii świec spadkowych (instrukcja 7a)")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--no-sync", action="store_true", help="bez synchronizacji bazy dobowej")
    ap.add_argument("--since", default=sw.DAILY_SINCE, help="początek portfela (domyślnie 2000-01-01)")
    ap.add_argument("--cost", type=float, default=sw.COST_PCT, help="spread %% na transakcję (raz; domyślnie 0,3)")
    ap.add_argument("--propozycja", action="store_true",
                    help="≥ 5 spadków, spadek ≥ 10%%, słaby rynek, 5 sesji, max 5 pozycji (pozostałe opcje ją zmieniają)")
    ap.add_argument("--n", type=int, help="seria DOKŁADNIE n świec spadkowych (domyślnie 7)")
    ap.add_argument("--min-n", type=int, help="seria CO NAJMNIEJ n świec — wejście w pierwszym dniu spełnienia warunków")
    ap.add_argument("--min-drop", type=float, help="spadek serii co najmniej tyle %% (np. 10)")
    ap.add_argument("--rynek", action="store_true", help=f"słaby rynek: SPY w serii ≤ −{SPY_DROP:g}%% albo SPY pod SMA200")
    ap.add_argument("--hold", type=int, help="sesji w pozycji (domyślnie 3; propozycja 5)")
    ap.add_argument("--max-poz", type=int, help="najwięcej pozycji naraz (propozycja 5)")
    ap.add_argument("--short", action="store_true", help="short zamiast long")
    ap.add_argument("--lev", type=float, help="dźwignia (domyślnie 5)")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    if a.selftest:
        selftest()
        return
    r = PROPOSAL if a.propozycja else Rule()
    ch = {}
    if a.n is not None:
        ch.update(n=a.n, min_n=None)
    if a.min_n is not None:
        ch["min_n"] = a.min_n
    if a.min_drop is not None:
        ch["min_drop"] = a.min_drop
    if a.rynek:
        ch["market"] = True
    if a.hold is not None:
        ch["hold"] = a.hold
    if a.max_poz is not None:
        ch["max_pos"] = a.max_poz or None
    if a.short:
        ch["d"] = -1
    if a.lev is not None:
        ch["lev"] = a.lev
    run(a.no_sync, a.out, a.since, a.cost, replace(r, **ch))


if __name__ == "__main__":
    main()
