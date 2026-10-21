"""
IA 4 — FX-POWRÓT v2: skąd biorą się stratne transakcje FX-POWRÓT, czy da się je uciąć (czas trzymania, cel, SL,
weekend) i sprawdzian zamrożonej strategii na 10 nowych parach fx1h_double_proof (instrukcja, sekcja 7b).
Wersja projektu: 1.58 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Pytanie człowieka (20.10.2026): większość zyskownych transakcji zamyka się szybko — czy krótszy czas trzymania (albo
inne wyjście) obetnie stratne? Wolumenu nie ma: Yahoo podaje dla par walutowych 0, baza fx1h go nie zapisuje.

Wejście bez zmian (FX-POWRÓT, 1.50): świeca zamyka się ≥ 2 ATR(100) wyżej niż poprzednia; pierwsza czerwona świeca
w ciągu 3 świec, świeca sygnału 7–16 UTC → short po open (open ≥ 0,5 ATR nad bazą = ceną sprzed wzrostu).
Rozliczenie pesymistyczne (zasada 5), spread raz na transakcję, jedna pozycja na parę.

  * v1 = FX-POWRÓT: TP w połowie drogi od szczytu do bazy, bez SL, najwyżej 96 świec.
  * v2 (zamrożona 20.10.2026, przed pobraniem par doubleProof): jak v1, ale
      - bez transakcji, gdy do TP jest mniej niż 1 × spread (cel nie pokrywa kosztu — pewna strata),
      - najwyżej 60 świec (po 60 świecach trzymanie dalej nie jest już lepsze od zamknięcia, a noc i weekend kosztują).
  * Analiza: wynik wg czasu trzymania, „trzymać czy zamknąć po k świecach”, ok. 40 wariantów wyjścia (limit czasu,
    wyjście warunkowe, cel bliżej po k świecach, cel ≤ x ATR, SL od wejścia, stop na wejściu, przed weekendem),
    lista stratnych i cechy wejścia transakcji trzymanych długo.
  * DOUBLEPROOF: v1 i v2 raz na parach fx1h_double_proof (EURCHF, EURAUD, EURCAD, GBPCHF, GBPAUD, AUDNZD, AUDJPY,
    CADJPY, CHFJPY, NZDJPY) — dane, których nikt nie oglądał. Warunek zaliczenia ustalony przed testem: v2 ≥ 30
    transakcji, średnia netto > 0 i t ≥ 2. Każde otwarcie zapisuje się w data/fxpowrot2/doubleproof.jsonl.
    Spread tych par — szacunek (do potwierdzenia zrzutem z XTB): DP_SPREAD.

    python -m ia4.fxpowrot2              # → ~/Desktop/IA4/FX-POWROT-V2.xlsx (doubleProof, gdy pary są już w Firestore)
    python -m ia4.fxpowrot2 --bez-sync   # bez odświeżenia kopii fx1h
    python -m ia4.fxpowrot2 --strona     # tylko pełna historia dla strony (branch fx, fx/pv2/; run() robi to też na końcu)
    python -m ia4.fxpowrot2 --selftest   # „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import numba as nb
import numpy as np
import pandas as pd

from . import __version__, config
from . import fxpowrot as fp
from . import fxshort as fs
from .fxshort import NPER, PAIRS

DP_PAIRS = ('EURCHF', 'EURAUD', 'EURCAD', 'GBPCHF', 'GBPAUD', 'AUDNZD', 'AUDJPY', 'CADJPY', 'CHFJPY', 'NZDJPY')
# spread w pipsach na transakcję — SZACUNEK dla XTB (typowy spread ECN × 1,1 / 0,8 jak przy pozostałych parach)
DP_SPREAD = {'EURCHF': 1.93, 'EURAUD': 2.61, 'EURCAD': 2.75, 'GBPCHF': 3.3, 'GBPAUD': 3.58, 'AUDNZD': 2.75,
             'AUDJPY': 2.06, 'CADJPY': 2.75, 'CHFJPY': 3.3, 'NZDJPY': 2.75}
CONV = {'EUR': ('EURUSD', False), 'GBP': ('GBPUSD', False), 'AUD': ('AUDUSD', False), 'NZD': ('NZDUSD', False),
        'CAD': ('USDCAD', True), 'CHF': ('USDCHF', True), 'JPY': ('USDJPY', True), 'USD': (None, False)}
MAJORS = ('EURUSD', 'GBPUSD', 'AUDUSD', 'NZDUSD', 'USDCAD', 'USDCHF', 'USDJPY')
WIDZIANE_DO = '2026-10-17'          # świece par bazowych oglądane przy badaniach 1.49–1.51; późniejsze = świeże
DP_MIN_CANDLES = 2000               # mniej świec na parę — doubleProof jeszcze się nie liczy (historia w toku)

V1 = dict(name='v1 = FX-POWRÓT', f=0.5, H=96, mintp=0.0)
V2 = dict(name='v2', f=0.5, H=60, mintp=1.0)
CRIT = dict(min_n=30, t=2.0)        # warunek zaliczenia doubleProof (ustalony przed testem)
WHY = {0: 'TP', 1: 'czas', 2: 'wyjście warunkowe', 3: 'przed weekendem', 4: 'SL'}


def pip_of(p):
    return 0.01 if p.endswith('JPY') else 0.0001


# ---------------------------------------------------------------- dane

class Rynek:
    """Świece wybranych par na wspólnej siatce czasu + pary do przeliczenia wartości pozycji na USD."""

    def __init__(self, frames: dict, pairs, spread: dict):
        self.pairs = tuple(pairs)
        need = set(self.pairs) | {CONV[p[:3]][0] for p in self.pairs if CONV[p[:3]][0]}
        common = None
        for p in need:
            tt = set(frames[p].time.values.tolist())
            common = tt if common is None else common & tt
        tgrid = np.array(sorted(common), np.int64)
        self.F = {}
        for p in need:
            F = fs.features(frames[p].set_index('time').loc[tgrid].reset_index())
            c = F['c']; F['run1'] = (c - np.r_[np.nan, c[:-1]]) / F['atr']
            F['ema10'] = fs._ema(c, 10)
            self.F[p] = F
        self.t = tgrid; self.n = len(tgrid); self.spread = spread
        T0, T1 = int(tgrid[0]), int(tgrid[-1]) + 3600
        self.per = np.minimum(((tgrid - T0) / (T1 - T0) * NPER).astype(np.int64), NPER - 1)
        self.periods = [(pd.Timestamp(T0 + (T1 - T0) * k / NPER, unit='s'), pd.Timestamp(T0 + (T1 - T0) * (k + 1) / NPER, unit='s'))
                        for k in range(NPER)]
        nxt = np.r_[tgrid[1:], tgrid[-1] + 3600]
        self.fri = ((nxt - tgrid) > 6 * 3600).astype(np.int64)          # ostatnia świeca przed przerwą weekendową

    def usd(self, p, j):
        """Wartość 1 jednostki waluty bazowej pary p w USD (open świecy j)."""
        q, inv = CONV[p[:3]]
        if q is None:
            return 1.0
        x = float(self.F[q]['o'][j])
        return 1.0 / x if inv else x


# ---------------------------------------------------------------- zdarzenia i wyjścia

def events(R: Rynek, hours=(7, 16), dist=0.5) -> pd.DataFrame:
    N, k, b, s, w = fp.main_ev()
    out = []
    for p in R.pairs:
        F = R.F[p]
        J, B, P, S0 = fp.gen_trades(F['o'], F['h'], F['l'], F['c'], F['atr'], F[f'run{N}'], k * np.sqrt(N), N, b, s, w, F['ema10'],
                                    F['valid'], fs.WARMUP)
        A = F['atr'][J]; en = F['o'][J + 1]
        m = (F['hour'][J] >= hours[0]) & (F['hour'][J] <= hours[1]) & ((en - B) / A >= dist)
        out.append(pd.DataFrame(dict(para=p, j=J[m], B=B[m], P=P[m], A=A[m], s0=S0[m])))
    return pd.concat(out, ignore_index=True) if out else pd.DataFrame(columns=['para', 'j', 'B', 'P', 'A', 's0'])


@nb.njit(cache=True)
def run_exit(o, h, l, c, fri, J, B, P, A, sp, f, H, Ts, u, wk, mintp, sl_atr, be, tpcap, td, g):
    """Short po open J+1. Zwraca wynik % netto, świece w pozycji, powód (0 TP, 1 czas, 2 warunkowe, 3 weekend, 4 SL; −1 brak)."""
    n = o.size; nt = J.size
    res = np.full(nt, np.nan); nb_ = np.zeros(nt, np.int64); why = np.full(nt, -1, np.int64)
    for t in range(nt):
        e0 = J[t] + 1
        if e0 + H > n:
            continue
        en = o[e0]; TP = P[t] - f * (P[t] - B[t])
        if tpcap > 0 and en - TP > tpcap * A[t]:
            TP = en - tpcap * A[t]
        if en - TP < mintp * sp or en - TP < 0.1 * A[t]:
            continue
        SL = en + sl_atr * A[t] if sl_atr > 0 else 1e18
        best = en; TP0 = TP; ex = np.nan; w = -1; k = 0
        for j in range(e0, e0 + H):
            k = j - e0 + 1
            if td > 0 and k == td + 1:
                TP = en - g * (en - TP0)
            if j > e0 and o[j] >= SL:
                ex = o[j]; w = 4; break
            if j > e0 and o[j] <= TP:
                ex = TP; w = 0; break
            if h[j] >= SL:
                ex = SL; w = 4; break
            if l[j] <= TP:
                ex = TP; w = 0; break
            if l[j] < best:
                best = l[j]
            if be > 0 and en - best >= be * (en - TP0):
                SL = min(SL, en)
            if Ts > 0 and k == Ts and (en - c[j]) / en * 100.0 - sp / en * 100.0 < u:
                ex = c[j]; w = 2; break
            if wk and fri[j] == 1:
                ex = c[j]; w = 3; break
            if k == H:
                ex = c[j]; w = 1
        res[t] = (en - ex) / en * 100.0 - sp / en * 100.0; nb_[t] = k; why[t] = w
    return res, nb_, why


def simulate(R: Rynek, E: pd.DataFrame, f=0.5, H=96, Ts=0, u=-99.0, wk=False, mintp=0.0, sl_atr=0.0, be=0.0, tpcap=0.0, td=0, g=1.0,
             **_) -> pd.DataFrame:
    parts = []
    for p in R.pairs:
        q = E[E.para == p]
        if not len(q):
            continue
        F = R.F[p]; sp = R.spread[p] * pip_of(p)
        r, x, w = run_exit(F['o'], F['h'], F['l'], F['c'], R.fri, q.j.values.astype(np.int64), q.B.values, q.P.values, q.A.values, sp,
                           f, H, Ts, u, wk, mintp, sl_atr, be, tpcap, td, g)
        d = q.assign(wynik=r, swiece=x, powod=w)
        d = d[~np.isnan(d.wynik)]
        keep = np.zeros(len(d), bool); last = -1                      # jedna pozycja na parę
        for k_, (jj, xx) in enumerate(zip(d.j.values, d.swiece.values)):
            if jj > last:
                keep[k_] = True; last = jj + xx
        parts.append(d[keep])
    if not parts:
        return pd.DataFrame(columns=['para', 'j', 'B', 'P', 'A', 's0', 'wynik', 'swiece', 'powod', 't', 'okres'])
    D = pd.concat(parts).sort_values(['j', 'para']).reset_index(drop=True)
    D['t'] = pd.to_datetime(R.t[D.j.values], unit='s'); D['okres'] = R.per[D.j.values] + 1
    return D


def stats(D: pd.DataFrame, granica='2025-10-15') -> dict:
    w = D.wynik.values.astype(float); n = len(w)
    if n == 0:
        return dict(n=0, srednia=np.nan, mediana=np.nan, suma=0.0, t=np.nan, ci=np.nan, trafne=np.nan, PF=np.nan, najgorsza=np.nan,
                    stratne_okresy=0, rok1=0.0, rok2=0.0, obs=0.0, med_sw=np.nan)
    per = D.groupby('okres').wynik.sum().reindex(range(1, NPER + 1), fill_value=0)
    y1 = (D.t < granica).values
    eq = np.cumsum(w); dd = float((np.maximum.accumulate(np.r_[0, eq])[1:] - eq).max())
    sd = w.std(ddof=1) if n > 1 else np.nan
    return dict(n=n, srednia=w.mean(), mediana=float(np.median(w)), suma=w.sum(), t=w.mean() / sd * np.sqrt(n) if sd and sd > 0 else np.nan,
                ci=1.96 * sd / np.sqrt(n) if n > 1 else np.nan, trafne=(w > 0).mean(),
                PF=w[w > 0].sum() / -w[w < 0].sum() if (w < 0).any() else 99.0, najgorsza=w.min(),
                stratne_okresy=int((per < 0).sum()), rok1=w[y1].sum(), rok2=w[~y1].sum(), obs=dd, med_sw=float(np.median(D.swiece)))


def variants():
    """Warianty wyjścia (nazwa, grupa, parametry) — wszystkie z tym samym wejściem."""
    V = [('v1 = FX-POWRÓT (TP ½ drogi, bez SL, 96 św.)', 'odniesienie', dict()),
         ('v2 (TP ≥ 1 × spread, 60 św.)', 'odniesienie', dict(mintp=1, H=60))]
    V += [(f'limit {H} św.', 'krótszy czas', dict(H=H)) for H in (6, 12, 18, 24, 36, 48, 60, 72)]
    V += [(f'po {Ts} św. zamknij, jeśli wynik < ' + f'{u:+.1f}%'.replace('.', ','), 'wyjście warunkowe', dict(Ts=Ts, u=u))
          for Ts in (12, 24, 36, 48) for u in (0.0, -0.2)]
    V += [(f'po {td} św. cel = ' + f'{g:.1f}'.replace('.', ',') + ' drogi do TP', 'cel bliżej po czasie', dict(td=td, g=g, mintp=1))
          for td in (12, 24, 48) for g in (0.0, 0.3)]
    V += [(f'cel ≤ {c:g} ATR od wejścia'.replace('.', ','), 'cel bliżej', dict(tpcap=c, mintp=1)) for c in (1.0, 1.5, 2.0)]
    V += [(f'SL {s} ATR od wejścia', 'stop loss', dict(sl_atr=s)) for s in (2, 3, 4, 6)]
    V += [('stop na wejściu po ' + f'{b:.1f}'.replace('.', ',') + ' drogi', 'stop loss', dict(be=b)) for b in (0.5, 0.8)]
    V += [('zamknij przed weekendem', 'weekend', dict(wk=True))]
    V += [(f'TP ≥ {m} × spread', 'koszt', dict(mintp=m)) for m in (1, 2, 3)]
    V += [(f'TP ≥ 1 × spread, limit {H} św.', 'v2 — wybór limitu', dict(mintp=1, H=H)) for H in (48, 72, 96, 144)]
    return V


# ---------------------------------------------------------------- rachunek 1000 $ na 1:30

def konto(D: pd.DataFrame, R: Rynek, f=0.20, lev=30.0, start=1000.0, stop_out=0.5, lot=1000):
    """Rachunek CFD jak fxshort.simulate (depozyt = f × saldo, 0,01 lota, stop out < 50%), z przeliczeniem waluty bazowej na USD."""
    ent = {}
    for k, r in enumerate(D.itertuples()):
        ent.setdefault(int(r.j) + 1, []).append(k)
    cash = start; open_ = {}; eq = np.full(R.n, np.nan); ev = 0; skipped = 0; done = 0; dead = False
    rows = list(D.itertuples())
    for j in range(R.n):
        if dead:
            eq[j] = cash; continue
        for k in ent.get(j, []):
            r = rows[k]; p = r.para
            if p in open_:
                skipped += 1; continue
            used = sum(q['margin'] for q in open_.values())
            margin = min(f * cash, cash - used)
            bu = R.usd(p, j)
            units = int(max(margin, 0) * lev / bu // lot) * lot
            if units < lot:
                skipped += 1; continue
            notional = units * bu; e0 = float(R.F[p]['o'][j])
            open_[p] = dict(j1=j + int(r.swiece) - 1, notional=notional, margin=notional / lev, entry=e0, pnl=float(r.wynik),
                            sp=R.spread[p] * pip_of(p))
        if open_:
            worst = sum(-q['notional'] * (float(R.F[p]['h'][j]) - q['entry'] + q['sp']) / q['entry'] for p, q in open_.items())
            used = sum(q['margin'] for q in open_.values())
            if cash + worst < stop_out * used:
                cash += worst; open_.clear(); ev += 1
                if cash <= 0:
                    cash = 0.0; dead = True
                eq[j] = cash
                continue
        for p in list(open_):
            q = open_[p]
            if q['j1'] == j:
                cash += q['notional'] * q['pnl'] / 100.0; done += 1
                del open_[p]
        eq[j] = cash + sum(-q['notional'] * (float(R.F[p]['c'][j]) - q['entry'] + q['sp']) / q['entry'] for p, q in open_.items())
    curve = pd.Series(eq, index=pd.to_datetime(R.t, unit='s')).ffill().fillna(start)
    peak = curve.cummax(); obs = float(((peak - curve) / peak).max() * 100)
    return dict(koniec=float(curve.iloc[-1]), zysk_pct=(curve.iloc[-1] / start - 1) * 100, obs_pct=obs, transakcje=done, pominiete=skipped,
                stop_out=ev), curve.resample('D').last().dropna()


# ---------------------------------------------------------------- analiza

def features_of(D: pd.DataFrame, R: Rynek) -> pd.DataFrame:
    rows = []
    for r in D.itertuples():
        F = R.F[r.para]; j = int(r.j); en = F['o'][j + 1]; A = r.A
        rows.append(dict(impuls=F['run1'][int(r.s0)], do_bazy=(en - r.B) / A, od_szczytu=(r.P - en) / A, zakres=(r.P - r.B) / A,
                         czekanie=j - int(r.s0), godzina=int(F['hour'][j]), dzien=int(pd.Timestamp(int(R.t[j]), unit='s').dayofweek),
                         nachylenie500=F['slope500'][j], zmiennosc=F['atr'][j] / F['atr500'][j], czerwona=(F['o'][j] - F['c'][j]) / A,
                         ruch24=(F['c'][j] - F['c'][j - 24]) / A, wejscie=en, cel=(r.P - 0.5 * (r.P - r.B)),
                         cel_atr=(en - (r.P - 0.5 * (r.P - r.B))) / A,
                         cel_spread=(en - (r.P - 0.5 * (r.P - r.B))) / (R.spread[r.para] * pip_of(r.para))))
    return pd.concat([D.reset_index(drop=True), pd.DataFrame(rows)], axis=1)


def hold_tables(D: pd.DataFrame, R: Rynek):
    b = pd.cut(D.swiece, [0, 1, 2, 3, 6, 12, 24, 48, 96, 1000], labels=['1', '2', '3', '4–6', '7–12', '13–24', '25–48', '49–96', '> 96'])
    B = D.groupby(b, observed=True).wynik.agg(n='size', srednia='mean', suma='sum', trafne=lambda s: (s > 0).mean()).reset_index()
    B = B.rename(columns={'swiece': 'swiece'})
    rows = []
    for k in (1, 3, 6, 12, 24, 36, 48, 60, 72):
        o = D[D.swiece > k]
        if not len(o):
            continue
        mk = np.array([(R.F[r.para]['o'][r.j + 1] - R.F[r.para]['c'][r.j + k]) / R.F[r.para]['o'][r.j + 1] * 100
                       - R.spread[r.para] * pip_of(r.para) / R.F[r.para]['o'][r.j + 1] * 100 for r in o.itertuples()])
        rows.append(dict(k=k, otwarte=len(o), udzial=len(o) / len(D), dojdzie_do_TP=(o.powod == 0).mean(), trzymaj=o.wynik.mean(),
                         zamknij=mk.mean(), roznica_suma=(o.wynik.values - mk).sum(), trzymanie_lepsze=(o.wynik.values > mk).mean()))
    return B, pd.DataFrame(rows)


def quartiles(X: pd.DataFrame):
    out = []
    X = X.assign(dlugie=X.swiece > 48)
    for c, opis in [('do_bazy', 'odległość wejścia od bazy (ATR)'), ('zakres', 'szczyt − baza (ATR)'), ('impuls', 'świeca impulsu (ATR)'),
                    ('od_szczytu', 'szczyt − wejście (ATR)'), ('nachylenie500', 'nachylenie SMA500'), ('zmiennosc', 'ATR100 / ATR500'),
                    ('czerwona', 'czerwona świeca sygnału (ATR)'), ('ruch24', 'ruch 24 świec przed sygnałem (ATR)')]:
        q = pd.qcut(X[c], 4, duplicates='drop')
        g = X.groupby(q, observed=True).agg(n=('wynik', 'size'), srednia=('wynik', 'mean'), suma=('wynik', 'sum'), dlugie=('dlugie', 'mean'))
        for iv, r in g.iterrows():
            out.append(dict(cecha=opis, przedzial=f'{iv.left:.2f} – {iv.right:.2f}'.replace('.', ','), **r.to_dict()))
    for c, opis, names in [('dzien', 'dzień tygodnia', ['pn', 'wt', 'śr', 'czw', 'pt', 'sb', 'nd']), ('godzina', 'godzina sygnału (UTC)', None)]:
        g = X.groupby(c).agg(n=('wynik', 'size'), srednia=('wynik', 'mean'), suma=('wynik', 'sum'), dlugie=('dlugie', 'mean'))
        for iv, r in g.iterrows():
            out.append(dict(cecha=opis, przedzial=names[iv] if names else str(iv), **r.to_dict()))
    return pd.DataFrame(out)


def pair_table(D: pd.DataFrame, pairs):
    rows = []
    for p in pairs:
        s = stats(D[D.para == p])
        rows.append(dict(para=p, **{k: s[k] for k in ['n', 'srednia', 'suma', 't', 'trafne', 'najgorsza', 'med_sw']}))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- doubleProof

LOG_DIR = config.RESEARCH_DIR / 'data' / 'fxpowrot2'          # dziennik otwarć doubleProof (zasada 4)


def frozen_hash() -> str:
    s = json.dumps(dict(V1=V1, V2=V2, CRIT=CRIT, ev=fp.main_ev(), DP=DP_PAIRS, SP=DP_SPREAD), sort_keys=True)
    return hashlib.sha256(s.encode()).hexdigest()[:16]


def dp_open(R: Rynek, log=print) -> list:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    path = LOG_DIR / 'doubleproof.jsonl'
    old = [json.loads(x) for x in path.read_text(encoding='utf-8').splitlines() if x.strip()] if path.exists() else []
    e = dict(otwarcie=len(old) + 1, czas=datetime.now(timezone.utc).isoformat(timespec='seconds'), reguly=frozen_hash(), wersja=__version__,
             od=str(pd.Timestamp(int(R.t[0]), unit='s')), do=str(pd.Timestamp(int(R.t[-1]), unit='s')), swiec=int(R.n))
    with path.open('a', encoding='utf-8') as fh:
        fh.write(json.dumps(e, ensure_ascii=False) + '\n')
    log(f'doubleProof otwarty (otwarcie nr {e["otwarcie"]}, reguły {e["reguly"]}) — zapis: {path}')
    return old + [e]


def compute(R: Rynek, Rdp: Rynek | None, f=0.20, lev=30.0, start=1000.0, log=print) -> dict:
    O = dict(f=f, lev=lev, start=start, od=pd.Timestamp(int(R.t[0]), unit='s'), do=pd.Timestamp(int(R.t[-1]), unit='s'), swiec=R.n,
             spread=R.spread, periods=R.periods, hash=frozen_hash())
    E = events(R)
    log(f'Zdarzenia (impuls → czerwona świeca, 7–16 UTC): {len(E)}')
    D1 = simulate(R, E, **V1); D2 = simulate(R, E, **V2)
    O['D1'], O['D2'] = D1, D2
    O['s1'], O['s2'] = stats(D1), stats(D2)
    sw = D2[D2.t >= WIDZIANE_DO]
    O['swieze'] = stats(sw) if len(sw) else None
    log(f"v1: {O['s1']['n']} transakcji, średnio {O['s1']['srednia']:+.3f}% (t {O['s1']['t']:.1f}); "
        f"v2: {O['s2']['n']}, {O['s2']['srednia']:+.3f}% (t {O['s2']['t']:.1f})")
    O['warianty'] = pd.DataFrame([dict(wariant=nm, grupa=gr, **stats(simulate(R, E, **kw))) for nm, gr, kw in variants()])
    O['czas'], O['trzymaj'] = hold_tables(D1, R)
    X = features_of(D1, R)
    O['stratne'] = X[X.wynik <= 0].sort_values('wynik').reset_index(drop=True)
    O['cechy'] = quartiles(X)
    O['pary'] = pair_table(D1, R.pairs).merge(pair_table(D2, R.pairs), on='para', suffixes=('_v1', '_v2'))
    O['konto'] = {}; O['krzywe'] = {}
    for nm, D in [('v1', D1), ('v2', D2)]:
        st, cv = konto(D, R, f=f, lev=lev, start=start)
        O['konto'][nm] = st; O['krzywe'][nm] = cv
    O['dp'] = None
    if Rdp is not None:
        O['dp_log'] = dp_open(Rdp, log)
        Edp = events(Rdp)
        A1 = simulate(Rdp, Edp, **V1); A2 = simulate(Rdp, Edp, **V2)
        s1, s2 = stats(A1), stats(A2)
        ok = s2['n'] >= CRIT['min_n'] and s2['srednia'] > 0 and (s2['t'] or 0) >= CRIT['t']
        k1, c1 = konto(A1, Rdp, f=f, lev=lev, start=start); k2, c2 = konto(A2, Rdp, f=f, lev=lev, start=start)
        O['dp'] = dict(D1=A1, D2=A2, s1=s1, s2=s2, zaliczony=ok, pary=pair_table(A1, Rdp.pairs).merge(pair_table(A2, Rdp.pairs), on='para', suffixes=('_v1', '_v2')),
                       konto={'v1': k1, 'v2': k2}, krzywe={'v1': c1, 'v2': c2}, od=pd.Timestamp(int(Rdp.t[0]), unit='s'),
                       do=pd.Timestamp(int(Rdp.t[-1]), unit='s'), swiec=Rdp.n, spread=Rdp.spread)
        log(f"doubleProof: v2 {s2['n']} transakcji, średnio {s2['srednia']:+.3f}% (t {s2['t']:.1f}) → {'ZALICZONY' if ok else 'NIE zaliczony'}; "
            f"v1 {s1['n']}, {s1['srednia']:+.3f}%")
    return O


# ---------------------------------------------------------------- strona: pełna historia (branch fx, fx/pv2/)

def history_docs(R: Rynek) -> tuple[dict, dict]:
    """Świece każdej pary (liczby całkowite w jednostkach ostatniego miejsca, czas jako odstępy w godzinach) i wszystkie
    transakcje v2 na całej historii — dla strony (FX → Powrót.v2 → historia / galeria)."""
    D = simulate(R, events(R), **V2)
    trades = []
    for r in D.itertuples():
        F = R.F[r.para]; j = int(r.j); e0 = j + 1; x = int(r.swiece); ex_i = j + x
        en = float(F['o'][e0]); TP = float(r.P - V2['f'] * (r.P - r.B))
        ex = TP if int(r.powod) == 0 else float(F['c'][ex_i])
        dp = 3 if r.para.endswith('JPY') else 5
        trades.append(dict(pair=r.para, sigT=int(R.t[j]), entryT=int(R.t[e0]), exitT=int(R.t[ex_i]), entry=round(en, dp), tp=round(TP, dp),
                           exit=round(ex, dp), base=round(float(r.B), dp), peak=round(float(r.P), dp), kind='TP' if int(r.powod) == 0 else 'czas (60 świec)',
                           ret=round(float(r.wynik), 4), bars=x))
    pairs = {}
    for p in R.pairs:
        F = R.F[p]; dp = 3 if p.endswith('JPY') else 5; k = 10 ** dp
        t = F['t'].astype(np.int64)
        pairs[p] = dict(pair=p, dp=dp, t0=int(t[0]), dt=[0] + (np.diff(t) // 3600).astype(int).tolist(),
                        **{q: np.rint(F[q] * k).astype(np.int64).tolist() for q in 'ohlc'})
    s = stats(D)
    index = dict(version=__version__, generatedAt=datetime.now(timezone.utc).isoformat(timespec='seconds'), od=int(R.t[0]), do=int(R.t[-1]),
                 rules=dict(V2), pairs=list(R.pairs), spread={p: R.spread[p] for p in R.pairs}, trades=trades,
                 summary=dict(n=s['n'], srednia=round(float(s['srednia']), 4), t=round(float(s['t']), 2), trafne=round(float(s['trafne']), 4)))
    return index, pairs


def publish_history(R: Rynek, push=True, log=print) -> bool:
    """Zapis na branch `fx` (fx/pv2/index.json + fx/pv2/{PARA}.json) — strona czyta stamtąd pełną historię."""
    from .fx.store import fx_repo
    index, pairs = history_docs(R)
    repo = fx_repo(push=push)
    repo.ensure()
    d = repo.dir / 'pv2'
    d.mkdir(parents=True, exist_ok=True)
    for p, doc in pairs.items():
        (d / f'{p}.json').write_text(json.dumps(doc, separators=(',', ':')), encoding='utf-8')
    (d / 'index.json').write_text(json.dumps(index, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    ok = repo.commit_push(f"pv2: historia {len(index['trades'])} transakcji, świece do {pd.Timestamp(index['do'], unit='s'):%Y-%m-%d %H:%M}")
    log(f"Strona: historia Powrót.v2 — {len(index['trades'])} transakcji, {R.n} świec × {len(pairs)} par → branch fx, fx/pv2/"
        + ('' if ok else f'  ! push nieudany: {repo.last_error}'))
    return ok


# ---------------------------------------------------------------- Excel

def write_xlsx(O: dict, path: Path, wnioski=None, dp_status='') -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    FN = 'Arial'
    fN = Font(name=FN, size=10); fB = Font(name=FN, size=10, bold=True); fT = Font(name=FN, size=14, bold=True)
    fH = Font(name=FN, size=10, bold=True, color='FFFFFF'); fS = Font(name=FN, size=9, italic=True, color='555555')
    HDR = PatternFill('solid', fgColor='1F3A5F'); HDR2 = PatternFill('solid', fgColor='7A4B00'); GRY = PatternFill('solid', fgColor='EEF2F7')
    GOOD = PatternFill('solid', fgColor='E2F0D9'); BAD = PatternFill('solid', fgColor='FBE5D6'); SEL = PatternFill('solid', fgColor='FFF2CC')
    thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')
    PCT = '0.0%'; PCT3 = '0.000%'; USD = '#,##0.00 [$$-409];[Red]-#,##0.00 [$$-409]'; USD0 = '#,##0 [$$-409]'; DT = 'yyyy-mm-dd hh:mm'; DD = 'yyyy-mm-dd'

    def pl(x, d=1):
        return f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')

    wb = Workbook()

    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name; ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def hdr(ws, r, c, names, widths=None, fill=HDR):
        for k, nm in enumerate(names):
            x = ws.cell(r, c + k, nm); x.font = fH; x.fill = fill; x.border = BOX
            x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for k, w in enumerate(widths or []):
            if w:
                ws.column_dimensions[CL(c + k)].width = w
        ws.row_dimensions[r].height = 42

    def put(ws, r, c, v, font=fN, fmt=None, fill=None, al=None):
        if isinstance(v, np.floating):
            v = float(v)
        if isinstance(v, np.integer):
            v = int(v)
        if isinstance(v, float) and not np.isfinite(v):
            v = None
        x = ws.cell(r, c, v); x.font = font; x.border = BOX
        if fmt: x.number_format = fmt
        if fill: x.fill = fill
        if al: x.alignment = al
        return x

    def note(ws, r, c, text, span=10, h=None):
        x = ws.cell(r, c, text); x.font = fS; x.alignment = WR
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        if h: ws.row_dimensions[r].height = h

    SC = ['n', 'srednia', 'mediana', 'ci', 't', 'trafne', 'PF', 'najgorsza', 'stratne_okresy', 'rok1', 'rok2', 'obs', 'med_sw']
    SN = ['Transakcje', 'Średnio % (netto)', 'Mediana %', '± 95%', 't', 'Trafne', 'PF', 'Najgorsza %', 'Okresy stratne / 20', 'Rok 1: suma %',
          'Rok 2: suma %', 'Obsunięcie (suma %)', 'Mediana świec']
    SF = ['0', PCT3, PCT3, PCT3, '0.0', PCT, '0.00', PCT3, '0', PCT, PCT, PCT, '0']

    def srow(ws, r, c, s, fill=None):
        for k, (key, fm) in enumerate(zip(SC, SF)):
            v = s.get(key)
            if fm in (PCT3, PCT) and key not in ('trafne',) and v is not None and np.isfinite(v):
                v = v / 100
            put(ws, r, c + k, v, fmt=fm, fill=fill)

    # PODSUMOWANIE
    ws = sheet('PODSUMOWANIE', first=True)
    ws.column_dimensions['A'].width = 34
    ws['A1'] = 'FX-POWRÓT v2 — skąd biorą się straty i sprawdzian na nowych parach (doubleProof)'; ws['A1'].font = fT
    ws['A2'] = (f"IA4 {__version__} · badanie informacyjne · 10 par: {O['od']:%Y-%m-%d} – {O['do']:%Y-%m-%d} ({O['swiec']} świec 1h) · "
                f"reguły zamrożone: {O['hash']} · policzono {pd.Timestamp.now():%Y-%m-%d %H:%M}")
    ws['A2'].font = fS
    r = 4
    put(ws, r, 1, '10 PAR BADANIA (w próbie — na tych danych powstał FX-POWRÓT)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    hdr(ws, r + 1, 1, ['Strategia'] + SN, [44] + [11] * len(SN))
    put(ws, r + 2, 1, V1['name'] + ' (½ drogi, bez SL, 96 św.)', fB); srow(ws, r + 2, 2, O['s1'])
    put(ws, r + 3, 1, 'v2 (TP ≥ 1 × spread, 60 św.)', fB, fill=SEL); srow(ws, r + 3, 2, O['s2'], fill=SEL)
    rr = r + 4
    if O.get('swieze'):
        put(ws, rr, 1, f'v2 — świeże świece (od {WIDZIANE_DO})', fB); srow(ws, rr, 2, O['swieze']); rr += 1
    r = rr + 1
    put(ws, r, 1, f"RACHUNEK {pl(O['start'], 0)} $, 1:{O['lev']:g}, depozyt {pl(O['f'] * 100, 0)}% salda na pozycję (10 par)", fH, fill=HDR)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    hdr(ws, r + 1, 1, ['Strategia', 'Saldo końcowe', 'Zysk', 'Największe obsunięcie', 'Transakcje', 'Pominięte', 'Stop out'])
    for k, nm in enumerate(['v1', 'v2']):
        st = O['konto'][nm]
        for c, (v, fm) in enumerate([(nm, None), (st['koniec'], USD), (st['zysk_pct'] / 100, PCT), (-st['obs_pct'] / 100, PCT), (st['transakcje'], '0'),
                                     (st['pominiete'], '0'), (st['stop_out'], '0')]):
            put(ws, r + 2 + k, 1 + c, v, fmt=fm)
    r += 5
    put(ws, r, 1, 'DOUBLEPROOF — 10 nowych par fx1h_double_proof (dane, których nikt nie oglądał)', fH, fill=HDR2)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    dp = O.get('dp')
    if dp:
        hdr(ws, r + 1, 1, ['Strategia'] + SN, fill=HDR2)
        put(ws, r + 2, 1, 'v1 (dla porównania)', fB); srow(ws, r + 2, 2, dp['s1'])
        put(ws, r + 3, 1, 'v2 — test', fB, fill=SEL); srow(ws, r + 3, 2, dp['s2'], fill=SEL)
        x = put(ws, r + 4, 1, ('ZALICZONY' if dp['zaliczony'] else 'NIE ZALICZONY') +
                f" (warunek ustalony przed testem: ≥ {CRIT['min_n']} transakcji, średnia > 0, t ≥ {pl(CRIT['t'], 0)})", fB,
                fill=GOOD if dp['zaliczony'] else BAD)
        ws.merge_cells(start_row=r + 4, start_column=1, end_row=r + 4, end_column=8)
        r += 6
    else:
        x = ws.cell(r + 1, 1, dp_status or 'Brak danych par fx1h_double_proof — Apps Script (Fx.gs 1.52) pobiera ich historię; '
                                        'potem na Macu: python -m ia4.fxpowrot2'); x.font = fB; x.alignment = WR
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=14); ws.row_dimensions[r + 1].height = 30
        r += 3
    put(ws, r, 1, 'WNIOSKI' if wnioski else 'JAK CZYTAĆ', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    txt = wnioski or ['CZAS TRZYMANIA — wynik wg świec w pozycji i „trzymać czy zamknąć po k świecach”.', 'WARIANTY WYJŚCIA — ok. 40 wyjść z tym samym wejściem.',
                      'STRATNE — każda stratna transakcja v1 z cechami wejścia; CECHY — cechy wejścia w ćwiartkach.', 'DOUBLEPROOF — test na nowych parach.']
    for k, s_ in enumerate(txt):
        x = ws.cell(r + 1 + k, 1, s_); x.font = fN; x.alignment = WR
        ws.merge_cells(start_row=r + 1 + k, start_column=1, end_row=r + 1 + k, end_column=14); ws.row_dimensions[r + 1 + k].height = 44 if wnioski else 16

    # CZAS TRZYMANIA
    ws = sheet('CZAS TRZYMANIA')
    ws['A1'] = 'CZAS TRZYMANIA — czy długo trzymane transakcje warto zamykać wcześniej? (v1, 10 par)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Lewa tabela: wynik transakcji wg liczby świec w pozycji. Prawa: transakcje wciąż otwarte po k świecach — średni wynik, jeśli trzymać do końca '
                   '(TP albo limit), wobec zamknięcia po close k-tej świecy. „Trzymanie lepsze” > 50% i dodatnia różnica = cięcie czasem szkodzi.', 14, 42)
    hdr(ws, 4, 1, ['Świec w pozycji', 'Transakcje', 'Średnio %', 'Suma %', 'Trafne'], [14, 11, 11, 11, 10])
    B = O['czas']
    for k, rr_ in enumerate(B.itertuples()):
        for c, (v, fm) in enumerate([(str(rr_.swiece), None), (rr_.n, '0'), (rr_.srednia / 100, PCT3), (rr_.suma / 100, PCT), (rr_.trafne, PCT)]):
            x = put(ws, 5 + k, 1 + c, v, fmt=fm)
            if c == 2: x.fill = GOOD if rr_.srednia > 0 else BAD
    ch = BarChart(); ch.title = 'Średni wynik % wg świec w pozycji'; ch.height = 7; ch.width = 14; ch.legend = None
    ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=4 + len(B)), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(B)))
    ch.y_axis.number_format = '0.00%'; ch.x_axis.delete = False; ch.y_axis.delete = False; ch.x_axis.tickLblPos = 'low'
    ws.add_chart(ch, f'A{7 + len(B)}')
    Tq = O['trzymaj']
    hdr(ws, 4, 7, ['Po k świecach', 'Wciąż otwarte', 'Udział', 'Dojdzie do TP', 'Trzymaj do końca: średnio %', 'Zamknij teraz: średnio %',
                   'Różnica (suma %)', 'Trzymanie lepsze'], [11, 10, 9, 10, 13, 13, 11, 11])
    for k, q in enumerate(Tq.itertuples()):
        for c, (v, fm) in enumerate([(q.k, '0'), (q.otwarte, '0'), (q.udzial, PCT), (q.dojdzie_do_TP, PCT), (q.trzymaj / 100, PCT3), (q.zamknij / 100, PCT3),
                                     (q.roznica_suma / 100, PCT), (q.trzymanie_lepsze, PCT)]):
            x = put(ws, 5 + k, 7 + c, v, fmt=fm)
            if c == 6: x.fill = GOOD if q.roznica_suma > 0 else BAD

    # WARIANTY WYJŚCIA
    ws = sheet('WARIANTY WYJŚCIA')
    ws['A1'] = 'WARIANTY WYJŚCIA — to samo wejście, różne wyjścia (10 par, w próbie)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Zielony = lepszy t niż v1. Rok 1 / rok 2 — suma % w każdym roku osobno (czy zmiana pomaga w obu). Wariant wybrany na tych danych jest '
                   'zawyżony — dlatego v2 zawiera tylko zmiany uzasadnione analizą (koszt, trzymanie ≈ zamknięcie po 60 świecach) i jest sprawdzany na nowych parach.', 15, 42)
    hdr(ws, 4, 1, ['Wariant', 'Grupa'] + SN, [42, 18] + [10] * len(SN))
    W = O['warianty']; t1 = W.t.iloc[0]
    for k, v in enumerate(W.itertuples()):
        put(ws, 5 + k, 1, v.wariant, fB if k < 2 else fN, fill=SEL if k == 1 else None); put(ws, 5 + k, 2, v.grupa)
        srow(ws, 5 + k, 3, v._asdict())
        ws.cell(5 + k, 7).fill = GOOD if (v.t or 0) > t1 + 1e-9 else (BAD if (v.t or 0) < t1 - 0.5 else PatternFill())
    ws.freeze_panes = 'C5'

    # STRATNE
    ws = sheet('STRATNE')
    ws['A1'] = 'STRATNE — każda transakcja v1 z wynikiem ≤ 0 i cechy jej wejścia'; ws['A1'].font = fT
    note(ws, 2, 1, 'Dwa rodzaje strat: (1) limit 96 świec — kurs nie wrócił, (2) TP trafiony w 1. świecy, ale cel bliżej niż spread (cel w spreadach < 1) — '
                   'pewna strata, v2 takich transakcji nie otwiera. ATR — średni zakres świecy (100).', 16, 30)
    cols = [('para', 'Para', None, 8), ('t', 'Sygnał (UTC)', DT, 16), ('swiece', 'Świec', '0', 7), ('powod', 'Wyjście', None, 10), ('wynik', 'Wynik %', PCT3, 10),
            ('cel_spread', 'Cel w spreadach', '0.0', 9), ('cel_atr', 'Cel (ATR)', '0.00', 8), ('do_bazy', 'Wejście − baza (ATR)', '0.00', 9),
            ('zakres', 'Szczyt − baza (ATR)', '0.00', 9), ('impuls', 'Impuls (ATR)', '0.00', 8), ('nachylenie500', 'Nachylenie SMA500', '0.00', 9),
            ('zmiennosc', 'ATR100 / ATR500', '0.00', 9), ('ruch24', 'Ruch 24 św. (ATR)', '0.0', 9), ('godzina', 'Godz. UTC', '0', 7), ('dzien', 'Dzień', None, 6)]
    hdr(ws, 4, 1, [c[1] for c in cols], [c[3] for c in cols])
    dn = ['pn', 'wt', 'śr', 'czw', 'pt', 'sb', 'nd']
    for k, v in enumerate(O['stratne'].itertuples()):
        for c, (key, _, fm, _) in enumerate(cols):
            val = getattr(v, key)
            if key == 't': val = val.to_pydatetime()
            if key == 'powod': val = WHY.get(int(val), '?')
            if key == 'wynik': val = val / 100
            if key == 'dzien': val = dn[int(val)]
            put(ws, 5 + k, 1 + c, val, fmt=fm)
    ws.freeze_panes = 'A5'

    # CECHY
    ws = sheet('CECHY')
    ws['A1'] = 'CECHY WEJŚCIA — czy da się z góry rozpoznać transakcję, która utknie (> 48 świec)? (v1, 10 par)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Każda cecha w ćwiartkach (po ok. 90 transakcji). „Długie” = udział transakcji trzymanych > 48 świec (tam są prawie wszystkie straty). '
                   'Długie skupiają się tam, gdzie cel jest daleko — ale tam też jest największy średni zysk, więc odcięcie ich obniża wynik.', 8, 42)
    hdr(ws, 4, 1, ['Cecha', 'Przedział', 'Transakcje', 'Średnio %', 'Suma %', 'Długie (> 48 św.)'], [32, 16, 10, 10, 10, 12])
    for k, v in enumerate(O['cechy'].itertuples()):
        for c, (val, fm) in enumerate([(v.cecha, None), (v.przedzial, None), (v.n, '0'), (v.srednia / 100, PCT3), (v.suma / 100, PCT), (v.dlugie, PCT)]):
            x = put(ws, 5 + k, 1 + c, val, fmt=fm)
            if c == 3: x.fill = GOOD if v.srednia > 0 else BAD

    # PARY
    ws = sheet('PARY')
    ws['A1'] = 'PARY — v1 i v2 na każdej z 10 par badania'; ws['A1'].font = fT
    hdr(ws, 3, 1, ['Para', 'v1: n', 'v1: średnio %', 'v1: t', 'v1: najgorsza %', 'v2: n', 'v2: średnio %', 'v2: t', 'v2: najgorsza %', 'v2: mediana świec'],
        [9, 8, 11, 8, 11, 8, 11, 8, 11, 10])
    for k, v in enumerate(O['pary'].itertuples()):
        for c, (val, fm) in enumerate([(v.para, None), (v.n_v1, '0'), (v.srednia_v1 / 100, PCT3), (v.t_v1, '0.0'), (v.najgorsza_v1 / 100, PCT3),
                                       (v.n_v2, '0'), (v.srednia_v2 / 100, PCT3), (v.t_v2, '0.0'), (v.najgorsza_v2 / 100, PCT3), (v.med_sw_v2, '0')]):
            put(ws, 4 + k, 1 + c, val, fmt=fm)

    # TRANSAKCJE V2
    def trades_sheet(name, D, title, spread):
        ws = sheet(name)
        ws['A1'] = title; ws['A1'].font = fT
        note(ws, 2, 1, f"Short po open świecy po sygnale; wynik % netto po spreadzie (pipsy: {', '.join(f'{p} {pl(v, 2)}' for p, v in spread.items())}).", 10, 30)
        hdr(ws, 4, 1, ['Para', 'Sygnał (UTC)', 'Baza (cena sprzed wzrostu)', 'Szczyt', 'ATR', 'Świec', 'Wyjście', 'Wynik %', 'Okres'], [8, 16, 13, 11, 10, 7, 10, 10, 7])
        for k, v in enumerate(D.itertuples()):
            for c, (val, fm) in enumerate([(v.para, None), (v.t.to_pydatetime(), DT), (v.B, '0.00000'), (v.P, '0.00000'), (v.A, '0.00000'), (v.swiece, '0'),
                                           (WHY.get(int(v.powod), '?'), None), (v.wynik / 100, PCT3), (v.okres, '0')]):
                x = put(ws, 5 + k, 1 + c, val, fmt=fm)
                if c == 7: x.fill = GOOD if v.wynik > 0 else BAD
        n = len(D)
        put(ws, 6 + n, 7, 'Razem', fB, fill=GRY); put(ws, 6 + n, 8, f'=SUM(H5:H{4 + n})', fB, PCT)
        put(ws, 7 + n, 7, 'Średnio', fB, fill=GRY); put(ws, 7 + n, 8, f'=AVERAGE(H5:H{4 + n})', fB, PCT3)
        put(ws, 8 + n, 7, 'Trafne', fB, fill=GRY); put(ws, 8 + n, 8, f'=COUNTIF(H5:H{4 + n},">0")/{max(n, 1)}', fB, PCT)
        ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:I{4 + n}'

    trades_sheet('TRANSAKCJE V2', O['D2'], 'TRANSAKCJE V2 — 10 par badania', O['spread'])

    # KAPITAŁ
    ws = sheet('KAPITAŁ')
    ws['A1'] = f"KAPITAŁ — {pl(O['start'], 0)} $, 1:{O['lev']:g}, depozyt {pl(O['f'] * 100, 0)}% salda na pozycję"; ws['A1'].font = fT
    series = [('10 par: v1', O['krzywe']['v1']), ('10 par: v2', O['krzywe']['v2'])]
    if dp:
        series += [('doubleProof: v1', dp['krzywe']['v1']), ('doubleProof: v2', dp['krzywe']['v2'])]
    col = 1
    for g, (a, b) in enumerate([(0, 2), (2, 4)] if dp else [(0, 2)]):
        ss = series[a:b]
        idx = ss[0][1].index.union(ss[1][1].index)
        hdr(ws, 3, col, ['Dzień'] + [s[0] for s in ss], [12, 13, 13])
        for k, d in enumerate(idx):
            put(ws, 4 + k, col, d.to_pydatetime(), fmt=DD)
            for q, (_, cv) in enumerate(ss):
                v = cv.get(d, np.nan)
                put(ws, 4 + k, col + 1 + q, None if pd.isna(v) else round(float(v), 2), fmt=USD0)
        ch = LineChart(); ch.title = ('10 par badania' if g == 0 else 'doubleProof') + f" — {pl(O['start'], 0)} $, 1:{O['lev']:g}"
        ch.height = 9; ch.width = 20; ch.y_axis.crossAx = 500; ch.x_axis = DateAxis(crossAx=100); ch.x_axis.number_format = 'yyyy-mm'
        ch.x_axis.majorTimeUnit = 'months'; ch.x_axis.delete = False; ch.y_axis.delete = False
        ch.add_data(Reference(ws, min_col=col + 1, max_col=col + 2, min_row=3, max_row=3 + len(idx)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=col, min_row=4, max_row=3 + len(idx))); ch.legend.position = 'b'
        for s_, colr in zip(ch.series, ['9DC3E6', 'C00000']):
            s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 18000
        ws.add_chart(ch, f'{CL(9 if g == 0 else 9)}{3 if g == 0 else 23}')
        col += 4

    # DOUBLEPROOF
    ws = sheet('DOUBLEPROOF')
    ws['A1'] = 'DOUBLEPROOF — v1 i v2 na 10 parach fx1h_double_proof (nowe dane, test jednorazowy)'; ws['A1'].font = fT
    note(ws, 2, 1, f"Reguły zamrożone przed pobraniem par (skrót {O['hash']}). Warunek zaliczenia v2 ustalony przed testem: ≥ {CRIT['min_n']} transakcji, średnia > 0, "
                   f"t ≥ {pl(CRIT['t'], 0)}. Spread tych par — szacunek dla XTB: {', '.join(f'{p} {pl(v, 2)}' for p, v in DP_SPREAD.items())} pipsa. "
                   'Każde otwarcie zapisane w data/fxpowrot2/doubleproof.jsonl (tabela niżej).', 12, 54)
    if dp:
        put(ws, 4, 1, f"Dane: {dp['od']:%Y-%m-%d} – {dp['do']:%Y-%m-%d %H:%M} UTC, {dp['swiec']} świec na parę", fB)
        hdr(ws, 6, 1, ['Strategia'] + SN, [26] + [10] * len(SN), fill=HDR2)
        put(ws, 7, 1, 'v1', fB); srow(ws, 7, 2, dp['s1']); put(ws, 8, 1, 'v2 — test', fB, fill=SEL); srow(ws, 8, 2, dp['s2'], fill=SEL)
        hdr(ws, 10, 1, ['Para', 'v1: n', 'v1: średnio %', 'v1: t', 'v1: najgorsza %', 'v2: n', 'v2: średnio %', 'v2: t', 'v2: najgorsza %', 'v2: mediana świec'], fill=HDR2)
        for k, v in enumerate(dp['pary'].itertuples()):
            for c, (val, fm) in enumerate([(v.para, None), (v.n_v1, '0'), (v.srednia_v1 / 100, PCT3), (v.t_v1, '0.0'), (v.najgorsza_v1 / 100, PCT3),
                                           (v.n_v2, '0'), (v.srednia_v2 / 100, PCT3), (v.t_v2, '0.0'), (v.najgorsza_v2 / 100, PCT3), (v.med_sw_v2, '0')]):
                put(ws, 11 + k, 1 + c, val, fmt=fm)
        r = 23
        hdr(ws, r, 1, ['Rachunek', 'Saldo końcowe', 'Zysk', 'Obsunięcie', 'Transakcje', 'Pominięte', 'Stop out'], fill=HDR2)
        for k, nm in enumerate(['v1', 'v2']):
            st = dp['konto'][nm]
            for c, (v, fm) in enumerate([(nm, None), (st['koniec'], USD), (st['zysk_pct'] / 100, PCT), (-st['obs_pct'] / 100, PCT), (st['transakcje'], '0'),
                                         (st['pominiete'], '0'), (st['stop_out'], '0')]):
                put(ws, r + 1 + k, 1 + c, v, fmt=fm)
        r += 5
    else:
        x = ws.cell(4, 1, dp_status or 'Brak danych — czeka na historię par z Apps Script.'); x.font = fB
        r = 6
    hdr(ws, r, 1, ['Otwarcie', 'Czas (UTC)', 'Reguły (skrót)', 'Dane od', 'Dane do', 'Świec', 'Wersja'], fill=HDR2)
    for k, e in enumerate(O.get('dp_log') or []):
        for c, key in enumerate(['otwarcie', 'czas', 'reguly', 'od', 'do', 'swiec', 'wersja']):
            put(ws, r + 1 + k, 1 + c, e.get(key))
    if dp:
        trades_sheet('DOUBLEPROOF TRANSAKCJE', dp['D2'], 'DOUBLEPROOF TRANSAKCJE — v2 na parach fx1h_double_proof', dp['spread'])

    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- uruchomienie

def run(no_sync=False, out=None, f=0.20, lev=30.0, start=1000.0, frames=None, dp_frames=None, log=print, keep=None, wnioski=None,
        publish=False) -> Path:
    t0 = time.time(); status = ''
    if frames is None:
        from .fx import sync as fxs
        if not no_sync:
            log('Kopia fx1h — odświeżenie 10 par badania i 10 par fx1h_double_proof (python -m ia4.fx.sync)')
            fxs.sync(quiet=True, syms=list(PAIRS))
            try:
                fxs.sync(quiet=True, syms=list(DP_PAIRS))
            except Exception as e:
                log(f'  ! pary doubleProof: {type(e).__name__}: {e}')
        frames = {p: fxs.load(p) for p in PAIRS}
        if dp_frames is None:
            dp_frames = {p: fxs.load(p) for p in DP_PAIRS if fxs.parquet_path(p).exists()}
    R = Rynek(frames, PAIRS, fs.spreads())
    log(f'10 par badania: {R.n} świec, {pd.Timestamp(int(R.t[0]), unit="s"):%Y-%m-%d} – {pd.Timestamp(int(R.t[-1]), unit="s"):%Y-%m-%d}')
    Rdp = None
    dp_frames = dp_frames or {}
    miss = [p for p in DP_PAIRS if p not in dp_frames or len(dp_frames[p]) < DP_MIN_CANDLES]
    if miss:
        status = (f"doubleProof jeszcze nie liczony: brak historii {len(miss)} z 10 par ({', '.join(miss)}) w kopii fx1h — Apps Script (Fx.gs 1.52) "
                  'pobiera po jednej parze; uruchom ponownie, gdy wszystkie mają ≥ 2000 świec.')
        log(status)
    else:
        allf = dict(frames); allf.update(dp_frames)
        Rdp = Rynek(allf, DP_PAIRS, DP_SPREAD)
        log(f'doubleProof: {Rdp.n} świec, {pd.Timestamp(int(Rdp.t[0]), unit="s"):%Y-%m-%d} – {pd.Timestamp(int(Rdp.t[-1]), unit="s"):%Y-%m-%d}')
    O = compute(R, Rdp, f=f, lev=lev, start=start, log=log)
    if publish:
        try:
            publish_history(R, log=log)
        except Exception as e:                                   # strona nie zatrzymuje badania
            log(f'  ! historia dla strony: {type(e).__name__}: {e}')
    if keep is not None:
        keep.update(O)
    path = write_xlsx(O, out or (config.RESEARCH_DIR.parent / 'FX-POWROT-V2.xlsx'), wnioski=wnioski, dp_status=status)
    log(f'Zapisano: {path}  ({time.time() - t0:.0f} s)')
    return path


def selftest() -> bool:
    ok = True

    def check(name, cond):
        nonlocal ok
        print(f"  {'✓' if cond else '✗'} {name}")
        ok &= bool(cond)

    n = 30
    o = np.full(n, 1.0); h = np.full(n, 1.001); l = np.full(n, 0.999); c = np.full(n, 1.0); fri = np.zeros(n, np.int64)
    J = np.array([0]); B = np.array([0.98]); P = np.array([1.02]); A = np.array([0.002])
    h2, l2 = h.copy(), l.copy(); l2[3] = 0.995                                 # TP = 1,02 − ½ · 0,04 = 1,00 → za blisko wejścia (1,00)
    r, x, w = run_exit(o, h2, l2, c, fri, J, B, P, A, 0.0001, 0.5, 20, 0, -99.0, False, 0.0, 0.0, 0.0, 0.0, 0, 1.0)
    check('cel ≤ 0,1 ATR od wejścia → bez transakcji', np.isnan(r[0]))
    P2 = np.array([1.03])                                                      # TP = 1,03 − ½ · 0,05 = 1,005
    o3 = np.full(n, 1.01); c = np.full(n, 1.01); h3 = np.full(n, 1.011); l3 = np.full(n, 1.009); l3[4] = 1.004   # wejście 1,01, TP w 4. świecy
    h = h3.copy()
    r, x, w = run_exit(o3, h3, l3, c, fri, J, B, P2, A, 0.0001, 0.5, 20, 0, -99.0, False, 0.0, 0.0, 0.0, 0.0, 0, 1.0)
    check('TP w 4. świecy: wynik (1,01 − 1,005)/1,01 − spread, 4 świece', abs(r[0] - ((1.01 - 1.005) / 1.01 * 100 - 0.0001 / 1.01 * 100)) < 1e-9 and x[0] == 4 and w[0] == 0)
    r, x, w = run_exit(o3, h3, l3, c, fri, J, B, P2, A, 0.0060, 0.5, 20, 0, -99.0, False, 1.0, 0.0, 0.0, 0.0, 0, 1.0)
    check('v2: cel 0,005 < 1 × spread 0,006 → bez transakcji', np.isnan(r[0]))
    l4 = np.full(n, 1.009); c4 = c.copy(); c4[5] = 1.008
    r, x, w = run_exit(o3, h3, l4, c4, fri, J, B, P2, A, 0.0001, 0.5, 5, 0, -99.0, False, 0.0, 0.0, 0.0, 0.0, 0, 1.0)
    check('limit 5 świec → wyjście po close 5. świecy (1,008)', x[0] == 5 and w[0] == 1 and abs(r[0] - ((1.01 - 1.008) / 1.01 * 100 - 0.0001 / 1.01 * 100)) < 1e-9)
    f2 = fri.copy(); f2[3] = 1
    r, x, w = run_exit(o3, h3, l4, c, f2, J, B, P2, A, 0.0001, 0.5, 20, 0, -99.0, True, 0.0, 0.0, 0.0, 0.0, 0, 1.0)
    check('przed weekendem → zamknięcie 3. świecy', x[0] == 3 and w[0] == 3)
    h5 = h.copy(); h5[2] = 1.02
    r, x, w = run_exit(o3, h5, l4, c, fri, J, B, P2, A, 0.0001, 0.5, 20, 0, -99.0, False, 0.0, 3.0, 0.0, 0.0, 0, 1.0)
    check('SL 3 ATR (1,016) w 2. świecy', w[0] == 4 and x[0] == 2 and abs(r[0] - ((1.01 - 1.016) / 1.01 * 100 - 0.0001 / 1.01 * 100)) < 1e-9)
    # jedna pozycja na parę i rachunek: syntetyczny rynek
    t = 1_700_000_000 + 3600 * np.arange(1500)
    rng = np.random.default_rng(3)
    fr = {}
    for p in set(DP_PAIRS) | set(MAJORS):
        base = 150.0 if p.endswith('JPY') else 1.2
        z = rng.normal(0, 0.002, t.size); z[rng.integers(0, t.size, 60)] += 0.012          # co jakiś czas gwałtowna świeca w górę
        cc = base * np.exp(np.cumsum(z))
        fr[p] = pd.DataFrame(dict(time=t, o=np.r_[cc[0], cc[:-1]], h=np.maximum(np.r_[cc[0], cc[:-1]], cc) * 1.001,
                                  l=np.minimum(np.r_[cc[0], cc[:-1]], cc) * 0.999, c=cc, filled=False))
    R = Rynek(fr, DP_PAIRS, DP_SPREAD)
    check('przeliczenie na USD: CHFJPY → 1 / USDCHF, EURAUD → EURUSD',
          abs(R.usd('CHFJPY', 700) - 1 / R.F['USDCHF']['o'][700]) < 1e-12 and abs(R.usd('EURAUD', 700) - R.F['EURUSD']['o'][700]) < 1e-12)
    E = events(R)
    D = simulate(R, E, **V1)
    over = any((g.j.values[1:] <= (g.j.values[:-1] + g.swiece.values[:-1])).any() for _, g in D.groupby('para'))
    check(f'jedna pozycja na parę ({len(D)} transakcji)', not over)
    if len(D):
        st, cv = konto(D, R)
        check('rachunek: saldo końcowe liczbą', np.isfinite(st['koniec']))
    print('\nWYNIK: OK' if ok else '\nWYNIK: BŁĄD')
    return ok


def main():
    ap = argparse.ArgumentParser(description='FX-POWRÓT v2 — analiza strat i doubleProof')
    ap.add_argument('--bez-sync', action='store_true', help='bez odświeżenia kopii fx1h')
    ap.add_argument('--depozyt', type=float, default=0.20, help='depozyt na pozycję jako część salda (domyślnie 0.20)')
    ap.add_argument('--strona', action='store_true', help='tylko historia dla strony (branch fx, fx/pv2/) — bez Excela i bez doubleProof')
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args()
    print(f'IA 4 {__version__} — FX-POWRÓT v2\n')
    if a.selftest:
        sys.exit(0 if selftest() else 1)
    if a.strona:
        from .fx import sync as fxs
        if not a.bez_sync:
            log_ = print
            log_('Kopia fx1h — odświeżenie 10 par (python -m ia4.fx.sync)')
            fxs.sync(quiet=True, syms=list(PAIRS))
        R = Rynek({p: fxs.load(p) for p in PAIRS}, PAIRS, fs.spreads())
        sys.exit(0 if publish_history(R) else 1)
    run(no_sync=a.bez_sync, f=a.depozyt, publish=True)


if __name__ == '__main__':
    main()
