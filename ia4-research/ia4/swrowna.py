"""
IA 4 — strategia RÓWNA: kupowanie spadków tylko przy rosnącym trendzie rynku (instrukcja, sekcja 7a).
Wersja projektu: 1.53 (2026-10-20) — musi zgadzać się z IA4_INSTRUKCJA.md

Cel człowieka (17.10.2026): krzywa kapitału konsekwentnie w górę, w bessie bez strat (najwyżej bez zysku), dużo
mniejsze obsunięcia, bez przeuczenia. Wniosek z badań 1.47: zysk kupowania paniki to głównie odbicie całego rynku,
więc zawsze przechodzi przez krach — równą krzywą daje kupowanie spadków TYLKO przy rosnącym trendzie.

Dwie części, obie z filtrem trendu (w bessie stoją z boku):
  1. QQQ — reguła człowieka z `ia4.swrsi` bez zmian (RSI(3) < 22, EMA100 rośnie, close ≥ 1% pod VWAP10, wyjście
     RSI(3) > 55). Wartość pozycji = min(3, 25% / zmienność QQQ) × stan portfela (zmienność — roczna, z 20 sesji
     do dnia sygnału): spokojny rynek — większa pozycja, nerwowy — mniejsza.
  2. Spółki (73) — co najmniej 5 świec spadkowych, spadek serii ≥ 10% (jak `ia4.swportfel`, wejście w pierwszym dniu
     spełnienia), tylko gdy EMA100 SPY rośnie (wyżej niż 5 sesji wcześniej); long po open, 5 sesji, bez SL; wartość
     pozycji 20% stanu portfela, najwyżej 5 pozycji naraz (najpierw najgłębszy spadek).
Skala (`--skala 1.5`) mnoży wszystkie pozycje. Portfel 10 000 $, rachunek CFD z dźwignią 1:5 (depozyt = wartość / 5,
suma depozytów ≤ stan portfela), spread 0,3% raz na transakcję, wycena dzienna; bez kosztu nocnego (swap).

Arkusze: PODSUMOWANIE (wynik całości i każdej części, obsunięcia, lata bessy), KAPITAŁ (wykresy), OBSUNIĘCIA,
ROK PO ROKU (z zyskiem każdej części), TRANSAKCJE, SYGNAŁY SPÓŁEK (także odrzucone przez trend i limit),
ODPORNOŚĆ (sąsiednie ustawienia, części osobno, spółki 53 / doubleProof, wynik transakcji w okresach).

    python -m ia4.swrowna               # poziom podstawowy (najpierw python -m ia4.daily)
    python -m ia4.swrowna --skala 1.5   # wszystkie pozycje 1,5 × większe
    python -m ia4.swrowna --selftest    # „WYNIK: OK”
Wynik: ~/Desktop/IA4/SWING-ROWNA.xlsx (z --skala: SWING-ROWNA-x1.5.xlsx), poza gitem.
"""

from __future__ import annotations

import argparse
import math
import sys
import time
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path

import numpy as np

from . import __version__, config
from . import swing as sw
from . import swportfel as sp
from . import swrsi

INDEX = "QQQ"
VOL_TARGET, VOL_CAP, VOL_N = 25.0, 3.0, 20
STOCK_SIZE, STOCK_MAX = 0.20, 5
STOCK_RULE = sp.Rule(min_n=5, min_drop=10.0, hold=5)
TREND_EMA, TREND_BACK = 100, 5
LEV = 5.0
BEARS = [("2000-03-24", "2002-10-09"), ("2007-10-09", "2009-03-09"), ("2020-02-19", "2020-03-23"),
         ("2022-01-03", "2022-10-12")]        # bessy S&P 500 (szczyt → dno) — do opisu wyników


@dataclass(frozen=True)
class Cfg:
    skala: float = 1.0
    vol_target: float = VOL_TARGET
    vol_cap: float = VOL_CAP
    stock_size: float = STOCK_SIZE
    stock_max: int = STOCK_MAX
    trend: bool = True               # filtr trendu dla spółek
    vol_sizing: bool = True          # wielkość pozycji QQQ wg zmienności
    index: bool = True
    stocks: bool = True
    groups: tuple = ("main", "proof", "doubleProof")

    def short(self) -> str:
        parts = []
        if self.index:
            parts.append(f"{INDEX} RSI(3)" + (f" (cel zmienności {self.vol_target:g}%, max {self.vol_cap:g}×)" if self.vol_sizing else " (cała kwota)"))
        if self.stocks:
            parts.append(f"spółki: ≥ 5 spadków, spadek ≥ 10%{', EMA100 SPY rośnie' if self.trend else ''}, "
                         f"{self.stock_size * 100:g}%/poz., max {self.stock_max}")
        return " + ".join(parts) + (f" · skala {self.skala:g}" if self.skala != 1 else "")


def realized_vol(c: np.ndarray, n: int = VOL_N) -> np.ndarray:
    """Zmienność roczna % z n ostatnich zmian zamknięć (do sesji włącznie)."""
    r = np.diff(np.log(c), prepend=np.nan)
    out = np.full(len(c), np.nan)
    for q in range(n, len(c)):
        out[q] = float(np.std(r[q - n + 1:q + 1], ddof=1)) * math.sqrt(252) * 100.0
    return out


def index_trades(D: sw.Data, cfg: Cfg) -> list[dict]:
    si = next((i for i, s in enumerate(D.stocks) if s.sym == INDEX), None)
    if si is None:
        return []
    s = D.stocks[si]
    tr, _ = swrsi.trades(s, si)
    vol = realized_vol(s.c)
    out = []
    for t in tr:
        v = vol[t["qe"] - 1]
        f = min(cfg.vol_cap, cfg.vol_target / v) if cfg.vol_sizing and np.isfinite(v) and v > 0 else 1.0
        out.append(dict(t, label=INDEX, size=f * cfg.skala, prio=-1e9, why=4))
    return out


def trend_up(D: sw.Data) -> np.ndarray:
    """Na kalendarzu: EMA100 SPY wyżej niż 5 sesji SPY wcześniej (stan po close)."""
    T = len(D.cal)
    out = np.zeros(T, bool)
    s = D.by_sym.get("SPY")
    if s is None:
        return out
    e = sw.ema(s.c, TREND_EMA)
    with np.errstate(invalid="ignore"):
        up = e > sw.prior(e, TREND_BACK)
    out[s.t] = np.nan_to_num(up, nan=False).astype(bool)
    last = False                             # dni bez SPY — ostatni znany stan
    have = np.zeros(T, bool)
    have[s.t] = True
    for t in range(T):
        if have[t]:
            last = out[t]
        else:
            out[t] = last
    return out


def stock_events(D: sw.Data, cfg: Cfg, since: str | None) -> list[dict]:
    rows = D.rows(*cfg.groups)
    up = trend_up(D)
    out = []
    for e in sp.events(D, rows, STOCK_RULE, sw.COST_PCT, sp.market(D)):
        if since and e["last"] < since or not e["ok"]:
            continue
        s = D.stocks[e["si"]]
        q = int(np.flatnonzero(s.dates == e["last"])[0])
        e["trend"] = bool(up[s.t[q]])
        if e["trade"] is not None and (e["trend"] or not cfg.trend):
            e["trade"] = dict(e["trade"], label="spółki", size=cfg.stock_size * cfg.skala, prio=e["drop"])
        else:
            if e["trade"] is not None:
                e["status"] = "odrzucona: EMA100 SPY nie rośnie (rynek bez trendu wzrostowego)"
            e["trade"] = None
        out.append(e)
    return out


def build(D: sw.Data, cfg: Cfg, since: str | None, cost: float, keep_rows: bool = True):
    idx = [t for t in index_trades(D, cfg) if not since or D.stocks[t["si"]].dates[t["qe"]] >= since] if cfg.index else []
    evs = stock_events(D, cfg, since) if cfg.stocks else []
    trades = idx + [e["trade"] for e in evs if e["trade"] is not None]
    pf = sw.portfolio([dict(t) for t in trades], D, lev=LEV, cost=cost, keep_rows=keep_rows,
                      max_by_label={"spółki": cfg.stock_max})
    return pf, idx, evs


def bear_returns(pf: dict) -> list[tuple]:
    cv = pf["curve"]
    out = []
    for a, b in BEARS:
        pts = [(d, v) for d, v in cv if a <= d <= b]
        before = [v for d, v in cv if d < a]
        if pts:
            v0 = before[-1] if before else pts[0][1]
            out.append((a, b, (pts[-1][1] / v0 - 1.0) * 100.0))
    return out


def smooth(pf: dict) -> dict:
    """Miary równości krzywej: R² log(stan) wobec czasu, odsetek 12-miesięcznych okien na plusie."""
    cv = pf["curve"]
    if len(cv) < 260:
        return {"r2": None, "roll12": None}
    v = np.array([x for _, x in cv])
    lg = np.log(np.maximum(v, 1e-9))
    tt = np.arange(len(lg))
    r2 = float(np.corrcoef(tt, lg)[0, 1] ** 2)
    roll = float(np.mean(v[252:] > v[:-252]) * 100.0)
    return {"r2": r2, "roll12": roll}


def run(no_sync: bool = False, out: Path | None = None, since: str | None = sw.DAILY_SINCE, cost: float = sw.COST_PCT,
        cfg: Cfg = Cfg(), data: sw.Data | None = None, robust: bool = True) -> Path:
    t0 = time.time()
    print(f"IA 4 — strategia RÓWNA: {cfg.short()}; spread {cost:g}% raz na transakcję\n")
    D = data or sw.load(no_sync, None, groups=("main", "proof", "doubleProof", "context"))
    print(f"  {D.note}\n  portfel od {since or D.cal[0]} do {D.cal[-1]}")
    pf, idx, evs = build(D, cfg, since, cost)
    parts = {}
    if cfg.index and cfg.stocks:
        parts[f"tylko {INDEX}"] = build(D, replace(cfg, stocks=False), since, cost, keep_rows=False)[0]
        parts["tylko spółki"] = build(D, replace(cfg, index=False), since, cost, keep_rows=False)[0]
    dds = sp.drawdowns(pf["curve"])
    sm = smooth(pf)
    cg = sw.cagr(pf, D)
    yrs = dict(pf["years"])
    print(f"  transakcji: {pf['trades']} ({sum(1 for r in pf['rows'] if r['label'] == INDEX)} {INDEX}, "
          f"{sum(1 for r in pf['rows'] if r['label'] == 'spółki')} spółki); pominięte przez limit {pf['skipped_limit']}")
    print(f"  portfel: {pf['ret']:+.1f}% (stan {pf['end']:,.0f} $), CAGR {cg or 0:.1f}%, największe obsunięcie {pf['mdd']:.1f}%, "
          f"R² {sm['r2'] or 0:.2f}, lata stratne {sum(v < 0 for v in yrs.values())}")
    print("  bessy: " + ", ".join(f"{a[:7]}–{b[:7]} {r:+.1f}%" for a, b, r in bear_returns(pf)))
    for x in sorted(dds, key=lambda x: -x["depth"])[:3]:
        print(f"   obsunięcie {x['depth']:.1f}%: szczyt {x['start']} → dno {x['trough']} → "
              + (f"odrobione {x['rec']}" if x["rec"] else "jeszcze nie odrobione"))
    rob = robustness(D, cfg, since, cost) if robust else []
    path = out or (config.RESEARCH_DIR.parent / ("SWING-ROWNA.xlsx" if cfg.skala == 1 else f"SWING-ROWNA-x{cfg.skala:g}.xlsx"))
    write_xlsx(path, D, cfg, pf, parts, idx, evs, dds, rob, cost)
    print(f"\nGotowe w {time.time() - t0:.0f} s → {path}")
    return path


def robustness(D, cfg: Cfg, since, cost) -> list[tuple]:
    """Sąsiednie ustawienia i części osobno — czy wynik nie zależy od jednej dokładnej liczby."""
    out = []
    variants = [("ustawienie z pliku", cfg)]
    for vt in (20.0, 30.0):
        variants.append((f"cel zmienności QQQ {vt:g}%", replace(cfg, vol_target=vt)))
    for ss in (0.10, 0.30):
        variants.append((f"spółki {ss * 100:g}% na pozycję", replace(cfg, stock_size=ss)))
    for mx in (3, 10):
        variants.append((f"spółki max {mx} pozycji", replace(cfg, stock_max=mx)))
    variants += [(f"{INDEX} bez skalowania zmiennością (cała kwota)", replace(cfg, vol_sizing=False)),
                 ("spółki BEZ filtru trendu", replace(cfg, trend=False)),
                 ("tylko spółki handlowane (53)", replace(cfg, groups=("main", "proof"))),
                 ("tylko spółki doubleProof (20) — sprawdzian", replace(cfg, groups=("doubleProof",), index=False)),
                 ("tylko spółki (73)", replace(cfg, index=False)),
                 (f"tylko {INDEX}", replace(cfg, stocks=False))]
    for name, c in variants:
        pf, idx, evs = build(D, c, since, cost, keep_rows=True)
        sm = smooth(pf)
        yrs = dict(pf["years"])
        out.append((name, pf["trades"], pf["ret"], sw.cagr(pf, D), pf["mdd"], sm["r2"], sum(v < 0 for v in yrs.values()),
                    min(yrs.values()) if yrs else None, [r for _, _, r in bear_returns(pf)]))
    # wynik transakcji w okresach (bez portfela): QQQ od 1999, spółki od 1980 (trend SPY od 1993)
    per = []
    for lab, trs in ((INDEX, index_trades(D, cfg)), ("spółki (z filtrem trendu)", [e["trade"] for e in stock_events(D, cfg, None) if e["trade"]])):
        x = [(D.stocks[t["si"]].dates[t["qe"]], (t["pout"] / t["pin"] - 1.0) * 100.0 - cost) for t in trs]
        for a, b in (("1980", "2000"), ("2000", "2013"), ("2013", "2100")):
            v = np.array([n for d, n in x if a <= d < b])
            per.append((lab, f"{a}–{int(b) - 1 if b != '2100' else 'dziś'}", len(v), float(v.mean()) if len(v) else None,
                        float((v > 0).mean() * 100) if len(v) else None))
    return [("ROB", out), ("PER", per)]


def write_xlsx(path: Path, D, cfg: Cfg, pf, parts, idx, evs, dds, rob, cost) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.utils import get_column_letter
    S = sw.styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "PODSUMOWANIE"
    sm = smooth(pf)
    cg = sw.cagr(pf, D)
    yrs = dict(pf["years"])
    lines = [
        ("IA 4 — strategia RÓWNA: kupowanie spadków tylko przy rosnącym trendzie rynku", ""),
        ("Wersja", __version__), ("Wygenerowano", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Dane", f"portfel od {pf['first'] or '—'} do {D.cal[-1]}; {D.note}"),
        ("Ustawienie", cfg.short()),
        (f"Część 1 — {INDEX}", f"reguła RSI(3) z ia4.swrsi bez zmian (RSI(3) < 22, EMA100 rośnie, close ≥ 1% pod VWAP10, wyjście "
                              f"RSI(3) > 55); wartość pozycji = min({cfg.vol_cap:g}, {cfg.vol_target:g}% / zmienność {INDEX} z "
                              f"{VOL_N} sesji) × stan portfela" if cfg.vol_sizing else "cała kwota"),
        ("Część 2 — spółki", f"≥ 5 świec spadkowych, spadek serii ≥ 10%, EMA{TREND_EMA} SPY wyżej niż {TREND_BACK} sesji wcześniej; "
                             f"long po open, 5 sesji, bez SL; {cfg.stock_size * 100:g}% stanu portfela na pozycję, najwyżej "
                             f"{cfg.stock_max} naraz (najpierw najgłębszy spadek)"),
        ("Rachunek", f"10 000 $, CFD 1:{LEV:g} (depozyt = wartość / {LEV:g}, suma depozytów ≤ stan portfela), skala {cfg.skala:g}; "
                     f"spread {cost:g}% raz na transakcję; bez kosztu nocnego (swap) i dywidend"),
        ("Uwaga", "73 spółki to te, które przetrwały; QQQ ma dane od 1999 r.; parametry 25% i 20% wybrane z kilku "
                  "sprawdzonych wartości — sąsiednie w arkuszu ODPORNOŚĆ"),
        ("", ""),
        ("Stan końcowy $", round(pf["end"], 2)), ("Wynik %", round(pf["ret"], 1)), ("CAGR %", round(cg, 2) if cg else None),
        ("Największe obsunięcie % (wycena dzienna)", round(pf["mdd"], 1)),
        ("R² krzywej (1 = linia prosta w skali log.)", round(sm["r2"], 3) if sm["r2"] else None),
        ("12-miesięczne okresy na plusie %", round(sm["roll12"], 1) if sm["roll12"] else None),
        ("Lata stratne", f"{sum(v < 0 for v in yrs.values())} z {len(yrs)}"),
        ("Najgorszy rok %", round(min(yrs.values()), 1) if yrs else None),
        ("Transakcji", f"{pf['trades']} ({sum(1 for r in pf['rows'] if r['label'] == INDEX)} {INDEX}, "
                       f"{sum(1 for r in pf['rows'] if r['label'] == 'spółki')} spółki)"),
        ("Zyskownych %", round(pf["wins"] / pf["trades"] * 100, 1) if pf["trades"] else None),
        ("Najwięcej pozycji naraz", pf["max_open"]), ("Pominięte przez limit spółek", pf["skipped_limit"]),
        ("Upadłość", "TAK" if pf["ruined"] else "nie"),
    ]
    for k, (a, b) in enumerate(lines, start=1):
        ws.cell(k, 1, a)
        ws.cell(k, 2, b)
    ws["A1"].font = S.bold
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    for col in "CDEF":
        ws.column_dimensions[col].width = 14
    r = len(lines) + 2
    ws.cell(r, 1, "Bessy S&P 500 (szczyt → dno)").font = S.bold
    sw.header(ws, r + 1, ["Bessa", "Wynik portfela w tym czasie %"] + list(parts), [44, 18, 14, 14], S, 30)
    r += 2
    pb = {k: bear_returns(v) for k, v in parts.items()}
    for i, (a, b, x) in enumerate(bear_returns(pf)):
        ws.cell(r, 1, f"{a} → {b}")
        c = ws.cell(r, 2, round(x, 1))
        c.fill = S.good if x >= 0 else S.bad
        for j, k in enumerate(parts, start=3):
            ws.cell(r, j, round(pb[k][i][2], 1) if i < len(pb[k]) else None)
        r += 1
    r += 1
    ws.cell(r, 1, "Części osobno (ten sam rachunek, tylko jedna część)").font = S.bold
    sw.header(ws, r + 1, ["Część", "Wynik %", "CAGR %", "Obsunięcie %", "R²"], [44, 18, 14, 14, 14], S, 30)
    r += 2
    for k, v in parts.items():
        smv = smooth(v)
        for j, x in enumerate((k, round(v["ret"], 1), round(sw.cagr(v, D) or 0, 2), round(v["mdd"], 1),
                               round(smv["r2"], 3) if smv["r2"] else None), start=1):
            ws.cell(r, j, x)
        r += 1
    r += 1
    ws.cell(r, 1, "Najgłębsze obsunięcia").font = S.bold
    sw.header(ws, r + 1, ["Szczyt (od)", "Dno", "Obsunięcie %", "Odrobione", "Dni"], [44, 18, 14, 14, 14], S, 30)
    r += 2
    for x in sorted(dds, key=lambda x: -x["depth"])[:8]:
        for j, v in enumerate((x["start"], x["trough"], round(x["depth"], 1), x["rec"] or "jeszcze nie",
                               int(x["days"]) if x["days"] is not None else None), start=1):
            ws.cell(r, j, v)
        r += 1

    # KAPITAŁ
    kp = wb.create_sheet("KAPITAŁ")
    curves = {"RÓWNA": dict(pf["curve"])}
    for k, v in parts.items():
        curves[k] = dict(v["curve"])
    names = list(curves)
    sw.header(kp, 1, ["Dzień"] + [f"{n} $" for n in names] + ["Obsunięcie RÓWNA %"], [12] + [14] * len(names) + [12], S, 30)
    peak = 0.0
    last = {n: sw.CAPITAL for n in names}
    for d, v in pf["curve"]:
        peak = max(peak, v)
        row = [d]
        for n in names:
            last[n] = curves[n].get(d, last[n])
            row.append(round(last[n], 2))
        row.append(round(-(peak - v) / peak * 100.0 if peak > 0 else 0.0, 2))
        kp.append(row)
    n = len(pf["curve"])
    if n > 1:
        ch = LineChart()
        ch.title = "Stan portfela 10 000 $ (skala logarytmiczna): RÓWNA i jej części osobno"
        ch.y_axis.scaling.logBase = 10
        ch.add_data(Reference(kp, min_col=2, max_col=1 + len(names), min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(kp, min_col=1, min_row=2, max_row=n + 1))
        ch.x_axis.tickLblSkip = ch.x_axis.tickMarkSkip = max(1, n // 12)
        ch.height, ch.width = 12, 30
        for sser, col, w in zip(ch.series, ("2A78D6", "1BAF7A", "EB6834"), (22000, 9000, 9000)):
            sser.graphicalProperties.line.solidFill = col
            sser.graphicalProperties.line.width = w
        kp.add_chart(ch, "H2")
        c2 = LineChart()
        c2.title = "Obsunięcie RÓWNA od szczytu %"
        c2.add_data(Reference(kp, min_col=2 + len(names), min_row=1, max_row=n + 1), titles_from_data=True)
        c2.set_categories(Reference(kp, min_col=1, min_row=2, max_row=n + 1))
        c2.x_axis.tickLblSkip = c2.x_axis.tickMarkSkip = max(1, n // 12)
        c2.x_axis.tickLblPos = "low"
        c2.legend = None
        c2.height, c2.width = 8, 30
        c2.series[0].graphicalProperties.line.solidFill = "2A78D6"
        kp.add_chart(c2, "H27")
    kp.freeze_panes = "A2"

    ob = wb.create_sheet("OBSUNIĘCIA")
    ob["A1"] = "Każde obsunięcie portfela ≥ 2% (wycena po close)"
    ob["A1"].font = S.bold
    sw.header(ob, 2, ["Szczyt (od)", "Stan w szczycie $", "Dno", "Stan na dnie $", "Obsunięcie %", "Odrobione", "Dni"],
              [12, 14, 12, 14, 12, 14, 10], S, 30)
    for x in dds:
        ob.append([x["start"], round(x["peak"], 2), x["trough"], round(x["low"], 2), round(x["depth"], 1),
                   x["rec"] or "jeszcze nie", int(x["days"]) if x["days"] is not None else None])
        ob.cell(ob.max_row, 5).fill = S.bad if x["depth"] >= 15 else S.sub

    yr = wb.create_sheet("ROK PO ROKU")
    sw.header(yr, 1, ["Rok", "Wynik roku %", f"Zysk $ {INDEX}", "Zysk $ spółki", f"Transakcji {INDEX}", "Transakcji spółki"],
              [8, 12, 13, 13, 12, 12], S, 30)
    agg: dict = {}
    for row in pf["rows"]:
        a = agg.setdefault(row["out"][:4], {INDEX: [0.0, 0], "spółki": [0.0, 0]})
        a[row["label"]][0] += row["pnl"]
        a[row["label"]][1] += 1
    for y, v in pf["years"]:
        a = agg.get(y, {INDEX: [0.0, 0], "spółki": [0.0, 0]})
        yr.append([y, round(v, 2), round(a[INDEX][0], 2), round(a["spółki"][0], 2), a[INDEX][1], a["spółki"][1]])
        yr.cell(yr.max_row, 2).fill = S.good if v >= 0 else S.bad

    sw.portfolio_sheet(wb, "TRANSAKCJE", pf, f"Przebieg portfela — {cfg.short()}", S)

    sg = wb.create_sheet("SYGNAŁY SPÓŁEK")
    sg["A1"] = ("Każda seria ≥ 5 świec spadkowych ze spadkiem ≥ 10% — czy rynek miał trend wzrostowy (EMA100 SPY rośnie), "
                "czy weszła do portfela; wynik netto (dla odrzuconych — hipotetyczny)")
    sg["A1"].font = S.bold
    cols = ["Spółka", "Grupa", "1. świeca", "Sygnał", "Świec", "Spadek serii %", "R spółki %", "Trend SPY rośnie",
            "Wejście", "Wyjście", "Wynik netto %", "W portfelu", "Zysk $", "Uwagi"]
    sw.header(sg, 2, cols, [9, 11, 11, 11, 7, 10, 9, 9, 11, 11, 10, 9, 10, 50], S)
    inpf = {(r["sym"], r["in"]): r for r in pf["rows"]}
    skip = {(D.stocks[si].sym, D.cal[t]): why for si, t, why in pf["skipped_list"]}
    rr = lambda v, p=2: None if v is None or (isinstance(v, float) and np.isnan(v)) else round(v, p)
    for e in evs:
        p = inpf.get((e["sym"], e.get("in")))
        note = "" if p else (e.get("status", "") if e["trade"] is None else skip.get((e["sym"], e.get("in")), ""))
        sg.append([e["sym"], e["group"], e["first"], e["last"], e["k"], rr(e["drop"]), rr(e["R"]),
                   "tak" if e.get("trend") else "nie", e.get("in"), e.get("out"), rr(e.get("net"), 3),
                   "tak" if p else "nie", rr(p["pnl"]) if p else None, note])
        if e.get("net") is not None:
            sg.cell(sg.max_row, 11).fill = (S.good if e["net"] > 0 else S.bad) if p else S.sub
    sg.freeze_panes = "B3"
    sg.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{sg.max_row}"

    if rob:
        od = wb.create_sheet("ODPORNOŚĆ")
        od["A1"] = ("Czy wynik zależy od jednej dokładnej liczby? Ten sam rachunek przy sąsiednich ustawieniach i dla części osobno. "
                    "Bessy: 2000–02, 2007–09, 2020, 2022.")
        od["A1"].font = S.bold
        sw.header(od, 3, ["Wariant", "Transakcji", "Wynik %", "CAGR %", "Obsunięcie %", "R²", "Lata stratne", "Najgorszy rok %",
                          "Bessa 2000–02 %", "Bessa 2007–09 %", "Bessa 2020 %", "Bessa 2022 %"],
                  [46, 10, 10, 9, 11, 8, 9, 10, 10, 10, 10, 10], S, 30)
        for name, n, ret, cg_, dd, r2, yneg, worst, bears in dict(rob)["ROB"]:
            od.append([name, n, round(ret, 1), round(cg_ or 0, 2), round(dd, 1), round(r2, 3) if r2 else None, yneg,
                       round(worst, 1) if worst is not None else None] + [round(b, 1) for b in bears])
            od.cell(od.max_row, 5).fill = S.good if dd < 25 else S.bad
        r = od.max_row + 3
        od.cell(r, 1, "Wynik pojedynczej transakcji w okresach (bez portfela)").font = S.bold
        sw.header(od, r + 1, ["Część", "Okres", "Transakcji", "Śr. netto %", "Zyskownych %"], [46, 10, 10, 9, 11], S, 30)
        r += 2
        for lab, per, n, m, w in dict(rob)["PER"]:
            for j, v in enumerate((lab, per, n, round(m, 3) if m is not None else None, round(w, 1) if w is not None else None), start=1):
                c = od.cell(r, j, v)
                if j == 4 and m is not None:
                    c.fill = S.good if m > 0 else S.bad
            r += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
def selftest() -> None:
    fails: list[str] = []

    def check(name, got, exp):
        ok = got == exp
        print(f"  {'OK  ' if ok else 'BŁĄD'} {name}: {got}" + ("" if ok else f" (oczekiwano {exp})"))
        if not ok:
            fails.append(name)

    print(f"IA 4 {__version__} — selftest ia4.swrowna\n")
    sw.selftest_core(check)
    rng = np.random.default_rng(3)
    v = realized_vol(100 * np.exp(np.cumsum(rng.normal(0, 0.01, 400))))
    check("zmienność roczna ok. 16% przy 1% dziennie", 12 < float(np.nanmean(v)) < 20, True)
    # Dane: SPY i QQQ rosną (trend), spółka A: 6 spadków po 2% (−11,4%), potem wzrosty; QQQ: spadek → sygnał RSI(3)
    n = 260
    days = np.array([str(np.datetime64("2010-01-04") + np.timedelta64(k, "D")) for k in range(n)])

    def mk(sym, grp, path):
        c = np.array(path, float)
        o = np.concatenate([[c[0]], c[:-1]])
        h, l = np.maximum(o, c) * 1.003, np.minimum(o, c) * 0.997
        return sw.Stock(sym, grp, days, o, h, l, c, np.full(n, 1e6))
    up = list(100 * 1.002 ** np.arange(n))
    spy = mk("SPY", "context", up)
    qp = list(100 * 1.003 ** np.arange(n))
    for k in range(200, 203):
        qp[k] = qp[k - 1] * 0.975
    for k in range(203, n):
        qp[k] = qp[k - 1] * 1.01
    qqq = mk("QQQ", "context", qp)
    a = [100.0] * 150
    p = 100.0
    for _ in range(6):
        p *= 0.98
        a.append(p)
    while len(a) < n:
        p *= 1.005
        a.append(p)
    ac = np.array(a)
    ac[:150] = 100.2
    ao = np.concatenate([[100.0], np.array(a[:-1])])
    ao[:150] = 100.0
    ao[150:156] = ac[150:156] / 0.98                     # świece spadkowe: open wyżej niż close
    A = sw.Stock("A", "main", days, ao, np.maximum(ao, ac) * 1.003, np.minimum(ao, ac) * 0.997, ac, np.full(n, 1e6))
    D = sw.finish([A, spy, qqq])
    cfg = Cfg()
    it = index_trades(D, cfg)
    check(f"{INDEX}: jedna transakcja RSI(3), wielkość ≤ {VOL_CAP:g}", (len(it), bool(it and 0 < it[0]["size"] <= VOL_CAP)), (1, True))
    ev = stock_events(D, cfg, None)
    check("spółka: seria 6 spadków, wejście przy trendzie", (len(ev), ev[0]["k"] if ev else None, ev[0]["trade"] is not None if ev else None),
          (1, 6, True))
    pf, idx, evs = build(D, cfg, None, 0.3)
    check("portfel: obie części", sorted({r["label"] for r in pf["rows"]}), ["QQQ", "spółki"])
    spy_dn = mk("SPY", "context", list(100 * 0.998 ** np.arange(n)))
    D2 = sw.finish([A, spy_dn, qqq])
    ev2 = stock_events(D2, cfg, None)
    check("spółka: bez trendu SPY → odrzucona", (ev2[0]["trade"] is None, "EMA100" in ev2[0]["status"]), (True, True))
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        pth = run(no_sync=True, out=Path(td) / "t.xlsx", data=D, since=None, robust=False)
        from openpyxl import load_workbook
        wb = load_workbook(pth)
        check("arkusze", wb.sheetnames, ["PODSUMOWANIE", "KAPITAŁ", "OBSUNIĘCIA", "ROK PO ROKU", "TRANSAKCJE", "SYGNAŁY SPÓŁEK"])
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    if fails:
        sys.exit(1)


def main() -> None:
    ap = argparse.ArgumentParser(description="Strategia RÓWNA — kupowanie spadków tylko przy trendzie wzrostowym (instrukcja 7a)")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--no-sync", action="store_true", help="bez synchronizacji bazy dobowej")
    ap.add_argument("--since", default=sw.DAILY_SINCE, help="początek portfela (domyślnie 2000-01-01)")
    ap.add_argument("--cost", type=float, default=sw.COST_PCT, help="spread %% na transakcję (raz; domyślnie 0,3)")
    ap.add_argument("--skala", type=float, default=1.0, help="mnożnik wszystkich pozycji (np. 1.5)")
    ap.add_argument("--bez-odpornosci", action="store_true", help="bez arkusza ODPORNOŚĆ (szybciej)")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    if a.selftest:
        selftest()
        return
    run(a.no_sync, a.out, a.since, a.cost, Cfg(skala=a.skala), robust=not a.bez_odpornosci)


if __name__ == "__main__":
    main()
