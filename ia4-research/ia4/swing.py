"""
IA 4 — wspólny silnik badań wielodniowych (instrukcja, sekcja 7a).
Wersja projektu: 1.51 (2026-10-20) — musi zgadzać się z IA4_INSTRUKCJA.md

Badania, w których pozycja może trwać dłużej niż jedną sesję (decyzja człowieka z 16.10.2026):
`ia4.swtrend` (strategie jednokierunkowe z SL i TP), `ia4.swrsi` (strategia RSI(3) / EMA100 / VWAP10),
`ia4.swserie` (mapa: wejście po serii świec). Ten moduł daje im wspólne:

  • dane — baza dobowa (`data/daily/`, `python -m ia4.daily`) od 2000 r., każda spółka jako ciąg pełnych sesji;
  • rozliczenie SL / TP na świecach dobowych — zawsze pesymistycznie (zasada 5): wejście po open, SL i TP
    w tej samej świecy → SL, luka przez SL → wyjście po open (gorzej), luka przez TP → wyjście po TP (bez premii);
  • spread 0,3% wartości pozycji, liczony RAZ na transakcję (wejście + wyjście razem; decyzja człowieka
    z 16.10.2026 — obowiązuje we wszystkich badaniach);
  • portfel 10 000 $ z dźwignią: depozyt pozycji = STAKE_PCT% stanu portfela (wycena po zamknięciu poprzedniej
    sesji), wartość pozycji = depozyt × dźwignia, wynik = wartość × (ruch % − spread %); suma depozytów
    ≤ stan portfela (nadmiar sygnałów — stawki proporcjonalnie mniejsze), jedna pozycja na spółkę,
    obsunięcie liczone z wyceny dziennej; ≤ 10% kapitału startowego = upadłość. Bez kosztu nocnego (swap CFD).
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field

import numpy as np

from . import oneday as od

COST_PCT = 0.3            # spread na transakcję, % wartości pozycji, raz (wejście + wyjście) — decyzja człowieka 16.10.2026
DAILY_SINCE = "2000-01-01"
CAPITAL = 10_000.0
STAKE_PCT = 10.0          # depozyt pozycji = 10% stanu portfela
RUIN_PCT = 10.0
SPLIT_YEAR = 2013         # część A < 2013, część B ≥ 2013

REASON = {0: "—", 1: "SL", 2: "TP", 3: "czas", 4: "sygnał", 5: "luka SL"}


# ---------------------------------------------------------------------------
#  Dane
# ---------------------------------------------------------------------------
@dataclass
class Stock:
    sym: str
    group: str
    dates: np.ndarray          # str RRRR-MM-DD, tylko pełne sesje
    o: np.ndarray
    h: np.ndarray
    l: np.ndarray
    c: np.ndarray
    v: np.ndarray
    t: np.ndarray = None       # indeks dnia w kalendarzu wspólnym


@dataclass
class Data:
    stocks: list
    cal: list
    note: str = ""
    by_sym: dict = field(default_factory=dict)

    def rows(self, *groups) -> list[int]:
        return [i for i, s in enumerate(self.stocks) if s.group in groups]


def clean(df, since: str | None):
    """Tabela bazy dobowej → pełne sesje (spójne ceny, nie płaskie) — jak odsearch.days_from_daily."""
    if since:
        df = df[df["date"] >= since]
    o, h, l, c = (df[k].to_numpy(float) for k in ("o", "h", "l", "c"))
    v = df["v"].to_numpy(float) if "v" in df else np.zeros(len(df))
    tol = 1e-4
    ok = ((o > 0) & (c > 0) & (l > 0) & (h >= l) & (o <= h * (1 + tol)) & (o >= l * (1 - tol))
          & (c <= h * (1 + tol)) & (c >= l * (1 - tol)) & ~((o == h) & (h == l) & (l == c)))
    d = df["date"].astype(str).to_numpy()
    o, h, l, c = o[ok], h[ok], l[ok], c[ok]
    return d[ok], o, np.maximum(h, np.maximum(o, c)), np.minimum(l, np.minimum(o, c)), c, v[ok]


def load(no_sync: bool = False, since: str | None = DAILY_SINCE, groups=("main", "proof", "doubleProof", "context")) -> Data:
    from . import daily as dl
    note = "bez synchronizacji (--no-sync)"
    if not no_sync:
        try:
            r = dl.sync(quiet=True)
            note = f"baza dobowa zsynchronizowana: {r['ready']}/{r['total']} instrumentów z historią"
        except Exception as e:                                  # bez sieci / bez klucza — kopia lokalna
            note = f"bez synchronizacji ({type(e).__name__}: {e})"
    have = set(dl.local_symbols())
    stocks = []
    for s, g in od.ordered(od.universe()):
        if g not in groups:
            continue
        if s not in have:
            print(f"  ! {s}: brak w data/daily/ — pominięta")
            continue
        d, o, h, l, c, v = clean(dl.load(s), since)
        if len(d) < 30:
            continue
        stocks.append(Stock(s, g, d, o, h, l, c, v))
    if not stocks:
        raise SystemExit("Brak bazy dobowej w kopii lokalnej — uruchom najpierw `python -m ia4.daily`.")
    return finish(stocks, note + f" · od {since or 'początku'}")


def finish(stocks: list, note: str = "") -> Data:
    cal = sorted(set().union(*(set(s.dates.tolist()) for s in stocks)))
    pos = {d: i for i, d in enumerate(cal)}
    for s in stocks:
        s.t = np.array([pos[d] for d in s.dates], dtype=np.int32)
    return Data(stocks, cal, note, {s.sym: s for s in stocks})


# ---------------------------------------------------------------------------
#  Wskaźniki (na ciągu pełnych sesji jednej spółki)
# ---------------------------------------------------------------------------
def sma(x: np.ndarray, n: int) -> np.ndarray:
    out = np.full(len(x), np.nan)
    if len(x) >= n:
        cs = np.cumsum(np.concatenate([[0.0], x]))
        out[n - 1:] = (cs[n:] - cs[:-n]) / n
    return out


def ema(x: np.ndarray, n: int) -> np.ndarray:
    """EMA (α = 2/(n+1)), start — SMA z pierwszych n wartości."""
    out = np.full(len(x), np.nan)
    if len(x) < n:
        return out
    a = 2.0 / (n + 1)
    e = float(np.mean(x[:n]))
    out[n - 1] = e
    for i in range(n, len(x)):
        e = a * x[i] + (1 - a) * e
        out[i] = e
    return out


def rsi(c: np.ndarray, n: int) -> np.ndarray:
    """RSI Wildera z zamknięć (start — średnia z pierwszych n zmian)."""
    out = np.full(len(c), np.nan)
    if len(c) <= n:
        return out
    d = np.diff(c)
    up, dn = np.maximum(d, 0.0), np.maximum(-d, 0.0)
    au, ad = float(up[:n].mean()), float(dn[:n].mean())
    out[n] = 100.0 if ad == 0 else 100.0 - 100.0 / (1.0 + au / ad)
    for i in range(n, len(d)):
        au = (au * (n - 1) + up[i]) / n
        ad = (ad * (n - 1) + dn[i]) / n
        out[i + 1] = 100.0 if ad == 0 else 100.0 - 100.0 / (1.0 + au / ad)
    return out


def rolling(x: np.ndarray, n: int, fn) -> np.ndarray:
    from numpy.lib.stride_tricks import sliding_window_view
    out = np.full(len(x), np.nan)
    if len(x) >= n:
        out[n - 1:] = fn(sliding_window_view(x, n), axis=1)
    return out


def prior(x: np.ndarray, k: int = 1) -> np.ndarray:
    """Wartość sprzed k sesji."""
    out = np.full(len(x), np.nan)
    if k < len(x):
        out[k:] = x[:-k]
    return out


def streak(flag: np.ndarray) -> np.ndarray:
    """Ile sesji z rzędu (kończąc na tej) flaga jest prawdziwa."""
    out = np.zeros(len(flag), dtype=np.int32)
    run = 0
    for i, f in enumerate(flag):
        run = run + 1 if f else 0
        out[i] = run
    return out


# ---------------------------------------------------------------------------
#  Rozliczenie SL / TP (pesymistycznie) — wejście po open sesji q
# ---------------------------------------------------------------------------
def sltp(o, h, l, c, d: int, slp, tpp, H: int):
    """Dla KAŻDEJ sesji wejścia q (open q): wynik pozycji d (+1 long, −1 short) z SL slp% i TP tpp% od wejścia
    (tablice albo liczby; inf — brak), najwyżej H sesji (wyjście po close sesji q+H−1).
    Zwraca (ruch % w kierunku pozycji przed spreadem, przesunięcie sesji wyjścia, powód); NaN — za mało sesji."""
    Q = len(o)
    q = np.arange(Q)
    slp = np.broadcast_to(np.asarray(slp, float), (Q,))
    tpp = np.broadcast_to(np.asarray(tpp, float), (Q,))
    ent = o
    SL = ent * (1.0 - d * slp / 100.0)
    TP = ent * (1.0 + d * tpp / 100.0)
    px = np.full(Q, np.nan)
    off = np.zeros(Q, dtype=np.int16)
    why = np.zeros(Q, dtype=np.int8)
    done = np.zeros(Q, bool)
    for j in range(H):
        k = q + j
        inb = k < Q
        kk = np.minimum(k, Q - 1)
        oo, hh, ll, cc = o[kk], h[kk], l[kk], c[kk]
        if j > 0:
            act = ~done & inb
            gs = act & ((oo <= SL) if d > 0 else (oo >= SL))
            gt = act & ~gs & ((oo >= TP) if d > 0 else (oo <= TP))
            px[gs], why[gs] = oo[gs], 5
            px[gt], why[gt] = TP[gt], 2
            off[gs | gt] = j
            done |= gs | gt
        act = ~done & inb
        hs = act & ((ll <= SL) if d > 0 else (hh >= SL))
        ht = act & ~hs & ((hh >= TP) if d > 0 else (ll <= TP))
        px[hs], why[hs] = SL[hs], 1
        px[ht], why[ht] = TP[ht], 2
        off[hs | ht] = j
        done |= hs | ht
        if j == H - 1:
            last = ~done & inb
            px[last], why[last], off[last] = cc[last], 3, j
            done |= last
    move = d * (px / ent - 1.0) * 100.0
    return move, off, why


# ---------------------------------------------------------------------------
#  Portfel z dźwignią
# ---------------------------------------------------------------------------
def portfolio(trades: list[dict], data: Data, capital: float = CAPITAL, stake_pct: float = STAKE_PCT,
              lev: float = 1.0, cost: float = COST_PCT, keep_rows: bool = True, max_pos: int | None = None,
              max_by_label: dict | None = None) -> dict:
    """trades: {si, qe, qx, d, pin, pout, why, label?, prio?} — qe/qx: sesja wejścia (open) i wyjścia w ciągu spółki si.
    Kolejność w dniu: najpierw wejścia po open (depozyt z wyceny po poprzednim close), potem wyjścia.
    max_pos — najwięcej pozycji naraz; gdy sygnałów jest więcej niż wolnych miejsc, wchodzą te o najmniejszym `prio`.
    max_by_label — limit pozycji naraz osobno dla każdej etykiety transakcji (`label`), np. {"spółki": 5}.
    `size` w transakcji — wartość pozycji jako ułamek stanu portfela (depozyt = size / dźwignia); bez niej depozyt
    = stake_pct% stanu portfela.
    Wynik `skipped_list`: (si, te, powód) każdego sygnału, który nie wszedł do portfela."""
    st = data.stocks
    T = len(data.cal)
    ents: dict[int, list] = {}
    for tr in trades:
        tr["te"] = int(st[tr["si"]].t[tr["qe"]])
        tr["tx"] = int(st[tr["si"]].t[tr["qx"]])
        ents.setdefault(tr["te"], []).append(tr)
    if not ents:
        return {"rows": [], "end": capital, "ret": 0.0, "mdd": 0.0, "trades": 0, "wins": 0, "ruined": False,
                "max_open": 0, "scaled": 0, "skipped": 0, "skipped_limit": 0, "skipped_list": [], "curve": [],
                "years": [], "first": None, "last": None, "capital": capital, "lev": lev, "stake": stake_pct,
                "max_pos": max_pos}
    closes = {}
    eq = capital
    peak = capital
    mdd, mtm_prev = 0.0, capital
    openp: dict[int, tuple] = {}               # si → (tr, margin)
    rows, curve, years = [], [], {}
    n, wins, scaled, skipped, max_open, skipped_limit = 0, 0, 0, 0, 0, 0
    skipped_list = []
    ruined = False
    t0 = min(ents)
    last_t = max(tr["tx"] for tr in trades)
    for t in range(t0, last_t + 1):
        new = []
        for tr in ents.get(t, []):
            if tr["si"] in openp:
                skipped += 1
                skipped_list.append((tr["si"], t, "spółka już w pozycji"))
            else:
                new.append(tr)
        if max_pos and new:
            new.sort(key=lambda tr: tr.get("prio", 0.0))
            room = max(0, max_pos - len(openp))
            for tr in new[room:]:
                skipped_list.append((tr["si"], t, f"limit {max_pos} pozycji naraz"))
            skipped_limit += max(0, len(new) - room)
            new = new[:room]
        if max_by_label and new:
            keep = []
            for tr in sorted(new, key=lambda tr: tr.get("prio", 0.0)):
                lab = tr.get("label", "")
                lim = max_by_label.get(lab)
                if lim is not None and sum(1 for x, _ in openp.values() if x.get("label", "") == lab) \
                        + sum(1 for x in keep if x.get("label", "") == lab) >= lim:
                    skipped_list.append((tr["si"], t, f"limit {lim} pozycji naraz ({lab})"))
                    skipped_limit += 1
                else:
                    keep.append(tr)
            new = keep
        if new:
            base = mtm_prev
            ms = [base * (tr["size"] / lev if "size" in tr else stake_pct / 100.0) for tr in new]
            free = base - sum(m for _, m in openp.values())
            need = sum(ms)
            sc = 1.0 if need <= free * (1 + 1e-9) else (max(0.0, free) / need if need > 0 else 0.0)
            scaled += sc < 1.0
            for tr, m0 in zip(new, ms):
                if m0 * sc > 0:
                    tr["scale"] = sc
                    openp[tr["si"]] = (tr, m0 * sc)
                else:
                    skipped_list.append((tr["si"], t, "brak wolnego depozytu"))
        max_open = max(max_open, len(openp))
        for si in [k for k, (tr, _) in openp.items() if tr["tx"] == t]:
            tr, m = openp.pop(si)
            pct = tr["d"] * (tr["pout"] / tr["pin"] - 1.0) * 100.0 - cost
            pnl = m * lev * pct / 100.0
            eq += pnl
            n += 1
            wins += pnl > 0
            if keep_rows:
                s = st[si]
                rows.append({"in": data.cal[tr["te"]], "out": data.cal[t], "sym": s.sym, "group": s.group,
                             "d": tr["d"], "pin": tr["pin"], "pout": tr["pout"], "why": REASON.get(tr["why"], tr["why"]),
                             "sessions": int(tr["qx"] - tr["qe"] + 1), "move": pct + cost, "net": pct,
                             "margin": m, "pnl": pnl, "eq": eq, "label": tr.get("label", ""),
                             "scale": tr.get("scale", 1.0)})
        mtm = eq
        for si, (tr, m) in openp.items():
            s = st[si]
            q = np.searchsorted(s.t, t, side="right") - 1
            cl = s.c[q] if q >= tr["qe"] else tr["pin"]
            mtm += m * lev * (tr["d"] * (cl / tr["pin"] - 1.0) * 100.0 - cost) / 100.0
        peak = max(peak, mtm)
        mdd = max(mdd, min(100.0, (peak - mtm) / peak * 100.0) if peak > 0 else 0.0)
        mtm_prev = mtm
        y = data.cal[t][:4]
        years.setdefault(y, [None, None])
        if years[y][0] is None:
            years[y][0] = curve[-1][1] if curve else capital
        years[y][1] = mtm
        curve.append((data.cal[t], mtm))
        if mtm <= capital * RUIN_PCT / 100.0:
            ruined = True                       # rachunek z ochroną przed ujemnym saldem: nie mniej niż 0
            curve[-1] = (data.cal[t], max(0.0, mtm))
            years[y][1] = max(0.0, mtm)
            mdd = max(mdd, min(100.0, (peak - max(0.0, mtm)) / peak * 100.0))
            break
    end = curve[-1][1] if curve else capital
    return {"rows": rows, "end": end, "ret": (end / capital - 1.0) * 100.0, "mdd": mdd, "trades": n, "wins": wins,
            "ruined": ruined, "max_open": max_open, "scaled": scaled, "skipped": skipped, "curve": curve,
            "years": [(y, (b / a - 1.0) * 100.0 if a else 0.0) for y, (a, b) in sorted(years.items())],
            "first": rows[0]["in"] if rows else None, "last": rows[-1]["out"] if rows else None,
            "capital": capital, "lev": lev, "stake": stake_pct, "max_pos": max_pos, "skipped_limit": skipped_limit,
            "skipped_list": skipped_list}


def cagr(pf: dict, data: Data) -> float | None:
    if not pf["curve"] or pf["end"] <= 0:
        return None
    from datetime import date
    a, b = date.fromisoformat(pf["curve"][0][0]), date.fromisoformat(pf["curve"][-1][0])
    yrs = (b - a).days / 365.25
    if yrs < 1.0:
        return None
    return ((pf["end"] / pf["capital"]) ** (1.0 / yrs) - 1.0) * 100.0


# ---------------------------------------------------------------------------
#  Excel — wspólne drobiazgi
# ---------------------------------------------------------------------------
def styles():
    from openpyxl.styles import Font, PatternFill
    return SimpleStyles(PatternFill("solid", fgColor="1F2A37"), PatternFill("solid", fgColor="DCFCE7"),
                        PatternFill("solid", fgColor="FEE2E2"), PatternFill("solid", fgColor="E5E7EB"),
                        Font(bold=True, color="FFFFFF"), Font(bold=True))


@dataclass
class SimpleStyles:
    head: object
    good: object
    bad: object
    sub: object
    white: object
    bold: object


def header(ws, row: int, cols: list, widths: list, S=None, height: int = 45):
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter
    S = S or styles()
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row, j, h)
        c.font, c.fill = S.white, S.head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = height


def heat(v: float | None, lim: float):
    """Kolor komórki mapy: zielony (zysk) — czerwony (strata), nasycenie wg |v| / lim."""
    from openpyxl.styles import PatternFill
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return PatternFill("solid", fgColor="F3F4F6")
    x = max(-1.0, min(1.0, v / lim if lim > 0 else 0.0))
    if x >= 0:
        r, g, b = int(255 - 155 * x), int(255 - 35 * x), int(255 - 155 * x)
    else:
        x = -x
        r, g, b = int(255 - 15 * x), int(255 - 155 * x), int(255 - 155 * x)
    return PatternFill("solid", fgColor=f"{r:02X}{g:02X}{b:02X}")


def portfolio_sheet(wb, title: str, pf: dict, intro: str, S=None):
    """Arkusz z przebiegiem portfela: podsumowanie, lata, transakcje."""
    S = S or styles()
    ws = wb.create_sheet(title[:31])
    ws["A1"] = intro
    ws["A1"].font = S.bold
    summ = [("Kapitał startowy $", pf["capital"]), ("Dźwignia", f"1:{pf['lev']:g}"),
            ("Depozyt na pozycję", f"{pf['stake']:g}% stanu portfela"),
            ("Stan końcowy $", round(pf["end"], 2)), ("Wynik %", round(pf["ret"], 2)),
            ("Największe obsunięcie % (wycena dzienna)", round(pf["mdd"], 2)),
            ("Transakcji", pf["trades"]), ("Zyskownych %", round(pf["wins"] / pf["trades"] * 100, 1) if pf["trades"] else None),
            ("Najwięcej pozycji naraz", pf["max_open"]), ("Dni z pomniejszoną stawką (brak wolnego depozytu)", pf["scaled"]),
            ("Sygnały pominięte (spółka już w pozycji)", pf["skipped"]),
            ("Limit pozycji naraz", pf.get("max_pos") or "bez limitu"),
            ("Sygnały pominięte (limit pozycji)", pf.get("skipped_limit", 0)),
            ("Upadłość (≤ 10% kapitału)", "TAK" if pf["ruined"] else "nie"),
            ("Okres", f"{pf['first'] or '—'} – {pf['last'] or '—'}")]
    for k, (a, b) in enumerate(summ, start=3):
        ws.cell(k, 1, a)
        ws.cell(k, 2, b)
    r0 = 3 + len(summ) + 1
    ws.cell(r0, 1, "Rok").font = S.bold
    ws.cell(r0, 2, "Wynik roku %").font = S.bold
    for k, (y, v) in enumerate(pf["years"], start=r0 + 1):
        ws.cell(k, 1, y)
        c = ws.cell(k, 2, round(v, 2))
        c.fill = S.good if v >= 0 else S.bad
    r1 = r0 + len(pf["years"]) + 2
    cols = ["Nr", "Wejście (open)", "Wyjście", "Spółka", "Kierunek", "Cena wejścia", "Cena wyjścia", "Powód wyjścia",
            "Sesji w pozycji", "Ruch % (przed spreadem)", "Wynik % po spreadzie", f"Wynik % depozytu (×{pf['lev']:g})",
            "Depozyt $", "Zysk / strata $", "Stan portfela $ (zrealizowany)"]
    header(ws, r1, cols, [6, 12, 12, 10, 9, 11, 11, 10, 9, 11, 11, 12, 11, 11, 14], S)
    for k, r in enumerate(pf["rows"], start=1):
        ws.append([k, r["in"], r["out"], r["sym"] + (" · dP" if r["group"] == "doubleProof" else ""),
                   "long" if r["d"] > 0 else "short", round(r["pin"], 4), round(r["pout"], 4), r["why"], r["sessions"],
                   round(r["move"], 3), round(r["net"], 3), round(r["net"] * pf["lev"], 3), round(r["margin"], 2),
                   round(r["pnl"], 2), round(r["eq"], 2)])
        ws.cell(ws.max_row, 14).fill = S.good if r["pnl"] >= 0 else S.bad
    ws.freeze_panes = ws.cell(r1 + 1, 1)
    ws.column_dimensions["A"].width = 44
    return ws


def selftest_core(check) -> None:
    """Testy silnika (wołane z selftestów modułów)."""
    o = np.array([100.0, 101, 99, 103, 104])
    h = np.array([102.0, 103, 100, 106, 105])
    l = np.array([99.0, 98, 95, 102, 103])
    c = np.array([101.0, 99, 96, 105, 104])
    mv, off, why = sltp(o, h, l, c, +1, 2.0, 4.0, 3)
    # wejście 100: SL 98, TP 104 — dzień 0: L 99, H 102 → nic; dzień 1: L 98 ≤ SL → SL (pesymistycznie, choć H 103 < TP)
    check("sltp long: SL w 2. sesji", (round(mv[0], 6), int(off[0]), int(why[0])), (-2.0, 1, 1))
    mv, off, why = sltp(o, h, l, c, -1, 2.0, 4.0, 3)
    # short od 100: SL 102, TP 96 — dzień 0: H 102 ≥ SL → SL
    check("sltp short: SL w 1. sesji", (round(mv[0], 6), int(off[0]), int(why[0])), (-2.0, 0, 1))
    mv, off, why = sltp(o, h, l, c, +1, 50.0, 50.0, 2)
    check("sltp: wyjście po czasie (close 2. sesji)", (round(mv[0], 6), int(off[0]), int(why[0])), (-1.0, 1, 3))
    check("sltp: za mało sesji → NaN", bool(np.isnan(mv[-1])), True)
    g = sltp(np.array([100.0, 95]), np.array([101.0, 96]), np.array([99.5, 94]), np.array([100.0, 95]), +1, 2.0, 9.0, 2)
    check("sltp: luka przez SL → wyjście po open", (round(float(g[0][0]), 6), int(g[2][0])), (-5.0, 5))
    check("ema/rsi bez NaN po rozgrzewce", (bool(np.isfinite(ema(np.arange(10.0), 3)[2])), bool(np.isfinite(rsi(np.arange(10.0), 3)[3]))), (True, True))
    check("rsi rosnącego = 100", round(float(rsi(np.arange(1.0, 20.0), 3)[-1]), 6), 100.0)
    # portfel: jedna transakcja +10% na dźwigni 5, depozyt 10% → +5% portfela minus spread
    st = Stock("A", "main", np.array(["2020-01-02", "2020-01-03"]), np.array([100.0, 100.0]), np.array([111.0, 111.0]),
               np.array([99.0, 99.0]), np.array([110.0, 110.0]), np.zeros(2))
    dd = finish([st])
    pf = portfolio([{"si": 0, "qe": 0, "qx": 0, "d": 1, "pin": 100.0, "pout": 110.0, "why": 3}], dd, lev=5.0)
    check("portfel 1:5, depozyt 10%: +10% ruchu → wynik", round(pf["ret"], 6), round(10.0 * 5 * (10.0 - COST_PCT) / 100.0, 6))
    two = [{"si": 0, "qe": 0, "qx": 0, "d": 1, "pin": 100.0, "pout": 110.0, "why": 3, "prio": 2.0},
           {"si": 1, "qe": 0, "qx": 0, "d": 1, "pin": 100.0, "pout": 90.0, "why": 3, "prio": 1.0}]
    st2 = Stock("B", "main", st.dates, st.o, st.h, st.l, st.c, st.v)
    pf = portfolio([dict(x) for x in two], finish([st, st2]), lev=1.0, max_pos=1)
    check("limit 1 pozycji: wchodzi sygnał o mniejszym prio (B, −10%)", (pf["trades"], pf["skipped_limit"], round(pf["ret"], 6)),
          (1, 1, round(10.0 * (-10.0 - COST_PCT) / 100.0, 6)))
    two2 = [dict(x, label="spółki" if x["si"] == 1 else "QQQ", size=2.0 if x["si"] == 0 else 0.5) for x in two]
    pf = portfolio([dict(x) for x in two2], finish([st, st2]), lev=5.0, max_by_label={"spółki": 1})
    check("wielkość pozycji (size) i limit dla etykiety", (pf["trades"], round(pf["ret"], 6)),
          (2, round(2.0 * (10.0 - COST_PCT) + 0.5 * (-10.0 - COST_PCT), 6)))
