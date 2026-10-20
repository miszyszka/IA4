"""
IA 4 — FX-POWRÓT: short po gwałtownym wzroście, wejście po sygnale zawrócenia, TP na cenie sprzed wzrostu
(instrukcja, sekcja 7b). Wersja projektu: 1.53 (2026-10-20) — musi zgadzać się z IA4_INSTRUKCJA.md

Pomysł człowieka (20.10.2026): po wybiciu cena bardzo często wraca do ceny sprzed gwałtownych świec — TP ustawić na
tej cenie, a wejściem ma być nie sama gwałtowna świeca, tylko sygnał, że kurs zawraca. Liczymy wszystkie konfiguracje.

Zdarzenie (na każdej parze osobno, świece 1h):
  1. IMPULS — wzrost N świec ≥ k × √N × ATR(100) (zamknięcie do zamknięcia). Baza B = cena sprzed impulsu
     (zamknięcie świecy przed impulsem albo najniższy punkt impulsu), szczyt P = najwyższy punkt od początku impulsu
     (rośnie, póki trwa). Kolejne świece spełniające próg przedłużają ten sam impuls.
  2. ZAWRÓCENIE — pierwszy sygnał z listy SIGNALS w ciągu W świec od ostatniej świecy impulsu (albo „natychmiast”:
     wejście zaraz po świecy impulsu, dla porównania). Impuls wygasa bez transakcji, gdy zamknięcie spadło do bazy
     albo minęło W świec.
  3. SHORT po open następnej świecy. TP = P − f × (P − B): f = 1 → dokładnie cena sprzed wzrostu, 0,5 / 0,75 —
     część drogi, 1,25 — poniżej bazy, albo bez TP. SL nad szczytem (P + 0,1 / 0,5 / 1 ATR) albo od wejścia
     (1 / 2 / 3 ATR) albo bez. Limit 6–96 świec. TP, który wejście już minęło (< 0,1 ATR do TP) — konfiguracja pomija
     transakcję. Rozliczenie pesymistyczne (zasada 5), spread raz na transakcję (EUR/USD 1,1 pipsa), jedna pozycja na parę.
Filtry: trend (SMA500), godzina świecy sygnału, odległość wejścia od bazy (≥ 0,5 / 1 / 2 ATR).
Ocena jak w `ia4.fxshort`: 20 okresów, połowy, placebo (ten sam „nawias” TP / SL w ATR, ale w losowym czasie —
sprawdza, czy liczy się moment, a nie kształt wyjścia), wybór chronologiczny (1. rok → 2. rok), rachunek 1000 $ na 1:30.

    python -m ia4.fxpowrot              # → ~/Desktop/IA4/FX-POWROT.xlsx
    python -m ia4.fxpowrot --bez-sync   # bez odświeżenia kopii fx1h
    python -m ia4.fxpowrot --selftest   # „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import numba as nb
import numpy as np
import pandas as pd

from . import __version__, config
from . import fxshort as fs
from .fxshort import NPER, PAIRS, PIP

IMPULSE_N = [1, 2, 3, 4, 6, 8, 12]
IMPULSE_K = [1.0, 1.25, 1.5, 2.0]
BASES = ['zamknięcie przed wzrostem', 'najniższy punkt wzrostu']
SIGNALS = ['natychmiast (bez potwierdzenia)', 'pierwsza czerwona świeca', 'czerwona świeca ≥ 0,5 ATR', 'czerwona świeca ≥ 1 ATR',
           'zamknięcie pod dołkiem poprzedniej świecy', 'zamknięcie pod dołkiem świecy szczytu', 'objęcie bessy',
           'dwie czerwone z rzędu', 'cofnięcie ≥ 24% wzrostu', 'cofnięcie ≥ 38% wzrostu', 'górny knot ≥ 50% świecy',
           'zamknięcie pod EMA10']
WAITS = [3, 6, 12]
TP_F = [0.5, 0.75, 1.0, 1.25, 0.0]                 # 0 = bez TP
SL_DEF = [(1, 0.1), (1, 0.5), (1, 1.0), (2, 1.0), (2, 2.0), (2, 3.0), (0, 0.0)]   # 1 = nad szczytem + v ATR, 2 = wejście + v ATR, 0 = bez
HOLD = [6, 12, 24, 48, 96]
EXITS = [(f, s, h) for f in TP_F for s in SL_DEF for h in HOLD]
NE = len(EXITS)
EX_F = np.array([e[0] for e in EXITS]); EX_SK = np.array([e[1][0] for e in EXITS], np.int64)
EX_SV = np.array([e[1][1] for e in EXITS]); EX_H = np.array([e[2] for e in EXITS], np.int64)
TRENDS = ['bez filtra trendu', 'SMA500 spada', 'SMA500 nie rośnie mocno (< +1)']
HOURS = [('wszystkie godziny', 0, 23), ('7–16 UTC', 7, 16), ('11–16 UTC', 11, 16)]
DISTS = [0.5, 1.0, 2.0]


def _n(x):
    return f'{x:g}'.replace('.', ',')


def exit_name(e: int) -> str:
    f, (sk, sv), h = EXITS[e]
    tp = 'bez TP' if f == 0 else ('TP = cena sprzed wzrostu' if f == 1 else f'TP = cofnięcie {_n(f * 100)}% wzrostu')
    sl = 'bez SL' if sk == 0 else (f'SL szczyt + {_n(sv)} ATR' if sk == 1 else f'SL wejście + {_n(sv)} ATR')
    return f'{tp} · {sl} · {h} św.'


def impulse_name(N, k):
    return f'wzrost {N} św. ≥ {_n(round(k * np.sqrt(N), 1))} ATR'


# ---------------------------------------------------------------- dane

class Kurs:
    """Świece 10 par na wspólnej siatce czasu (bez macierzy wyjść z fxshort)."""

    def __init__(self, frames: dict, spread: dict):
        common = None
        for p in PAIRS:
            tt = set(frames[p].time.values.tolist())
            common = tt if common is None else common & tt
        tgrid = np.array(sorted(common), np.int64)
        self.F = {}
        for p in PAIRS:
            F = fs.features(frames[p].set_index('time').loc[tgrid].reset_index())
            c = F['c']; F['run1'] = (c - np.r_[np.nan, c[:-1]]) / F['atr']
            F['ema10'] = fs._ema(c, 10)
            self.F[p] = F
        self.t = tgrid; self.n = len(tgrid); self.spread = spread
        T0, T1 = int(tgrid[0]), int(tgrid[-1]) + 3600
        self.per = np.minimum(((tgrid - T0) / (T1 - T0) * NPER).astype(np.int64), NPER - 1)
        self.periods = [(pd.Timestamp(T0 + (T1 - T0) * k / NPER, unit='s'), pd.Timestamp(T0 + (T1 - T0) * (k + 1) / NPER, unit='s'))
                        for k in range(NPER)]


# ---------------------------------------------------------------- zdarzenia i wyniki

@nb.njit(cache=True)
def gen_trades(o, h, l, c, atr, run, thr, N, base_mode, sig, W, ema10, valid, start):
    """Impuls → (pierwszy) sygnał zawrócenia. Zwraca indeksy świec sygnału, bazę, szczyt, początek impulsu."""
    n = c.size
    J = np.empty(n, np.int64); B_ = np.empty(n); P_ = np.empty(n); S0 = np.empty(n, np.int64)
    k = 0; active = False; s0 = 0; last = 0; pk = 0; B = 0.0; P = 0.0
    for s in range(max(start, N + 1), n - 1):
        if atr[s] != atr[s] or run[s] != run[s]:
            continue
        trig = run[s] >= thr
        if not active:
            if trig:
                active = True; s0 = s; last = s
                if base_mode == 0:
                    B = c[s - N]
                else:
                    B = l[s - N + 1]
                    for q in range(s - N + 2, s + 1):
                        if l[q] < B: B = l[q]
                P = h[s - N + 1]; pk = s - N + 1
                for q in range(s - N + 2, s + 1):
                    if h[q] >= P: P = h[q]; pk = q
                if sig == 0:
                    if valid[s] and o[s + 1] > B:
                        J[k] = s; B_[k] = B; P_[k] = P; S0[k] = s0; k += 1
                    active = False
            continue
        if h[s] > P:
            P = h[s]; pk = s
        if trig:
            last = s
            continue
        if c[s] <= B or s - last > W:
            active = False
            continue
        a = atr[s]; red = c[s] < o[s]; fire = False
        if sig == 1:
            fire = red
        elif sig == 2:
            fire = o[s] - c[s] >= 0.5 * a
        elif sig == 3:
            fire = o[s] - c[s] >= 1.0 * a
        elif sig == 4:
            fire = c[s] < l[s - 1]
        elif sig == 5:
            fire = pk != s and c[s] < l[pk]
        elif sig == 6:
            fire = red and c[s - 1] > o[s - 1] and o[s] >= c[s - 1] and c[s] <= o[s - 1]
        elif sig == 7:
            fire = red and c[s - 1] < o[s - 1]
        elif sig == 8:
            fire = c[s] <= P - 0.236 * (P - B)
        elif sig == 9:
            fire = c[s] <= P - 0.382 * (P - B)
        elif sig == 10:
            rng = h[s] - l[s]
            fire = rng >= 0.5 * a and (h[s] - max(o[s], c[s])) >= 0.5 * rng
        elif sig == 11:
            fire = c[s] < ema10[s]
        if fire:
            if valid[s] and o[s + 1] > B:
                J[k] = s; B_[k] = B; P_[k] = P; S0[k] = s0; k += 1
            active = False
    return J[:k], B_[:k], P_[:k], S0[:k]


@nb.njit(cache=True)
def outcomes_tr(o, h, l, c, J, B, P, A, sp, exf, exsk, exsv, exh):
    """Short po open świecy J+1. TP = P − f(P − B); SL nad szczytem albo od wejścia. Wynik % netto, świece, MAE %."""
    n = o.size; nt = J.size; ne = exf.size
    pnl = np.full((nt, ne), np.nan, np.float32)
    xb = np.zeros((nt, ne), np.int16)
    mae = np.zeros((nt, ne), np.float32)
    for t in range(nt):
        e0 = J[t] + 1
        if e0 >= n:
            continue
        en = o[e0]; a = A[t]
        for e in range(ne):
            f = exf[e]
            if f == 0.0:
                TP = -1e18
            else:
                TP = P[t] - f * (P[t] - B[t])
                if TP > en - 0.1 * a:
                    continue
            sk = exsk[e]
            if sk == 0:
                SL = 1e18
            elif sk == 1:
                SL = max(P[t], en) + exsv[e] * a
            else:
                SL = en + exsv[e] * a
            H = exh[e]
            if e0 + H - 1 >= n:
                continue
            ex = np.nan; nb_ = 0; hi = en
            for j in range(e0, e0 + H):
                nb_ = j - e0 + 1
                if j > e0 and o[j] >= SL:
                    ex = o[j]; hi = max(hi, o[j]); break
                if j > e0 and o[j] <= TP:
                    ex = TP; break
                if h[j] > hi:
                    hi = min(h[j], SL)
                if h[j] >= SL:
                    ex = SL; break
                if l[j] <= TP:
                    ex = TP; break
                if j == e0 + H - 1:
                    ex = c[j]
            pnl[t, e] = (en - ex) / en * 100.0 - sp / en * 100.0
            xb[t, e] = nb_
            mae[t, e] = (hi - en) / en * 100.0 + sp / en * 100.0
    return pnl, xb, mae


@nb.njit(cache=True)
def evaluate_tr(pp, jj, PN, XB, per, nper, eu):
    """Jak fxshort.evaluate, ale wyniki na transakcję (PN[k, e]) zamiast na świecę. Kolejność pp/jj — czas."""
    ne = PN.shape[1]; npair = 10; half = nper // 2
    out = np.zeros((ne, 2, 22), np.float64)
    m = pp.size
    buf = np.empty(m, np.float64); bufe = np.empty(m, np.float64)
    last = np.empty(npair, np.int64)
    ps = np.zeros((2, nper))
    for e in range(ne):
        for p in range(npair):
            last[p] = -1
        ps[:, :] = 0.0
        cnt = np.zeros(2, np.int64); cum = np.zeros(2); peak = np.zeros(2); dd = np.zeros(2)
        gp = np.zeros(2); gl = np.zeros(2); hold = np.zeros(2)
        hn = np.zeros((2, 2)); hcum = np.zeros((2, 2)); hpk = np.zeros((2, 2)); hdd = np.zeros((2, 2))
        for k in range(m):
            p = pp[k]; i = jj[k]
            if i <= last[p]:
                continue
            v = PN[k, e]
            if v != v:
                continue
            x = XB[k, e]
            last[p] = i + x
            hf = 0 if per[i] < half else 1
            for g in range(2):
                if g == 1 and p != eu:
                    continue
                if g == 0:
                    buf[cnt[0]] = v
                else:
                    bufe[cnt[1]] = v
                cnt[g] += 1
                cum[g] += v
                if cum[g] > peak[g]: peak[g] = cum[g]
                if peak[g] - cum[g] > dd[g]: dd[g] = peak[g] - cum[g]
                if v > 0: gp[g] += v
                else: gl[g] -= v
                hold[g] += x
                ps[g, per[i]] += v
                hn[g, hf] += 1; hcum[g, hf] += v
                if hcum[g, hf] > hpk[g, hf]: hpk[g, hf] = hcum[g, hf]
                if hpk[g, hf] - hcum[g, hf] > hdd[g, hf]: hdd[g, hf] = hpk[g, hf] - hcum[g, hf]
        for g in range(2):
            n_ = cnt[g]
            out[e, g, 0] = n_
            if n_ == 0:
                continue
            b = buf[:n_] if g == 0 else bufe[:n_]
            s = b.sum(); mu = s / n_; sd = b.std()
            out[e, g, 1] = s; out[e, g, 2] = mu; out[e, g, 3] = np.median(b)
            out[e, g, 4] = mu / sd * np.sqrt(n_) if sd > 0 else 0.0
            out[e, g, 5] = (b > 0).sum() / n_
            out[e, g, 6] = gp[g] / gl[g] if gl[g] > 0 else 99.0
            out[e, g, 7] = dd[g]
            lose = 0; empty = 0
            for q in range(nper):
                if ps[g, q] < 0: lose += 1
                if ps[g, q] == 0: empty += 1
            out[e, g, 8] = lose; out[e, g, 9] = empty
            out[e, g, 10] = ps[g].mean(); out[e, g, 11] = np.median(ps[g]); out[e, g, 12] = ps[g].min()
            out[e, g, 13] = hold[g] / n_
            for hf in range(2):
                l2 = 0
                for q in range(half * hf, half * hf + half):
                    if ps[g, q] < 0: l2 += 1
                out[e, g, 14 + 4 * hf] = hn[g, hf]; out[e, g, 15 + 4 * hf] = hcum[g, hf]
                out[e, g, 16 + 4 * hf] = l2; out[e, g, 17 + 4 * hf] = hdd[g, hf]
    return out


FIELDS = fs.FIELDS


# ---------------------------------------------------------------- przeszukanie

def event_sets():
    """Wszystkie definicje zdarzeń: (N, k, baza, sygnał, okno)."""
    out = []
    for N in IMPULSE_N:
        for k in IMPULSE_K:
            for b in range(len(BASES)):
                for s in range(len(SIGNALS)):
                    for w in ([0] if s == 0 else WAITS):
                        out.append((N, k, b, s, w))
    return out


def ev_name(ev):
    N, k, b, s, w = ev
    return impulse_name(N, k), BASES[b], SIGNALS[s], (f'≤ {w} św.' if s else '—')


class Trades:
    """Transakcje jednej definicji zdarzenia na wszystkich parach (prawdziwe albo przesunięte w czasie — placebo)."""

    def __init__(self, K: Kurs, ev, shifts=None):
        N, k, b, s, w = ev
        rows = {f: [] for f in ['p', 'j', 'B', 'P', 'A', 'trend', 'slope', 'hour', 'dist', 's0']}
        PN, XB, MA = [], [], []
        for pi, p in enumerate(PAIRS):
            F = K.F[p]
            J, B, P, S0 = gen_trades(F['o'], F['h'], F['l'], F['c'], F['atr'], F[f'run{N}'], k * np.sqrt(N), N, b, s, w,
                                     F['ema10'], F['valid'], fs.WARMUP)
            if J.size == 0:
                continue
            A = F['atr'][J]; en = F['o'][J + 1]
            slope = np.nan_to_num(F['slope500'][J], nan=9.0); hour = F['hour'][J]; dist = (en - B) / A
            if shifts is not None:                                  # placebo: ten sam nawias w losowym czasie
                brel = (en - B) / A; prel = (P - en) / A
                J2 = (J + shifts[p] - WARM) % (K.n - WARM - 2) + WARM
                ok = F['valid'][J2] & ~np.isnan(F['atr'][J2])
                J, brel, prel, slope, hour, dist, S0 = J2[ok], brel[ok], prel[ok], slope[ok], hour[ok], dist[ok], S0[ok]
                A = F['atr'][J]; en = F['o'][J + 1]; B = en - brel * A; P = en + prel * A
            pn, x, ma = outcomes_tr(F['o'], F['h'], F['l'], F['c'], J, B, P, A, K.spread[p] * PIP[p], EX_F, EX_SK, EX_SV, EX_H)
            PN.append(pn); XB.append(x); MA.append(ma)
            for f_, v in [('p', np.full(J.size, pi)), ('j', J), ('B', B), ('P', P), ('A', A), ('slope', slope), ('hour', hour),
                          ('dist', dist), ('s0', S0)]:
                rows[f_].append(v)
        if PN:
            cat = {f_: np.concatenate(v) for f_, v in rows.items() if v}
            o = np.lexsort((cat['p'], cat['j']))
            self.d = {f_: v[o] for f_, v in cat.items()}
            self.PN = np.concatenate(PN)[o]; self.XB = np.concatenate(XB)[o]; self.MA = np.concatenate(MA)[o]
        else:
            self.d = None

    def masks(self):
        d = self.d
        for ti, tn in enumerate(TRENDS):
            mt = np.ones(d['j'].size, bool) if ti == 0 else (d['slope'] < 0 if ti == 1 else d['slope'] < 1)
            for hn, a, b in HOURS:
                mh = (d['hour'] >= a) & (d['hour'] <= b)
                for dv in DISTS:
                    yield tn, hn, dv, mt & mh & (d['dist'] >= dv)


WARM = fs.WARMUP


def scan(K: Kurs, seed=None, collect=None, log=print):
    rng = np.random.default_rng(seed) if seed is not None else None
    shifts = {p: int(rng.integers(1500, K.n - 1500)) for p in PAIRS} if rng is not None else None
    evs = event_sets(); t0 = time.time()
    for q, ev in enumerate(evs):
        T = Trades(K, ev, shifts)
        if T.d is not None:
            for tn, hn, dv, m in T.masks():
                if m.sum() < 20:
                    continue
                res = evaluate_tr(T.d['p'][m].astype(np.int64), T.d['j'][m].astype(np.int64), T.PN[m], T.XB[m], K.per, NPER, PAIRS.index('EURUSD'))
                collect(ev, tn, hn, dv, res)
        if log and (q % 200 == 199 or q == len(evs) - 1):
            log(f'    {q + 1}/{len(evs)} definicji zdarzeń, {time.time() - t0:.0f} s')


def rule_trades(K: Kurs, ev, tn, hn, dv, e) -> pd.DataFrame:
    """Lista transakcji jednej reguły (jedna pozycja na parę)."""
    T = Trades(K, ev)
    m = next(mm for a, b, c, mm in T.masks() if (a, b, c) == (tn, hn, dv))
    d = {f_: v[m] for f_, v in T.d.items()}; PN = T.PN[m][:, e]; XB = T.XB[m][:, e]; MA = T.MA[m][:, e]
    f, (sk, sv), H = EXITS[e]
    rows = []; last = {}
    for k in range(d['j'].size):
        p = PAIRS[int(d['p'][k])]; j = int(d['j'][k])
        if j <= last.get(p, -1) or np.isnan(PN[k]):
            continue
        x = int(XB[k]); last[p] = j + x
        F = K.F[p]; en = float(F['o'][j + 1]); A = float(d['A'][k]); B = float(d['B'][k]); P = float(d['P'][k])
        TP = P - f * (P - B) if f else np.nan
        SL = (max(P, en) + sv * A) if sk == 1 else ((en + sv * A) if sk == 2 else np.nan)
        sp = K.spread[p] * PIP[p]; ex = en - float(PN[k]) * en / 100.0 - sp
        if not np.isnan(SL) and ex >= SL - 0.05 * PIP[p]:
            why = 'SL' if abs(ex - SL) < 0.05 * PIP[p] else 'SL (luka)'
        elif not np.isnan(TP) and abs(ex - TP) < 0.05 * PIP[p]:
            why = 'TP'
        else:
            why = 'czas (zamknięcie świecy)'
        rows.append(dict(para=p, i=j, impuls_t=int(F['t'][int(d['s0'][k])]), sygnal_t=int(F['t'][j]), wejscie_t=int(F['t'][j + 1]),
                         wyjscie_t=int(F['t'][j + x]) + 3600, baza=B, szczyt=P, wejscie=en, atr=A, SL=SL, TP=TP, wyjscie=ex, powod=why,
                         swiece=x, wynik=float(PN[k]), pipsy=(en - ex - sp) / PIP[p], mae=float(MA[k]), okres=int(K.per[j]) + 1,
                         cap=((SL - en + sp) / en) if not np.isnan(SL) else 9e9))
    df = pd.DataFrame(rows)
    for c_ in ['impuls_t', 'sygnal_t', 'wejscie_t', 'wyjscie_t']:
        if len(df):
            df[c_] = pd.to_datetime(df[c_], unit='s')
    return df


# ---------------------------------------------------------------- statystyka powrotów i pojedyncze reguły

@nb.njit(cache=True)
def race(h, l, J, B, P, A, Hmax, m):
    """Dla każdego zdarzenia: po ilu świecach (od wejścia) cena dotknęła bazy i po ilu przebiła szczyt o m ATR (9999 = nigdy)."""
    n = h.size; rb = np.full(J.size, 9999); rp = np.full(J.size, 9999)
    for t in range(J.size):
        for j in range(J[t] + 1, min(n, J[t] + 1 + Hmax)):
            if rp[t] == 9999 and h[j] >= P[t] + m * A[t]:
                rp[t] = j - J[t]
            if rb[t] == 9999 and l[j] <= B[t]:
                rb[t] = j - J[t]
    return rb, rp


def return_stats(K: Kurs) -> pd.DataFrame:
    rows = []
    for N in IMPULSE_N:
        for k in IMPULSE_K:
            for sig, w in [(0, 0), (1, 3), (9, 6)]:
                rb_all, rp_all = [], []
                for p in PAIRS:
                    F = K.F[p]
                    J, B, P, S0 = gen_trades(F['o'], F['h'], F['l'], F['c'], F['atr'], F[f'run{N}'], k * np.sqrt(N), N, 0, sig, w,
                                             F['ema10'], F['valid'], WARM)
                    rb, rp = race(F['h'], F['l'], J, B, P, F['atr'][J], 96, 0.5)
                    rb_all.append(rb); rp_all.append(rp)
                rb = np.concatenate(rb_all); rp = np.concatenate(rp_all)
                if rb.size == 0:
                    continue
                rows.append(dict(impuls=impulse_name(N, k), N=N, k=k, sygnal=SIGNALS[sig], n=rb.size,
                                 baza_6=(rb <= 6).mean(), baza_24=(rb <= 24).mean(), baza_96=(rb <= 96).mean(),
                                 przed_24=((rb <= 24) & (rb < rp)).mean(), przed_96=((rb <= 96) & (rb < rp)).mean(),
                                 szczyt_24=(rp <= 24).mean()))
    return pd.DataFrame(rows)


def rule_stats(K: Kurs, ev, tn, hn, dv, f, sk, sv, H):
    """Statystyki jednej reguły z dowolnym wyjściem (poza siatką EXITS)."""
    T = Trades(K, ev)
    if T.d is None:
        return None
    m = next(mm for a, b, c, mm in T.masks() if (a, b, c) == (tn, hn, dv))
    d = {k_: v[m] for k_, v in T.d.items()}
    PN = np.full(m.sum(), np.nan, np.float32); XB = np.zeros(m.sum(), np.int16)
    for pi, p in enumerate(PAIRS):
        q = d['p'] == pi
        if not q.any():
            continue
        F = K.F[p]
        pn, x, _ = outcomes_tr(F['o'], F['h'], F['l'], F['c'], d['j'][q], d['B'][q], d['P'][q], d['A'][q], K.spread[p] * PIP[p],
                               np.array([float(f)]), np.array([sk], np.int64), np.array([float(sv)]), np.array([H], np.int64))
        PN[q] = pn[:, 0]; XB[q] = x[:, 0]
    return evaluate_tr(d['p'].astype(np.int64), d['j'].astype(np.int64), PN[:, None], XB[:, None], K.per, NPER, PAIRS.index('EURUSD'))[0]


MAIN = dict(N=1, k=2.0, baza=0, sygnal=1, okno=3, trend='bez filtra trendu', godziny='7–16 UTC', odl=0.5, f=0.5, sl=(0, 0.0), H=96)


def main_ev():
    return (MAIN['N'], MAIN['k'], MAIN['baza'], MAIN['sygnal'], MAIN['okno'])


def main_exit():
    return EXITS.index((MAIN['f'], MAIN['sl'], MAIN['H']))


def compute(K: Kurs, f=0.20, lev=30.0, start=1000.0, log=print) -> dict:
    O = {'MAIN': MAIN, 'periods': K.periods, 'spread': K.spread, 'f': f, 'lev': lev, 'start': start, 'swiec': K.n,
         'od': pd.Timestamp(int(K.t[0]), unit='s'), 'do': pd.Timestamp(int(K.t[-1]), unit='s')}
    cols = ['N', 'k', 'baza', 'sygnal', 'okno', 'trend', 'godziny', 'odl', 'e'] + FIELDS[:14] + ['suma1', 'suma2', 'eu_n', 'eu_suma']
    O['placebo'] = {}; O['wf'] = {}; nr = [0]
    for tag, seed in [('prawdziwe', None), ('placebo 1', 101), ('placebo 2', 202), ('placebo 3', 303)]:
        log(f'  przeszukanie — {tag}')
        cnt = dict(all200=0, all300=0, halves=0, eu40=0, eu60=0, max_t=0.0, t5=0, max_zo=0.0)
        wf = {0: [], 1: []}; keep = []; dims = O.setdefault('dims_acc', {}) if tag == 'prawdziwe' else {}

        def col(ev, tn, hn, dv, res, cnt=cnt, wf=wf, keep=keep, tag=tag, dims=dims):
            A, E = res[:, 0, :], res[:, 1, :]
            ok = A[:, 0] >= 200
            m3 = ok & (A[:, 8] <= 3) & (A[:, 1] > 0)
            cnt['all200'] += int(m3.sum()); cnt['all300'] += int((m3 & (A[:, 0] >= 300)).sum())
            cnt['halves'] += int(((A[:, 14] >= 100) & (A[:, 18] >= 100) & (A[:, 16] <= 2) & (A[:, 20] <= 2) & (A[:, 15] > 0) & (A[:, 19] > 0)).sum())
            cnt['eu40'] += int(((E[:, 0] >= 40) & (E[:, 8] <= 3) & (E[:, 1] > 0)).sum())
            cnt['eu60'] += int(((E[:, 0] >= 60) & (E[:, 8] <= 3) & (E[:, 1] > 0)).sum())
            if ok.any():
                cnt['max_t'] = max(cnt['max_t'], float(A[ok, 4].max())); cnt['t5'] += int((A[ok, 4] >= 5).sum())
            if m3.any():
                cnt['max_zo'] = max(cnt['max_zo'], float((A[m3, 1] / np.maximum(A[m3, 7], 0.05)).max()))
            for fold, (a, b) in enumerate([(14, 18), (18, 14)]):
                mm = (A[:, a] >= 100) & (A[:, a + 2] <= 1) & (A[:, a + 1] > 0)
                for e in np.nonzero(mm)[0]:
                    wf[fold].append((A[e, a + 1] / max(A[e, a + 3], 0.05), A[e, b], A[e, b + 1], ev, tn, hn, dv, e))
            if tag == 'prawdziwe':
                nr[0] += int(np.isfinite(A[:, 1]).sum())
                for e in np.nonzero(m3 | (ok & (EX_F == 1.0) & (A[:, 1] > 0)))[0]:
                    keep.append((*ev, tn, hn, dv, e, *A[e, :14], A[e, 15], A[e, 19], E[e, 0], E[e, 1]))
                if ok.any():
                    tt = A[:, 4]; ss = A[:, 1]
                    for dim, vals in [('sygnał', None), ('impuls N', None), ('godziny', None), ('trend', None), ('TP', EX_F), ('SL', EX_SK * 100 + EX_SV), ('limit świec', EX_H)]:
                        if vals is None:
                            v0 = {'sygnał': SIGNALS[ev[3]], 'impuls N': ev[0], 'godziny': hn, 'trend': tn}[dim]
                            groups = [(v0, ok)]
                        else:
                            groups = [(v, ok & (vals == v)) for v in np.unique(vals[ok])]
                        for v, mk in groups:
                            if not mk.any():
                                continue
                            a = dims.setdefault((dim, v), [0, 0, -99.0, -99.0])
                            a[0] += int(mk.sum()); a[1] += int((mk & m3).sum()); a[2] = max(a[2], float(tt[mk].max())); a[3] = max(a[3], float(ss[mk].max()))
        scan(K, seed, col, log=None)
        O['placebo'][tag] = cnt
        w = {}
        for fold in [0, 1]:
            L = sorted(wf[fold], key=lambda x: -x[0])
            if L:
                s2 = np.array([x[2] for x in L]); n2 = np.array([x[1] for x in L])
                w[fold] = dict(spelnia=len(L), top1=s2[0], top10=s2[:10].mean(), top10_poz=(s2[:10] > 0).mean(),
                               top50_tr=s2[:50].sum() / max(n2[:50].sum(), 1), top=L[0][3:])
        O['wf'][tag] = w
        if tag == 'prawdziwe':
            r = pd.DataFrame(keep, columns=cols)
            r['zysk/obs'] = r.suma / r.obsuniecie.clip(lower=0.05)
            r['impuls'] = [impulse_name(a, b) for a, b in zip(r.N, r.k)]; r['sygnał'] = [SIGNALS[s] for s in r.sygnal]
            r['wyjście'] = [exit_name(int(e)) for e in r.e]
            ok = (r.n >= 200) & (r.stratne_okresy <= 3) & (r.suma > 0)
            O['rank'] = r[ok].sort_values('zysk/obs', ascending=False).head(40)
            O['rank_t'] = r[ok].sort_values('t', ascending=False).head(40)
            O['rank_base'] = r[r.e.map(lambda e: EXITS[int(e)][0] == 1.0)].sort_values('t', ascending=False).head(25)
            def lab(dim, v):
                if dim == 'TP':
                    return 'bez TP' if v == 0 else ('cena sprzed wzrostu' if v == 1 else f'cofnięcie {_n(v * 100)}% wzrostu')
                if dim == 'SL':
                    k_, x = int(v // 100), v - 100 * int(v // 100)
                    return 'bez SL' if k_ == 0 else (f'szczyt + {_n(round(x, 2))} ATR' if k_ == 1 else f'wejście + {_n(round(x, 2))} ATR')
                return str(v)
            O['dims'] = pd.DataFrame([dict(wymiar=k_[0], wartosc=lab(*k_), reguly=v[0], warunek=v[1], max_t=v[2], max_suma=v[3])
                                      for k_, v in sorted(dims.items(), key=lambda kv: (kv[0][0], str(kv[0][1])))])
    O['n_rules'] = nr[0]
    log('  statystyka powrotów, sąsiednie ustawienia, strategia')
    O['returns'] = return_stats(K)
    M = MAIN; ev = main_ev()
    nb_ = {}

    def st(lab, **kw):
        a = dict(ev=ev, tn=M['trend'], hn=M['godziny'], dv=M['odl'], f=M['f'], sk=M['sl'][0], sv=M['sl'][1], H=M['H']); a.update(kw)
        r_ = rule_stats(K, a['ev'], a['tn'], a['hn'], a['dv'], a['f'], a['sk'], a['sv'], a['H'])
        return dict(ustawienie=lab, **({FIELDS[i]: r_[0, i] for i in range(22)} if r_ is not None else {}),
                    eu_n=r_[1, 0] if r_ is not None else 0, eu_suma=r_[1, 1] if r_ is not None else 0)
    nb_['A. Wielkość świecy impulsu i okno na czerwoną świecę'] = [
        st(f'świeca ≥ {_n(k)} ATR, czerwona w ≤ {w} św.', ev=(1, k, 0, 1, w)) for k in [1.5, 1.75, 2.0, 2.25, 2.5, 3.0] for w in [1, 2, 3]]
    nb_['B. Sygnał zawrócenia (świeca ≥ 2 ATR)'] = [st(SIGNALS[s], ev=(1, 2.0, 0, s, 0 if s == 0 else 3)) for s in range(len(SIGNALS))]
    nb_['C. Impuls z kilku świec (pierwsza czerwona ≤ 3 św.)'] = [st(impulse_name(N, k), ev=(N, k, 0, 1, 3)) for N in [2, 3, 4, 6] for k in [1.25, 1.5, 2.0]]
    nb_['D. Take profit'] = [st(lab, f=fv) for fv, lab in [(0.5, 'cofnięcie 50% wzrostu'), (0.75, 'cofnięcie 75%'),
                                                          (1.0, 'cena sprzed wzrostu (pomysł człowieka)'), (1.25, '25% poniżej bazy'), (0.0, 'bez TP')]]
    nb_['E. Stop loss'] = [st(lab, sk=a, sv=b) for (a, b), lab in [((0, 0.0), 'bez SL'), ((1, 0.5), 'szczyt + 0,5 ATR'), ((1, 1.0), 'szczyt + 1 ATR'),
                                                                  ((1, 2.0), 'szczyt + 2 ATR'), ((1, 4.0), 'szczyt + 4 ATR'), ((2, 2.0), 'wejście + 2 ATR'),
                                                                  ((2, 3.0), 'wejście + 3 ATR'), ((2, 4.0), 'wejście + 4 ATR'), ((2, 6.0), 'wejście + 6 ATR'),
                                                                  ((2, 8.0), 'wejście + 8 ATR')]]
    nb_['F. Limit świec'] = [st(f'{h} świec', H=h) for h in [6, 12, 24, 48, 96]]
    nb_['G. Godzina świecy sygnału i trend'] = [st(f'{hn} · {tn}', hn=hn, tn=tn) for hn, _, _ in HOURS for tn in TRENDS]
    O['neighbors'] = {k_: pd.DataFrame(v) for k_, v in nb_.items()}
    main = rule_trades(K, ev, M['trend'], M['godziny'], M['odl'], main_exit())
    main['noce'] = [fs.rollovers(a, b) for a, b in zip(main.wejscie_t, main.wyjscie_t)]
    main['bazowa_usd'] = [1.0 if fs.BASE_USD[p] is None else float(K.F[fs.BASE_USD[p]]['o'][i + 1]) for p, i in zip(main.para, main.i)]
    O['main'] = main; O['main_sum'] = fs.summary(main)
    w_ = main.wynik.values
    O['main_t'] = w_.mean() / w_.std() * np.sqrt(len(w_)); O['main_ci'] = 1.96 * w_.std() / np.sqrt(len(w_))
    var = []; curves = {}
    for fv in sorted({0.05, 0.10, 0.20, 1 / 3, 0.5, 1.0, f}):
        for sw in [0.0, 3.0]:
            L, cv, evs = fs.simulate(main, K, lev=lev, f=fv, start=start, swap_pa=sw)
            var.append(dict(f=fv, swap=sw, **fs.curve_stats(cv, start), zawarte=int((L.status == 'zamknięta').sum()) if len(L) else 0,
                            pominiete=int((L.status == 'pominięta').sum()) if len(L) else 0, stop_out=len(evs)))
            if sw == 0.0:
                curves[fv] = fs.daily(cv)
    O['variants'] = pd.DataFrame(var); O['curves'] = curves
    eu = main[main.para == 'EURUSD'].reset_index(drop=True)
    O['eu_sum'] = fs.summary(eu); ev_ = []; O['eu_curves'] = {}
    for fv in [0.2, 1 / 3, 1.0]:
        L, cv, evs = fs.simulate(eu, K, lev=lev, f=fv, start=start)
        ev_.append(dict(f=fv, **fs.curve_stats(cv, start), zawarte=int((L.status == 'zamknięta').sum()) if len(L) else 0, stop_out=len(evs)))
        O['eu_curves'][fv] = fs.daily(cv)
    O['eu_var'] = pd.DataFrame(ev_)
    y2 = K.periods[NPER // 2][0]
    top = O['wf']['prawdziwe'].get(0, {}).get('top')
    if top:
        tev, ttn, thn, tdv, te = top
        wtr = rule_trades(K, tev, ttn, thn, tdv, int(te)); wtr = wtr[wtr.okres > NPER // 2].reset_index(drop=True)
        O['wf_rule'] = dict(opis=f"{' · '.join(ev_name(tev)[:3])} · {ttn} · {thn} · odl. ≥ {_n(tdv)} ATR · {exit_name(int(te))}")
        O['wf_sum'] = fs.summary(wtr)
        _, cv, _ = fs.simulate(wtr, K, lev=lev, f=0.20, start=start)
        O['wf_stats'] = fs.curve_stats(cv, start, since=y2); O['wf_curve'] = fs.daily(cv[cv.index >= y2])
    m2 = main[main.okres > NPER // 2].reset_index(drop=True)
    _, cv, _ = fs.simulate(m2, K, lev=lev, f=0.20, start=start)
    O['main_y2'] = fs.curve_stats(cv, start, since=y2); O['main_y2_curve'] = fs.daily(cv[cv.index >= y2]); O['main_y2_tr'] = m2
    return O


# ---------------------------------------------------------------- Excel

def write_xlsx(O: dict, path: Path, wnioski=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    M = O['MAIN']; FN = 'Arial'
    fN = Font(name=FN, size=10); fB = Font(name=FN, size=10, bold=True); fT = Font(name=FN, size=14, bold=True)
    fH = Font(name=FN, size=10, bold=True, color='FFFFFF'); fIn = Font(name=FN, size=10, color='0000FF')
    fS = Font(name=FN, size=9, italic=True, color='555555')
    HDR = PatternFill('solid', fgColor='1F3A5F'); YEL = PatternFill('solid', fgColor='FFFF00'); GRY = PatternFill('solid', fgColor='EEF2F7')
    GOOD = PatternFill('solid', fgColor='E2F0D9'); BAD = PatternFill('solid', fgColor='FBE5D6'); SEL = PatternFill('solid', fgColor='92D050')
    thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')
    PCT = '0.0%'; PCT3 = '0.000%'; USD = '#,##0.00 [$$-409];[Red]-#,##0.00 [$$-409]'; USD0 = '#,##0 [$$-409]'; DT = 'yyyy-mm-dd hh:mm'; DD = 'yyyy-mm-dd'

    def pl(x, d=1):
        return f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')

    wb = Workbook()

    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name; ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def hdr(ws, r, c, names, widths=None):
        for k, nm in enumerate(names):
            x = ws.cell(r, c + k, nm); x.font = fH; x.fill = HDR; x.border = BOX
            x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for k, w in enumerate(widths or []):
            ws.column_dimensions[CL(c + k)].width = w
        ws.row_dimensions[r].height = 42

    def put(ws, r, c, v, font=fN, fmt=None, fill=None, al=None):
        x = ws.cell(r, c, v); x.font = font; x.border = BOX
        if fmt: x.number_format = fmt
        if fill: x.fill = fill
        if al: x.alignment = al
        return x

    def note(ws, r, c, text, span=10, h=None):
        x = ws.cell(r, c, text); x.font = fS; x.alignment = WR
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        if h: ws.row_dimensions[r].height = h

    def lchart(title, h, w):
        ch = LineChart(); ch.title = title; ch.height = h; ch.width = w
        ch.y_axis.crossAx = 500; ch.x_axis = DateAxis(crossAx=100); ch.x_axis.number_format = 'yyyy-mm'; ch.x_axis.majorTimeUnit = 'months'
        ch.x_axis.delete = False; ch.y_axis.delete = False; ch.display_blanks = 'span'
        return ch

    def py(ts):
        return pd.Timestamp(ts).to_pydatetime()

    def stats_row(ws, r, c, d, eu=True):
        vals = [(int(d['n']), '0'), (d['suma'] / 100, PCT3), (d['srednia'] / 100, PCT3), (d['trafne'], PCT), (min(d['PF'], 99), '0.00'),
                (d['t'], '0.0'), (d['obsuniecie'] / 100, PCT3), (int(d['stratne_okresy']), '0'), (d['suma1'] / 100, PCT3), (d['suma2'] / 100, PCT3)]
        if eu:
            vals += [(int(d['eu_n']), '0'), (d['eu_suma'] / 100, PCT3)]
        for j, (v, fm) in enumerate(vals):
            x = put(ws, r, c + j, v, fmt=fm)
        ws.cell(r, c + 1).fill = GOOD if d['suma'] > 0 else BAD
    STATS_H = ['Transakcje', 'Wynik % (1:1, suma)', 'Średnia', 'Trafne', 'PF', 't', 'Obsunięcie (suma %)', 'Okresy stratne / 20', '1. połowa', '2. połowa',
               'EUR/USD: transakcje', 'EUR/USD: wynik %']

    # TRANSAKCJE
    tr = O['main']; N = len(tr); R0 = 11; R1 = R0 + max(N, 1) - 1
    wsP = sheet('PODSUMOWANIE', first=True); ws = sheet('TRANSAKCJE')
    ws['A1'] = 'TRANSAKCJE — FX-POWRÓT, 10 par, świece 1h (rozliczenie pesymistyczne, spread XTB)'; ws['A1'].font = fT
    inp = [('Kapitał początkowy ($)', O['start'], USD0, 'Założenie człowieka: start 1000 $.'),
           ('Dźwignia (1:x)', O['lev'], '0', 'Założenie człowieka: dźwignia 1:30. Depozyt pozycji = wartość / dźwignia.'),
           ('Depozyt na pozycję (% salda)', O['f'], PCT, '20% → wartość pozycji = 6 × saldo; 5 pozycji naraz = cała dźwignia 1:30. Strategia nie ma SL — przy większym depozycie rośnie ryzyko stop outu (arkusz WARIANTY).'),
           ('Koszt nocny (% rocznie wartości pozycji)', 0.0, '0.0%', 'Swap CFD — szacunek, nie dane XTB. Pozycje trwają do 4 dni, więc koszt nocny ma tu większe znaczenie niż w FX-SHORT-NY. Środa ×3.'),
           ('Jednostki w 0,01 lota', 1000, '#,##0', 'Minimalna pozycja XTB: 0,01 lota = 1000 jednostek waluty bazowej.')]
    for k, (lab, v, fm, com) in enumerate(inp):
        put(ws, 3 + k, 1, lab, fB, fill=GRY); x = put(ws, 3 + k, 2, v, fIn, fm, YEL); x.comment = Comment(com, 'IA4')
    ws.column_dimensions['A'].width = 34
    note(ws, 8, 1, 'Żółte komórki można zmieniać — saldo, wykresy w KAPITAŁ i okresy przeliczą się same. Wynik % transakcji (kolumna O) z symulacji Python (ia4.fxpowrot): baza = cena sprzed gwałtownej świecy, szczyt = najwyższy punkt do sygnału, wejście po open świecy po sygnale zawrócenia; SL i TP w tej samej świecy → SL, spread odjęty raz.', 16, 40)
    cols = ['Nr', 'Para', 'Impuls — świeca (UTC)', 'Sygnał zawrócenia — świeca (UTC)', 'Wejście (UTC, open)', 'Wyjście (UTC)', 'Świec w pozycji',
            'Baza (cena sprzed wzrostu)', 'Szczyt', 'Cena wejścia', 'SL (cena)', 'TP (cena)', 'Cena wyjścia', 'Powód wyjścia', 'Wynik % (netto)',
            'Wynik w pipsach (netto)', 'Największy ruch przeciw (MAE %)', 'Okres (1–20)', 'Noce (rolowania, środa ×3)', 'Kurs waluty bazowej w USD',
            'Saldo przy wejściu ($)', 'Depozyt zajęty przez otwarte pozycje ($)', 'Jednostki', 'Loty', 'Wartość pozycji ($)', 'Depozyt pozycji ($)',
            'Wynik ($)', 'Saldo po wyjściach do tej chwili ($)', 'Status']
    hdr(ws, 10, 1, cols, [5, 8, 16, 16, 16, 16, 7, 10, 10, 10, 10, 10, 10, 13, 10, 9, 10, 7, 9, 10, 12, 13, 10, 7, 12, 11, 11, 13, 10])
    for k, r in tr.reset_index(drop=True).iterrows():
        row = R0 + k; pf = '0.000' if r.para.endswith('JPY') else '0.00000'
        vals = [(k + 1, None), (r.para, None), (py(r.impuls_t), DT), (py(r.sygnal_t), DT), (py(r.wejscie_t), DT), (py(r.wyjscie_t), DT), (int(r.swiece), None),
                (round(float(r.baza), 6), pf), (round(float(r.szczyt), 6), pf), (round(float(r.wejscie), 6), pf),
                (None if pd.isna(r.SL) else round(float(r.SL), 6), pf), (None if pd.isna(r.TP) else round(float(r.TP), 6), pf),
                (round(float(r.wyjscie), 6), pf), (r.powod, None), (float(r.wynik) / 100, PCT3), (round(float(r.pipsy), 1), '0.0'),
                (float(r.mae) / 100, PCT3), (int(r.okres), None), (int(r.noce), None), (round(float(r.bazowa_usd), 6), '0.0000')]
        for c, (v, fm) in enumerate(vals):
            put(ws, row, 1 + c, v, fmt=fm)
        put(ws, row, 21, f'=$B$3+SUMIFS(AA$10:AA{row-1},F$10:F{row-1},"<="&E{row})', fmt=USD)
        put(ws, row, 22, f'=SUMIFS(Z$10:Z{row-1},E$10:E{row-1},"<="&E{row},F$10:F{row-1},">"&E{row})', fmt=USD)
        put(ws, row, 23, f'=IF(MIN($B$5*U{row},U{row}-V{row})<=0,0,INT(MIN($B$5*U{row},U{row}-V{row})*$B$4/T{row}/$B$7)*$B$7)', fmt='#,##0')
        put(ws, row, 24, f'=W{row}/100000', fmt='0.00')
        put(ws, row, 25, f'=W{row}*T{row}', fmt=USD)
        put(ws, row, 26, f'=Y{row}/$B$4', fmt=USD)
        put(ws, row, 27, f'=Y{row}*O{row}-Y{row}*$B$6/365*S{row}', fmt=USD)
        put(ws, row, 28, f'=$B$3+SUMIFS(AA${R0}:AA${R1},F${R0}:F${R1},"<="&F{row})', fmt=USD)
        put(ws, row, 29, f'=IF(W{row}=0,"pominięta","zawarta")')
    ws.freeze_panes = 'C11'; ws.auto_filter.ref = f'A10:AC{R1}'

    def TR(c):
        return f'TRANSAKCJE!${c}${R0}:${c}${R1}'

    # KAPITAŁ
    ws = sheet('KAPITAŁ')
    ws['A1'] = 'KAPITAŁ — saldo dzień po dniu (ustawienia z TRANSAKCJE!B3:B7)'; ws['A1'].font = fT
    note(ws, 2, 1, f"Saldo = kapitał początkowy + wynik transakcji zamkniętych do końca dnia (formuła). Kolumna F — wycena z otwartymi pozycjami co godzinę z symulacji Python dla ustawień domyślnych ({pl(O['f'] * 100, 0)}% depozytu, bez kosztu nocnego). Strategia trzyma pozycję do 4 dni bez SL — obsunięcie z wyceny godzinowej jest wyraźnie głębsze niż z salda.", 8, 42)
    hdr(ws, 4, 1, ['Dzień (koniec, UTC)', 'Saldo ($)', 'Najwyższe saldo ($)', 'Obsunięcie od szczytu', 'Liczba transakcji zamkniętych', 'Wycena godzinowa — Python, ustawienia domyślne ($)'], [13, 13, 13, 12, 12, 16])
    cur = O['curves'][O['f']]
    days = pd.date_range(cur.index.min().normalize(), cur.index.max().normalize(), freq='D')
    for k, d in enumerate(days):
        row = 5 + k
        put(ws, row, 1, py(d), fmt=DD)
        put(ws, row, 2, f'=TRANSAKCJE!$B$3+SUMIFS({TR("AA")},{TR("F")},"<"&(A{row}+1))', fmt=USD)
        put(ws, row, 3, f'=MAX(B$5:B{row})', fmt=USD); put(ws, row, 4, f'=B{row}/C{row}-1', fmt=PCT)
        put(ws, row, 5, f'=COUNTIFS({TR("F")},"<"&(A{row}+1),{TR("AC")},"zawarta")', fmt='0')
        v = cur.get(d, np.nan); put(ws, row, 6, None if np.isnan(v) else round(float(v), 2), fmt=USD)
    KR1 = 5 + len(days) - 1
    ch = lchart(f"Kapitał: {pl(O['start'], 0)} $ na dźwigni 1:{O['lev']:g} (saldo, ustawienia z TRANSAKCJE)", 9, 24)
    ch.y_axis.title = '$'; ch.y_axis.scaling.min = 0.8 * O['start']
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=KR1), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=6, min_row=4, max_row=KR1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=KR1))
    ch.series[0].graphicalProperties.line.width = 22000; ch.series[0].graphicalProperties.line.solidFill = '1F3A5F'
    ch.series[1].graphicalProperties.line.width = 9000; ch.series[1].graphicalProperties.line.solidFill = 'A6A6A6'
    ch.legend.position = 'b'; ws.add_chart(ch, 'H4')
    ch2 = lchart('Obsunięcie salda od szczytu', 6, 24)
    ch2.add_data(Reference(ws, min_col=4, min_row=4, max_row=KR1), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=5, max_row=KR1)); ch2.y_axis.number_format = '0%'; ch2.x_axis.tickLblPos = 'low'
    ch2.series[0].graphicalProperties.line.solidFill = 'C00000'; ch2.legend = None; ws.add_chart(ch2, 'H23')
    ws.freeze_panes = 'A5'

    # OKRESY
    ws = sheet('OKRESY')
    ws['A1'] = 'OKRESY — dostępne świece podzielone na 20 równych części'; ws['A1'].font = fT
    note(ws, 2, 1, 'Transakcja należy do okresu świecy sygnału. Wynik % = suma wyników transakcji bez dźwigni. Zmiana kapitału — saldo z TRANSAKCJE na początku i końcu okresu. Okres 1 prawie pusty: pierwsze 600 świec to rozbieg średnich.', 11, 42)
    hdr(ws, 4, 1, ['Okres', 'Od (UTC)', 'Do (UTC)', 'Transakcje', 'Wynik % (1:1, suma)', 'Średnia na transakcję', 'Trafne', 'Wynik ($)', 'Saldo na początku ($)', 'Saldo na końcu ($)', 'Zmiana kapitału'],
        [7, 16, 16, 10, 11, 11, 9, 12, 13, 13, 11])
    for k, (a, b) in enumerate(O['periods']):
        row = 5 + k
        put(ws, row, 1, k + 1); put(ws, row, 2, py(a), fmt=DT); put(ws, row, 3, py(b), fmt=DT)
        put(ws, row, 4, f'=COUNTIF({TR("R")},A{row})', fmt='0')
        put(ws, row, 5, f'=SUMIFS({TR("O")},{TR("R")},A{row})', fmt=PCT3)
        put(ws, row, 6, f'=IF(D{row}=0,0,E{row}/D{row})', fmt=PCT3)
        put(ws, row, 7, f'=IF(D{row}=0,0,COUNTIFS({TR("R")},A{row},{TR("O")},">0")/D{row})', fmt=PCT)
        put(ws, row, 8, f'=SUMIFS({TR("AA")},{TR("R")},A{row})', fmt=USD)
        put(ws, row, 9, f'=TRANSAKCJE!$B$3+SUMIFS({TR("AA")},{TR("F")},"<="&B{row})', fmt=USD)
        put(ws, row, 10, f'=TRANSAKCJE!$B$3+SUMIFS({TR("AA")},{TR("F")},"<="&C{row})', fmt=USD)
        put(ws, row, 11, f'=J{row}/I{row}-1', fmt=PCT)
    ws.cell(25, 4, 'Podsumowanie').font = fB; ws.cell(25, 5, 'wynik % (1:1)').font = fB; ws.cell(25, 11, 'zmiana kapitału').font = fB
    for k, (lab, f1, fm1, f2, fm2) in enumerate([('Okresy stratne (< 0)', '=COUNTIF(E5:E24,"<0")', '0', '=COUNTIF(K5:K24,"<0")', '0'),
                                                 ('Okresy puste', '=COUNTIF(D5:D24,0)', '0', None, None),
                                                 ('Średnia okresu', '=AVERAGE(E5:E24)', PCT3, '=AVERAGE(K5:K24)', PCT),
                                                 ('Mediana okresu', '=MEDIAN(E5:E24)', PCT3, '=MEDIAN(K5:K24)', PCT),
                                                 ('Najgorszy okres', '=MIN(E5:E24)', PCT3, '=MIN(K5:K24)', PCT),
                                                 ('Najlepszy okres', '=MAX(E5:E24)', PCT3, '=MAX(K5:K24)', PCT)]):
        put(ws, 26 + k, 1, lab, fB, fill=GRY); ws.merge_cells(start_row=26 + k, start_column=1, end_row=26 + k, end_column=4)
        put(ws, 26 + k, 5, f1, fmt=fm1)
        if f2:
            put(ws, 26 + k, 11, f2, fmt=fm2)
    bc = BarChart(); bc.title = 'Zmiana kapitału w każdym z 20 okresów'; bc.height = 8; bc.width = 22
    bc.x_axis.delete = False; bc.y_axis.delete = False
    bc.add_data(Reference(ws, min_col=11, min_row=4, max_row=24), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=24)); bc.y_axis.number_format = '0%'; bc.legend = None
    bc.series[0].graphicalProperties.solidFill = '1F3A5F'; bc.series[0].invertIfNegative = False
    ws.add_chart(bc, 'M4')

    # PARY
    ws = sheet('PARY')
    ws['A1'] = 'PARY — wynik strategii na każdej parze'; ws['A1'].font = fT
    hdr(ws, 3, 1, ['Para', 'Spread (pipsy)', 'Transakcje', 'Wynik % (1:1, suma)', 'Średnia na transakcję', 'Średnio pipsów', 'Trafne', 'Wynik ($)'], [10, 9, 10, 12, 12, 10, 9, 12])
    for k, p in enumerate(PAIRS):
        row = 4 + k
        put(ws, row, 1, p); put(ws, row, 2, O['spread'][p], fIn, '0.00')
        put(ws, row, 3, f'=COUNTIF({TR("B")},A{row})', fmt='0')
        put(ws, row, 4, f'=SUMIFS({TR("O")},{TR("B")},A{row})', fmt=PCT3)
        put(ws, row, 5, f'=IF(C{row}=0,0,D{row}/C{row})', fmt=PCT3)
        put(ws, row, 6, f'=IF(C{row}=0,0,SUMIFS({TR("P")},{TR("B")},A{row})/C{row})', fmt='0.0')
        put(ws, row, 7, f'=IF(C{row}=0,0,COUNTIFS({TR("B")},A{row},{TR("O")},">0")/C{row})', fmt=PCT)
        put(ws, row, 8, f'=SUMIFS({TR("AA")},{TR("B")},A{row})', fmt=USD)
    put(ws, 14, 1, 'Razem', fB, fill=GRY); put(ws, 14, 3, '=SUM(C4:C13)', fB, '0'); put(ws, 14, 4, '=SUM(D4:D13)', fB, PCT3)
    put(ws, 14, 5, '=IF(C14=0,0,D14/C14)', fB, PCT3); put(ws, 14, 7, f'=IF(C14=0,0,COUNTIF({TR("O")},">0")/C14)', fB, PCT); put(ws, 14, 8, '=SUM(H4:H13)', fB, USD)
    note(ws, 16, 1, 'Spread: EUR/USD 1,1 pipsa (zrzut XTB człowieka), pozostałe pary — fx_spread_pips z ia4/ai/config.py × 1,1 / 0,8.', 8, 30)

    # WARIANTY
    ws = sheet('WARIANTY')
    ws['A1'] = 'WARIANTY — ile depozytu na pozycję przy dźwigni (symulacja godzinowa Python)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Wycena otwartych pozycji co godzinę po najgorszej cenie świecy, stop out przy poziomie depozytu < 50%, lot 0,01. Strategia nie ma stop lossa — przy dużym depozycie jedna zła transakcja może dać stop out (kolumna „Stop out”). Koszt nocny 3% rocznie — szacunek ostrożny.', 10, 42)
    hdr(ws, 4, 1, ['Depozyt na pozycję (% salda)', 'Wartość pozycji (× saldo)', 'Koszt nocny (% rocznie)', 'Saldo końcowe ($)', 'Zysk', 'Zysk roczny (CAGR)', 'Największe obsunięcie', 'Transakcje zawarte', 'Pominięte (brak depozytu)', 'Stop out'],
        [13, 11, 10, 13, 10, 11, 12, 11, 12, 8])
    for k, v in O['variants'].reset_index(drop=True).iterrows():
        row = 5 + k
        for c, (val, fm) in enumerate([(v.f, PCT), (v.f * O['lev'], '0.0"×"'), (v.swap / 100, '0%'), (round(v.koniec, 2), USD), (v.zysk_pct / 100, PCT),
                                       (v.cagr_pct / 100, PCT), (-v.max_obs_pct / 100, PCT), (int(v.zawarte), '0'), (int(v.pominiete), '0'), (int(v.stop_out), '0')]):
            x = put(ws, row, 1 + c, val, fmt=fm)
            if abs(v.f - O['f']) < 1e-9 and v.swap == 0:
                x.fill = GOOD
            if c == 9 and v.stop_out > 0:
                x.fill = BAD
    c0 = 13; fs_ = [f_ for f_ in [0.10, 0.20, 1 / 3, 1.0] if f_ in O['curves']]
    hdr(ws, 4, c0, ['Dzień'] + [f"Depozyt {pl(f_ * 100, 0)}% ({pl(f_ * O['lev'], 0)}×)" for f_ in fs_], [12] + [12] * len(fs_))
    cv = pd.DataFrame({f_: O['curves'][f_] for f_ in fs_})
    for k, (d, rr) in enumerate(cv.iterrows()):
        put(ws, 5 + k, c0, py(d), fmt=DD)
        for j, f_ in enumerate(fs_):
            put(ws, 5 + k, c0 + 1 + j, None if np.isnan(rr[f_]) else round(float(rr[f_]), 2), fmt=USD0)
    ch = lchart(f"{pl(O['start'], 0)} $ przy dźwigni 1:{O['lev']:g} — różny depozyt na pozycję (bez kosztu nocnego)", 10, 24)
    ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + len(fs_), min_row=4, max_row=4 + len(cv)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=c0, min_row=5, max_row=4 + len(cv))); ch.y_axis.title = '$'
    for s_, colr in zip(ch.series, ['9DC3E6', '1F3A5F', 'ED7D31', 'C00000']):
        s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 15000
    ch.legend.position = 'b'; ws.add_chart(ch, 'A20')

    # ODPORNOŚĆ
    ws = sheet('ODPORNOŚĆ')
    ws['A1'] = 'ODPORNOŚĆ — sąsiednie ustawienia (wszystkie pary, wynik % bez dźwigni)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Każda tabela zmienia jedną rzecz, reszta jak w strategii FX-POWRÓT (zielony wiersz). Zysk ma być „płaskowyżem”, nie pojedynczym szczytem. t ≥ 3 — wynik wyraźnie różny od zera (bez poprawki na liczbę prób).', 13, 30)
    r = 4
    for title, df in O['neighbors'].items():
        ws.cell(r, 1, title).font = fB
        hdr(ws, r + 1, 1, ['Ustawienie'] + STATS_H, [44] + [10] * 12)
        for k, (_, d) in enumerate(df.iterrows()):
            put(ws, r + 2 + k, 1, d.ustawienie, fB)
            if not pd.isna(d.get('n', np.nan)) and d['n'] > 0:
                stats_row(ws, r + 2 + k, 2, d)
            base = ('≥ 2 ATR, czerwona w ≤ 3' in d.ustawienie or d.ustawienie in ('pierwsza czerwona świeca', 'cofnięcie 50% wzrostu', 'bez SL', '96 świec',
                    f"{M['godziny']} · {M['trend']}"))
            if base:
                ws.cell(r + 2 + k, 1).fill = SEL
        r += 3 + len(df)

    # POWRÓT DO BAZY
    ws = sheet('POWRÓT DO BAZY')
    ws['A1'] = 'Jak często cena wraca do ceny sprzed gwałtownego wzrostu?'; ws['A1'].font = fT
    note(ws, 2, 1, 'Wszystkie impulsy na 10 parach. Liczone od świecy po sygnale. „Przed nowym szczytem” = baza dotknięta, zanim cena przebiła szczyt impulsu o 0,5 ATR — tylko wtedy short z SL tuż nad szczytem zarobiłby pełny ruch. Bez kosztów, bez SL.', 10, 42)
    hdr(ws, 4, 1, ['Impuls', 'Wejście', 'Zdarzeń', 'Baza w ≤ 6 św.', 'Baza w ≤ 24 św.', 'Baza w ≤ 96 św.', 'Baza przed nowym szczytem, ≤ 24 św.', 'Baza przed nowym szczytem, ≤ 96 św.', 'Nowy szczyt (+0,5 ATR) w ≤ 24 św.'],
        [22, 30, 9, 10, 10, 10, 13, 13, 13])
    for k, (_, d) in enumerate(O['returns'].iterrows()):
        row = 5 + k
        put(ws, row, 1, d.impuls); put(ws, row, 2, d.sygnal); put(ws, row, 3, int(d.n), fmt='#,##0')
        for j, c in enumerate(['baza_6', 'baza_24', 'baza_96', 'przed_24', 'przed_96', 'szczyt_24']):
            put(ws, row, 4 + j, d[c], fmt='0%')

    # UCZCIWY TEST
    ws = sheet('UCZCIWY TEST')
    ws['A1'] = 'UCZCIWY TEST — czy to przewaga, czy przypadek z przeszukania?'; ws['A1'].font = fT
    note(ws, 2, 1, f"Przeszukanie: {pl(O['n_rules'], 0)} reguł z wynikiem. Placebo — te same transakcje z tym samym „nawiasem” (odległość do bazy, szczytu, TP i SL w ATR), ale w losowym czasie (przesunięcie ≥ 1500 świec, osobno na parę). Liczba reguł spełniających warunek okresów zależy w placebo mocno od trendu walut — lepiej rozróżnia najlepsze t i zysk / obsunięcie.", 12, 54)
    hdr(ws, 4, 1, ['Proces', 'Reguł: ≥ 200 transakcji, ≤ 3 stratne okresy', 'j.w. ≥ 300 transakcji', 'Obie połowy ≤ 2 stratne z 10', 'EUR/USD ≥ 60, ≤ 3 stratne',
                   'Najlepsze t (≥ 200 transakcji)', 'Reguł z t ≥ 5', 'Najlepszy zysk / obsunięcie', 'Wybór na 1. roku: najlepsza w 2. roku (%)',
                   'Wybór na 1. roku: najlepsze 10 w 2. roku (%)', 'Najlepsze 50 w 2. roku: średnio na transakcję'], [26, 12, 11, 11, 11, 11, 9, 11, 13, 13, 13])
    for k, (tag, c) in enumerate(O['placebo'].items()):
        row = 5 + k; w = O['wf'][tag].get(0, {})
        vals = [(tag, None), (c['all200'], '#,##0'), (c['all300'], '#,##0'), (c['halves'], '#,##0'), (c['eu60'], '#,##0'), (c['max_t'], '0.0'), (c['t5'], '0'),
                (c['max_zo'], '0.0'), (w.get('top1', np.nan) / 100, PCT3), (w.get('top10', np.nan) / 100, PCT3), (w.get('top50_tr', np.nan) / 100, PCT3)]
        for j, (v, fm) in enumerate(vals):
            put(ws, row, 1 + j, v, fmt=fm, fill=GRY if k == 0 else None)
    r = 7 + len(O['placebo'])
    ws.cell(r, 1, f"FX-POWRÓT w 2. roku i reguła wybrana wyłącznie na 1. roku ({pl(O['start'], 0)} $, 1:{O['lev']:g}, depozyt 20%)").font = fB
    hdr(ws, r + 1, 1, ['Strategia', 'Transakcje', 'Średnia na transakcję', 'Trafne', 'Saldo po 2. roku ($)', 'Zysk', 'Największe obsunięcie'])
    if 'wf_rule' in O:
        put(ws, r + 2, 1, f"Wybrana na 1. roku (poza próbą): {O['wf_rule']['opis']}", al=WR)
        put(ws, r + 2, 2, int(O['wf_sum']['n'])); put(ws, r + 2, 3, O['wf_sum']['srednia'] / 100, fmt=PCT3); put(ws, r + 2, 4, O['wf_sum']['trafne'], fmt=PCT)
        put(ws, r + 2, 5, round(O['wf_stats']['koniec'], 2), fmt=USD); put(ws, r + 2, 6, O['wf_stats']['zysk_pct'] / 100, fmt=PCT); put(ws, r + 2, 7, -O['wf_stats']['max_obs_pct'] / 100, fmt=PCT)
    m2, my2 = O['main_y2_tr'], O['main_y2']
    put(ws, r + 3, 1, 'FX-POWRÓT w 2. roku (w próbie — ustawienia znane z całych danych)', al=WR)
    put(ws, r + 3, 2, len(m2)); put(ws, r + 3, 3, (m2.wynik.mean() if len(m2) else 0) / 100, fmt=PCT3); put(ws, r + 3, 4, (m2.wynik > 0).mean() if len(m2) else 0, fmt=PCT)
    put(ws, r + 3, 5, round(my2['koniec'], 2), fmt=USD); put(ws, r + 3, 6, my2['zysk_pct'] / 100, fmt=PCT); put(ws, r + 3, 7, -my2['max_obs_pct'] / 100, fmt=PCT)
    ws.row_dimensions[r + 2].height = 54; ws.row_dimensions[r + 3].height = 30
    if 'wf_curve' in O:
        c0 = 9
        hdr(ws, r + 1, c0, ['Dzień', 'Wybrana na 1. roku (poza próbą)', 'FX-POWRÓT (w próbie)'], [12, 14, 14])
        cc = pd.DataFrame({'a': O['wf_curve'], 'b': O['main_y2_curve']}).dropna()
        for k, (d, rr) in enumerate(cc.iterrows()):
            put(ws, r + 2 + k, c0, py(d), fmt=DD); put(ws, r + 2 + k, c0 + 1, round(rr.a, 2), fmt=USD0); put(ws, r + 2 + k, c0 + 2, round(rr.b, 2), fmt=USD0)
        ch = lchart('2. rok: reguła wybrana bez wiedzy o nim vs FX-POWRÓT', 9, 20)
        ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + 2, min_row=r + 1, max_row=r + 1 + len(cc)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=c0, min_row=r + 2, max_row=r + 1 + len(cc)))
        ch.series[0].graphicalProperties.line.solidFill = 'ED7D31'; ch.series[1].graphicalProperties.line.solidFill = '1F3A5F'; ch.legend.position = 'b'
        ws.add_chart(ch, f'A{r + 6}')
    r += 26
    ws.cell(r, 1, 'Wymiary przeszukania (prawdziwe sygnały): ile reguł z ≥ 200 transakcjami, ile spełnia warunek okresów, najlepsze t i suma').font = fB
    hdr(ws, r + 1, 1, ['Wymiar', 'Wartość', 'Reguł (≥ 200 transakcji)', 'Spełnia warunek', 'Najlepsze t', 'Najlepsza suma %'])
    for k, (_, d) in enumerate(O['dims'].iterrows()):
        row = r + 2 + k
        put(ws, row, 1, d.wymiar); put(ws, row, 2, d.wartosc); put(ws, row, 3, int(d.reguly), fmt='#,##0'); put(ws, row, 4, int(d.warunek), fmt='#,##0')
        put(ws, row, 5, d.max_t, fmt='0.0'); put(ws, row, 6, d.max_suma / 100, fmt=PCT3)

    # RANKING
    ws = sheet('RANKING')
    ws['A1'] = 'RANKING — tylko do wglądu (w próbie, zawyżony)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Warunek: ≥ 200 transakcji, ≤ 3 stratne okresy z 20, suma > 0. Wyniki % bez dźwigni. Trzecia tabela — tylko TP dokładnie na cenie sprzed wzrostu (pomysł człowieka), najlepsze t bez warunku okresów.', 16, 30)
    rc = ['Impuls', 'Baza', 'Sygnał', 'Okno', 'Trend', 'Godziny', 'Odl. od bazy (ATR)', 'Wyjście', 'Transakcje', 'Suma %', 'Średnia', 'Trafne', 'PF', 't', 'Obsunięcie',
          'Okresy stratne', '1. połowa', '2. połowa', 'EUR/USD n', 'EUR/USD suma', 'Zysk / obsunięcie']
    r = 4
    for src, title in [(O['rank'], 'Ranking: zysk / obsunięcie'), (O['rank_t'], 'Ranking: t'), (O['rank_base'], 'TP = cena sprzed wzrostu (najlepsze t)')]:
        ws.cell(r, 1, title).font = fB
        hdr(ws, r + 1, 1, rc, [22, 20, 30, 6, 22, 14, 9, 40] + [9] * 13)
        for k, (_, d) in enumerate(src.iterrows()):
            vals = [d.impuls, BASES[int(d.baza)], d['sygnał'], int(d.okno), d.trend, d.godziny, d.odl, d['wyjście'], int(d.n), d.suma / 100, d.srednia / 100,
                    d.trafne, min(d.PF, 99), d.t, d.obsuniecie / 100, int(d.stratne_okresy), d.suma1 / 100, d.suma2 / 100, int(d.eu_n), d.eu_suma / 100, d['zysk/obs']]
            fms = [None] * 6 + ['0.0', None, '0', PCT3, PCT3, PCT, '0.00', '0.0', PCT3, '0', PCT3, PCT3, '0', PCT3, '0.0']
            for j, (v, fm) in enumerate(zip(vals, fms)):
                put(ws, r + 2 + k, 1 + j, v, fmt=fm)
        r += 4 + len(src)

    # EURUSD
    ws = sheet('EURUSD')
    ws['A1'] = 'EUR/USD — ta sama strategia tylko na EUR/USD'; ws['A1'].font = fT
    note(ws, 2, 1, 'Formuły czytają TRANSAKCJE. Saldo przy dźwigni — symulacja Python tylko z transakcjami EUR/USD (jedna pozycja naraz). Mało transakcji → szeroki przedział ufności (zasada 6).', 10, 42)
    hdr(ws, 4, 1, ['Okres', 'Transakcje', 'Wynik % (1:1)', 'Trafne'], [8, 10, 12, 9])
    for k in range(NPER):
        row = 5 + k; put(ws, row, 1, k + 1)
        put(ws, row, 2, f'=COUNTIFS({TR("R")},A{row},{TR("B")},"EURUSD")', fmt='0')
        put(ws, row, 3, f'=SUMIFS({TR("O")},{TR("R")},A{row},{TR("B")},"EURUSD")', fmt=PCT3)
        put(ws, row, 4, f'=IF(B{row}=0,0,COUNTIFS({TR("R")},A{row},{TR("B")},"EURUSD",{TR("O")},">0")/B{row})', fmt=PCT)
    put(ws, 25, 1, 'Razem', fB, fill=GRY); put(ws, 25, 2, '=SUM(B5:B24)', fB, '0'); put(ws, 25, 3, '=SUM(C5:C24)', fB, PCT3)
    put(ws, 26, 1, 'Okresy stratne', fB, fill=GRY); put(ws, 26, 3, '=COUNTIF(C5:C24,"<0")', fB, '0')
    put(ws, 27, 1, 'Okresy puste', fB, fill=GRY); put(ws, 27, 3, '=COUNTIF(B5:B24,0)', fB, '0')
    hdr(ws, 4, 6, ['Depozyt na pozycję', 'Wartość pozycji', 'Saldo końcowe ($)', 'Zysk', 'Największe obsunięcie', 'Transakcje', 'Stop out'], [11, 10, 13, 9, 11, 9, 8])
    for k, v in O['eu_var'].reset_index(drop=True).iterrows():
        for c, (val, fm) in enumerate([(v.f, PCT), (v.f * O['lev'], '0.0"×"'), (round(v.koniec, 2), USD), (v.zysk_pct / 100, PCT), (-v.max_obs_pct / 100, PCT),
                                       (int(v.zawarte), '0'), (int(v.stop_out), '0')]):
            put(ws, 5 + k, 6 + c, val, fmt=fm)
    c0 = 14
    hdr(ws, 4, c0, ['Dzień', 'EUR/USD, depozyt 33%', 'EUR/USD, depozyt 100%'], [12, 13, 13])
    ec = pd.DataFrame({'a': O['eu_curves'][1 / 3], 'b': O['eu_curves'][1.0]})
    for k, (d, rr) in enumerate(ec.iterrows()):
        put(ws, 5 + k, c0, py(d), fmt=DD); put(ws, 5 + k, c0 + 1, round(rr.a, 2), fmt=USD0); put(ws, 5 + k, c0 + 2, round(rr.b, 2), fmt=USD0)
    ch = lchart(f"EUR/USD: {pl(O['start'], 0)} $ przy 1:{O['lev']:g}", 8, 20)
    ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + 2, min_row=4, max_row=4 + len(ec)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=c0, min_row=5, max_row=4 + len(ec))); ch.legend.position = 'b'
    ws.add_chart(ch, 'A30')

    # PODSUMOWANIE
    ws = wsP
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 70
    ws['A1'] = 'FX-POWRÓT — short po gwałtownej świecy wzrostowej, wejście po zawróceniu, TP w połowie drogi do ceny sprzed wzrostu'; ws['A1'].font = fT
    ws['A2'] = (f"IA4 {__version__} · badanie informacyjne · kopia fx1h: {O['od']:%Y-%m-%d} – {O['do']:%Y-%m-%d} ({pl(O['swiec'], 0)} świec na parę) · "
                f"policzono {pd.Timestamp.now():%Y-%m-%d %H:%M}")
    ws['A2'].font = fS
    r = 4; put(ws, r, 1, 'REGUŁA (wszystkie liczby stałe)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    rule = [('Rynek', 'EUR/USD, USD/JPY, GBP/USD, AUD/USD, USD/CAD, USD/CHF, NZD/USD, EUR/JPY, GBP/JPY, EUR/GBP — świece 1h; jedna pozycja na parę'),
            ('Impuls', f"zamknięcie świecy wyżej niż zamknięcie poprzedniej o ≥ {_n(M['k'])} × ATR(100) — gwałtowna zielona świeca"),
            ('Baza i szczyt', 'baza = zamknięcie świecy przed impulsem (cena sprzed wzrostu); szczyt = najwyższy punkt od impulsu do sygnału'),
            ('Sygnał wejścia', f"pierwsza czerwona świeca (close < open) w ciągu {M['okno']} świec od ostatniej świecy impulsu; świeca sygnału zaczyna się {M['godziny']} (w Polsce 9:00–18:00 latem); impuls wygasa, gdy kurs wrócił do bazy wcześniej"),
            ('Wejście', 'SHORT po open następnej świecy (gdy open jest wyżej niż baza o ≥ 0,5 ATR)'),
            ('Take profit', 'szczyt − 50% × (szczyt − baza) — połowa drogi do ceny sprzed wzrostu'),
            ('Stop loss', 'brak — każdy SL (także 3–8 ATR) obniża wynik; stopem jest limit czasu (arkusz ODPORNOŚĆ, tabela E)'),
            ('Limit czasu', f"po {M['H']} świecach (ok. 4 dni) wyjście po close"),
            ('Koszty', 'spread raz na transakcję: EUR/USD 1,1 pipsa (zrzut XTB), pozostałe pary — arkusz PARY; koszt nocny — TRANSAKCJE!B6'),
            ('Rozliczenie', 'pesymistyczne (zasada 5): SL i TP w tej samej świecy → SL; luka przez TP → TP bez premii')]
    for k, (a, b) in enumerate(rule):
        put(ws, r + 1 + k, 1, a, fB, fill=GRY, al=WR); put(ws, r + 1 + k, 2, b, al=WR)
        ws.merge_cells(start_row=r + 1 + k, start_column=2, end_row=r + 1 + k, end_column=3); ws.row_dimensions[r + 1 + k].height = 30
    r += 2 + len(rule)
    put(ws, r, 1, 'WYNIK (formuły z arkuszy TRANSAKCJE, OKRESY, KAPITAŁ)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    reasons = O['main'].powod.value_counts() if len(O['main']) else {}
    vdef = O['variants'][(abs(O['variants'].f - O['f']) < 1e-9) & (O['variants'].swap == 0)].iloc[0]
    res = [('Transakcje (wszystkie pary)', f'=COUNTA({TR("B")})', '0', f"EUR/USD: {O['eu_sum']['n']} (arkusz EURUSD)"),
           ('Średni wynik transakcji (bez dźwigni)', f'=AVERAGE({TR("O")})', PCT3, f"przedział 95%: ± {pl(O['main_ci'], 3)}% (t = {pl(O['main_t'])})"),
           ('Mediana transakcji', f'=MEDIAN({TR("O")})', PCT3, ''),
           ('Trafne transakcje', f'=COUNTIF({TR("O")},">0")/COUNTA({TR("O")})', PCT, ' · '.join(f'{k} {v}' for k, v in dict(reasons).items())),
           ('Profit factor', f'=SUMIF({TR("O")},">0")/-SUMIF({TR("O")},"<0")', '0.00', ''),
           ('Najgorsza transakcja (bez dźwigni)', f'=MIN({TR("O")})', PCT3, 'brak SL — strata ograniczona tylko limitem czasu'),
           ('Mediana świec w pozycji', f'=MEDIAN({TR("G")})', '0', 'godziny'),
           ('Okresy stratne z 20 (wynik %)', '=OKRESY!E26', '0', 'warunek człowieka: najmniej stratnych okresów'),
           ('Średnia / mediana okresu (wynik % bez dźwigni)', '=OKRESY!E28', PCT3, ''), ('', '=OKRESY!E29', PCT3, ''),
           ('Saldo końcowe przy ustawieniach z TRANSAKCJE ($)', f'=KAPITAŁ!B{KR1}', USD, 'zmień TRANSAKCJE!B3:B7'),
           ('Zysk', f'=KAPITAŁ!B{KR1}/TRANSAKCJE!B3-1', PCT, ''),
           ('Największe obsunięcie salda (dzienne)', f'=-MIN(KAPITAŁ!D5:D{KR1})', PCT, f"z wyceną otwartych pozycji co godzinę (Python, ustawienia domyślne): {pl(vdef.max_obs_pct)}%"),
           ('Zmiana kapitału: okresy stratne / mediana okresu', '=OKRESY!K26', '0', ''), ('', '=OKRESY!K29', PCT, '')]
    for k, (a, f_, fm, c) in enumerate(res):
        row = r + 1 + k
        put(ws, row, 1, a, fB, fill=GRY); put(ws, row, 2, f_, fB, fm); put(ws, row, 3, c, fS, al=WR)
    r += 2 + len(res)
    put(ws, r, 1, 'WNIOSKI' if wnioski else 'JAK TO CZYTAĆ', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    txt = wnioski or ['Jak często cena wraca do bazy — arkusz POWRÓT DO BAZY.', 'Sąsiednie ustawienia, sygnały, TP i SL — ODPORNOŚĆ.',
                      'Kapitał przy dźwigni — KAPITAŁ i WARIANTY.', 'Uczciwość wyniku — UCZCIWY TEST (placebo, wybór na 1. roku).']
    for k, s_ in enumerate(txt):
        x = ws.cell(r + 1 + k, 1, s_); x.font = fN; x.alignment = WR
        ws.merge_cells(start_row=r + 1 + k, start_column=1, end_row=r + 1 + k, end_column=3); ws.row_dimensions[r + 1 + k].height = 56 if wnioski else 18
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- uruchomienie

def run(no_sync=False, out=None, f=0.20, lev=30.0, start=1000.0, frames=None, log=print, keep=None) -> Path:
    t0 = time.time()
    if frames is None:
        from .fx import sync as fxs
        if not no_sync:
            log('Kopia fx1h — odświeżenie (python -m ia4.fx.sync)')
            fxs.sync(quiet=True, syms=list(PAIRS))
        frames = {p: fxs.load(p) for p in PAIRS}
    K = Kurs(frames, fs.spreads())
    log(f'Dane: {K.n} świec na parę, {pd.Timestamp(int(K.t[0]), unit="s"):%Y-%m-%d} – {pd.Timestamp(int(K.t[-1]), unit="s"):%Y-%m-%d}')
    O = compute(K, f=f, lev=lev, start=start, log=log)
    if keep is not None:
        keep.update(O)
    path = write_xlsx(O, out or (config.RESEARCH_DIR.parent / 'FX-POWROT.xlsx'))
    ms = O['main_sum']; v = O['variants'][(abs(O['variants'].f - f) < 1e-9) & (O['variants'].swap == 0)].iloc[0]
    log(f"\nFX-POWRÓT: {ms['n']} transakcji, średnio {ms['srednia']:+.3f}% (t = {O['main_t']:.1f}), trafne {ms['trafne']:.0%}, okresy stratne "
        f"{ms['stratne']}/20; {start:.0f} $ przy 1:{lev:g} i depozycie {f:.0%} → {v.koniec:,.0f} $, obsunięcie {v.max_obs_pct:.1f}%, stop out {int(v.stop_out)}")
    if 'wf_stats' in O:
        log(f"Poza próbą (reguła wybrana na 1. roku, 2. rok, 20%): {O['wf_stats']['zysk_pct']:+.1f}%")
    log(f'Zapisano: {path}  ({time.time() - t0:.0f} s)')
    return path


def selftest() -> bool:
    ok = True

    def check(name, cond):
        nonlocal ok
        print(f"  {'✓' if cond else '✗'} {name}")
        ok &= bool(cond)

    # impuls: świeca 5 rośnie o 3 ATR, świeca 6 czerwona → sygnał w świecy 6, baza = close świecy 4, szczyt = high świecy 5
    n = 30
    c = np.full(n, 1.0); o = c.copy(); h = c + 0.001; l = c - 0.001; atr = np.full(n, 0.01); valid = np.ones(n, bool)
    o[5], c[5], h[5] = 1.0, 1.03, 1.032
    o[6], c[6], h[6], l[6] = 1.03, 1.025, 1.031, 1.024
    o[7:] = 1.025; c[7:] = 1.025; h[7:] = 1.026; l[7:] = 1.024
    o[9], c[9], h[9], l[9] = 1.025, 1.012, 1.025, 1.010           # spadek do 1,010 — TP 50% = 1,016
    run = np.r_[np.nan, np.diff(c)] / atr
    J, B, P, S0 = gen_trades(o, h, l, c, atr, run, 2.0, 1, 0, 1, 3, c, valid, 2)
    check('impuls → sygnał w pierwszej czerwonej świecy', J.size == 1 and J[0] == 6 and abs(B[0] - 1.0) < 1e-12 and abs(P[0] - 1.032) < 1e-12)
    pn, xb, ma = outcomes_tr(o, h, l, c, J, B, P, atr[J], 0.0001, np.array([0.5]), np.array([0], np.int64), np.array([0.0]), np.array([10], np.int64))
    check('TP = połowa drogi do bazy (1,016), wynik = (1,025 − 1,016) / 1,025 − spread', abs(pn[0, 0] - ((1.025 - 1.016) / 1.025 * 100 - 0.0001 / 1.025 * 100)) < 1e-4 and xb[0, 0] == 3)
    pn, xb, ma = outcomes_tr(o, h, l, c, J, B, P, atr[J], 0.0001, np.array([1.0]), np.array([1], np.int64), np.array([0.1]), np.array([10], np.int64))
    check('TP = baza 1,000 nieosiągnięty, wyjście po czasie (close 1,025)', abs(pn[0, 0] - (-0.0001 / 1.025 * 100)) < 1e-4 and xb[0, 0] == 10)
    J2, _, _, _ = gen_trades(o, h, l, c, atr, run, 2.0, 1, 0, 0, 0, c, valid, 2)
    check('„natychmiast” → wejście po świecy impulsu', J2.size == 1 and J2[0] == 5)
    # całość na losowym spacerze
    fr = fs._synthetic()
    K = Kurs(fr, fs.spreads())
    T = Trades(K, (1, 1.5, 0, 1, 3))
    check(f'transakcje na danych syntetycznych ({0 if T.d is None else T.d["j"].size})', T.d is not None and T.d['j'].size > 20)
    tr = rule_trades(K, (1, 1.5, 0, 1, 3), 'bez filtra trendu', 'wszystkie godziny', 0.5, EXITS.index((0.5, (0, 0.0), 24)))
    tr['bazowa_usd'] = 1.0
    L, cv, ev = fs.simulate(tr, K, f=0.2)
    check('rachunek z dźwignią przebiega', len(L) > 0)
    print(f"\nWYNIK: {'OK' if ok else 'BŁĄD'}")
    return ok


def main():
    ap = argparse.ArgumentParser(description='IA 4 — FX-POWRÓT: short po gwałtownym wzroście, TP do ceny sprzed wzrostu (sekcja 7b)')
    ap.add_argument('--bez-sync', action='store_true', help='bez odświeżenia kopii fx1h')
    ap.add_argument('--depozyt', type=float, default=0.20)
    ap.add_argument('--dzwignia', type=float, default=30.0)
    ap.add_argument('--start', type=float, default=1000.0)
    ap.add_argument('--out', type=Path)
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args()
    if a.selftest:
        sys.exit(0 if selftest() else 1)
    print(f'IA 4 {__version__} — FX-POWRÓT (impuls → zawrócenie → TP w stronę ceny sprzed wzrostu, 10 par, 1h)\n')
    run(no_sync=a.bez_sync, out=a.out, f=a.depozyt, lev=a.dzwignia, start=a.start)


if __name__ == '__main__':
    main()
