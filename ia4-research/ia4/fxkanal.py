"""
IA 4 — FX-KANAŁ: short przy górnej linii kanału, long przy dolnej (instrukcja, sekcja 7b). Wersja projektu: 1.56 (2026-10-21)
— musi zgadzać się z IA4_INSTRUKCJA.md

Obserwacja człowieka (21.10.2026, zrzut z wykresu EURUSD): cena faluje między dwiema „średnimi” — górną (zielona: za każdym
razem, gdy cena do niej dobija, idealny moment na short) i dolną (żółta: long). Pytanie: czy da się to złapać regułą?

1. ZŁUDZENIE — linie narysowane na gotowym wykresie „znają przyszłość” (każdy punkt linii leży między ceną sprzed i po).
   Liczymy ten sam kanał wycentrowany (z przyszłością) i zwykły (tylko przeszłość) — różnica to cena opóźnienia.
2. PRZESZUKANIE — kanały, które da się policzyć na żywo:
     Bollinger SMA(n) ± k·σ(n) · Keltner EMA(n) ± k·ATR(100) · Donchian (najwyższy / najniższy punkt n świec) ·
     obwiednia: EMA(24) z najwyższego / najniższego punktu n świec;
   wejście: dotknięcie linii (pierwsze), zamknięcie za linią (pierwsze), powrót do kanału (zamknięcie znów w środku),
   dotknięcie + świeca w przeciwną stronę; short przy górnej, long przy dolnej, po open następnej świecy;
   filtry: nachylenie środka kanału (płaski / zgodny / przeciwny), godzina, szerokość kanału;
   wyjścia: TP na środku kanału (aktualnym), w połowie drogi do środka, na przeciwnej linii, 0,5 / 1 / 2 ATR; SL brak / 1 / 2 / 3 ATR;
   limit 12 / 24 / 48 / 96 świec. Rozliczenie pesymistyczne (zasada 5), spread raz na transakcję, jedna pozycja na parę.
3. UCZCIWY TEST — placebo (te same wyjścia, wejście w losowej chwili), wybór wyłącznie na 1. roku → wynik 2. roku,
   sąsiednie ustawienia najlepszej reguły, rachunek 1000 $ na 1:30; doubleProof — raz, na 10 parach fx1h_double_proof
   (`--double-proof`, reguła zamrożona w pliku data/fxkanal/strategia.json, otwarcia w data/fxkanal/doubleproof.jsonl).

    python -m ia4.fxkanal                  # ok. 10–30 min (wszystkie rdzenie) → ~/Desktop/IA4/FX-KANAL.xlsx
    python -m ia4.fxkanal --double-proof   # dodatkowo sprawdzian zamrożonej strategii na parach fx1h_double_proof
    python -m ia4.fxkanal --selftest       # „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import numba as nb
import numpy as np
import pandas as pd

from . import __version__, config
from . import fxpowrot as fp
from . import fxpowrot2 as F2
from . import fxshort as fs
from .fxshort import NPER, PAIRS

OUT_DIR = config.RESEARCH_DIR / 'data' / 'fxkanal'
GRANICA_PER = NPER // 2                 # okresy 1–10 = rok 1, 11–20 = rok 2

# ---------------------------------------------------------------- kanały, wejścia, filtry, wyjścia
BANDS = ([('BB', n, k) for n in (20, 50, 100, 200, 400) for k in (1.5, 2.0, 2.5, 3.0, 3.5)] +
         [('KC', n, k) for n in (20, 50, 100, 200, 400) for k in (1.0, 1.5, 2.0, 3.0, 4.0, 5.0)] +
         [('DC', n, 0.0) for n in (12, 24, 48, 120, 240, 480)] +
         [('EN', n, 24.0) for n in (24, 48, 120, 240, 480)] + [('EN', n, 72.0) for n in (120, 240)])
TRIG = ['dotknięcie linii', 'zamknięcie za linią', 'powrót do kanału', 'dotknięcie + świeca w drugą stronę']
DIRS = [(-1, 'short przy górnej'), (1, 'long przy dolnej')]
SLOPE = ['każde nachylenie', 'środek płaski', 'środek w stronę transakcji', 'środek przeciw transakcji']
HOURS = [('wszystkie godziny', 0, 23, False), ('7–16 UTC', 7, 16, False), ('poza 7–16 UTC', 7, 16, True)]
WIDTH = ['każda szerokość', 'kanał wąski (dolna ⅓)', 'kanał szeroki (górna ⅓)']
TREND = ['każdy trend długi', 'SMA500 w stronę transakcji', 'SMA500 przeciw transakcji']
TP_M = [('środek kanału', 0, 0.0), ('½ drogi do środka', 1, 0.5), ('przeciwna linia', 2, 0.0),
        ('0,5 ATR', 3, 0.5), ('1 ATR', 3, 1.0), ('2 ATR', 3, 2.0)]
SL_V = [0.0, 1.0, 2.0, 3.0]
HOLD = [12, 24, 48, 96]
EXITS = [(t, s, h) for t in range(len(TP_M)) for s in SL_V for h in HOLD]
NE = len(EXITS)
EX_TM = np.array([TP_M[t][1] for t, s, h in EXITS], np.int64); EX_TV = np.array([TP_M[t][2] for t, s, h in EXITS])
EX_SL = np.array([s for t, s, h in EXITS]); EX_H = np.array([h for t, s, h in EXITS], np.int64)
POOL = dict(n=300, stratne=4, t=3.0, pf=1.2)                    # warunek „stabilna i dużo transakcji”
RULE_KEYS = ['b', 'trig', 'dir', 'slope', 'hours', 'width', 'trend', 'e']
NF = len(SLOPE) * len(HOURS) * len(WIDTH) * len(TREND)
NR = len(BANDS) * len(TRIG) * len(DIRS) * NF * NE


def band_name(b):
    kind, n, k = b
    kk = f'{k:g}'.replace('.', ',')
    return {'BB': f'Bollinger {n}, ±{kk}σ', 'KC': f'Keltner EMA{n} ±{kk} ATR', 'DC': f'Donchian {n} św.',
            'EN': f'obwiednia {n} św. (EMA{int(k)})'}[kind]


def exit_name(e):
    t, s, h = EXITS[e]
    return f"TP {TP_M[t][0]} · {'bez SL' if not s else 'SL ' + f'{s:g}'.replace('.', ',') + ' ATR'} · {h} św."


def _sma(x, n):
    return pd.Series(x).rolling(n).mean().values


def _ema(x, n):
    return pd.Series(x).ewm(span=n, adjust=False).mean().values


def bands(Fp: dict, b, centered=False):
    """(dolna, środek, górna) na świecy s — z danych do zamknięcia s włącznie (Donchian: n świec PRZED s).
    centered=True: ten sam kanał przesunięty o połowę okna wstecz — „zna przyszłość” (tylko do pokazania złudzenia)."""
    kind, n, k = b
    c, h, l = Fp['c'], Fp['h'], Fp['l']
    if kind == 'BB':
        m = _sma(c, n); sd = pd.Series(c).rolling(n).std(ddof=0).values
        lo, up = m - k * sd, m + k * sd
    elif kind == 'KC':
        m = _ema(c, n); a = Fp['atr']
        lo, up = m - k * a, m + k * a
    elif kind == 'DC':
        up = pd.Series(h).rolling(n).max().shift(1).values; lo = pd.Series(l).rolling(n).min().shift(1).values
        m = (up + lo) / 2
    else:
        up = np.array(_ema(np.nan_to_num(pd.Series(h).rolling(n).max().values, nan=h[0]), int(k)), dtype=np.float64)
        lo = np.array(_ema(np.nan_to_num(pd.Series(l).rolling(n).min().values, nan=l[0]), int(k)), dtype=np.float64)
        up[:n] = np.nan; lo[:n] = np.nan
        m = (up + lo) / 2
    if centered:
        sh = n // 2
        lo, m, up = [np.r_[a[sh:], np.full(sh, np.nan)] for a in (lo, m, up)]
    return lo.astype(np.float64), m.astype(np.float64), up.astype(np.float64)


@nb.njit(cache=True)
def events(o, h, l, c, lo, mid, up, valid, trig, d, start):
    """Świece sygnału (indeksy). d = −1: górna linia (short), +1: dolna (long). Pierwsze wystąpienie (świeca wcześniej — nie)."""
    n = c.size
    out = np.empty(n, np.int64); k = 0
    for s in range(max(start, 2), n - 1):
        if up[s] != up[s] or lo[s] != lo[s] or up[s - 1] != up[s - 1]:
            continue
        if d < 0:
            a0 = h[s] >= up[s]; a1 = h[s - 1] >= up[s - 1]
            b0 = c[s] > up[s]; b1 = c[s - 1] > up[s - 1]
            rev = c[s] < o[s]
        else:
            a0 = l[s] <= lo[s]; a1 = l[s - 1] <= lo[s - 1]
            b0 = c[s] < lo[s]; b1 = c[s - 1] < lo[s - 1]
            rev = c[s] > o[s]
        if trig == 0:
            ok = a0 and not a1
        elif trig == 1:
            ok = b0 and not b1
        elif trig == 2:
            ok = b1 and not b0
        else:
            prev_rev = (c[s - 1] < o[s - 1]) if d < 0 else (c[s - 1] > o[s - 1])
            ok = a0 and rev and not (a1 and prev_rev)
        if ok and valid[s]:
            out[k] = s; k += 1
    return out[:k]


@nb.njit(cache=True)
def outcomes(o, h, l, c, J, d, A, lo, mid, up, sp, tm, tv, slv, hh):
    """Wynik % netto każdej transakcji przy każdym wyjściu (NaN — wyjście nie pasuje, np. cena już za celem)."""
    n = o.size; nt = J.size; ne = tm.size
    pnl = np.full((nt, ne), np.nan, np.float32); xb = np.zeros((nt, ne), np.int16)
    for t in range(nt):
        s = J[t]; e0 = s + 1
        en = o[e0]; a = A[t]
        for e in range(ne):
            H = hh[e]
            if e0 + H > n:
                continue
            m = tm[e]
            # cel stały (1, 3) albo ruchomy (0 — środek, 2 — przeciwna linia; wartość z poprzedniej świecy)
            if m == 1:
                tgt0 = en + tv[e] * (mid[s] - en)
            elif m == 3:
                tgt0 = en + d * tv[e] * a
            elif m == 0:
                tgt0 = mid[s]
            else:
                tgt0 = lo[s] if d < 0 else up[s]
            if tgt0 != tgt0 or d * (tgt0 - en) <= 0.05 * a:
                continue
            SL = en - d * slv[e] * a if slv[e] > 0 else (-1e18 if d > 0 else 1e18)
            ex = np.nan; k = 0
            for j in range(e0, e0 + H):
                k = j - e0 + 1
                if m == 0:
                    tgt = mid[j - 1]
                elif m == 2:
                    tgt = lo[j - 1] if d < 0 else up[j - 1]
                else:
                    tgt = tgt0
                if d < 0:
                    if j > e0 and o[j] >= SL:
                        ex = o[j]; break
                    if j > e0 and o[j] <= tgt:
                        ex = tgt; break
                    if h[j] >= SL:
                        ex = SL; break
                    if l[j] <= tgt:
                        ex = tgt; break
                else:
                    if j > e0 and o[j] <= SL:
                        ex = o[j]; break
                    if j > e0 and o[j] >= tgt:
                        ex = tgt; break
                    if l[j] <= SL:
                        ex = SL; break
                    if h[j] >= tgt:
                        ex = tgt; break
                if k == H:
                    ex = c[j]
            pnl[t, e] = d * (ex - en) / en * 100.0 - sp / en * 100.0
            xb[t, e] = k
    return pnl, xb


# ---------------------------------------------------------------- dane i przeszukanie

class Kanaly:
    """Świece 10 par na wspólnej siatce (fxpowrot2.Rynek) + cechy do filtrów."""

    def __init__(self, frames: dict, pairs=PAIRS, spread=None):
        self.R = F2.Rynek(frames, pairs, spread or fs.spreads())
        self.pairs = list(pairs)
        self.n = self.R.n; self.t = self.R.t; self.per = self.R.per


def _filters(R, p, b, J, d, lo, mid, up, q_width):
    Fp = R.F[p]; A = Fp['atr'][J]
    slope = (mid[J] - mid[np.maximum(J - 24, 0)]) / A * d                  # > 0: środek idzie w stronę transakcji
    hour = Fp['hour'][J]
    width = (up[J] - lo[J]) / A
    m_s = [np.ones(J.size, bool), np.abs(slope) <= 0.3, slope > 0.3, slope < -0.3]
    m_h = [np.ones(J.size, bool)] + [((hour >= a) & (hour <= b_)) ^ inv for _, a, b_, inv in HOURS[1:]]
    m_w = [np.ones(J.size, bool), width <= q_width[0], width >= q_width[1]]
    tr = np.nan_to_num(Fp['slope500'][J], nan=0.0) * d
    m_t = [np.ones(J.size, bool), tr > 0, tr < 0]
    return m_s, m_h, m_w, m_t


def scan_band(R, bi, shifts=None):
    """Wszystkie reguły jednego kanału: wejścia × kierunki × filtry × wyjścia → tablica statystyk (float32, kolumny COLS)."""
    b = BANDS[bi]
    out = []
    pair_idx = {p: i for i, p in enumerate(R.pairs)}
    eu = R.pairs.index('EURUSD') if 'EURUSD' in R.pairs else 0
    bandv = {p: bands(R.F[p], b) for p in R.pairs}
    # progi szerokości z roku 1 (wszystkie pary) — bez podglądania roku 2
    W = np.concatenate([((bandv[p][2] - bandv[p][0]) / R.F[p]['atr'])[(R.per < GRANICA_PER) & ~np.isnan(bandv[p][2])] for p in R.pairs])
    W = W[np.isfinite(W)]
    qw = (np.quantile(W, 1 / 3), np.quantile(W, 2 / 3)) if W.size else (0, 1e9)
    ek = np.arange(NE, dtype=np.float32)
    for ti in range(len(TRIG)):
        for di, (d, _) in enumerate(DIRS):
            PP, JJ, PN, XB = [], [], [], []
            MK = [[[] for _ in range(n_)] for n_ in (len(SLOPE), len(HOURS), len(WIDTH), len(TREND))]
            for p in R.pairs:
                Fp = R.F[p]; lo, mid, up = bandv[p]
                J = events(Fp['o'], Fp['h'], Fp['l'], Fp['c'], lo, mid, up, Fp['valid'], ti, d, fs.WARMUP)
                if shifts is not None and J.size:            # placebo: tyle samo wejść, losowe chwile
                    rng = np.random.default_rng(shifts * 1_000_003 + 7919 * bi + 31 * ti + di + 101 * pair_idx[p])
                    okv = np.nonzero(Fp['valid'] & ~np.isnan(up) & (np.arange(R.n) < R.n - 100) & (np.arange(R.n) >= fs.WARMUP))[0]
                    J = np.sort(rng.choice(okv, size=min(J.size, okv.size), replace=False))
                if J.size == 0:
                    continue
                pn, xb = outcomes(Fp['o'], Fp['h'], Fp['l'], Fp['c'], J, float(d), Fp['atr'][J], lo, mid, up,
                                  R.spread[p] * F2.pip_of(p), EX_TM, EX_TV, EX_SL, EX_H)
                masks = _filters(R, p, b, J, d, lo, mid, up, qw)
                PP.append(np.full(J.size, pair_idx[p])); JJ.append(J); PN.append(pn); XB.append(xb)
                for g in range(4):
                    for q in range(len(masks[g])):
                        MK[g][q].append(masks[g][q])
            if not PP:
                continue
            pp = np.concatenate(PP); jj = np.concatenate(JJ); o_ = np.lexsort((pp, jj))
            pp, jj = pp[o_].astype(np.int64), jj[o_].astype(np.int64); pn = np.concatenate(PN)[o_]; xb = np.concatenate(XB)[o_]
            MK = [[np.concatenate(x)[o_] for x in grp] for grp in MK]
            for si in range(len(SLOPE)):
                for hi in range(len(HOURS)):
                    for wi in range(len(WIDTH)):
                        for ri in range(len(TREND)):
                            m = MK[0][si] & MK[1][hi] & MK[2][wi] & MK[3][ri]
                            if m.sum() < 30:
                                continue
                            res = fp.evaluate_tr(pp[m], jj[m], pn[m], xb[m], R.per, NPER, eu)
                            keep = res[:, 0, 0] >= 30
                            if not keep.any():
                                continue
                            k = int(keep.sum())
                            key = np.column_stack([np.full((k, 6), [bi, ti, di, si, hi, wi], np.float32), ek[keep], np.full(k, ri, np.float32)])
                            out.append(np.column_stack([key, res[keep, 0, :].astype(np.float32), res[keep, 1, :2].astype(np.float32)]))
    return np.concatenate(out) if out else np.zeros((0, len(COLS)), np.float32)


COLS = ['b', 'trig', 'dir', 'slope', 'hours', 'width', 'e', 'trend'] + fs.FIELDS + ['eu_n', 'eu_suma']

_G = {}


def _init(frames, pairs, spread):
    _G['R'] = F2.Rynek(frames, pairs, spread)


def _task(args):
    bi, shift = args
    return scan_band(_G['R'], bi, shifts=shift)


def search(frames, pairs=PAIRS, spread=None, procesy=None, shift=None, log=print) -> pd.DataFrame:
    import multiprocessing as mp
    spread = spread or fs.spreads()
    procesy = procesy or max(1, (os.cpu_count() or 2) - 1)
    t0 = time.time(); parts = []; nrow = 0
    tasks = [(bi, shift) for bi in range(len(BANDS))]

    def take(k, r):
        nonlocal nrow
        parts.append(r); nrow += len(r)
        if log and (k % 10 == 9 or k == len(tasks) - 1):
            log(f'    kanały {k + 1}/{len(BANDS)}, {nrow:,} reguł z ≥ 30 transakcjami, {time.time() - t0:.0f} s')
    if procesy == 1:
        _init(frames, pairs, spread)
        for k, r in enumerate(map(_task, tasks)):
            take(k, r)
    else:
        with mp.get_context('spawn').Pool(procesy, initializer=_init, initargs=(frames, pairs, spread)) as pool:
            for k, r in enumerate(pool.imap_unordered(_task, tasks)):
                take(k, r)
    D = pd.DataFrame(np.concatenate(parts) if parts else np.zeros((0, len(COLS)), np.float32), columns=COLS)
    for c_ in ['b', 'trig', 'dir', 'slope', 'hours', 'width', 'e', 'trend']:
        D[c_] = D[c_].astype(np.int32)
    return D.sort_values(['b', 'trig', 'dir', 'slope', 'hours', 'width', 'trend', 'e'], kind='stable').reset_index(drop=True)


def pool_mask(D):
    return ((D.n >= POOL['n']) & (D.stratne_okresy <= POOL['stratne']) & (D.t >= POOL['t']) & (D.PF >= POOL['pf']) &
            (D.suma1 > 0) & (D.suma2 > 0))


def chrono(D, top=(1, 10, 50)):
    """Wybór wyłącznie na roku 1 (okresy 1–10): ≥ 150 transakcji, ≤ 2 stratne okresy z 10, zysk; ranking suma1 / obsunięcie1 →
    wynik tych reguł w roku 2."""
    s = D[(D.n1 >= 150) & (D.stratne1 <= 2) & (D.suma1 > 0)].copy()
    if not len(s):
        return dict(spelnia=0), s
    s['ocena1'] = s.suma1 / np.maximum(s.obs1, 0.05)
    s = s.sort_values('ocena1', ascending=False)
    out = dict(spelnia=len(s))
    for k in top:
        h = s.head(k)
        out[f'top{k}_suma2'] = float(h.suma2.mean()); out[f'top{k}_na_tr2'] = float(h.suma2.sum() / max(h.n2.sum(), 1))
        out[f'top{k}_dodatnie'] = float((h.suma2 > 0).mean())
    return out, s


def describe(r) -> dict:
    return dict(kanal=band_name(BANDS[int(r.b)]), wejscie=TRIG[int(r.trig)], kierunek=DIRS[int(r.dir)][1], nachylenie=SLOPE[int(r.slope)],
                godziny=HOURS[int(r.hours)][0], szerokosc=WIDTH[int(r.width)], trend=TREND[int(r.trend)], wyjscie=exit_name(int(r.e)))


# ---------------------------------------------------------------- transakcje jednej reguły i rachunek

def rule_trades(R, rule: dict, centered=False) -> pd.DataFrame:
    """Lista transakcji reguły (jedna pozycja na parę). rule: b, trig, dir, slope, hours, width, e (indeksy)."""
    b = BANDS[rule['b']]; d = DIRS[rule['dir']][0]; e = rule['e']
    W = np.concatenate([((bands(R.F[p], b)[2] - bands(R.F[p], b)[0]) / R.F[p]['atr'])[(R.per < GRANICA_PER)] for p in R.pairs])
    W = W[np.isfinite(W)]
    qw = (np.quantile(W, 1 / 3), np.quantile(W, 2 / 3))
    rows = []
    for p in R.pairs:
        Fp = R.F[p]; lo, mid, up = bands(Fp, b, centered)
        J = events(Fp['o'], Fp['h'], Fp['l'], Fp['c'], lo, mid, up, Fp['valid'], rule['trig'], d, fs.WARMUP)
        if not J.size:
            continue
        ms, mh, mw, mt = _filters(R, p, b, J, d, lo, mid, up, qw)
        J = J[ms[rule['slope']] & mh[rule['hours']] & mw[rule['width']] & mt[rule.get('trend', 0)]]
        if not J.size:
            continue
        pn, xb = outcomes(Fp['o'], Fp['h'], Fp['l'], Fp['c'], J, float(d), Fp['atr'][J], lo, mid, up, R.spread[p] * F2.pip_of(p),
                          EX_TM[e:e + 1], EX_TV[e:e + 1], EX_SL[e:e + 1], EX_H[e:e + 1])
        last = -1
        for k, j in enumerate(J):
            if j <= last or np.isnan(pn[k, 0]):
                continue
            x = int(xb[k, 0]); last = j + x
            en = float(Fp['o'][j + 1])
            rows.append(dict(para=p, j=int(j), t=pd.Timestamp(int(R.t[j]), unit='s'), kierunek='short' if d < 0 else 'long', wejscie=en,
                             wynik=float(pn[k, 0]), swiece=x, okres=int(R.per[j]) + 1, linia=float(up[j] if d < 0 else lo[j]), srodek=float(mid[j])))
    D = pd.DataFrame(rows)
    return D.sort_values(['j', 'para']).reset_index(drop=True) if len(D) else D


def konto(D, R, f=0.20, lev=30.0, start=1000.0, stop_out=0.5):
    """1000 $ na 1:30, depozyt = f × saldo, jedna pozycja na parę, wynik po zamknięciu (bez zaokrąglania do lota)."""
    cash = start; open_ = {}; eq = np.full(R.n, np.nan); by = {}
    for k, r in enumerate(D.itertuples()):
        by.setdefault(r.j + 1, []).append(r)
    so = 0
    for j in range(R.n):
        for r in by.get(j, []):
            if r.para in open_:
                continue
            used = sum(q['m'] for q in open_.values())
            m = min(f * cash, cash - used)
            if m <= 0:
                continue
            open_[r.para] = dict(m=m, notional=m * lev, j1=j + r.swiece - 1, pnl=r.wynik, en=r.wejscie, d=-1 if r.kierunek == 'short' else 1)
        if open_:
            worst = sum(q['notional'] * min(0.0, q['d'] * ((R.F[p]['l'][j] if q['d'] > 0 else R.F[p]['h'][j]) - q['en']) / q['en']) for p, q in open_.items())
            used = sum(q['m'] for q in open_.values())
            if cash + worst < stop_out * used:
                cash += worst; open_.clear(); so += 1
        for p in list(open_):
            if open_[p]['j1'] <= j:
                cash += open_[p]['notional'] * open_[p]['pnl'] / 100; del open_[p]
        eq[j] = cash + sum(q['notional'] * q['d'] * (R.F[p]['c'][j] - q['en']) / q['en'] for p, q in open_.items())
    cv = pd.Series(eq, index=pd.to_datetime(R.t, unit='s')).ffill().fillna(start)
    pk = cv.cummax()
    return dict(koniec=float(cv.iloc[-1]), zysk_pct=float(cv.iloc[-1] / start * 100 - 100), obs_pct=float(((pk - cv) / pk).max() * 100), stop_out=so), \
        cv.resample('D').last().dropna()


def tstats(D) -> dict:
    if not len(D):
        return dict(n=0)
    w = D.wynik.values; n = len(w); sd = w.std(ddof=1) if n > 1 else 0
    per = D.groupby('okres').wynik.sum().reindex(range(1, NPER + 1), fill_value=0)
    return dict(n=n, srednia=w.mean(), mediana=float(np.median(w)), t=w.mean() / sd * np.sqrt(n) if sd > 0 else 0, trafne=(w > 0).mean(),
                PF=w[w > 0].sum() / max(-w[w < 0].sum(), 1e-9), najgorsza=w.min(), stratne_okresy=int((per < 0).sum()),
                rok1=float(w[D.okres.values <= GRANICA_PER].sum()), rok2=float(w[D.okres.values > GRANICA_PER].sum()),
                swiece=float(np.median(D.swiece)), long=int((D.kierunek == 'long').sum()), short=int((D.kierunek == 'short').sum()))


# ---------------------------------------------------------------- złudzenie: kanał „z przyszłością”

def illusion(R, log=print) -> pd.DataFrame:
    """Proste reguły (dotknięcie → TP środek, bez SL, 48 św.) na kanale wycentrowanym (zna przyszłość) i zwykłym."""
    rows = []
    e = EXITS.index((0, 0.0, 48))
    for bi, b in enumerate(BANDS):
        if b not in [('KC', 50, 2.0), ('KC', 100, 3.0), ('BB', 50, 2.0), ('BB', 100, 2.0), ('EN', 120, 24.0), ('EN', 240, 24.0), ('EN', 240, 72.0), ('DC', 120, 0.0)]:
            continue
        for di in range(2):
            rule = dict(b=bi, trig=0, dir=di, slope=0, hours=0, width=0, trend=0, e=e)
            for cen in (True, False):
                s = tstats(rule_trades(R, rule, centered=cen))
                rows.append(dict(kanal=band_name(b), kierunek=DIRS[di][1], wersja='wycentrowany — zna przyszłość' if cen else 'zwykły — tylko przeszłość',
                                 **{k: s.get(k) for k in ['n', 'srednia', 't', 'trafne', 'najgorsza', 'stratne_okresy']}))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- całość

def compute(frames, procesy=None, log=print, placebo_seeds=(1, 2), dp_frames=None) -> dict:
    t0 = time.time()
    R = F2.Rynek(frames, PAIRS, fs.spreads())
    O = dict(od=pd.Timestamp(int(R.t[0]), unit='s'), do=pd.Timestamp(int(R.t[-1]), unit='s'), swiec=R.n, spread=R.spread, periods=R.periods)
    log(f'Złudzenie: kanał wycentrowany (zna przyszłość) i zwykły…')
    O['zludzenie'] = illusion(R, log)
    log(f'Przeszukanie: {len(BANDS)} kanałów × {len(TRIG)} wejść × 2 kierunki × {NF} filtrów × {NE} wyjść = {NR:,} reguł')
    D = search(frames, procesy=procesy, log=log)
    O['n_regul'] = NR
    O['n_ocenione'] = len(D)
    pm = pool_mask(D)
    O['pula'] = int(pm.sum())
    log(f'  reguł z ≥ 30 transakcjami: {len(D):,}; stabilnych (≥ {POOL["n"]} transakcji, ≤ {POOL["stratne"]} stratne okresy z 20, t ≥ {POOL["t"]:g}, '
        f'PF ≥ {POOL["pf"]:g}, oba lata na plusie): {O["pula"]:,}')
    P = D[pm].sort_values('t', ascending=False)
    O['ranking'] = P.head(300).reset_index(drop=True)
    O['ranking_n'] = P[P.t >= POOL['t']].sort_values('n', ascending=False).head(50).reset_index(drop=True)
    O['chrono'], ch = chrono(D)
    O['chrono_top'] = ch.head(50).reset_index(drop=True)
    # najlepsze wg kierunku i kanału (przegląd)
    O['przeglad'] = (D[pm].sort_values('t', ascending=False).groupby(['b', 'dir']).head(1).sort_values('t', ascending=False).reset_index(drop=True)
                     if pm.any() else D.head(0))
    # statystyki dla placebo
    O['placebo'] = []
    for sd in placebo_seeds:
        log(f'Placebo {sd}: te same wyjścia, wejście w losowej chwili…')
        Dp = search(frames, procesy=procesy, shift=sd, log=None)
        cp, _ = chrono(Dp)
        O['placebo'].append(dict(seed=sd, ocenione=len(Dp), pula=int(pool_mask(Dp).sum()), max_t=float(Dp.t.max()) if len(Dp) else 0,
                                 t_ge5=int((Dp.t >= 5).sum()), **cp))
        del Dp
    O['real'] = dict(ocenione=len(D), pula=O['pula'], max_t=float(D.t.max()) if len(D) else 0, t_ge5=int((D.t >= 5).sum()), **O['chrono'])
    # strategia: najwyższe t w puli (wybór na całych 2 latach — w próbie); sąsiedzi, rachunek, pary
    if len(P):
        best = P.iloc[0]
        rule = {k: int(best[k]) for k in RULE_KEYS}
        O['rule'] = rule; O['rule_desc'] = describe(best)
        T = rule_trades(R, rule)
        O['trades'] = T; O['stats'] = tstats(T)
        O['konto'] = {}; O['krzywe'] = {}
        for f in (0.10, 0.20, 0.33):
            st, cv = konto(T, R, f=f)
            O['konto'][f] = st; O['krzywe'][f] = cv
        O['pary'] = pd.DataFrame([dict(para=p, **tstats(T[T.para == p])) for p in PAIRS])
        O['okresy'] = T.groupby('okres').wynik.agg(['size', 'sum']).reindex(range(1, NPER + 1), fill_value=0).reset_index()
        O['sasiedzi'] = neighbors(D, best)
        save_rule(rule, O)
        if dp_frames:
            O['dp'] = double_proof(dp_frames, rule, log)
    O['czas_min'] = (time.time() - t0) / 60
    return O


def neighbors(D, best) -> pd.DataFrame:
    """Odporność: ta sama reguła z jedną zmianą (kanał tego samego typu, wejście, filtr, wyjście)."""
    rows = []
    key = {k: int(best[k]) for k in RULE_KEYS}
    for k in key:
        q = D
        for k2, v in key.items():
            if k2 != k:
                q = q[q[k2] == v]
        if k == 'b':
            q = q[[BANDS[int(x)][0] == BANDS[key['b']][0] for x in q.b]]
        for r in q.itertuples():
            rows.append(dict(zmiana=k, **describe(r), n=r.n, srednia=r.srednia, t=r.t, stratne_okresy=r.stratne_okresy, suma1=r.suma1, suma2=r.suma2,
                             to_ona=all(int(getattr(r, k2)) == v for k2, v in key.items())))
    return pd.DataFrame(rows)


def save_rule(rule, O):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = dict(rule=rule, opis=O['rule_desc'], wersja=__version__, zamrozona=datetime.now(timezone.utc).isoformat(timespec='seconds'),
               dane=f"{O['od']:%Y-%m-%d} – {O['do']:%Y-%m-%d}")
    path = OUT_DIR / 'strategia.json'
    if not path.exists():                                   # pierwsza zamrożona zostaje (zmiana — tylko ręcznie, za zgodą)
        path.write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding='utf-8')


def double_proof(dp_frames, rule_now, log=print):
    """Sprawdzian zamrożonej reguły (data/fxkanal/strategia.json) na 10 parach fx1h_double_proof — każde otwarcie zapisane."""
    path = OUT_DIR / 'strategia.json'
    frozen = json.loads(path.read_text(encoding='utf-8')) if path.exists() else dict(rule=rule_now)
    rule = frozen['rule']
    Rd = F2.Rynek(dp_frames, F2.DP_PAIRS, F2.DP_SPREAD)
    logp = OUT_DIR / 'doubleproof.jsonl'
    n_old = sum(1 for _ in logp.open(encoding='utf-8')) if logp.exists() else 0
    with logp.open('a', encoding='utf-8') as fh:
        fh.write(json.dumps(dict(otwarcie=n_old + 1, czas=datetime.now(timezone.utc).isoformat(timespec='seconds'), regula=rule,
                                 do=str(pd.Timestamp(int(Rd.t[-1]), unit='s')), wersja=__version__), ensure_ascii=False) + '\n')
    T = rule_trades(Rd, rule)
    s = tstats(T)
    ok = s.get('n', 0) >= 30 and s.get('srednia', 0) > 0 and s.get('t', 0) >= 2
    log(f"doubleProof (otwarcie nr {n_old + 1}): {s.get('n', 0)} transakcji, średnio {s.get('srednia', 0):+.3f}% (t {s.get('t', 0):.1f}) → "
        f"{'ZALICZONY' if ok else 'NIE zaliczony'}")
    st, cv = konto(T, Rd) if len(T) else ({}, None)
    return dict(stats=s, ok=ok, trades=T, konto=st, krzywa=cv, otwarcie=n_old + 1, rule=rule, opis=frozen.get('opis'),
                pary=pd.DataFrame([dict(para=p, **tstats(T[T.para == p])) for p in F2.DP_PAIRS]) if len(T) else None)


# ---------------------------------------------------------------- wykres-ilustracja (PNG)

def picture(frames, path, pair='EURUSD', b=('EN', 120, 24.0), days=110):
    """Kanał wycentrowany (jak rysunek ręką) i zwykły na tym samym fragmencie wykresu."""
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    R = F2.Rynek(frames, PAIRS, fs.spreads())
    Fp = R.F[pair]; n = R.n; i0 = n - days * 24
    lo_c, mid_c, up_c = bands(Fp, b, centered=True); lo, mid, up = bands(Fp, b)
    x = np.arange(i0, n)
    fig, ax = plt.subplots(figsize=(14, 5.2), dpi=110)
    fig.patch.set_facecolor('#0b1220'); ax.set_facecolor('#0b1220')
    ax.plot(x, Fp['c'][i0:n], color='#c9d4e3', lw=0.7, label='kurs (zamknięcia 1h)')
    ax.plot(x, up_c[i0:n], color='#3ddc4a', lw=2.4, label='górna — wycentrowana (zna przyszłość, jak rysunek ręką)')
    ax.plot(x, lo_c[i0:n], color='#ffd400', lw=2.4, label='dolna — wycentrowana')
    ax.plot(x, up[i0:n], color='#3ddc4a', lw=1.2, ls='--', label='górna — zwykła (tylko przeszłość, na żywo)')
    ax.plot(x, lo[i0:n], color='#ffd400', lw=1.2, ls='--', label='dolna — zwykła')
    ax.set_title(f'{pair}: ten sam kanał ({band_name(b)}) — z przyszłością i bez', color='#e8eef6', fontsize=12)
    ax.tick_params(colors='#7f93a8'); [s.set_color('#24364a') for s in ax.spines.values()]
    ticks = np.linspace(i0, n - 1, 8).astype(int)
    ax.set_xticks(ticks); ax.set_xticklabels([pd.Timestamp(int(R.t[i]), unit='s').strftime('%d.%m') for i in ticks])
    ax.legend(loc='upper left', fontsize=8, facecolor='#0b1220', edgecolor='#24364a', labelcolor='#c9d4e3')
    fig.tight_layout(); fig.savefig(path, facecolor=fig.get_facecolor()); plt.close(fig)
    return path


# ---------------------------------------------------------------- Excel

def write_xlsx(O: dict, path: Path, wnioski=None, png=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.drawing.image import Image as XImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    FN = 'Arial'
    fN = Font(name=FN, size=10); fB = Font(name=FN, size=10, bold=True); fT = Font(name=FN, size=14, bold=True)
    fH = Font(name=FN, size=10, bold=True, color='FFFFFF'); fS = Font(name=FN, size=9, italic=True, color='555555')
    HDR = PatternFill('solid', fgColor='1F3A5F'); GRY = PatternFill('solid', fgColor='EEF2F7'); SEL = PatternFill('solid', fgColor='FFF2CC')
    GOOD = PatternFill('solid', fgColor='E2F0D9'); BAD = PatternFill('solid', fgColor='FBE5D6')
    thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')
    PCT = '0.0%'; PCT3 = '0.000%'; USD = '#,##0.00 [$$-409];[Red]-#,##0.00 [$$-409]'; DT = 'yyyy-mm-dd hh:mm'; DD = 'yyyy-mm-dd'
    wb = Workbook()

    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name; ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def hdr(ws, r, c, names, widths=None):
        for k, nm in enumerate(names):
            x = ws.cell(r, c + k, nm); x.font = fH; x.fill = HDR; x.border = BOX
            x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for k, w in enumerate(widths or []):
            if w: ws.column_dimensions[CL(c + k)].width = w
        ws.row_dimensions[r].height = 40

    def put(ws, r, c, v, font=fN, fmt=None, fill=None, al=None):
        if isinstance(v, np.floating): v = float(v)
        if isinstance(v, np.integer): v = int(v)
        if isinstance(v, float) and not np.isfinite(v): v = None
        x = ws.cell(r, c, v); x.font = font; x.border = BOX
        if fmt: x.number_format = fmt
        if fill: x.fill = fill
        if al: x.alignment = al
        return x

    def note(ws, r, c, text, span=10, h=None):
        x = ws.cell(r, c, text); x.font = fS; x.alignment = WR
        if span > 1: ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        if h: ws.row_dimensions[r].height = h

    DESC = ['kanal', 'wejscie', 'kierunek', 'nachylenie', 'godziny', 'szerokosc', 'trend', 'wyjscie']
    DESCN = ['Kanał', 'Wejście', 'Kierunek', 'Nachylenie środka', 'Godziny', 'Szerokość', 'Trend długi', 'Wyjście']
    ST = [('n', 'Transakcje', '0'), ('srednia', 'Średnio % (netto)', PCT3), ('t', 't', '0.0'), ('trafne', 'Trafne', PCT), ('PF', 'PF', '0.00'),
          ('stratne_okresy', 'Okresy stratne / 20', '0'), ('suma1', 'Rok 1: suma %', PCT), ('suma2', 'Rok 2: suma %', PCT), ('obsuniecie', 'Obsunięcie (suma %)', PCT),
          ('swiece', 'Średnio świec', '0.0'), ('eu_n', 'EURUSD: n', '0'), ('eu_suma', 'EURUSD: suma %', PCT)]

    def rule_table(ws, r0, D, title, note_txt=None, maxn=300):
        ws.cell(r0, 1, title).font = fB
        if note_txt: note(ws, r0 + 1, 1, note_txt, 16, 30); r0 += 1
        hdr(ws, r0 + 1, 1, ['Nr'] + DESCN + [x[1] for x in ST], [5, 26, 22, 16, 22, 15, 18, 22, 34] + [10] * len(ST))
        for k, r in enumerate(D.head(maxn).itertuples()):
            row = r0 + 2 + k; dsc = describe(r)
            put(ws, row, 1, k + 1, fmt='0')
            for c, key in enumerate(DESC): put(ws, row, 2 + c, dsc[key])
            for c, (key, _, fm) in enumerate(ST):
                v = getattr(r, key)
                if fm in (PCT, PCT3) and key not in ('trafne',): v = float(v) / 100
                put(ws, row, 2 + len(DESC) + c, v, fmt=fm)
        return r0 + 3 + min(len(D), maxn)

    s = O.get('stats') or {}
    # PODSUMOWANIE
    ws = sheet('PODSUMOWANIE', first=True)
    ws.column_dimensions['A'].width = 44
    for col in 'BCDEFGH': ws.column_dimensions[col].width = 15
    ws['A1'] = 'FX-KANAŁ — short przy górnej linii, long przy dolnej (obserwacja z 21.10.2026)'; ws['A1'].font = fT
    ws['A2'] = (f"IA4 {__version__} · badanie informacyjne · 10 par 1h: {O['od']:%Y-%m-%d} – {O['do']:%Y-%m-%d} ({O['swiec']} świec) · "
                f"policzono {pd.Timestamp.now():%Y-%m-%d %H:%M} ({O.get('czas_min', 0):.0f} min)"); ws['A2'].font = fS
    r = 4
    put(ws, r, 1, 'WNIOSKI' if wnioski else 'JAK CZYTAĆ', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    txt = wnioski or ['ZŁUDZENIE — dlaczego linie narysowane na wykresie wyglądają idealnie, a na żywo nie.', 'STRATEGIA — najlepsza stabilna reguła (w próbie).',
                      'UCZCIWY TEST — placebo i wybór na 1. roku.', 'RANKING — reguły stabilne wg t.']
    for k, s_ in enumerate(txt):
        x = ws.cell(r + 1 + k, 1, s_); x.font = fN; x.alignment = WR
        ws.merge_cells(start_row=r + 1 + k, start_column=1, end_row=r + 1 + k, end_column=8); ws.row_dimensions[r + 1 + k].height = 58 if wnioski else 16
    r += 2 + len(txt)
    put(ws, r, 1, 'PRZESZUKANIE', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    re_, pl = O['real'], O['placebo']
    lines = [('Reguł policzonych', f"{O['n_regul']:,}".replace(',', ' ')), ('Reguł z ≥ 30 transakcjami', f"{O['n_ocenione']:,}".replace(',', ' ')),
             (f"Stabilnych (≥ {POOL['n']} tr., ≤ {POOL['stratne']} stratne okresy z 20, t ≥ {POOL['t']:g}, PF ≥ {POOL['pf']:g}, oba lata +)",
              f"{re_['pula']:,}".replace(',', ' ') + ' — placebo: ' + ', '.join(f"{p_['pula']:,}".replace(',', ' ') for p_ in pl)),
             ('Najwyższe t', f"{re_['max_t']:.1f} — placebo: " + ', '.join(f"{p_['max_t']:.1f}" for p_ in pl)),
             ('Wybór na 1. roku → 2. rok (10 najlepszych, średnia suma %)', f"{re_.get('top10_suma2', float('nan')):+.2f} — placebo: " +
              ', '.join(f"{p_.get('top10_suma2', float('nan')):+.2f}" for p_ in pl))]
    for k, (a, b_) in enumerate(lines):
        put(ws, r + 1 + k, 1, a, fB, fill=GRY, al=WR); put(ws, r + 1 + k, 2, b_); ws.merge_cells(start_row=r + 1 + k, start_column=2, end_row=r + 1 + k, end_column=8)
        ws.row_dimensions[r + 1 + k].height = 28
    r += 2 + len(lines)
    if O.get('rule_desc'):
        put(ws, r, 1, 'STRATEGIA FX-KANAŁ (najwyższe t wśród stabilnych — wybór na całych 2 latach, w próbie)', fH, fill=HDR)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        for k, (key, lab) in enumerate(zip(DESC, DESCN)):
            put(ws, r + 1 + k, 1, lab, fB, fill=GRY); put(ws, r + 1 + k, 2, O['rule_desc'][key]); ws.merge_cells(start_row=r + 1 + k, start_column=2, end_row=r + 1 + k, end_column=8)
        r += 2 + len(DESC)
        vals = [('Transakcje (long / short)', f"{s['n']} ({s['long']} / {s['short']})"), ('Średnio na transakcję', f"{s['srednia']:+.3f}% (mediana {s['mediana']:+.3f}%)"),
                ('t', f"{s['t']:.1f}"), ('Trafne / PF', f"{s['trafne']:.0%} / {s['PF']:.2f}"), ('Najgorsza transakcja', f"{s['najgorsza']:+.2f}%"),
                ('Okresy stratne z 20', str(s['stratne_okresy'])), ('Rok 1 / rok 2 (suma %, bez dźwigni)', f"{s['rok1']:+.1f}% / {s['rok2']:+.1f}%"),
                ('Mediana świec w pozycji', f"{s['swiece']:.0f}")]
        vals += [(f"1000 $, 1:30, depozyt {int(f * 100)}%", f"{st['koniec']:,.0f} $ (obsunięcie {st['obs_pct']:.1f}%, stop out {st['stop_out']})".replace(',', ' '))
                 for f, st in O['konto'].items()]
        if O.get('dp'):
            ds = O['dp']['stats']
            vals.append(('DOUBLEPROOF (10 nowych par)', f"{'ZALICZONY' if O['dp']['ok'] else 'NIE ZALICZONY'}: {ds.get('n', 0)} transakcji, "
                                                       f"średnio {ds.get('srednia', 0):+.3f}%, t {ds.get('t', 0):.1f}"))
        for k, (a, b_) in enumerate(vals):
            put(ws, r + k, 1, a, fB, fill=GRY); put(ws, r + k, 2, b_); ws.merge_cells(start_row=r + k, start_column=2, end_row=r + k, end_column=8)

    # ZŁUDZENIE
    ws = sheet('ZŁUDZENIE')
    ws['A1'] = 'ZŁUDZENIE — kanał narysowany na gotowym wykresie zna przyszłość'; ws['A1'].font = fT
    note(ws, 2, 1, 'Ta sama reguła (dotknięcie linii → TP na środku kanału, bez SL, najwyżej 48 świec) na kanale wycentrowanym (każdy punkt linii liczony z ceny przed '
                   'I PO nim — tak wygląda linia rysowana ręką) i na zwykłym (tylko przeszłość — tak liczy się na żywo). Różnica to koszt opóźnienia średniej.', 12, 44)
    hdr(ws, 4, 1, ['Kanał', 'Kierunek', 'Wersja', 'Transakcje', 'Średnio %', 't', 'Trafne', 'Najgorsza %', 'Okresy stratne / 20'], [28, 18, 30, 11, 11, 8, 9, 11, 11])
    for k, v in enumerate(O['zludzenie'].itertuples()):
        cen = v.wersja.startswith('wycentr')
        for c, (val, fm) in enumerate([(v.kanal, None), (v.kierunek, None), (v.wersja, None), (v.n, '0'), ((v.srednia or 0) / 100, PCT3), (v.t, '0.0'),
                                       (v.trafne, PCT), ((v.najgorsza or 0) / 100, PCT3), (v.stratne_okresy, '0')]):
            x = put(ws, 5 + k, 1 + c, val, fmt=fm, fill=SEL if cen else None)
            if c == 4: x.fill = GOOD if (v.srednia or 0) > 0 else BAD
    if png and Path(png).exists():
        img = XImage(str(png)); img.width, img.height = 1100, 410
        ws.add_image(img, f'A{7 + len(O["zludzenie"])}')

    if O.get('trades') is not None and len(O['trades']):
        T = O['trades']
        ws = sheet('STRATEGIA')
        ws['A1'] = 'STRATEGIA — wszystkie transakcje (10 par, w próbie)'; ws['A1'].font = fT
        note(ws, 2, 1, ' · '.join(f'{lab}: {O["rule_desc"][key]}' for key, lab in zip(DESC, DESCN)), 10, 44)
        hdr(ws, 4, 1, ['Para', 'Sygnał (UTC)', 'Kierunek', 'Wejście', 'Linia kanału', 'Środek kanału', 'Świec', 'Wynik %', 'Okres'], [9, 17, 9, 11, 12, 12, 7, 10, 7])
        for k, v in enumerate(T.itertuples()):
            for c, (val, fm) in enumerate([(v.para, None), (v.t.to_pydatetime(), DT), (v.kierunek, None), (v.wejscie, '0.00000'), (v.linia, '0.00000'),
                                           (v.srodek, '0.00000'), (v.swiece, '0'), (v.wynik / 100, PCT3), (v.okres, '0')]):
                x = put(ws, 5 + k, 1 + c, val, fmt=fm)
                if c == 7: x.fill = GOOD if v.wynik > 0 else BAD
        n = len(T)
        put(ws, 6 + n, 7, 'Razem', fB, fill=GRY); put(ws, 6 + n, 8, f'=SUM(H5:H{4 + n})', fB, PCT)
        put(ws, 7 + n, 7, 'Średnio', fB, fill=GRY); put(ws, 7 + n, 8, f'=AVERAGE(H5:H{4 + n})', fB, PCT3)
        put(ws, 8 + n, 7, 'Trafne', fB, fill=GRY); put(ws, 8 + n, 8, f'=COUNTIF(H5:H{4 + n},">0")/{n}', fB, PCT)
        ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:I{4 + n}'

        ws = sheet('KAPITAŁ')
        ws['A1'] = 'KAPITAŁ — 1000 $, dźwignia 1:30, depozyt 10 / 20 / 33% salda na pozycję (bez zaokrąglania do lota)'; ws['A1'].font = fT
        cv = pd.DataFrame(O['krzywe']); cv.columns = [f'depozyt {int(f * 100)}%' for f in cv.columns]
        hdr(ws, 3, 1, ['Dzień'] + list(cv.columns), [12, 14, 14, 14])
        for k, (d, rr) in enumerate(cv.iterrows()):
            put(ws, 4 + k, 1, d.to_pydatetime(), fmt=DD)
            for q, c_ in enumerate(cv.columns): put(ws, 4 + k, 2 + q, round(float(rr[c_]), 2), fmt=USD)
        ch = LineChart(); ch.title = 'Saldo (1000 $ na start)'; ch.height = 10; ch.width = 24
        ch.y_axis.crossAx = 500; ch.x_axis = DateAxis(crossAx=100); ch.x_axis.number_format = 'yyyy-mm'; ch.x_axis.majorTimeUnit = 'months'
        ch.x_axis.delete = False; ch.y_axis.delete = False
        ch.add_data(Reference(ws, min_col=2, max_col=1 + len(cv.columns), min_row=3, max_row=3 + len(cv)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + len(cv))); ch.legend.position = 'b'
        ws.add_chart(ch, 'F3')

        ws = sheet('OKRESY')
        ws['A1'] = 'OKRESY — 20 równych części danych (1–10 = rok 1, 11–20 = rok 2)'; ws['A1'].font = fT
        hdr(ws, 3, 1, ['Okres', 'Od', 'Transakcje', 'Suma %'], [8, 12, 11, 11])
        for k, v in enumerate(O['okresy'].itertuples()):
            put(ws, 4 + k, 1, int(v.okres), fmt='0'); put(ws, 4 + k, 2, O['periods'][k][0].to_pydatetime(), fmt=DD)
            put(ws, 4 + k, 3, int(v.size), fmt='0'); x = put(ws, 4 + k, 4, float(v.sum) / 100, fmt=PCT); x.fill = GOOD if v.sum > 0 else BAD
        bc = BarChart(); bc.title = 'Suma % w okresie'; bc.height = 8; bc.width = 18; bc.legend = None
        bc.add_data(Reference(ws, min_col=4, min_row=3, max_row=23), titles_from_data=True); bc.set_categories(Reference(ws, min_col=1, min_row=4, max_row=23))
        bc.y_axis.number_format = '0%'; bc.x_axis.delete = False; bc.y_axis.delete = False; bc.x_axis.tickLblPos = 'low'
        ws.add_chart(bc, 'F3')

        ws = sheet('PARY')
        ws['A1'] = 'PARY — strategia na każdej parze'; ws['A1'].font = fT
        hdr(ws, 3, 1, ['Para', 'Transakcje', 'Średnio %', 't', 'Trafne', 'Najgorsza %', 'Rok 1 %', 'Rok 2 %'], [9, 10, 10, 8, 9, 11, 10, 10])
        for k, v in enumerate(O['pary'].itertuples()):
            if not getattr(v, 'n', 0): put(ws, 4 + k, 1, v.para); continue
            for c, (val, fm) in enumerate([(v.para, None), (v.n, '0'), (v.srednia / 100, PCT3), (v.t, '0.0'), (v.trafne, PCT), (v.najgorsza / 100, PCT3),
                                           (v.rok1 / 100, PCT), (v.rok2 / 100, PCT)]):
                put(ws, 4 + k, 1 + c, val, fmt=fm)

        ws = sheet('ODPORNOŚĆ')
        ws['A1'] = 'ODPORNOŚĆ — ta sama reguła z jedną zmianą (żółty = strategia)'; ws['A1'].font = fT
        note(ws, 2, 1, 'Strategia, która „działa”, ma podobne wyniki u sąsiadów (płaskowyż). Pojedynczy szczyt wśród słabych sąsiadów = dopasowanie do szumu.', 12, 30)
        cols = ['zmiana'] + DESC + ['n', 'srednia', 't', 'stratne_okresy', 'suma1', 'suma2']
        hdr(ws, 4, 1, ['Zmienione'] + DESCN + ['Transakcje', 'Średnio %', 't', 'Okresy stratne', 'Rok 1 %', 'Rok 2 %'], [11, 24, 20, 15, 20, 14, 16, 20, 32, 9, 9, 7, 8, 9, 9])
        NB = O['sasiedzi']
        for k, v in enumerate(NB.itertuples()):
            for c, key in enumerate(cols):
                val = getattr(v, key); fm = None
                if key in ('srednia', 'suma1', 'suma2'): val = float(val) / 100; fm = PCT3 if key == 'srednia' else PCT
                if key in ('n', 'stratne_okresy'): fm = '0'
                if key == 't': fm = '0.0'
                put(ws, 5 + k, 1 + c, val, fmt=fm, fill=SEL if v.to_ona else None)

    ws = sheet('UCZCIWY TEST')
    ws['A1'] = 'UCZCIWY TEST — czy przewaga to nie przypadek'; ws['A1'].font = fT
    note(ws, 2, 1, 'Placebo: dokładnie te same reguły i wyjścia, ale wejście w losowej chwili (tyle samo wejść na parę). Jeśli placebo daje podobnie dużo „stabilnych” '
                   'reguł albo podobne t — wynik bierze się z kształtu wyjścia (np. małe TP bez SL), nie z kanału. Wybór chronologiczny: reguły wybrane wyłącznie na roku 1 '
                   '(≥ 150 transakcji, ≤ 2 stratne okresy z 10, ranking zysk / obsunięcie) i ich wynik w roku 2.', 10, 58)
    hdr(ws, 4, 1, ['', 'Prawdziwe wejścia'] + [f'Placebo {p_["seed"]}' for p_ in O['placebo']], [52, 16, 16, 16])
    keys = [('ocenione', 'Reguł z ≥ 30 transakcjami', '#,##0'), ('pula', 'Stabilnych (warunek puli)', '#,##0'), ('max_t', 'Najwyższe t', '0.0'),
            ('t_ge5', 'Reguł z t ≥ 5', '#,##0'), ('spelnia', 'Wybór na roku 1: reguł spełniających warunek', '#,##0'),
            ('top1_suma2', 'Najlepsza z roku 1 → rok 2: suma %', '0.00'), ('top10_suma2', '10 najlepszych z roku 1 → rok 2: średnia suma %', '0.00'),
            ('top10_dodatnie', '10 najlepszych → rok 2: na plusie', PCT), ('top50_suma2', '50 najlepszych z roku 1 → rok 2: średnia suma %', '0.00'),
            ('top50_na_tr2', '50 najlepszych → rok 2: średnio na transakcję %', '0.000')]
    for k, (key, lab, fm) in enumerate(keys):
        put(ws, 5 + k, 1, lab, fB, fill=GRY)
        for c, src in enumerate([O['real']] + O['placebo']):
            put(ws, 5 + k, 2 + c, src.get(key), fmt=fm)
    rule_table(ws, 7 + len(keys), O['chrono_top'], 'Reguły wybrane na roku 1 (ranking) — kolumna „Rok 2” to sprawdzian', None, 50)

    ws = sheet('RANKING')
    ws['A1'] = 'RANKING — reguły stabilne (warunek puli), od najwyższego t — w próbie, zawyżone'; ws['A1'].font = fT
    rule_table(ws, 3, O['ranking'], f"{O['pula']:,} reguł spełnia warunek; pokazane {min(300, O['pula'])}".replace(',', ' '))
    ws.freeze_panes = 'B6'
    ws = sheet('NAJWIĘCEJ TRANSAKCJI')
    ws['A1'] = 'NAJWIĘCEJ TRANSAKCJI — reguły stabilne, od największej liczby transakcji'; ws['A1'].font = fT
    rule_table(ws, 3, O['ranking_n'], 'Warunek puli; sortowanie po liczbie transakcji', None, 50)
    ws = sheet('KANAŁY')
    ws['A1'] = 'KANAŁY — najlepsza stabilna reguła każdego kanału i kierunku'; ws['A1'].font = fT
    rule_table(ws, 3, O['przeglad'], 'Czy przewaga jest w wielu kanałach (dobry znak), czy w jednym (przypadek)?', None, 200)

    if O.get('dp'):
        dp = O['dp']; ds = dp['stats']
        ws = sheet('DOUBLEPROOF')
        ws['A1'] = f"DOUBLEPROOF — zamrożona reguła na 10 parach fx1h_double_proof (otwarcie nr {dp['otwarcie']})"; ws['A1'].font = fT
        note(ws, 2, 1, 'Warunek zaliczenia ustalony przed testem: ≥ 30 transakcji, średnia > 0, t ≥ 2. Reguła: data/fxkanal/strategia.json. Spread nowych par — szacunek.', 10, 30)
        put(ws, 4, 1, 'ZALICZONY' if dp['ok'] else 'NIE ZALICZONY', fB, fill=GOOD if dp['ok'] else BAD)
        for k, (a, b_) in enumerate([('Transakcje', ds.get('n')), ('Średnio %', ds.get('srednia')), ('t', ds.get('t')), ('Trafne', ds.get('trafne')),
                                     ('Okresy stratne / 20', ds.get('stratne_okresy')), ('1000 $ na 1:30, 20%', (dp.get('konto') or {}).get('koniec'))]):
            put(ws, 6 + k, 1, a, fB, fill=GRY); put(ws, 6 + k, 2, b_, fmt='0.000' if isinstance(b_, float) else None)
        if dp.get('pary') is not None:
            hdr(ws, 14, 1, ['Para', 'Transakcje', 'Średnio %', 't', 'Trafne'], [10, 10, 10, 8, 9])
            for k, v in enumerate(dp['pary'].itertuples()):
                if getattr(v, 'n', 0):
                    for c, (val, fm) in enumerate([(v.para, None), (v.n, '0'), (v.srednia / 100, PCT3), (v.t, '0.0'), (v.trafne, PCT)]):
                        put(ws, 15 + k, 1 + c, val, fmt=fm)
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- uruchomienie

def run(no_sync=False, out=None, procesy=None, double=False, log=print, keep=None, wnioski=None) -> Path:
    t0 = time.time()
    from .fx import sync as fxs
    if not no_sync:
        log('Kopia fx1h — odświeżenie 10 par (python -m ia4.fx.sync)')
        fxs.sync(quiet=True, syms=list(PAIRS) + (list(F2.DP_PAIRS) if double else []))
    frames = {p: fxs.load(p) for p in PAIRS}
    dp = None
    if double:
        dp = {p: fxs.load(p) for p in F2.DP_PAIRS if fxs.parquet_path(p).exists()}
        if len(dp) < len(F2.DP_PAIRS):
            log(f'doubleProof: brak kopii {len(F2.DP_PAIRS) - len(dp)} par — pomijam'); dp = None
        else:
            dp.update(frames)
    O = compute(frames, procesy=procesy, log=log, dp_frames=dp)
    if keep is not None:
        keep.update(O)
    png = OUT_DIR / 'zludzenie.png'
    try:
        picture(frames, png)
    except Exception as e:
        log(f'  ! obrazek: {e}'); png = None
    path = write_xlsx(O, out or (config.RESEARCH_DIR.parent / 'FX-KANAL.xlsx'), wnioski=wnioski, png=png)
    if O.get('stats'):
        s = O['stats']
        log(f"\nFX-KANAŁ: {s['n']} transakcji, średnio {s['srednia']:+.3f}% (t {s['t']:.1f}), okresy stratne {s['stratne_okresy']}/20 — "
            f"{O['rule_desc']['kanal']}, {O['rule_desc']['kierunek']}, {O['rule_desc']['wyjscie']}")
    log(f'Zapisano: {path}  ({(time.time() - t0) / 60:.0f} min)')
    return path


def selftest() -> bool:
    ok = True

    def check(name, cond):
        nonlocal ok
        print(f"  {'✓' if cond else '✗'} {name}")
        ok &= bool(cond)

    n = 40
    o = np.full(n, 1.0); h = np.full(n, 1.001); l = np.full(n, 0.999); c = np.full(n, 1.0)
    lo = np.full(n, 0.99); up = np.full(n, 1.01); mid = np.full(n, 1.0); valid = np.ones(n, bool)
    h2 = h.copy(); h2[10] = 1.012
    J = events(o, h2, l, c, lo, mid, up, valid, 0, -1, 0)
    check('dotknięcie górnej linii → sygnał short na tej świecy', list(J) == [10])
    c3 = c.copy(); c3[10] = 1.011; c3[11] = 1.005
    check('powrót do kanału → sygnał na świecy powrotu', list(events(o, h, l, c3, lo, mid, up, valid, 2, -1, 0)) == [11])
    l4 = l.copy(); l4[5] = 0.985
    check('dotknięcie dolnej linii → long', list(events(o, h, l4, c, lo, mid, up, valid, 0, 1, 0)) == [5])
    # wynik: short z wejściem 1,0 (open 11), TP środek 0,995 w 3. świecy
    o5 = np.full(n, 1.0); h5 = np.full(n, 1.001); l5 = np.full(n, 0.999); c5 = np.full(n, 1.0); l5[13] = 0.994
    mid5 = np.full(n, 0.995)
    pn, xb = outcomes(o5, h5, l5, c5, np.array([10]), -1.0, np.array([0.002]), lo, mid5, up, 0.0001,
                      np.array([0], np.int64), np.array([0.0]), np.array([0.0]), np.array([24], np.int64))
    check('short: TP na środku kanału w 3. świecy = (1 − 0,995) / 1 − spread', abs(pn[0, 0] - (0.5 - 0.01)) < 1e-4 and xb[0, 0] == 3)
    h6 = h5.copy(); h6[12] = 1.005
    pn, xb = outcomes(o5, h6, l5, c5, np.array([10]), -1.0, np.array([0.002]), lo, mid5, up, 0.0001,
                      np.array([0], np.int64), np.array([0.0]), np.array([2.0]), np.array([24], np.int64))
    check('short: SL 2 ATR (1,004) w 2. świecy, zanim TP', abs(pn[0, 0] - (-0.4 - 0.01)) < 1e-4 and xb[0, 0] == 2)
    l7 = l5.copy(); l7[12] = 0.99
    pn, xb = outcomes(o5, h6, l7, c5, np.array([10]), -1.0, np.array([0.002]), lo, mid5, up, 0.0001,
                      np.array([0], np.int64), np.array([0.0]), np.array([2.0]), np.array([24], np.int64))
    check('SL i TP w jednej świecy → SL (zasada 5)', abs(pn[0, 0] - (-0.41)) < 1e-4)
    pn, xb = outcomes(o5, h5, l5, c5, np.array([10]), 1.0, np.array([0.002]), lo, np.full(n, 1.004), up, 0.0001,
                      np.array([3], np.int64), np.array([1.0]), np.array([0.0]), np.array([5], np.int64))
    check('long: bez TP przez 5 świec → wyjście po close 5. świecy (spread)', abs(pn[0, 0] + 0.01) < 1e-4 and xb[0, 0] == 5)
    # kanał wycentrowany = zwykły przesunięty wstecz
    Fp = {'c': np.linspace(1, 2, 300), 'h': np.linspace(1, 2, 300) + 0.01, 'l': np.linspace(1, 2, 300) - 0.01, 'atr': np.full(300, 0.02)}
    a = bands(Fp, ('KC', 20, 2.0)); b_ = bands(Fp, ('KC', 20, 2.0), centered=True)
    check('kanał wycentrowany = zwykły przesunięty o połowę okna (zna przyszłość)', np.allclose(b_[1][:-10], a[1][10:], equal_nan=True))
    print('\nWYNIK: OK' if ok else '\nWYNIK: BŁĄD')
    return ok


def main():
    ap = argparse.ArgumentParser(description='FX-KANAŁ — short przy górnej linii kanału, long przy dolnej')
    ap.add_argument('--bez-sync', action='store_true', help='bez odświeżenia kopii fx1h')
    ap.add_argument('--procesy', type=int, default=None, help='liczba procesów (domyślnie wszystkie rdzenie − 1)')
    ap.add_argument('--double-proof', action='store_true', help='sprawdzian zamrożonej strategii na 10 parach fx1h_double_proof (otwarcie zapisane)')
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args()
    print(f'IA 4 {__version__} — FX-KANAŁ\n')
    if a.selftest:
        sys.exit(0 if selftest() else 1)
    run(no_sync=a.bez_sync, procesy=a.procesy, double=a.double_proof)


if __name__ == '__main__':
    main()
