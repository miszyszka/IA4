"""
IA 4 — One-Day: portfel strategii wybieranych krok po kroku w czasie (instrukcja, sekcja 7).
Wersja projektu: 1.59 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Decyzja człowieka z 15.10.2026 po wyniku 1.42 (spread 0,6%: żadna reguła nie spełniła warunku 20 okresów):
„10 razy więcej strategii, rzadsze strategie dozwolone, potem portfel wielu strategii naraz, bardziej rygorystycznie”.

1. Poszukiwanie (szerokie, z przycinaniem): kierunek + setup (82) + do 3 filtrów z różnych grup (33).
   Kombinacja z < 30 transakcjami nie jest liczona dalej (dodanie filtra tylko zmniejsza liczbę transakcji).
   Dla każdej reguły zapisywane są wyniki ROK PO ROKU (transakcje, suma i suma kwadratów dziennych wyników
   netto po spreadzie) — 53 spółki handlowane.
2. Wybór krok po kroku w czasie (walk-forward): na początku każdego roku Y program wybiera najwyżej 20 reguł
   WYŁĄCZNIE na danych sprzed Y (≥ 30 transakcji, suma netto > 0, t-stat ≥ 3, dodatni wynik w ≥ 60% lat
   z transakcjami, najwyżej 2 reguły z tego samego setupu) i gra nimi przez rok Y. Wynik portfela jest więc
   w całości poza próbą — żaden dzień nie był widziany przy wyborze strategii, które w nim handlują.
3. Portfel: w danym dniu i spółce głosy wybranych reguł się sumują (long +1, short −1); pozycja = znak sumy,
   jedna transakcja na spółkę i dzień, spread od każdej.
4. Ocena portfela: 20 okresów kalendarza (najwyżej 3 stratne — warunek człowieka), rok po roku, doubleProof
   (te same wybory, 20 spółek, których poszukiwanie nie widzi), placebo (cały proces z losowym znakiem dni),
   portfel 10 000 $ (5% stanu portfela na transakcję).

    python -m ia4.odport --cost 0.6          # baza dobowa od 2000 (najpierw python -m ia4.daily); ok. 30–60 min
    python -m ia4.odport --no-placebo        # bez placebo (o połowę krócej)
    python -m ia4.odport --selftest          # „WYNIK: OK”
Wynik: ~/Desktop/IA4/ONE-DAY-PORTFEL.xlsx (poza gitem).
"""

from __future__ import annotations

import argparse
import math
import os
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from numpy.lib.stride_tricks import sliding_window_view

from . import __version__, config
from . import odsearch as osr
from . import oneday as od

# ---------------------------------------------------------------------------
#  Parametry
# ---------------------------------------------------------------------------
COST_PCT = osr.COST_PCT     # spread na transakcję (wejście + wyjście); --cost
MIN_TRADES = 30             # reguła liczona dalej / wybierana tylko przy ≥ 30 transakcjach
MAX_FILTERS = 3
WF_MIN_YEARS = 5            # pierwszy rok portfela: po ≥ 5 latach historii w oknie
WF_K = 20                   # najwyżej tyle reguł w portfelu na rok
WF_T = 3.0                  # t-stat dziennych wyników netto w historii
WF_POS_YEARS = 0.6          # dodatnie lata (z ≥ 3 transakcjami) w historii
WF_FAMILY = 2               # najwyżej tyle reguł z tym samym setupem i kierunkiem
N_PERIODS = 20              # ocena portfela: 20 okresów …
MAX_LOSING = 3              # … najwyżej 3 stratne
TOP_RULES = 300             # arkusz REGUŁY (pełna próba — informacyjnie, zawyżone)
PF_CAPITAL = 10_000.0
PLACEBO_SEED = osr.PLACEBO_SEED


# ---------------------------------------------------------------------------
#  Dodatkowe cechy (znane przed open dnia t)
# ---------------------------------------------------------------------------
def extend(P: osr.Panel) -> None:
    f = P.feat
    f["body_p"] = f["body_r"] * f["r"]
    f["chg3p"] = f["chg3_r"] * f["r"]
    f["rel1p"] = f["chg1p"] - P.spy["chg1"][None, :]
    S, T = P.o.shape
    for k in ("low252", "high252"):
        f[k] = np.full((S, T), np.nan)
    for i in range(S):
        vi = np.flatnonzero(P.valid[i])
        if len(vi) < 252:
            continue
        c = P.c[i, vi]
        mn = np.full(len(c), np.nan)
        mx = np.full(len(c), np.nan)
        mn[251:] = sliding_window_view(c, 252).min(axis=1)
        mx[251:] = sliding_window_view(c, 252).max(axis=1)
        keep = vi + 1 < T
        src, tgt = np.flatnonzero(keep), vi[keep] + 1
        lo = np.where(np.isnan(mn), np.nan, (c <= mn).astype(float))
        hi = np.where(np.isnan(mx), np.nan, (c >= mx).astype(float))
        f["low252"][i, tgt], f["high252"][i, tgt] = lo[src], hi[src]
    # szerokość rynku: udział spółek ze spadkową świecą w poprzednim dniu kalendarza (≥ 10 spółek)
    dn = (P.c < P.o) & P.valid
    cnt = P.valid.sum(axis=0)
    share = np.where(cnt >= 10, dn.sum(axis=0) / np.maximum(cnt, 1), np.nan)
    P.breadth = np.concatenate([[np.nan], share[:-1]])
    # kalendarz (sesje znane z góry): początek / koniec miesiąca, dzień przed świętem
    ds = [date.fromisoformat(d) for d in P.dates]
    T_ = len(ds)
    ym = [(d.year, d.month) for d in ds]
    tom_s, tom_e, pre = np.zeros(T_, bool), np.zeros(T_, bool), np.zeros(T_, bool)
    k = 0
    for t in range(T_):
        k = k + 1 if t and ym[t] == ym[t - 1] else 0
        tom_s[t] = k <= 2
    k = 0
    for t in range(T_ - 1, -1, -1):
        k = k + 1 if t < T_ - 1 and ym[t] == ym[t + 1] else 0
        tom_e[t] = k <= 1 and t < T_ - 1
    for t in range(T_ - 1):
        gap = (ds[t + 1] - ds[t]).days
        wd = ds[t].weekday()
        pre[t] = gap > (3 if wd == 4 else 1)                    # między sesjami wypada dzień roboczy bez sesji
    P.cal = {"tom_s": tom_s, "tom_e": tom_e, "pre": pre}


def setups() -> list:
    C, g, l = osr.Cond, osr._gt, osr._lt
    out = list(osr.setups())
    for x in (3, 5, 8):
        out.append((C(f"bdp{x}", f"wczoraj spadek korpusu ≥ {x}%", "s"), lambda P, x=x: l(P.feat["body_p"], -x)))
        out.append((C(f"bup{x}", f"wczoraj wzrost korpusu ≥ {x}%", "s"), lambda P, x=x: g(P.feat["body_p"], x)))
        out.append((C(f"c1d{x}", f"wczoraj spadek ≥ {x}% (zamknięcie → zamknięcie)", "s"), lambda P, x=x: l(P.feat["chg1p"], -x)))
        out.append((C(f"c1u{x}", f"wczoraj wzrost ≥ {x}% (zamknięcie → zamknięcie)", "s"), lambda P, x=x: g(P.feat["chg1p"], x)))
    for x in (8, 12):
        out.append((C(f"c3d{x}", f"spadek 3 sesji ≥ {x}%", "s"), lambda P, x=x: l(P.feat["chg3p"], -x)))
        out.append((C(f"c3u{x}", f"wzrost 3 sesji ≥ {x}%", "s"), lambda P, x=x: g(P.feat["chg3p"], x)))
    for x in (3, 5):
        out.append((C(f"rw{x}", f"wczoraj słabsza od SPY o ≥ {x} pkt proc.", "s"), lambda P, x=x: l(P.feat["rel1p"], -x)))
        out.append((C(f"rs{x}", f"wczoraj mocniejsza od SPY o ≥ {x} pkt proc.", "s"), lambda P, x=x: g(P.feat["rel1p"], x)))
    out.append((C("low252", "zamknięcie najniższe z 252 sesji (dołek roczny)", "s"), lambda P: g(P.feat["low252"], 1)))
    out.append((C("high252", "zamknięcie najwyższe z 252 sesji (szczyt roczny)", "s"), lambda P: g(P.feat["high252"], 1)))
    return out


def filters() -> list:
    C, g, l, day = osr.Cond, osr._gt, osr._lt, osr._day
    out = list(osr.filters())
    out.append((C("r>5", "zmienność bezwzględna: r ≥ 5%", "va"), lambda P: g(P.feat["r"], 5)))
    out.append((C("br_dn", "wczoraj szeroka wyprzedaż (≥ 80% spółek spadkowych)", "br"), lambda P: day(P, g(P.breadth, 0.8))))
    out.append((C("br_up", "wczoraj szeroki wzrost (≤ 20% spółek spadkowych)", "br"),
                lambda P: day(P, l(np.nan_to_num(P.breadth, nan=1.0), 0.2))))
    out.append((C("spy-1.5", "SPY wczoraj ≤ −1,5%", "sm"), lambda P: day(P, l(P.spy["chg1"], -1.5))))
    out.append((C("spy-3", "SPY wczoraj ≤ −3%", "sm"), lambda P: day(P, l(P.spy["chg1"], -3))))
    out.append((C("spy+1.5", "SPY wczoraj ≥ +1,5%", "sm"), lambda P: day(P, g(P.spy["chg1"], 1.5))))
    out.append((C("spyr>1.5", "zmienność SPY ≥ 1,5% (średni zakres 20 sesji)", "sv"), lambda P: day(P, g(P.spy["r"], 1.5))))
    out.append((C("spyr>2.5", "zmienność SPY ≥ 2,5%", "sv"), lambda P: day(P, g(P.spy["r"], 2.5))))
    out.append((C("tom_s", "pierwsze 3 sesje miesiąca", "cal"), lambda P: day(P, P.cal["tom_s"])))
    out.append((C("tom_e", "ostatnie 2 sesje miesiąca", "cal"), lambda P: day(P, P.cal["tom_e"])))
    out.append((C("pre_hol", "sesja przed świętem", "cal"), lambda P: day(P, P.cal["pre"])))
    return out


@dataclass(frozen=True)
class Spec:
    dir: int
    s: int
    f: tuple


# ---------------------------------------------------------------------------
#  1. Poszukiwanie: wyniki rok po roku każdej reguły
# ---------------------------------------------------------------------------
def years_of(P: osr.Panel, g0: int) -> tuple[list[int], np.ndarray, np.ndarray]:
    ys = np.array([int(d[:4]) for d in P.dates[g0:]])
    uniq = sorted(set(ys.tolist()))
    starts = np.array([int(np.argmax(ys == y)) for y in uniq])
    days = np.diff(np.concatenate([starts, [len(ys)]]))
    return uniq, starts, days


def search(P: osr.Panel, rows: np.ndarray, g0: int, cost: float, sign: np.ndarray | None = None, log=print) -> dict:
    S, F = setups(), filters()
    R = (P.ret if sign is None else P.ret * sign[None, :])[rows][:, g0:].astype(np.float32)
    A = P.base[rows][:, g0:]
    smask = [fn(P)[rows][:, g0:] & A for _, fn in S]
    fmask = [np.ascontiguousarray(fn(P)[rows][:, g0:]) for _, fn in F]
    fgrp = [c.group for c, _ in F]
    years, ystart, ydays = years_of(P, g0)
    specs, YN, YS, YQ = [], [], [], []
    stats = {"masks": 0, "rules": 0}
    t0 = time.time()

    def record(m, si, fs):
        stats["rules"] += 2
        dd = (R * m).sum(axis=0, dtype=np.float64)
        dn = m.sum(axis=0)
        yn = np.add.reduceat(dn, ystart)
        for d in (1, -1):
            net = d * dd - cost * dn
            specs.append(Spec(d, si, fs))
            YN.append(yn)
            YS.append(np.add.reduceat(net, ystart))
            YQ.append(np.add.reduceat(net * net, ystart))

    def rec(m, si, fs, start):
        if len(fs) >= MAX_FILTERS:
            return
        used = {fgrp[j] for j in fs}
        for j in range(start, len(F)):
            if fgrp[j] in used:
                continue
            m2 = m & fmask[j]
            stats["masks"] += 1
            if m2.sum() < MIN_TRADES:
                continue
            record(m2, si, fs + (j,))
            rec(m2, si, fs + (j,), j + 1)

    for si in range(len(S)):
        m = smask[si]
        stats["masks"] += 1
        if m.sum() < MIN_TRADES:
            continue
        record(m, si, ())
        rec(m, si, (), 0)
        if log and (si + 1) % 10 == 0:
            log(f"    setupy {si + 1}/{len(S)} · reguł {stats['rules']} · {time.time() - t0:.0f} s")
    return {"specs": specs, "YN": np.array(YN, np.float32), "YS": np.array(YS, np.float64),
            "YQ": np.array(YQ, np.float64), "years": years, "ydays": ydays, "S": S, "F": F,
            "masks": stats["masks"], "rules": stats["rules"], "sec": time.time() - t0}


def tstat(s, q, d):
    """t-stat średniego dziennego wyniku (d dni, s — suma, q — suma kwadratów)."""
    mean = s / d
    var = np.maximum((q - s * s / d) / np.maximum(d - 1, 1), 1e-12)
    return mean / np.sqrt(var) * np.sqrt(d)


def select(res: dict, upto: int, counts: dict | None = None) -> list[int]:
    """Wybór na początku roku o indeksie `upto` (lata 0 … upto−1 = historia). Zwraca indeksy reguł."""
    YN, YS, YQ = res["YN"][:, :upto], res["YS"][:, :upto], res["YQ"][:, :upto]
    D = float(res["ydays"][:upto].sum())
    n, s, q = YN.sum(1), YS.sum(1), YQ.sum(1)
    t = tstat(s, q, D)
    act = YN >= 3
    pos = ((YS > 0) & act).sum(1) / np.maximum(act.sum(1), 1)
    ok = (n >= MIN_TRADES) & (s > 0) & (t >= WF_T) & (pos >= WF_POS_YEARS)
    if counts is not None:                     # ile reguł przechodzi kolejne progi (dla raportu)
        counts.update({"trades": int((n >= MIN_TRADES).sum()), "net>0": int(((n >= MIN_TRADES) & (s > 0)).sum()),
                       "t": int(((n >= MIN_TRADES) & (s > 0) & (t >= WF_T)).sum()), "all": int(ok.sum()),
                       "best_t": float(t[(n >= MIN_TRADES) & (s > 0)].max()) if ((n >= MIN_TRADES) & (s > 0)).any() else None})
    idx = np.flatnonzero(ok)
    idx = idx[np.argsort(-t[idx], kind="stable")]
    out, fam = [], {}
    for i in idx:
        sp = res["specs"][i]
        key = (sp.s, sp.dir)
        if fam.get(key, 0) >= WF_FAMILY:
            continue
        fam[key] = fam.get(key, 0) + 1
        out.append(int(i))
        if len(out) >= WF_K:
            break
    return out


def rule_of(res: dict, i: int) -> osr.Rule:
    sp = res["specs"][i]
    return osr.Rule(sp.dir, res["S"][sp.s][0], tuple(res["F"][j][0] for j in sp.f))


def mask_of(P: osr.Panel, res: dict, i: int, cache: dict) -> np.ndarray:
    """Pełna maska reguły (wszystkie spółki, cały kalendarz)."""
    sp = res["specs"][i]
    key = (sp.s, sp.f)
    if key not in cache:
        m = res["S"][sp.s][1](P) & P.base
        for j in sp.f:
            m = m & res["F"][j][1](P)
        cache[key] = m
    return cache[key]


# ---------------------------------------------------------------------------
#  2–3. Walk-forward i portfel
# ---------------------------------------------------------------------------
def walk_forward(P: osr.Panel, res: dict, g0: int, sign: np.ndarray | None = None) -> dict:
    years = res["years"]
    first = next((k for k, y in enumerate(years) if k >= WF_MIN_YEARS), None)
    S_, T = P.o.shape
    sgn = np.zeros((S_, T), np.int8)
    picks, counts = [], []
    cache: dict = {}
    if first is None:
        return {"sgn": sgn, "picks": picks, "counts": counts, "start": T, "first_year": None}
    ys = np.array([int(d[:4]) for d in P.dates])
    for k in range(first, len(years)):
        cnt: dict = {}
        sel = select(res, k, cnt)
        if len(cache) > 120:                                       # pamięć: maski 73 × T
            cache.clear()
        cols = np.flatnonzero((ys == years[k]) & (np.arange(T) >= g0))
        votes = np.zeros((S_, len(cols)), np.int16)
        for i in sel:
            votes += res["specs"][i].dir * mask_of(P, res, i, cache)[:, cols].astype(np.int16)
        sgn[:, cols] = np.sign(votes).astype(np.int8)
        picks.append((years[k], sel))
        counts.append((years[k], cnt))
    start = int(np.argmax(ys >= years[first]))
    return {"sgn": sgn, "picks": picks, "counts": counts, "start": max(start, g0), "first_year": years[first]}


def assess(P: osr.Panel, sgn: np.ndarray, rows: np.ndarray, start: int, cost: float, sign=None) -> dict:
    """Portfel na wierszach `rows` od `start`: 20 okresów, rok po roku, statystyki."""
    R = P.ret if sign is None else P.ret * sign[None, :]
    sub = sgn[rows].astype(float)
    m = sub != 0
    x = np.where(m, sub * R[rows] - cost, 0.0)
    T = P.o.shape[1]
    idx = np.array_split(np.arange(start, T), N_PERIODS)
    win = (start, [(int(a[0]), int(a[-1]) + 1) for a in idx if len(a)])
    old = osr.MAX_LOSING, osr.N_PERIODS
    osr.MAX_LOSING, osr.N_PERIODS = MAX_LOSING, N_PERIODS
    st = osr.evaluate_x(x, m, win, [-cost] * N_PERIODS, cost)
    osr.MAX_LOSING, osr.N_PERIODS = old
    ys = np.array([int(d[:4]) for d in P.dates])
    per_year = []
    for y in sorted(set(ys[start:].tolist())):
        cols = (ys == y) & (np.arange(T) >= start)
        n = int(m[:, cols].sum())
        s = float(x[:, cols].sum())
        w = int(((x > 0) & m)[:, cols].sum())
        per_year.append((y, n, s, s / n if n else 0.0, w / n * 100 if n else 0.0))
    return {"st": st, "win": win, "years": per_year}


# ---------------------------------------------------------------------------
#  Uruchomienie i raport
# ---------------------------------------------------------------------------
def pipeline(P: osr.Panel, cost: float, sign=None, log=print) -> dict:
    rows = np.array([i for i, g in enumerate(P.groups) if g in ("main", "proof")])
    dprows = np.array([i for i, g in enumerate(P.groups) if g == "doubleProof"], dtype=int)
    g0 = osr.periods(P, rows)[0]
    res = search(P, rows, g0, cost, sign, log)
    wf = walk_forward(P, res, g0, sign)
    out = {"res": res, "wf": wf, "rows": rows, "dprows": dprows, "g0": g0}
    if wf["first_year"] is not None:
        out["main"] = assess(P, wf["sgn"], rows, wf["start"], cost, sign)
        out["dp"] = assess(P, wf["sgn"], dprows, wf["start"], cost, sign) if len(dprows) else None
    return out


def verdict(a: dict | None) -> str:
    if not a:
        return "brak portfela"
    st = a["st"]
    return ("SPEŁNIA" if st.ok else "nie spełnia") + f" warunku {N_PERIODS} okresów (stratne {st.losing}, puste {st.empty})"


def full_sample_top(res: dict) -> list[tuple[int, float, float, float, int]]:
    """Najlepsze reguły na pełnej próbie (informacyjnie, zawyżone): (indeks, t, śr. netto, udział dodatnich lat, n)."""
    D = float(res["ydays"].sum())
    n, s, q = res["YN"].sum(1), res["YS"].sum(1), res["YQ"].sum(1)
    t = tstat(s, q, D)
    act = res["YN"] >= 3
    pos = ((res["YS"] > 0) & act).sum(1) / np.maximum(act.sum(1), 1)
    ok = (n >= MIN_TRADES) & (s > 0)
    idx = np.flatnonzero(ok)
    idx = idx[np.argsort(-t[idx], kind="stable")][:TOP_RULES]
    return [(int(i), float(t[i]), float(s[i] / n[i]), float(pos[i]), int(n[i])) for i in idx]


def write_xlsx(path: Path, P: osr.Panel, out: dict, plc: dict | None, cost: float, note: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="1F2A37")
    good, bad, sub = (PatternFill("solid", fgColor=c) for c in ("DCFCE7", "FEE2E2", "E5E7EB"))
    white, bold = Font(bold=True, color="FFFFFF"), Font(bold=True)
    pct = "0.00;[Red]-0.00"
    res, wf = out["res"], out["wf"]

    def header(ws, row, cols, widths=None):
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row, j, h)
            c.font, c.fill = white, head
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for j, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(j)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "PODSUMOWANIE"
    ws["A1"] = "One-Day — portfel strategii wybieranych krok po kroku w czasie (walk-forward)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"spread {cost:g}% na transakcję · reguł sprawdzonych {res['rules']} · portfel od {wf['first_year'] or '—'} · "
                f"{note}")
    r = 4
    blocks = [("53 spółki handlowane (wybór i gra)", out.get("main")),
              ("doubleProof — te same wybory, 20 spółek, których poszukiwanie nie widzi", out.get("dp"))]
    if plc is not None:
        blocks.append(("PLACEBO — cały proces z losowym znakiem dni (zero przewagi)", plc.get("main")))
    for title, a in blocks:
        ws.cell(r, 1, title).font = bold
        ws.cell(r, 4, verdict(a)).font = Font(bold=True, color="166534" if a and a["st"].ok else "991B1B")
        r += 1
        if not a:
            r += 1
            continue
        st = a["st"]
        header(ws, r, ["Transakcji", "Zyskownych %", "Śr. netto %", "Śr. brutto %", "Suma netto %", "t-stat",
                       "R² krzywej", "Maks. obsunięcie (suma %)", "Okresy stratne", "Okresy puste"])
        r += 1
        for j, v in enumerate([st.n, round(st.wins / st.n * 100, 2) if st.n else None, round(st.avg, 4),
                               round(st.gross, 4), round(st.total, 2), round(st.tstat, 2), round(st.r2, 3),
                               round(st.mdd, 2), st.losing, st.empty], start=1):
            ws.cell(r, j, v)
        r += 1
        ws.cell(r, 1, "okres")
        for p, (a0, b0) in enumerate(a["win"][1]):
            c = ws.cell(r, 2 + p, f"{P.dates[a0][:7]}")
            c.alignment = Alignment(text_rotation=90)
        r += 1
        ws.cell(r, 1, "suma netto %")
        for p in range(len(a["win"][1])):
            c = ws.cell(r, 2 + p, round(st.per_sum[p], 1))
            c.fill = sub if st.per_n[p] < 3 else (good if st.per_sum[p] >= 0 else bad)
        r += 1
        ws.cell(r, 1, "transakcji")
        for p in range(len(a["win"][1])):
            ws.cell(r, 2 + p, st.per_n[p])
        r += 2
    ws.column_dimensions["A"].width = 30
    for j in range(2, 23):
        ws.column_dimensions[get_column_letter(j)].width = 11

    # ROK PO ROKU
    yr = wb.create_sheet("ROK PO ROKU")
    cols = ["Rok", "Strategii w portfelu", "Kandydaci: ≥ 30 transakcji i netto > 0", "Kandydaci: + t ≥ 3",
            "Najlepszy t w historii", "Transakcji (53)", "Suma netto % (53)", "Śr. netto % (53)", "Zyskownych % (53)",
            "Transakcji (dP)", "Suma netto % (dP)", "Śr. netto % (dP)"]
    if plc is not None and plc.get("main"):
        cols += ["Suma netto % (placebo)"]
    header(yr, 1, cols, [8, 12, 16, 12, 12, 12, 14, 12, 12, 12, 14, 12, 16])
    cnts = dict(wf.get("counts", []))
    picks = dict(wf["picks"])
    dpy = {y: v for y, *v in (out["dp"]["years"] if out.get("dp") else [])}
    ply = {y: v for y, *v in (plc["main"]["years"] if plc and plc.get("main") else [])}
    for y, n, s, avg, w in (out["main"]["years"] if out.get("main") else []):
        d = dpy.get(y, [0, 0.0, 0.0, 0.0])
        c_ = cnts.get(y, {})
        bt = c_.get("best_t")
        row = [y, len(picks.get(y, [])), c_.get("net>0"), c_.get("t"), round(bt, 2) if bt is not None else None,
               n, round(s, 2), round(avg, 4), round(w, 1), d[0], round(d[1], 2), round(d[2], 4)]
        if ply:
            row.append(round(ply.get(y, [0, 0.0])[1], 2))
        yr.append(row)
        for col in (7, 11, 13):
            c = yr.cell(yr.max_row, col)
            if isinstance(c.value, (int, float)):
                c.fill = good if c.value >= 0 else bad
    yr.freeze_panes = "A2"

    # STRATEGIE W PORTFELU
    sp = wb.create_sheet("STRATEGIE W PORTFELU")
    header(sp, 1, ["Rok", "Kod reguły", "Opis", "Kierunek", "Transakcji w historii", "Śr. netto % w historii",
                   "t-stat w historii", "Transakcji w roku", "Suma netto % w roku"], [7, 30, 70, 9, 12, 12, 10, 10, 12])
    years = res["years"]
    for y, sel in wf["picks"]:
        k = years.index(y)
        D = float(res["ydays"][:k].sum())
        for i in sel:
            ru = rule_of(res, i)
            n = float(res["YN"][i, :k].sum())
            s = float(res["YS"][i, :k].sum())
            t = float(tstat(s, float(res["YQ"][i, :k].sum()), D))
            sp.append([y, ru.code, ru.desc, "long" if ru.dir > 0 else "short", int(n), round(s / n, 4) if n else None,
                       round(t, 2), int(res["YN"][i, k]), round(float(res["YS"][i, k]), 2)])
            c = sp.cell(sp.max_row, 9)
            c.fill = good if c.value >= 0 else bad
    sp.freeze_panes = "A2"
    if sp.max_row > 1:
        sp.auto_filter.ref = f"A1:I{sp.max_row}"

    # REGUŁY — pełna próba
    rl = wb.create_sheet("REGUŁY (pełna próba)")
    rl.cell(1, 1, "Najlepsze reguły na CAŁEJ historii — wybrane po fakcie, więc zawyżone; tylko do wglądu. Uczciwy wynik: PODSUMOWANIE.")
    header(rl, 2, ["Nr", "Kod reguły", "Opis", "Kierunek", "Transakcji", "Śr. netto %", "Śr. brutto %", "t-stat",
                   "Dodatnie lata"], [5, 30, 70, 9, 11, 11, 11, 9, 11])
    for k, (i, t, avg, pos, n) in enumerate(full_sample_top(res), start=1):
        ru = rule_of(res, i)
        rl.append([k, ru.code, ru.desc, "long" if ru.dir > 0 else "short", n, round(avg, 4), round(avg + cost, 4),
                   round(t, 2), f"{pos:.0%}"])
    rl.freeze_panes = "A3"

    # PORTFEL 10 000 $
    if out.get("main"):
        allrows = np.concatenate([out["rows"], out["dprows"]]).astype(int)
        old = osr.COST_PCT
        osr.COST_PCT = cost
        pf = osr.portfolio(P, wf["sgn"], allrows, g0=wf["start"])
        osr.COST_PCT = old
        od.write_portfolio(wb, SimpleNamespace(id="WF", desc=f"strategie wybierane co rok (walk-forward), spread {cost:g}%, "
                                                             "stawka 5% portfela na transakcję"), pf, len(allrows))
        out["pf"] = pf

    op = wb.create_sheet("OPIS")
    op.column_dimensions["A"].width = 30
    op.column_dimensions["B"].width = 120
    S_, F_ = res["S"], res["F"]
    meta = [
        ("Wersja", f"IA 4 {__version__} · policzono {datetime.now().strftime('%Y-%m-%d %H:%M')} · {note}"),
        ("Dane", f"{P.dates[0]} – {P.dates[-1]}; poszukiwanie od {P.dates[out['g0']]}; portfel od {wf['first_year']}"),
        ("Transakcja", f"open → close tego samego dnia; spread {cost:g}% na każdą transakcję (wejście + wyjście); wyniki netto."),
        ("1. Poszukiwanie", f"kierunek + setup ({len(S_)}) + do {MAX_FILTERS} filtrów z różnych grup ({len(F_)}); kombinacja "
                            f"z < {MIN_TRADES} transakcjami nie jest rozwijana; sprawdzonych reguł {res['rules']} "
                            f"({res['masks']} masek, {res['sec']:.0f} s). Wyniki rok po roku, 53 spółki."),
        ("2. Wybór co rok", f"na początku każdego roku — tylko na danych sprzed tego roku: ≥ {MIN_TRADES} transakcji, suma netto > 0, "
                            f"t-stat ≥ {WF_T:g}, dodatnie ≥ {WF_POS_YEARS:.0%} lat z transakcjami, najwyżej {WF_FAMILY} reguły z tego "
                            f"samego setupu i kierunku, najwyżej {WF_K} reguł; pierwszy rok portfela — po {WF_MIN_YEARS} latach historii."),
        ("3. Portfel", "głosy wybranych reguł sumowane w każdej spółce i dniu (long +1, short −1); pozycja = znak sumy; "
                       "jedna transakcja na spółkę i dzień."),
        ("4. Ocena", f"{N_PERIODS} równych okresów od początku portfela, najwyżej {MAX_LOSING} stratne (warunek człowieka); "
                     "doubleProof — te same wybory na 20 spółkach, których poszukiwanie nie widzi; placebo — cały proces "
                     "z losowym znakiem dni (powinien tracić ok. spread na transakcję)."),
        ("Dlaczego to uczciwe", "każdy dzień portfela jest grany regułami wybranymi na danych sprzed tego roku — wynik "
                                "nie jest dopasowany do przeszłości (inaczej niż arkusz REGUŁY)."),
        ("Portfel 10 000 $", "arkusz PORTFEL WF: wszystkie 73 spółki (dP oznaczone), 5% stanu portfela na transakcję, suma "
                             "stawek dnia ≤ stan portfela."),
    ]
    for a, b in meta:
        op.append([a, b])
        op.cell(op.max_row, 1).font = bold
        op.cell(op.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".ONE-DAY-PORTFEL-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


def summary_line(name: str, a: dict | None) -> str:
    if not a:
        return f"  {name}: brak portfela"
    st = a["st"]
    return (f"  {name}: transakcji {st.n} · śr. netto {st.avg:+.3f}% (brutto {st.gross:+.3f}%) · suma {st.total:+.1f}% · "
            f"t {st.tstat:+.2f} · {verdict(a)}")


def run(cost: float, since: str, no_sync: bool, placebo: bool, out: Path | None) -> Path:
    t0 = time.time()
    print(f"IA 4 — One-Day: portfel strategii (walk-forward) · spread {cost:g}% · od {since}\n")
    P, note = osr.load_panel("daily", no_sync, since)
    extend(P)
    print(f"  {note}\n  dane {P.dates[0]} – {P.dates[-1]}, {len(P.dates)} sesji, {len(P.syms)} spółek")
    print("  1/2 poszukiwanie i portfel:")
    res = pipeline(P, cost)
    print(f"  reguł sprawdzonych {res['res']['rules']} · {res['res']['sec']:.0f} s · portfel od {res['wf']['first_year']}")
    plc = None
    if placebo:
        print("  2/2 placebo (cały proces z losowym znakiem dni):")
        sign = np.random.default_rng(PLACEBO_SEED).choice([-1.0, 1.0], size=len(P.dates))
        plc = pipeline(P, cost, sign=sign, log=None)
    path = write_xlsx(out or config.RESEARCH_DIR.parent / "ONE-DAY-PORTFEL.xlsx", P, res, plc, cost, note)
    print()
    cn = res["wf"].get("counts", [])
    if cn:
        print("  kandydaci w kolejnych latach (≥ 30 transakcji i netto > 0 / + t ≥ 3 / wybrane):")
        print("   " + " · ".join(f"{y}: {c.get('net>0', 0)}/{c.get('t', 0)}/{len(dict(res['wf']['picks']).get(y, []))}"
                                  for y, c in cn))
    print(summary_line("PORTFEL (53 spółki)", res.get("main")))
    print(summary_line("doubleProof", res.get("dp")))
    if plc is not None:
        print(summary_line("PLACEBO", plc.get("main")))
    if res.get("pf"):
        pf = res["pf"]
        print(f"  10 000 $ → {pf['end']:,.2f} $ ({pf['ret']:+.2f}%), obsunięcie {pf['mdd']:.2f}%, transakcji {pf['trades']}"
              .replace(",", " "))
    print(f"\n  {time.time() - t0:.0f} s · plik: {path}")
    return path


# ---------------------------------------------------------------------------
#  Selftest
# ---------------------------------------------------------------------------
def _synth_daily(n_stocks: int, n_days: int, seed: int, edge: float = 0.0, mkt: float = 1.0):
    """Świece dobowe od 2010 r.; edge > 0: po spadku korpusu ≥ 2,5% następny dzień daje średnio +edge %."""
    import pandas as pd
    rng = np.random.default_rng(seed)
    cal = pd.bdate_range("2010-01-04", periods=n_days).strftime("%Y-%m-%d").tolist()
    mk = rng.normal(0.0, mkt, n_days)
    days, spy = {}, None
    for s in range(n_stocks + 1):
        p, prev_big = 100.0, False
        o, h, l, c = [], [], [], []
        for t in range(n_days):
            mv = (mk[t] + rng.normal(0, 1.4)) if s < n_stocks else mk[t]
            if edge and prev_big and s < n_stocks:
                mv += edge
            op = p * (1 + rng.normal(0, 0.003))
            cl = op * (1 + mv / 100)
            o.append(op), h.append(max(op, cl) * 1.004), l.append(min(op, cl) * 0.996), c.append(cl)
            prev_big = (cl / op - 1) * 100 <= -2.5
            p = cl
        dd = od.Days(cal, o, h, l, c, [True] * n_days, [], [], [], [])
        if s == n_stocks:
            spy = dd
        else:
            days[f"S{s}"] = dd
    return days, spy, cal


def selftest() -> None:
    global WF_MIN_YEARS
    fails = []

    def check(name, ok, info=""):
        print(f"  {'OK ' if ok else 'BŁĄD'} {name}{': ' + str(info) if info != '' else ''}")
        if not ok:
            fails.append(name)

    print("IA 4 — One-Day: portfel strategii (walk-forward) — selftest\n")
    old_min = WF_MIN_YEARS
    WF_MIN_YEARS = 3
    try:
        S, F = setups(), filters()
        check("słownik", (len(S), len(F)) == (82, 33), (len(S), len(F)))
        days, spy, cal = _synth_daily(24, 1800, seed=5, edge=2.5)
        groups = {s: ("main" if i < 16 else "doubleProof") for i, s in enumerate(days)}
        P = osr.build(days, groups, spy, calendar=cal, source="daily")
        extend(P)
        # cechy
        t = 600
        i = 0
        vi = np.flatnonzero(P.valid[i])
        c = P.c[i, vi]
        q = int(np.flatnonzero(vi == t - 1)[0])
        check("dołek roczny = pętla", P.feat["low252"][i, t] == float(c[q] <= c[q - 251:q + 1].min()))
        dn = ((P.c[:, t - 1] < P.o[:, t - 1]) & P.valid[:, t - 1]).sum() / P.valid[:, t - 1].sum()
        check("szerokość rynku = udział spadkowych wczoraj", abs(P.breadth[t] - dn) < 1e-12)
        d0 = [k for k, d in enumerate(P.dates) if d.endswith("-01") or (k and P.dates[k][:7] != P.dates[k - 1][:7])]
        check("początek miesiąca", P.cal["tom_s"][d0[3]] and not P.cal["tom_s"][d0[3] + 3])
        check("koniec miesiąca", P.cal["tom_e"][d0[3] - 1] and P.cal["tom_e"][d0[3] - 2] and not P.cal["tom_e"][d0[3] - 3])
        cal2 = [d for d in cal if d != "2010-05-31"]                       # „święto” w poniedziałek
        P2 = osr.build({k: od.Days(cal2, *[[v for d, v in zip(cal, getattr(dd, a)) if d != "2010-05-31"] for a in "ohlc"],
                                   [True] * len(cal2), [], [], [], []) for k, dd in list(days.items())[:2]},
                       {"S0": "main", "S1": "main"}, None, calendar=cal2, source="daily")
        extend(P2)
        check("dzień przed świętem", bool(P2.cal["pre"][cal2.index("2010-05-28")]) and not P2.cal["pre"][cal2.index("2010-05-27")])
        # poszukiwanie: przycinanie = pełne wyliczenie (mały słownik)
        rows = np.array([k for k, g in enumerate(P.groups) if g == "main"])
        g0 = osr.periods(P, rows)[0]
        t1 = time.time()
        res = search(P, rows, g0, 0.6, log=None)
        check("poszukiwanie: reguły i czas", res["rules"] > 1000, f"{res['rules']} reguł, {time.time() - t1:.0f} s")
        sp0 = res["specs"][0]
        m = res["S"][sp0.s][1](P)[rows][:, g0:] & P.base[rows][:, g0:]
        net = (sp0.dir * P.ret[rows][:, g0:] - 0.6) * m
        check("wyniki rok po roku = suma transakcji netto", abs(res["YS"][0].sum() - net.sum()) < 1e-6 * max(1.0, abs(net.sum())) and
              res["YN"][0].sum() == m.sum())
        check("każda zapisana reguła ma ≥ 30 transakcji", bool((res["YN"].sum(1) >= MIN_TRADES).all()))
        # walk-forward bez zaglądania w przyszłość: zmiana danych od roku Y nie zmienia wyboru na rok Y
        k = WF_MIN_YEARS + 1
        sel1 = select(res, k)
        res2 = dict(res)
        res2["YS"] = res["YS"].copy()
        res2["YS"][:, k:] *= -5.0
        check("wybór na rok Y widzi tylko lata < Y", select(res2, k) == sel1 and len(sel1) > 0, len(sel1))
        # pełny proces: zaszyta przewaga przechodzi do portfela i doubleProof; placebo i szum — nie
        out = pipeline(P, 0.6, log=None)
        a, d = out["main"], out["dp"]
        codes = {rule_of(out["res"], i).code for _, sel in out["wf"]["picks"] for i in sel}
        check("portfel: wybiera reguły long (zaszyta przewaga long)", any(cd.startswith("L|") for cd in codes), sorted(codes)[:6])
        check("portfel: zysk netto i warunek okresów", a["st"].avg > 0 and a["st"].ok, verdict(a))
        check("doubleProof: zysk netto", d["st"].avg > 0 and d["st"].total > 0, round(d["st"].avg, 3))
        sign = np.random.default_rng(1).choice([-1.0, 1.0], size=len(P.dates))
        pl = pipeline(P, 0.6, sign=sign, log=None)
        check("placebo: portfel nie zarabia", pl.get("main") is None or pl["main"]["st"].total <= 0 or
              not pl["main"]["st"].ok, verdict(pl.get("main")))
        days_n, spy_n, cal_n = _synth_daily(24, 1800, seed=6)
        Pn = osr.build(days_n, groups, spy_n, calendar=cal_n, source="daily")
        extend(Pn)
        nz = pipeline(Pn, 0.6, log=None)
        check("szum: portfel bez warunku albo bez transakcji", nz.get("main") is None or not nz["main"]["st"].ok
              or nz["main"]["st"].n == 0, verdict(nz.get("main")))
        # plik
        f = Path(tempfile.mkdtemp()) / "p.xlsx"
        write_xlsx(f, P, out, pl, 0.6, "selftest")
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True)
        check("Excel: arkusze", wb.sheetnames == ["PODSUMOWANIE", "ROK PO ROKU", "STRATEGIE W PORTFELU",
                                                  "REGUŁY (pełna próba)", "PORTFEL WF", "OPIS"], wb.sheetnames)
        check("portfel 10 000 $: transakcje = sygnały", out["pf"]["trades"] ==
              int((out["wf"]["sgn"][np.concatenate([out["rows"], out["dprows"]])][:, out["wf"]["start"]:] != 0).sum()))
    finally:
        WF_MIN_YEARS = old_min
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    sys.exit(0 if not fails else 1)


def main() -> None:
    global COST_PCT
    ap = argparse.ArgumentParser(description="IA 4 — One-Day: portfel strategii wybieranych krok po kroku w czasie")
    ap.add_argument("--cost", type=float, default=COST_PCT, help=f"spread na transakcję w %% (domyślnie {COST_PCT:g})")
    ap.add_argument("--since", default=osr.DAILY_SINCE, help=f"od daty (domyślnie {osr.DAILY_SINCE})")
    ap.add_argument("--no-sync", action="store_true")
    ap.add_argument("--no-placebo", action="store_true", help="bez placebo (o połowę krócej)")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        selftest()
    COST_PCT = max(0.0, a.cost)
    run(COST_PCT, a.since, a.no_sync, not a.no_placebo, a.out)


if __name__ == "__main__":
    main()
