"""
IA 4 — strategia RSI(3) / EMA100 / VWAP10: portfel QQQ i to samo dla wszystkich spółek (instrukcja, sekcja 7a).
Wersja projektu: 1.57 (2026-10-21) — musi zgadzać się z IA4_INSTRUKCJA.md

Reguła podana przez człowieka (16.10.2026), liczona jednorazowo — bez poszukiwania, bez zmiany parametrów:
  1. RSI(3) zamyka się poniżej 22 (RSI Wildera z zamknięć),
  2. EMA(100) z zamknięć jest wyżej niż 5 sesji wcześniej,
  3. zamknięcie co najmniej 1% poniżej 10-sesyjnego kroczącego VWAP (Σ typowa cena × wolumen / Σ wolumen,
     typowa cena = (H + L + C) / 3, 10 ostatnich sesji łącznie z bieżącą),
  4. kupno (long) po open następnej sesji,
  5. wyjście po close sesji, w której RSI(3) zamknie się powyżej 55,
  6. jedna pozycja naraz (sygnały w czasie pozycji są pomijane).
Spread 0,3% wartości pozycji raz na transakcję. Bez SL i TP (reguła ich nie ma), bez kosztu nocnego.

Portfele 10 000 $: QQQ — cała kwota w pozycji (1:1) i cała kwota jako depozyt przy dźwigni 1:5;
wszystkie spółki (73: główne, kontrolne, doubleProof) osobno + wspólny portfel (depozyt 10% stanu na pozycję).

    python -m ia4.swrsi                # baza dobowa od 2000 (najpierw python -m ia4.daily); kilkanaście sekund
    python -m ia4.swrsi --selftest     # „WYNIK: OK”
Wynik: ~/Desktop/IA4/SWING-RSI3.xlsx (poza gitem).
"""

from __future__ import annotations

import argparse
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np

from . import __version__, config
from . import swing as sw

RSI_N, RSI_IN, RSI_OUT = 3, 22.0, 55.0
EMA_N, EMA_BACK = 100, 5
VWAP_N, VWAP_GAP = 10, 1.0
MAIN = "QQQ"


def signals(s: sw.Stock) -> dict:
    r3 = sw.rsi(s.c, RSI_N)
    e = sw.ema(s.c, EMA_N)
    tp = (s.h + s.l + s.c) / 3.0
    num = sw.rolling(tp * s.v, VWAP_N, np.sum)
    den = sw.rolling(s.v, VWAP_N, np.sum)
    vwap = np.where(den > 0, num / np.where(den > 0, den, 1.0), np.nan)
    with np.errstate(invalid="ignore"):
        sig = (r3 < RSI_IN) & (e > sw.prior(e, EMA_BACK)) & (s.c <= vwap * (1.0 - VWAP_GAP / 100.0))
    return {"rsi": r3, "ema": e, "vwap": vwap, "sig": np.nan_to_num(sig, nan=False).astype(bool)}


def trades(s: sw.Stock, si: int) -> tuple[list[dict], dict | None]:
    """Transakcje jednej spółki (jedna pozycja naraz); druga wartość — pozycja otwarta na końcu danych."""
    f = signals(s)
    sig, r3 = f["sig"], f["rsi"]
    out, Q, q = [], len(s.c), 0
    while q < Q - 1:
        if not sig[q]:
            q += 1
            continue
        e = q + 1
        x = e
        while x < Q and not (r3[x] > RSI_OUT):
            x += 1
        if x >= Q:
            return out, {"sym": s.sym, "in": str(s.dates[e]), "pin": float(s.o[e]), "last": float(s.c[-1]),
                         "move": (s.c[-1] / s.o[e] - 1.0) * 100.0}
        out.append({"si": si, "qe": e, "qx": x, "d": 1, "pin": float(s.o[e]), "pout": float(s.c[x]), "why": 4})
        q = x                       # następny sygnał najwcześniej na close sesji wyjścia
    return out, None


def stats(tr: list[dict], s: sw.Stock, cost: float) -> dict:
    net = np.array([(t["pout"] / t["pin"] - 1.0) * 100.0 - cost for t in tr])
    ses = np.array([t["qx"] - t["qe"] + 1 for t in tr])
    n = len(net)
    pos, neg = net[net > 0].sum(), -net[net < 0].sum()
    return {"n": n, "win": float((net > 0).mean() * 100) if n else None, "avg": float(net.mean()) if n else None,
            "med": float(np.median(net)) if n else None, "pf": float(pos / neg) if neg > 0 else (99.0 if pos > 0 else None),
            "sum": float(net.sum()), "comp": float((np.prod(1 + net / 100.0) - 1.0) * 100.0) if n else 0.0,
            "ses": float(ses.mean()) if n else None, "inmkt": float(ses.sum() / len(s.c) * 100.0),
            "bh": float((s.c[-1] / s.o[0] - 1.0) * 100.0), "worst": float(net.min()) if n else None,
            "t": float(net.mean() / net.std(ddof=1) * np.sqrt(n)) if n > 2 and net.std(ddof=1) > 0 else None}


def run(no_sync: bool = False, out: Path | None = None, since: str | None = sw.DAILY_SINCE, cost: float = sw.COST_PCT,
        data: sw.Data | None = None) -> Path:
    t0 = time.time()
    print(f"IA 4 — strategia RSI(3) < {RSI_IN:g} · EMA{EMA_N} rosnąca · close ≥ {VWAP_GAP:g}% pod VWAP{VWAP_N} · "
          f"wyjście RSI(3) > {RSI_OUT:g} — spread {cost:g}% raz na transakcję\n")
    D = data or sw.load(no_sync, since)
    print(f"  {D.note}\n  dane {D.cal[0]} – {D.cal[-1]}, {len(D.stocks)} instrumentów")
    per, alltr, opens = [], [], []
    for si, s in enumerate(D.stocks):
        if s.group == "context" and s.sym != MAIN:
            continue
        tr, op = trades(s, si)
        if op:
            opens.append(op)
        st = stats(tr, s, cost)
        p1 = sw.portfolio([dict(t) for t in tr], D, stake_pct=100.0, lev=1.0, cost=cost, keep_rows=False)
        p5 = sw.portfolio([dict(t) for t in tr], D, stake_pct=100.0, lev=5.0, cost=cost, keep_rows=False)
        per.append({"sym": s.sym, "group": s.group, "first": str(s.dates[0]), "st": st, "p1": p1, "p5": p5,
                    "c1": sw.cagr(p1, D), "c5": sw.cagr(p5, D)})
        if s.group != "context":
            alltr += tr
        if s.sym == MAIN:
            qtr = tr
    main = next((p for p in per if p["sym"] == MAIN), None)
    if main is None:
        raise SystemExit(f"Brak {MAIN} w bazie dobowej.")
    q1 = sw.portfolio([dict(t) for t in qtr], D, stake_pct=100.0, lev=1.0, cost=cost)
    q5 = sw.portfolio([dict(t) for t in qtr], D, stake_pct=100.0, lev=5.0, cost=cost)
    a1 = sw.portfolio([dict(t) for t in alltr], D, stake_pct=sw.STAKE_PCT, lev=1.0, cost=cost)
    a5 = sw.portfolio([dict(t) for t in alltr], D, stake_pct=sw.STAKE_PCT, lev=5.0, cost=cost)
    st = main["st"]
    print(f"  {MAIN}: {st['n']} transakcji, zyskownych {st['win']:.0f}%, śr. netto {st['avg']:+.2f}%, PF {st['pf']:.2f}, "
          f"śr. {st['ses']:.1f} sesji; portfel 1:1 {q1['ret']:+.0f}% (obsunięcie {q1['mdd']:.0f}%), "
          f"1:5 {q5['ret']:+.0f}% (obsunięcie {q5['mdd']:.0f}%{', UPADŁOŚĆ' if q5['ruined'] else ''}); kup i trzymaj {st['bh']:+.0f}%")
    pos = [p for p in per if p["st"]["n"] and p["st"]["avg"] > 0]
    print(f"  spółki: {len(pos)}/{len(per)} z dodatnią średnią netto; wspólny portfel 1:1 {a1['ret']:+.0f}% "
          f"(obsunięcie {a1['mdd']:.0f}%), 1:5 {a5['ret']:+.0f}% (obsunięcie {a5['mdd']:.0f}%{', UPADŁOŚĆ' if a5['ruined'] else ''})")
    path = out or (config.RESEARCH_DIR.parent / "SWING-RSI3.xlsx")
    write_xlsx(path, D, per, q1, q5, a1, a5, opens, cost)
    print(f"\nGotowe w {time.time() - t0:.0f} s → {path}")
    return path


def write_xlsx(path: Path, D, per, q1, q5, a1, a5, opens, cost) -> Path:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    S = sw.styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "OPIS"
    m = next(p for p in per if p["sym"] == MAIN)
    lines = [
        ("IA 4 — strategia RSI(3) / EMA100 / VWAP10", ""),
        ("Wersja", __version__), ("Wygenerowano", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Dane", f"{D.cal[0]} – {D.cal[-1]} (baza dobowa, ceny skorygowane o splity, bez dywidend); {D.note}"),
        ("Reguła 1", f"RSI({RSI_N}) zamyka się poniżej {RSI_IN:g} (RSI Wildera z zamknięć)"),
        ("Reguła 2", f"EMA({EMA_N}) wyżej niż {EMA_BACK} sesji wcześniej"),
        ("Reguła 3", f"close co najmniej {VWAP_GAP:g}% poniżej {VWAP_N}-sesyjnego kroczącego VWAP (typowa cena (H+L+C)/3 × wolumen)"),
        ("Reguła 4", "kupno po open następnej sesji"), ("Reguła 5", f"wyjście po close sesji, gdy RSI({RSI_N}) > {RSI_OUT:g}"),
        ("Reguła 6", "jedna pozycja naraz (na instrument)"),
        ("Spread", f"{cost:g}% wartości pozycji, raz na transakcję (decyzja człowieka z 16.10.2026)"),
        ("Nie liczone", "koszt nocny (swap) CFD — przy pozycjach wielodniowych na dźwigni jest realnym kosztem; dywidendy"),
        ("Portfel QQQ", "10 000 $; 1:1 — cała kwota w pozycji; 1:5 — cała kwota jako depozyt (pozycja = 5 × stan portfela)"),
        ("Wspólny portfel", f"wszystkie spółki (bez SPY/QQQ), depozyt {sw.STAKE_PCT:g}% stanu na pozycję, 1:1 i 1:5"),
        ("", ""),
        (f"{MAIN} — transakcji", m["st"]["n"]), (f"{MAIN} — zyskownych %", round(m["st"]["win"], 1)),
        (f"{MAIN} — śr. wynik netto %", round(m["st"]["avg"], 3)), (f"{MAIN} — PF", round(m["st"]["pf"], 2)),
        (f"{MAIN} — śr. sesji w pozycji", round(m["st"]["ses"], 1)), (f"{MAIN} — czas w rynku %", round(m["st"]["inmkt"], 1)),
        (f"{MAIN} — portfel 1:1: wynik % / obsunięcie %", f"{q1['ret']:+.1f} / {q1['mdd']:.1f}"),
        (f"{MAIN} — portfel 1:5: wynik % / obsunięcie %", f"{q5['ret']:+.1f} / {q5['mdd']:.1f}" + (" — UPADŁOŚĆ" if q5["ruined"] else "")),
        (f"{MAIN} — kup i trzymaj % (ten sam okres)", round(m["st"]["bh"], 1)),
        ("Wspólny portfel 1:1: wynik % / obsunięcie %", f"{a1['ret']:+.1f} / {a1['mdd']:.1f}"),
        ("Wspólny portfel 1:5: wynik % / obsunięcie %", f"{a5['ret']:+.1f} / {a5['mdd']:.1f}" + (" — UPADŁOŚĆ" if a5["ruined"] else "")),
        ("Pozycje otwarte na końcu danych", ", ".join(f"{o['sym']} od {o['in']} ({o['move']:+.1f}%)" for o in opens) or "—"),
        ("Uwaga", "Reguła przyszła z zewnątrz i jest liczona raz — wynik nie jest zawyżony poszukiwaniem. Ale to wynik na "
                  "historii: czy przewaga trwa, pokazują lata (arkusze portfeli) i spółki, na których jej nie dobierano."),
    ]
    for k, (a, b) in enumerate(lines, start=1):
        ws.cell(k, 1, a)
        ws.cell(k, 2, b)
        if k == 1 or a.startswith(MAIN):
            ws.cell(k, 1).font = S.bold
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 120

    sw.portfolio_sheet(wb, f"{MAIN} 1-1", q1, f"{MAIN} — portfel 10 000 $, cała kwota w pozycji (1:1), spread {cost:g}%", S)
    sw.portfolio_sheet(wb, f"{MAIN} 1-5", q5, f"{MAIN} — portfel 10 000 $, cała kwota jako depozyt, dźwignia 1:5, spread {cost:g}%", S)

    sp = wb.create_sheet("SPÓŁKI")
    cols = ["Spółka", "Grupa", "Dane od", "Transakcji", "Zyskownych %", "Śr. netto %", "Mediana netto %", "PF", "t-stat",
            "Najgorsza %", "Śr. sesji", "Czas w rynku %", "Portfel 1:1 %", "CAGR 1:1 %", "Obsunięcie 1:1 %",
            "Portfel 1:5 %", "Obsunięcie 1:5 %", "Upadłość 1:5", "Kup i trzymaj %"]
    sw.header(sp, 1, cols, [9, 12, 11] + [10] * (len(cols) - 3), S)
    gname = {"main": "główna", "proof": "kontrolna", "doubleProof": "doubleProof", "context": "tło rynku"}
    for p in sorted(per, key=lambda p: -(p["p1"]["ret"])):
        s = p["st"]
        r = lambda v, k=3: None if v is None else round(v, k)
        sp.append([p["sym"], gname.get(p["group"], p["group"]), p["first"], s["n"], r(s["win"], 1), r(s["avg"]), r(s["med"]),
                   r(s["pf"], 2), r(s["t"], 2), r(s["worst"], 2), r(s["ses"], 1), r(s["inmkt"], 1), r(p["p1"]["ret"], 1),
                   r(p["c1"], 2), r(p["p1"]["mdd"], 1), r(p["p5"]["ret"], 1), r(p["p5"]["mdd"], 1),
                   "TAK" if p["p5"]["ruined"] else "nie", r(s["bh"], 1)])
        for j in (6, 13, 16):
            c = sp.cell(sp.max_row, j)
            if isinstance(c.value, (int, float)):
                c.fill = S.good if c.value > 0 else S.bad
    sp.freeze_panes = "B2"
    sp.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{sp.max_row}"

    sw.portfolio_sheet(wb, "WSZYSTKIE 1-1", a1, f"Wszystkie spółki — wspólny portfel 10 000 $, depozyt {sw.STAKE_PCT:g}% na pozycję, 1:1", S)
    sw.portfolio_sheet(wb, "WSZYSTKIE 1-5", a5, f"Wszystkie spółki — wspólny portfel 10 000 $, depozyt {sw.STAKE_PCT:g}% na pozycję, 1:5", S)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
def selftest() -> None:
    fails: list[str] = []

    def check(name, got, exp):
        ok = got == exp
        print(f"  {'OK  ' if ok else 'BŁĄD'} {name}: {got}" + ("" if ok else f" (oczekiwano {exp})"))
        if not ok:
            fails.append(name)

    print(f"IA 4 {__version__} — selftest ia4.swrsi\n")
    sw.selftest_core(check)
    # Przebieg: wzrost (EMA100 rośnie), potem 3 spadki z wolumenem → RSI(3) < 22 i close < VWAP − 1% → wejście
    n = 160
    c = list(np.linspace(100, 140, n))
    c += [136.0, 131.0, 126.0]                  # spadki → sygnał na ostatnim
    c += [127.0, 133.0, 139.0, 141.0]           # odbicie → RSI(3) > 55
    c = np.array(c)
    o = np.concatenate([[c[0]], c[:-1]])
    h, l = np.maximum(o, c) * 1.002, np.minimum(o, c) * 0.998
    d = np.array([f"2020-{1 + i // 28:02d}-{1 + i % 28:02d}" for i in range(len(c))])
    s = sw.Stock("X", "main", d, o, h, l, c, np.full(len(c), 1e6))
    D = sw.finish([s])
    f = signals(s)
    first = int(np.flatnonzero(f["sig"])[0]) if f["sig"].any() else -1
    check("sygnał w serii spadków", first in (n, n + 1, n + 2), True)
    tr, op = trades(s, 0)
    check("jedna transakcja, bez otwartej pozycji", (len(tr), op is None), (1, True))
    if tr:
        t = tr[0]
        check("wejście po open następnej sesji", t["qe"], first + 1)
        check("wyjście na close, gdy RSI(3) > 55", (t["qx"], bool(f["rsi"][t["qx"]] > RSI_OUT), bool(f["rsi"][t["qx"] - 1] <= RSI_OUT)),
              (t["qx"], True, True))
        st = stats(tr, s, 0.3)
        check("wynik netto = ruch − spread", round(st["avg"], 6), round((c[t["qx"]] / o[t["qe"]] - 1) * 100 - 0.3, 6))
        pf = sw.portfolio([dict(x) for x in tr], D, stake_pct=100.0, lev=5.0)
        check("portfel 1:5 = 5 × wynik netto", round(pf["ret"], 6), round(5 * st["avg"], 6))
    # VWAP: stały wolumen = średnia typowej ceny
    tp = (h + l + c) / 3
    check("VWAP10 przy stałym wolumenie = średnia typowej ceny", round(float(f["vwap"][50]), 6), round(float(tp[41:51].mean()), 6))
    # Excel
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        s2 = sw.Stock(MAIN, "context", d, o, h, l, c, np.full(len(c), 1e6))
        D2 = sw.finish([s, s2])
        p = run(no_sync=True, out=Path(td) / "t.xlsx", data=D2)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        check("arkusze", wb.sheetnames, ["OPIS", f"{MAIN} 1-1", f"{MAIN} 1-5", "SPÓŁKI", "WSZYSTKIE 1-1", "WSZYSTKIE 1-5"])
    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    if fails:
        sys.exit(1)


def main() -> None:
    ap = argparse.ArgumentParser(description="Strategia RSI(3)/EMA100/VWAP10 — portfel QQQ i wszystkie spółki (instrukcja 7a)")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--no-sync", action="store_true", help="bez synchronizacji bazy dobowej")
    ap.add_argument("--since", default=sw.DAILY_SINCE, help="początek danych (domyślnie 2000-01-01)")
    ap.add_argument("--cost", type=float, default=sw.COST_PCT, help="spread %% na transakcję (raz; domyślnie 0,3)")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    if a.selftest:
        selftest()
        return
    run(a.no_sync, a.out, a.since, a.cost)


if __name__ == "__main__":
    main()
