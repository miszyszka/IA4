"""
IA 4 — test One-Day: kaskady na świecach dobowych (instrukcja, sekcja 6c).
Wersja projektu: 1.39 (2026-10-14) — musi zgadzać się z IA4_INSTRUKCJA.md

Jednorazowy test informacyjny (decyzja człowieka z 12.10.2026). Nie jest strategią `ia4-rule/1`
(zasada 9), nie zmienia poszukiwania, paper tradingu ani AI Investor.

- Świeca dobowa = sesja NYSE: open pierwszej świecy 1h, close ostatniej (7. albo 4. w dzień skrócony),
  high/low — maksimum i minimum sesji. Sesja bez pierwszej albo ostatniej świecy 1h (albo brakująca
  w kopii) jest „niepełna”: nie jest handlowana i przerywa serię oraz kaskadę.
- Żelazna zasada: każda transakcja otwiera się na open dnia i zamyka na close tego samego dnia.
- Kaskada: po N świecach spadkowych z rzędu (close < open) long w dniu N+1 ze stawką k₁; dopóki świece
  są spadkowe (strata), następnego dnia k₂, k₃… Świeca wzrostowa (zysk) albo doji (close = open)
  kończy kaskadę; nowa zaczyna się dopiero, gdy seria znów dojdzie do N. Short — lustrzanie.
- Każda spółka × strategia: osobny portfel 100 000 $, bez kosztów, stawka ≤ bieżący kapitał (bez dźwigni),
  kapitał ≤ 10% startowego = upadłość (koniec handlu na tej spółce).
- Wynik: transakcje, zyskowne, wynik % portfela, maksymalne obsunięcie (szczyt → dołek krzywej kapitału
  po zamknięciu każdego dnia). Plik: ~/Desktop/IA4/ONE-DAY-STRATEGY.xlsx (poza gitem).
- Od 1.36 arkusz PORTFEL OD-080: jedna strategia na wszystkich spółkach naraz, wspólny portfel 10 000 $,
  kolejne transakcje dzień po dniu (portfolio()).

    python -m ia4.oneday              # synchronizacja (gdy nie działa `ia4 start`), test, plik Excel
    python -m ia4.oneday --no-sync    # tylko kopia lokalna
    python -m ia4.oneday --portfolio OD-080 OD-034 --capital 10000   # arkusze PORTFEL (domyślnie OD-080, 10 000 $)
    python -m ia4.oneday --selftest   # test na danych syntetycznych — „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import math
import os
import statistics
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from . import __version__, config
from .lab.check import US_HALF_DAYS, trading_days

# ---------------------------------------------------------------------------
#  Parametry (instrukcja 6c — zmiana tylko razem z instrukcją)
# ---------------------------------------------------------------------------
CAPITAL = 100_000.0        # kapitał startowy każdego portfela
X_USD = 5_000.0            # kwota x = 5% kapitału startowego
R_USD = 500.0              # 1R = stawka, przy której ruch o r to 500 $ (0,5% kapitału startowego)
R_DAYS = 20                # r = średni zakres (H−L)/O w % z 20 poprzednich pełnych sesji
SMA_DAYS = 50              # filtr trendu: SMA50 z dziennych zamknięć
SLOPE_DAYS = 5             # „rosnąca” = SMA50 wczoraj > SMA50 sprzed 5 pełnych sesji
WARMUP = SMA_DAYS + SLOPE_DAYS   # wspólny start: pierwsza transakcja po 55 pełnych sesjach
RUIN_PCT = 10.0            # kapitał ≤ 10% startowego → upadłość
MAX_STEPS_E = 4            # blok E: najwyżej 4 kroki kaskady

# Kolejność kolumn arkusza = instrukcja, sekcja 4 (kopia listy z Code.gs; nadpisuje ją system/universe).
FALLBACK_UNIVERSE = {
    "main": ["AAPL", "TSLA", "NVDA"],
    "proof": ["DELL", "AMAT", "PLTR", "ORCL", "XOM", "V", "WMT", "JPM", "MU", "META", "AVGO", "MSFT", "GOOGL",
              "JNJ", "MA", "ABBV", "BAC", "CVX", "MRK", "PG", "HD", "PM", "WFC", "CRM", "CAT", "HON", "UNP", "RTX",
              "AMZN", "MCD", "NKE", "SBUX", "T", "VZ", "NFLX", "DIS", "UNH", "LLY", "PFE", "MDT", "PLD", "AMT",
              "LIN", "FCX", "NEE", "DUK", "GS", "AXP", "KO", "PEP"],
    "doubleProof": list(config.DOUBLE_PROOF_FALLBACK),
    "context": ["SPY", "QQQ"],
}

PROG_LABEL = {
    "const": ("stała (1-1-1…)", "stała (1R-1R-1R…)"),
    "lin": ("liniowa (1-2-3…)", "+1R (1R-2R-3R…)"),
    "x2": ("×2 (1-2-4-8…)", "×2 (1R-2R-4R…)"),
    "x3": ("×3 (1-3-9…)", "×3 (1R-3R-9R…)"),
    "fib": ("Fibonacci (1-1-2-3-5…)", "Fibonacci (1R-1R-2R-3R…)"),
}


# ---------------------------------------------------------------------------
#  Lista 100 strategii
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Strat:
    id: str
    block: str
    dir: int                 # +1 long, −1 short
    n: int = 0               # długość serii przed wejściem (0 = codziennie, blok F)
    size: str = "x"          # "x" albo "R"
    prog: str = "const"      # const / lin / x2 / x3 / fib
    filt: str | None = None  # None / "up" (SMA50 rosnąca) / "down" (SMA50 spadkowa)
    dsize: bool = False      # blok D: stawka × D / r
    max_steps: int | None = None
    daily: bool = False      # blok F: codziennie

    @property
    def desc(self) -> str:
        side = "LONG" if self.dir > 0 else "SHORT"
        if self.daily:
            return f"{side} codziennie · x = 5 000 $ · punkt odniesienia"
        series = f"po {self.n} {'spadkach' if self.dir > 0 else 'wzrostach'}"
        size = "x = 5 000 $" if self.size == "x" else "1R (ruch o r = 500 $)"
        parts = [f"{side} {series}", size]
        if self.dsize:
            parts.append("stawka × D / r (D — suma korpusów serii w %)")
        else:
            parts.append(PROG_LABEL[self.prog][0 if self.size == "x" else 1])
        if self.filt:
            parts.append("tylko gdy SMA50 " + ("rosnąca" if self.filt == "up" else "spadkowa"))
        if self.max_steps:
            parts.append(f"najwyżej {self.max_steps} kroki")
        return " · ".join(parts)


def strategies() -> list[Strat]:
    out: list[Strat] = []

    def add(**kw):
        out.append(Strat(id=f"OD-{len(out) + 1:03d}", **kw))

    dirs = (1, -1)
    for d in dirs:                                       # A: OD-001…030
        for n in (2, 3, 4):
            for p in ("const", "lin", "x2", "x3", "fib"):
                add(block="A", dir=d, n=n, size="x", prog=p)
    for d in dirs:                                       # B: OD-031…054
        for n in (2, 3, 4):
            for p in ("const", "lin", "x2", "x3"):
                add(block="B", dir=d, n=n, size="R", prog=p)
    for d in dirs:                                       # C: OD-055…078
        for n in (2, 3, 4):
            for size in ("x", "R"):
                for f in ("up", "down"):
                    add(block="C", dir=d, n=n, size=size, prog="x2", filt=f)
    for d in dirs:                                       # D: OD-079…090
        for n in (2, 3, 4):
            for size in ("x", "R"):
                add(block="D", dir=d, n=n, size=size, dsize=True)
    for d in dirs:                                       # E: OD-091…098
        for n in (2, 3):
            for size in ("x", "R"):
                add(block="E", dir=d, n=n, size=size, prog="x2", max_steps=MAX_STEPS_E)
    add(block="F", dir=1, daily=True)                    # OD-099
    add(block="F", dir=-1, daily=True)                   # OD-100
    return out


def mult(prog: str, k: int) -> float:
    """Mnożnik stawki w kroku k (k = 1 — pierwsze wejście kaskady)."""
    if prog == "const":
        return 1.0
    if prog == "lin":
        return float(k)
    if prog == "x2":
        return 2.0 ** (k - 1)
    if prog == "x3":
        return 3.0 ** (k - 1)
    if prog == "fib":
        a, b = 1, 1
        for _ in range(k - 1):
            a, b = b, a + b
        return float(a)
    raise ValueError(prog)


# ---------------------------------------------------------------------------
#  Świece dobowe
# ---------------------------------------------------------------------------
@dataclass
class Days:
    dates: list[str]
    o: list[float]
    h: list[float]
    l: list[float]
    c: list[float]
    valid: list[bool]
    color: list[int]         # +1 wzrostowa, −1 spadkowa, 0 doji albo niepełna
    pos: list[int]           # numer pełnej sesji (−1 dla niepełnej)
    r: list[float]           # r przed dniem (na numerze pełnej sesji)
    sma_up: list[int]        # +1 rosnąca, −1 spadkowa, 0 płaska / brak — znane przed open dnia


def daily(df) -> Days:
    """Tabela świec 1h jednego instrumentu → świece dobowe na każdy dzień sesyjny od pierwszej do ostatniej daty."""
    dates = [str(x) for x in df["date"].tolist()]
    slots = [int(x) for x in df["slot"].tolist()]
    O, H, L, C = (df[k].astype(float).tolist() for k in ("o", "h", "l", "c"))
    by: dict[str, dict[int, int]] = {}
    for i, (d, s) in enumerate(zip(dates, slots)):
        by.setdefault(d, {})[s] = i
    if not by:
        return Days([], [], [], [], [], [], [], [], [], [])
    days = trading_days(min(by), max(by))
    out = Days(days, [], [], [], [], [], [], [], [], [])
    for d in days:
        idx = by.get(d, {})
        last = 4 if d in US_HALF_DAYS else 7
        ok = 1 in idx and last in idx
        if ok:
            rows = list(idx.values())
            o, c = O[idx[1]], C[idx[last]]
            h, lo = max(H[i] for i in rows), min(L[i] for i in rows)
            ok = o > 0 and c > 0 and all(map(math.isfinite, (o, h, lo, c)))
        if not ok:
            o = h = lo = c = float("nan")
        out.o.append(o)
        out.h.append(h)
        out.l.append(lo)
        out.c.append(c)
        out.valid.append(ok)
        out.color.append(0 if not ok or c == o else (1 if c > o else -1))
    _indicators(out)
    return out


def _indicators(dd: Days) -> None:
    vi = [t for t, v in enumerate(dd.valid) if v]
    rng = [(dd.h[t] - dd.l[t]) / dd.o[t] * 100.0 for t in vi]
    cls = [dd.c[t] for t in vi]
    sma: list[float] = []
    acc = 0.0
    for q, x in enumerate(cls):
        acc += x
        if q >= SMA_DAYS:
            acc -= cls[q - SMA_DAYS]
        sma.append(acc / SMA_DAYS if q >= SMA_DAYS - 1 else float("nan"))
    n = len(dd.dates)
    dd.pos, dd.r, dd.sma_up = [-1] * n, [float("nan")] * n, [0] * n
    pre = [0.0]
    for x in rng:
        pre.append(pre[-1] + x)
    for p, t in enumerate(vi):
        dd.pos[t] = p
        if p >= R_DAYS:                                   # r przed dniem: pełne sesje p−20 … p−1
            dd.r[t] = (pre[p] - pre[p - R_DAYS]) / R_DAYS
        if p - 1 - SLOPE_DAYS >= SMA_DAYS - 1:            # SMA50 wczoraj vs sprzed 5 pełnych sesji
            a, b = sma[p - 1], sma[p - 1 - SLOPE_DAYS]
            dd.sma_up[t] = 1 if a > b else (-1 if a < b else 0)


# ---------------------------------------------------------------------------
#  Symulacja
# ---------------------------------------------------------------------------
def events(s: Strat, dd: Days) -> list[tuple[int, int, float, float]]:
    """Transakcje strategii s na jednym instrumencie, niezależnie od kapitału:
    [(dzień t, krok kaskady k, mnożnik stawki, r z początku kaskady)]. Stawka = jednostka × mnożnik,
    jednostka: x = 5% kapitału odniesienia, 1R = 0,5% kapitału odniesienia / (r / 100)."""
    out = []
    trig = -s.dir                      # kolor świecy serii: long — spadkowe, short — wzrostowe
    streak = 0                         # długość bieżącej serii (do wczoraj)
    body = 0.0                         # D: suma korpusów serii w % (kierunek serii)
    k = 0                              # 0 = brak kaskady; inaczej numer kroku do zagrania dziś
    r0 = float("nan")
    for t, ok in enumerate(dd.valid):
        if not ok:                     # niepełna sesja: bez handlu, koniec serii i kaskady
            streak, body, k = 0, 0.0, 0
            continue
        col = dd.color[t]
        if dd.pos[t] >= WARMUP:
            if s.daily:
                out.append((t, 0, 1.0, dd.r[t]))
            else:
                if k == 0 and streak == s.n and (s.filt is None or dd.sma_up[t] == (1 if s.filt == "up" else -1)):
                    if dd.r[t] > 0 or (s.size == "x" and not s.dsize):
                        k, r0 = 1, dd.r[t]
                if k:
                    out.append((t, k, body / r0 if s.dsize else mult(s.prog, k), r0))
                    k = k + 1 if col == trig else 0          # strata → następny krok; zysk/doji → koniec
                    if s.max_steps and k > s.max_steps:
                        k = 0
        if col == trig:
            streak += 1
            body += abs(dd.c[t] - dd.o[t]) / dd.o[t] * 100.0
        else:
            streak, body = 0, 0.0
    return out


def unit(s: Strat, ref: float, r0: float) -> float:
    """Jednostka stawki przy kapitale odniesienia ref: x albo 1R (instrukcja 6c)."""
    if s.daily or s.size == "x":
        return ref * X_USD / CAPITAL
    return ref * R_USD / CAPITAL / (r0 / 100.0)


def first_trade_day(dd: Days) -> str | None:
    for t, ok in enumerate(dd.valid):
        if ok and dd.pos[t] >= WARMUP:
            return dd.dates[t]
    return None


def simulate(s: Strat, dd: Days, trace: list | None = None) -> dict:
    """Jeden portfel 100 000 $: strategia s na świecach dobowych jednego instrumentu (stawki stałe w $)."""
    eq = peak = CAPITAL
    mdd = 0.0
    trades = wins = 0
    ruined = False
    ruin_at = CAPITAL * RUIN_PCT / 100.0
    for t, k, m, r0 in events(s, dd):
        stake = min(unit(s, CAPITAL, r0) * m, eq)
        if stake <= 0:
            continue
        pnl = stake * s.dir * (dd.c[t] / dd.o[t] - 1.0)
        eq += pnl
        trades += 1
        wins += pnl > 0
        if trace is not None:
            trace.append((dd.dates[t], k, round(stake, 6), round(pnl, 6), round(eq, 6)))
        peak = max(peak, eq)                       # krzywa po każdym dniu z transakcją (bez niej się nie zmienia)
        mdd = max(mdd, (peak - eq) / peak * 100.0)
        if eq <= ruin_at:
            ruined = True
            break
    return {"trades": trades, "wins": wins, "ret": (eq / CAPITAL - 1.0) * 100.0, "mdd": mdd,
            "ruined": ruined, "first": first_trade_day(dd)}


# ---------------------------------------------------------------------------
#  Portfel — jedna strategia na wszystkich spółkach naraz (instrukcja 6c, „Portfel”)
# ---------------------------------------------------------------------------
PORTFOLIO_CAPITAL = 10_000.0
PORTFOLIO_STRATS = ("OD-080",)


def portfolio(s: Strat, data: dict[str, Days], symbols: list[str], capital: float = PORTFOLIO_CAPITAL) -> dict:
    """Wspólny portfel: wszystkie transakcje strategii s na spółkach `symbols`, dzień po dniu.
    Jednostka stawki liczona od stanu portfela na początku kaskady (dla F — na początku dnia);
    suma stawek dnia > stan portfela na początku dnia → wszystkie stawki dnia zmniejszone proporcjonalnie
    (bez dźwigni). Stan portfela ≤ 10% startowego → upadłość, koniec handlu."""
    by_day: dict[str, list] = {}
    order = {sym: i for i, sym in enumerate(symbols)}
    for sym in symbols:
        dd = data.get(sym)
        if dd is None:
            continue
        for t, k, m, r0 in events(s, dd):
            by_day.setdefault(dd.dates[t], []).append((order[sym], sym, t, k, m, r0))
    eq = peak = capital
    mdd, ruined = 0.0, False
    ruin_at = capital * RUIN_PCT / 100.0
    ref: dict[str, float] = {}          # stan portfela na początku bieżącej kaskady spółki
    rows, n_scaled, max_load = [], 0, 0.0
    trades = wins = 0
    for day in sorted(by_day):
        e0 = eq
        evs = sorted(by_day[day])
        stakes = []
        for _o, sym, t, k, m, r0 in evs:
            if k <= 1:
                ref[sym] = e0
            stakes.append(unit(s, ref[sym], r0) * m)
        total = sum(stakes)
        scale = min(1.0, e0 / total) if total > 0 else 1.0
        max_load = max(max_load, total / e0 if e0 > 0 else 0.0)
        if scale < 1.0:
            n_scaled += 1
        for (_o, sym, t, k, m, r0), st0 in zip(evs, stakes):
            dd = data[sym]
            stake = st0 * scale
            move = (dd.c[t] / dd.o[t] - 1.0) * 100.0
            pnl = stake * s.dir * move / 100.0
            eq += pnl
            trades += 1
            wins += pnl > 0
            rows.append({"day": day, "sym": sym, "candle": (s.n + k) if not s.daily else 1, "k": k,
                         "o": dd.o[t], "c": dd.c[t], "move": move, "stake": stake, "pnl": pnl,
                         "trade_pct": s.dir * move, "ret": (eq / capital - 1.0) * 100.0, "eq": eq,
                         "scale": scale})
        peak = max(peak, eq)
        mdd = max(mdd, (peak - eq) / peak * 100.0)
        if eq <= ruin_at:
            ruined = True
            break
    days = sorted(by_day)
    return {"rows": rows, "capital": capital, "end": eq, "ret": (eq / capital - 1.0) * 100.0, "mdd": mdd,
            "trades": trades, "wins": wins, "ruined": ruined, "days": len({r["day"] for r in rows}),
            "scaled": n_scaled, "max_load": max_load, "first": days[0] if days else None,
            "last": rows[-1]["day"] if rows else None, "max_k": max((r["k"] for r in rows), default=0)}


# ---------------------------------------------------------------------------
#  Instrumenty i dane
# ---------------------------------------------------------------------------
def universe() -> dict:
    try:
        u = config.universe()
        if sum(len(v) for v in u.values()) >= 70:
            return u
    except Exception as e:                                     # bez klucza / sieci — lista z instrukcji
        print(f"  (system/universe niedostępny: {type(e).__name__} — lista z instrukcji, sekcja 4)")
    return FALLBACK_UNIVERSE


def ordered(u: dict) -> list[tuple[str, str]]:
    return ([(s, "main") for s in u["main"]] + [(s, "proof") for s in u["proof"]]
            + [(s, "doubleProof") for s in u["doubleProof"]] + [(s, "context") for s in u["context"]])


def maybe_sync(symbols: list[str]) -> str:
    from .ai import lock
    h = lock.holder()
    if h:
        return f"bez synchronizacji — działa `ia4 start` (PID {h.get('pid')}), jego feed trzyma kopię na bieżąco"
    from . import sync
    sync.sync(symbols, quiet=True)
    return "synchronizacja przyrostowa wykonana"


# ---------------------------------------------------------------------------
#  Excel
# ---------------------------------------------------------------------------
GROUP_TAG = {"main": "", "proof": "", "doubleProof": " · dP", "context": " · ETF"}
SUMMARY = ["Transakcji łącznie", "Zyskownych łącznie", "Średni wynik %", "Mediana wyniku %",
           "Spółek na plusie", "Upadłości", "Średnie obsunięcie %", "Mediana obsunięcia %",
           "Średni wynik % — 53 handlowane", "Średni wynik % — 20 doubleProof"]
PER = ["Transakcje", "Zyskowne", "Wynik %", "Obsunięcie %"]
PCT = "0.00;[Red]-0.00"


def summarize(res: dict, insts: list[tuple[str, str]]) -> list:
    stocks = [s for s, g in insts if g != "context" and s in res]
    traded = [s for s, g in insts if g in ("main", "proof") and s in res]
    dp = [s for s, g in insts if g == "doubleProof" and s in res]

    def mean(xs):
        return statistics.fmean(xs) if xs else None

    rets = [res[s]["ret"] for s in stocks]
    mdds = [res[s]["mdd"] for s in stocks]
    return [sum(res[s]["trades"] for s in stocks), sum(res[s]["wins"] for s in stocks),
            mean(rets), statistics.median(rets) if rets else None,
            f"{sum(x > 0 for x in rets)} / {len(rets)}", sum(res[s]["ruined"] for s in stocks),
            mean(mdds), statistics.median(mdds) if mdds else None,
            mean([res[s]["ret"] for s in traded]), mean([res[s]["ret"] for s in dp])]


PF_HEAD = ["Nr", "Dzień", "Spółka", "Świeca w serii", "Krok kaskady", "Stawka $", "Wynik transakcji %",
           "Zysk / strata $", "Zarobiony % (portfel od startu)", "Stan portfela $", "Uwagi"]


def write_portfolio(wb, s: Strat, pf: dict, n_syms: int) -> None:
    """Arkusz PORTFEL <id>: kolejne transakcje wspólnego portfela (instrukcja 6c, „Portfel”)."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(f"PORTFEL {s.id}")
    head = PatternFill("solid", fgColor="1F2A37")
    sumf = PatternFill("solid", fgColor="FEF3C7")
    bold = Font(bold=True)
    usd, pct = '#,##0.00;[Red]-#,##0.00', "0.00;[Red]-0.00"
    ws["A1"] = f"PORTFEL {s.id} — {s.desc}"
    ws["A1"].font = Font(bold=True, size=13)
    summ = [("Start", pf["capital"], usd), ("Koniec", pf["end"], usd), ("Zarobiony %", pf["ret"], pct),
            ("Maks. obsunięcie %", pf["mdd"], pct), ("Transakcji", pf["trades"], "0"),
            ("Zyskownych", pf["wins"], "0"), ("Dni z transakcjami", pf["days"], "0"),
            ("Dni przycięte do kapitału", pf["scaled"], "0"), ("Najdłuższa kaskada (kroki)", pf["max_k"], "0")]
    for j, (lab, v, fmt) in enumerate(summ):
        a = ws.cell(2, 1 + j, lab)
        a.font, a.fill = bold, sumf
        a.alignment = Alignment(wrap_text=True, vertical="center")
        b = ws.cell(3, 1 + j, round(v, 4) if isinstance(v, float) else v)
        b.number_format, b.fill = fmt, sumf
    ws.cell(4, 1, f"Spółki: {n_syms} (53 handlowane + doubleProof, bez SPY i QQQ) · "
                  f"{pf['first'] or '—'} – {pf['last'] or '—'}"
                  + (" · UPADŁOŚĆ — portfel spadł do 10% startu, koniec handlu" if pf["ruined"] else ""))
    hr = 6
    for j, h in enumerate(PF_HEAD, start=1):
        c = ws.cell(hr, j, h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(pf["rows"], start=1):
        note = f"stawki dnia ×{r['scale']:.3f} (suma > stan portfela)" if r["scale"] < 1 else None
        vals = [i, r["day"], r["sym"], r["candle"], r["k"], round(r["stake"], 2), round(r["trade_pct"], 4),
                round(r["pnl"], 2), round(r["ret"], 4), round(r["eq"], 2), note]
        row = hr + i                                   # bez ws.max_row w pętli (czas kwadratowy)
        for col, v in enumerate(vals, start=1):
            if v is not None:
                ws.cell(row, col, v)
        for col, fmt in ((6, usd), (7, pct), (8, usd), (9, pct), (10, usd)):
            ws.cell(row, col).number_format = fmt
    widths = [7, 11, 9, 9, 9, 13, 11, 13, 13, 14, 34]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[hr].height = 32
    ws.freeze_panes = ws.cell(hr + 1, 1)
    if pf["rows"]:
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(PF_HEAD))}{hr + len(pf['rows'])}"


def write_xlsx(path: Path, strats: list[Strat], results: dict, insts: list[tuple[str, str]],
               info: dict, meta: list[tuple[str, str]], portfolios: list | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head = PatternFill("solid", fgColor="1F2A37")
    sub = PatternFill("solid", fgColor="E5E7EB")
    sumf = PatternFill("solid", fgColor="FEF3C7")
    dpf = PatternFill("solid", fgColor="EDE9FE")
    white = Font(bold=True, color="FFFFFF")
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sep = Border(left=Side(style="thin", color="9CA3AF"))

    wb = Workbook()
    ws = wb.active
    ws.title = "ONE-DAY-STRATEGY"
    present = [(s, g) for s, g in insts if s in info]
    fixed = ["Id", "Opis", "Blok"]
    c0 = len(fixed) + 1                       # pierwsza kolumna podsumowania
    i0 = c0 + len(SUMMARY)                    # pierwsza kolumna spółek

    for j, h in enumerate(fixed, start=1):
        ws.cell(2, j, h)
        ws.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
        ws.cell(1, j, h)
    ws.cell(1, c0, f"PODSUMOWANIE — spółki: {sum(g != 'context' for _, g in present)} (bez SPY, QQQ)")
    ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=i0 - 1)
    for j, h in enumerate(SUMMARY):
        ws.cell(2, c0 + j, h)
    for n, (sym, g) in enumerate(present):
        c = i0 + 4 * n
        ws.cell(1, c, sym + GROUP_TAG[g])
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 3)
        for j, h in enumerate(PER):
            ws.cell(2, c + j, h)
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = center
            cell.font = white if cell.row == 1 else bold
            cell.fill = head if cell.row == 1 else sub
    for j in range(c0, i0):
        ws.cell(2, j).fill = sumf
    for n, (sym, g) in enumerate(present):
        if g == "doubleProof":
            for j in range(4):
                ws.cell(2, i0 + 4 * n + j).fill = dpf

    for r, s in enumerate(strats, start=3):
        res = results[s.id]
        ws.cell(r, 1, s.id).font = bold
        ws.cell(r, 2, s.desc)
        ws.cell(r, 3, s.block)
        for j, v in enumerate(summarize(res, present)):
            cell = ws.cell(r, c0 + j, None if v is None else (round(v, 4) if isinstance(v, float) else v))
            cell.fill = sumf
            if j in (2, 3, 6, 7, 8, 9):
                cell.number_format = PCT
        for n, (sym, _g) in enumerate(present):
            x = res[sym]
            c = i0 + 4 * n
            ws.cell(r, c, x["trades"]).border = sep
            ws.cell(r, c + 1, x["wins"])
            ws.cell(r, c + 2, round(x["ret"], 4)).number_format = PCT
            ws.cell(r, c + 3, round(x["mdd"], 4)).number_format = PCT
            if x["ruined"]:
                ws.cell(r, c + 2).font = Font(bold=True, color="B91C1C")
    ws.freeze_panes = ws.cell(3, c0)
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 6
    for j in range(c0, i0):
        ws.column_dimensions[get_column_letter(j)].width = 12
    for j in range(i0, i0 + 4 * len(present)):
        ws.column_dimensions[get_column_letter(j)].width = 9.5
    ws.row_dimensions[2].height = 45
    ws.auto_filter.ref = f"A2:{get_column_letter(i0 + 4 * len(present) - 1)}{len(strats) + 2}"

    stock_n = sum(g != "context" for _, g in present)
    for s, pf in portfolios or []:
        write_portfolio(wb, s, pf, stock_n)

    op = wb.create_sheet("OPIS")
    op.column_dimensions["A"].width = 34
    op.column_dimensions["B"].width = 110
    for k, v in meta:
        op.append([k, v])
        op.cell(op.max_row, 1).font = bold
        op.cell(op.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    op.append([])
    hdr = ["Spółka", "Grupa", "Pierwszy dzień", "Ostatni dzień", "Dni sesyjne", "Pełne", "Niepełne / brak",
           "Pierwszy dzień handlu (po 55 pełnych sesjach)"]
    op.append(hdr)
    for j in range(1, len(hdr) + 1):
        op.cell(op.max_row, j).font = white
        op.cell(op.max_row, j).fill = head
    for sym, g in present:
        i = info[sym]
        op.append([sym, g, i["first"], i["last"], i["days"], i["valid"], i["days"] - i["valid"], i["start"]])

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".ONE-DAY-", suffix=".xlsx", dir=str(path.parent))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return path


def default_path() -> Path:
    return config.RESEARCH_DIR.parent / "ONE-DAY-STRATEGY.xlsx"      # ~/Desktop/IA4 (katalog repo)


def meta_rows(strats, data_note: str, first: str, last: str) -> list[tuple[str, str]]:
    return [
        ("Test One-Day", f"IA 4 wersja {__version__} · instrukcja, sekcja 6c · policzono "
                         f"{datetime.now().strftime('%Y-%m-%d %H:%M')} · {data_note}"),
        ("Zakres danych", f"{first} – {last} (sesje NYSE z kopii lokalnej)"),
        ("Świeca dobowa", "open = open 1. świecy 1h, close = close ostatniej (7. albo 4. w dzień skrócony), "
                          "high/low z sesji. Sesja niepełna albo brakująca: bez handlu, przerywa serię i kaskadę."),
        ("Transakcja", "zawsze open → close tego samego dnia; mediana trzymania = 1 świeca dobowa (zasada 13); "
                       "bez kosztów transakcyjnych."),
        ("Świeca spadkowa / wzrostowa", "close < open / close > open; doji (close = open) przerywa serię i kończy kaskadę."),
        ("Kaskada", "po N świecach spadkowych z rzędu long w dniu N+1 (krok 1); strata (dzień znów spadkowy) → "
                    "następnego dnia krok 2, 3…; zysk albo doji kończy kaskadę, nowa — gdy seria znów dojdzie do N. "
                    "Short lustrzanie (po N świecach wzrostowych)."),
        ("Portfel", "każda spółka × strategia osobno: 100 000 $; stawka ≤ bieżący kapitał (bez dźwigni); "
                    "kapitał ≤ 10 000 $ = upadłość (koniec handlu, wynik pogrubiony na czerwono)."),
        ("Kwota x", "5 000 $ (5% kapitału startowego)."),
        ("r i 1R", f"r = średni zakres (H−L)/open w % z {R_DAYS} poprzednich pełnych sesji, ustalony na początku "
                   "kaskady; 1R = 500 $ / (r/100) — stawka, przy której ruch o r to 500 $."),
        ("Blok D", "stawka = podstawa × D / r, D = suma korpusów |close−open|/open w % wszystkich świec serii "
                   "(N świec + stratne dni kaskady)."),
        ("Filtr SMA50 (blok C)", f"SMA{SMA_DAYS} z dziennych zamknięć; rosnąca = SMA wczoraj > SMA sprzed {SLOPE_DAYS} "
                                 "pełnych sesji; sprawdzany tylko przy pierwszym wejściu kaskady."),
        ("Blok E", f"kaskada ×2 najwyżej {MAX_STEPS_E} kroki; po 4. stracie czeka na nową serię."),
        ("Wspólny start", f"pierwsza transakcja po {WARMUP} pełnych sesjach instrumentu (rozgrzewka r i SMA50) "
                          "— dla wszystkich strategii ta sama."),
        ("Kolumny spółki", "Transakcje · Zyskowne (wynik > 0) · Wynik % (kapitał końcowy / 100 000 − 1) · "
                           "Obsunięcie % (największy spadek krzywej kapitału od szczytu, po zamknięciu każdego dnia)."),
        ("Podsumowanie", "spółki bez SPY i QQQ (ETF-y tylko informacyjnie); dP = doubleProof."),
        ("Arkusze PORTFEL …", f"jedna strategia na wszystkich spółkach naraz, wspólny portfel {PORTFOLIO_CAPITAL:,.0f} $"
                              .replace(",", " ") + "; kolejne transakcje dzień po dniu (w dniu — kolejność spółek z sekcji 4). "
                              "Jednostka stawki (x = 5%, 1R = 0,5% / r) od stanu portfela na początku kaskady spółki "
                              "(rośnie i maleje z portfelem); suma stawek dnia > stan portfela na początku dnia → wszystkie "
                              "stawki dnia zmniejszone proporcjonalnie (bez dźwigni); ≤ 10% startu = upadłość. "
                              "Świeca w serii = N + krok (np. 3 = trzecia świeca po dwóch spadkowych). "
                              "Wynik transakcji % = ruch open → close w kierunku pozycji; obsunięcie — po zamknięciu każdego dnia."),
        ("Uwaga", "Wynik informacyjny na ok. 2 latach danych. Strategie One-Day nie trafiają do paper tradingu "
                  "ani do AI Investor. doubleProof policzone za zgodą człowieka (6a, 6c)."),
        ("Strategii", str(len(strats))),
    ]


def run(no_sync: bool = False, out: Path | None = None, pf_ids: tuple[str, ...] = PORTFOLIO_STRATS,
        pf_capital: float = PORTFOLIO_CAPITAL) -> Path:
    t0 = time.time()
    print("IA 4 — test One-Day (instrukcja, sekcja 6c)\n")
    u = universe()
    insts = ordered(u)
    syms = [s for s, _ in insts]
    note = "bez synchronizacji (--no-sync)" if no_sync else maybe_sync(syms)
    print(f"  {note}")
    from .sync import load, parquet_path
    data: dict[str, Days] = {}
    info: dict[str, dict] = {}
    for s in syms:
        if not parquet_path(s).exists():
            print(f"  ! {s}: brak kopii lokalnej — pominięta")
            continue
        dd = daily(load(s))
        if not dd.dates:
            continue
        data[s] = dd
        vi = [t for t, v in enumerate(dd.valid) if v]
        info[s] = {"first": dd.dates[0], "last": dd.dates[-1], "days": len(dd.dates), "valid": len(vi),
                   "start": dd.dates[vi[WARMUP]] if len(vi) > WARMUP else "—"}
    if not data:
        raise SystemExit("Brak danych w kopii lokalnej — uruchom najpierw `python -m ia4.sync`.")
    strats = strategies()
    results = {st.id: {s: simulate(st, dd) for s, dd in data.items()} for st in strats}
    first = min(i["first"] for i in info.values())
    last = max(i["last"] for i in info.values())
    by_id = {st.id: st for st in strats}
    stock_syms = [s for s, g in insts if g != "context" and s in data]
    pfs = []
    for pid in pf_ids:
        if pid not in by_id:
            raise SystemExit(f"Nieznana strategia {pid} (OD-001 … OD-100).")
        pfs.append((by_id[pid], portfolio(by_id[pid], data, stock_syms, pf_capital)))
    path = write_xlsx(out or default_path(), strats, results, insts, info,
                      meta_rows(strats, note, first, last), pfs)
    present = [(s, g) for s, g in insts if s in info]
    rank = sorted(strats, key=lambda st: -(summarize(results[st.id], present)[3] or -1e9))
    print(f"\n  {len(data)} instrumentów · {len(strats)} strategii · dane {first} – {last} · {time.time() - t0:.0f} s")
    print("  Najwyższa mediana wyniku % (spółki):")
    for st in rank[:5]:
        sm = summarize(results[st.id], present)
        print(f"    {st.id}  mediana {sm[3]:+8.2f}%  średnia {sm[2]:+8.2f}%  na plusie {sm[4]:>7}  · {st.desc}")
    for st, pf in pfs:
        print(f"\n  PORTFEL {st.id} ({pf['capital']:,.0f} $, {len(stock_syms)} spółek): koniec {pf['end']:,.2f} $ · "
              f"{pf['ret']:+.2f}% · obsunięcie {pf['mdd']:.2f}% · transakcji {pf['trades']} "
              f"(zyskownych {pf['wins']}) · dni przycięte {pf['scaled']}".replace(",", " "))
    print(f"\nPlik: {path}")
    return path


# ---------------------------------------------------------------------------
#  Selftest (dane syntetyczne, wyniki liczone ręcznie)
# ---------------------------------------------------------------------------
def _frame(days: list[tuple[str, float, float]], drop: dict | None = None):
    """[(data, open, close)] → świece 1h; high/low z marginesem 1%; drop: {data: [sloty do usunięcia]}."""
    import pandas as pd
    rows = []
    for d, o, c in days:
        last = 4 if d in US_HALF_DAYS else 7
        for s in range(1, last + 1):
            if drop and s in drop.get(d, []):
                continue
            a = o + (c - o) * (s - 1) / last
            b = o + (c - o) * s / last
            rows.append({"symbol": "T", "date": d, "slot": s, "o": a, "h": max(a, b) * 1.01,
                         "l": min(a, b) * 0.99, "c": b, "v": 0})
    return pd.DataFrame(rows)


def selftest() -> None:
    fails: list[str] = []

    def check(name, got, exp, tol=1e-6):
        if isinstance(exp, float):
            ok = got is not None and abs(got - exp) <= tol * max(1.0, abs(exp))
        elif isinstance(exp, list) and exp and isinstance(exp[0], float):
            ok = len(got) == len(exp) and all(abs(g - e) <= tol * max(1.0, abs(e)) for g, e in zip(got, exp))
        else:
            ok = got == exp
        print(f"  {'OK ' if ok else 'BŁĄD'} {name}: {got!r}" + ("" if ok else f" (oczekiwano {exp!r})"))
        if not ok:
            fails.append(name)

    print("IA 4 — test One-Day: selftest\n")
    st = {s.id: s for s in strategies()}
    check("100 strategii", len(st), 100)
    check("OD-003 = long 2 ×2 x", (st["OD-003"].dir, st["OD-003"].n, st["OD-003"].prog, st["OD-003"].size), (1, 2, "x2", "x"))
    check("OD-045 = short 2 ×2 1R", (st["OD-045"].dir, st["OD-045"].n, st["OD-045"].prog, st["OD-045"].size), (-1, 2, "x2", "R"))
    check("OD-058 = long 2 1R spadkowa", (st["OD-058"].dir, st["OD-058"].n, st["OD-058"].size, st["OD-058"].filt), (1, 2, "R", "down"))
    check("OD-084 = long 4 D 1R", (st["OD-084"].n, st["OD-084"].dsize, st["OD-084"].size), (4, True, "R"))
    check("OD-098 = short 3 1R limit", (st["OD-098"].dir, st["OD-098"].n, st["OD-098"].size, st["OD-098"].max_steps), (-1, 3, "R", 4))
    check("OD-100 = short codziennie", (st["OD-100"].dir, st["OD-100"].daily), (-1, True))
    check("Fibonacci 1-1-2-3-5", [mult("fib", k) for k in range(1, 6)], [1.0, 1.0, 2.0, 3.0, 5.0])

    # Rozgrzewka: 60 sesji naprzemiennie −1% / +1% (seria nigdy ≥ 2, ostatnia wzrostowa), cena 100.
    days = trading_days("2025-01-02", "2025-12-31")
    seq: list[tuple[str, float, float]] = []
    price = 100.0
    for i in range(60):
        c = price * (0.99 if i % 2 == 0 else 1.01)
        seq.append((days[i], price, c))
        price = c
    # Ogon: −2%, −2% (seria 2), −3% (krok 1 — strata), −1% (krok 2 — strata), +4% (krok 3 — zysk), +1%, −0,5%.
    tail = [-0.02, -0.02, -0.03, -0.01, 0.04, 0.01, -0.005]
    for j, m in enumerate(tail):
        c = price * (1 + m)
        seq.append((days[60 + j], price, c))
        price = c
    dd = daily(_frame(seq))
    check("dni pełne", sum(dd.valid), 67)
    check("kolory ogona", dd.color[59:67], [1, -1, -1, -1, -1, 1, 1, -1])

    def rng_before(t):                                    # r niezależnie od daily(): zakres z (open, close) ±1%
        xs = [(max(o, c) * 1.01 - min(o, c) * 0.99) / o * 100 for _, o, c in seq[t - R_DAYS:t]]
        return statistics.fmean(xs)

    check("r przed dniem 62", dd.r[62], rng_before(62))

    def replay(stakes, moves, d=1):                       # oczekiwany kapitał po stawkach (z limitem kapitału)
        eq, out = CAPITAL, []
        for st_, m in zip(stakes, moves):
            st_ = min(st_, eq)
            out.append(round(st_, 6))
            eq += st_ * d * m
        return out, eq

    tr: list = []
    res = simulate(st["OD-003"], dd, tr)                  # long po 2, ×2, x
    check("OD-003 dni", [x[0] for x in tr], [seq[62][0], seq[63][0], seq[64][0]])
    check("OD-003 stawki", [x[2] for x in tr], [5000.0, 10000.0, 20000.0])
    check("OD-003 transakcje / zyskowne", (res["trades"], res["wins"]), (3, 1))
    exp_eq = CAPITAL + 5000 * -0.03 + 10000 * -0.01 + 20000 * 0.04
    check("OD-003 wynik %", res["ret"], (exp_eq / CAPITAL - 1) * 100)
    check("OD-003 obsunięcie %", res["mdd"], (5000 * 0.03 + 10000 * 0.01) / CAPITAL * 100)

    res = simulate(st["OD-008"], dd)                      # long po 3: wejście −1%, potem +4%
    check("OD-008 wynik %", res["ret"], (5000 * -0.01 + 10000 * 0.04) / CAPITAL * 100)
    res = simulate(st["OD-013"], dd)                      # long po 4: wejście +4%
    check("OD-013 wynik %", res["ret"], 5000 * 0.04 / CAPITAL * 100)
    res = simulate(st["OD-005"], dd, tr := [])            # Fibonacci 1, 1, 2
    check("OD-005 stawki", [x[2] for x in tr], [5000.0, 5000.0, 10000.0])
    res = simulate(st["OD-002"], dd, tr := [])            # liniowa 1, 2, 3
    check("OD-002 stawki", [x[2] for x in tr], [5000.0, 10000.0, 15000.0])
    res = simulate(st["OD-004"], dd, tr := [])            # ×3: 1, 3, 9
    check("OD-004 stawki", [x[2] for x in tr], [5000.0, 15000.0, 45000.0])
    res = simulate(st["OD-018"], dd, tr := [])            # short po 2 wzrostach: wejście w −0,5% (zysk)
    check("OD-018 dni", [x[0] for x in tr], [seq[66][0]])
    check("OD-018 wynik %", res["ret"], 5000 * 0.005 / CAPITAL * 100)

    r62 = rng_before(62)
    base = 500 / (r62 / 100)
    exp, eq = replay([base, 2 * base, 4 * base], [-0.03, -0.01, 0.04])
    res = simulate(st["OD-033"], dd, tr := [])            # long po 2, ×2, 1R (r z początku kaskady)
    check("OD-033 stawki", [x[2] for x in tr], exp, tol=1e-6)
    check("OD-033 wynik %", res["ret"], (eq / CAPITAL - 1) * 100)
    # D: krok 1 — seria 2% + 2% = 4; krok 2 — + 3% = 7; krok 3 — + 1% = 8 (r z początku kaskady).
    res = simulate(st["OD-079"], dd, tr := [])
    check("OD-079 stawki", [x[2] for x in tr], [5000 * 4 / r62, 5000 * 7 / r62, 5000 * 8 / r62], tol=1e-4)

    # Blok E: 8 spadków po −1% po rozgrzewce, potem +2%.
    seq2, p2 = list(seq[:60]), seq[59][2]
    for j in range(8):
        seq2.append((days[60 + j], p2, p2 * 0.99))
        p2 *= 0.99
    seq2.append((days[68], p2, p2 * 1.02))
    dd2 = daily(_frame(seq2))
    res = simulate(st["OD-091"], dd2, tr := [])           # najwyżej 4 kroki, potem czeka na nową serię
    check("OD-091 kroki", [x[1] for x in tr], [1, 2, 3, 4])
    res = simulate(st["OD-003"], dd2, tr := [])           # bez limitu: 7 kroków
    check("OD-003 kroki bez limitu", [x[1] for x in tr], [1, 2, 3, 4, 5, 6, 7])
    check("stawka ≤ kapitał", all(x[2] <= x[4] - x[3] + 1e-6 for x in tr), True)

    # Niepełna sesja (brak 7. świecy w drugim dniu spadkowym) przerywa serię.
    dd3 = daily(_frame(seq, drop={seq[61][0]: [7]}))
    check("niepełna sesja", dd3.valid[61], False)
    res = simulate(st["OD-003"], dd3, tr := [])           # seria od nowa: dni 62, 63 → wejście dnia 64 (+4%)
    check("OD-003 po niepełnej", [x[0] for x in tr], [seq[64][0]])
    dd3b = daily(_frame(seq, drop={seq[61][0]: [3]}))     # brak świecy w środku sesji — dzień pełny
    check("brak świecy w środku", dd3b.valid[61], True)
    hd = sorted(d for d in US_HALF_DAYS if d >= "2025-01-01")[0]
    dd4 = daily(_frame([(hd, 100.0, 101.0)]))
    check("dzień skrócony (4 świece)", (dd4.valid, round(dd4.c[0], 6)), ([True], 101.0))

    # Filtr SMA50: trend wzrostowy → OD-055 (rosnąca) handluje, OD-056 (spadkowa) nie; short OD-069/070 lustrzanie.
    seq5, p5 = [], 100.0
    for i in range(60):
        c = p5 * (0.995 if i % 2 == 0 else 1.02)
        seq5.append((days[i], p5, c))
        p5 = c
    for j, m in enumerate([-0.01, -0.01, 0.02]):
        seq5.append((days[60 + j], p5, p5 * (1 + m)))
        p5 *= 1 + m
    dd5 = daily(_frame(seq5))
    check("SMA50 rosnąca", dd5.sma_up[62], 1)
    check("OD-055 handluje", simulate(st["OD-055"], dd5)["trades"], 1)
    check("OD-056 nie handluje", simulate(st["OD-056"], dd5)["trades"], 0)

    # Codziennie (F) i obsunięcie szczyt → dołek.
    res = simulate(st["OD-099"], dd, tr := [])
    check("OD-099 transakcje", res["trades"], 67 - WARMUP)
    pk, md = CAPITAL, 0.0
    for e in [CAPITAL] + [x[4] for x in tr]:
        pk = max(pk, e)
        md = max(md, (pk - e) / pk * 100)
    check("OD-099 obsunięcie", res["mdd"], md)
    res100 = simulate(st["OD-100"], dd)
    check("OD-099 + OD-100 = 0", round(res["ret"] + res100["ret"], 9), 0.0)

    # Upadłość: long ×2 na ciągłych spadkach −30%.
    seq6, p6 = list(seq[:60]), seq[59][2]
    for j in range(14):
        seq6.append((days[60 + j], p6, p6 * 0.7))
        p6 *= 0.7
    res = simulate(st["OD-003"], daily(_frame(seq6)), tr := [])
    check("upadłość", (res["ruined"], res["ret"] <= RUIN_PCT - 100), (True, True))
    check("po upadłości bez handlu", len(tr), 10)

    # Wspólny portfel 10 000 $: trzy spółki z tym samym przebiegiem, OD-003 (x, ×2) i OD-004 (x, ×3).
    data3 = {"A": dd, "B": dd, "C": dd}
    pf = portfolio(st["OD-003"], data3, ["A", "B", "C"])
    check("portfel: transakcje", (pf["trades"], pf["wins"], pf["days"]), (9, 3, 3))
    check("portfel: kolejność", [(r["day"], r["sym"]) for r in pf["rows"][:3]], [(seq[62][0], "A"), (seq[62][0], "B"), (seq[62][0], "C")])
    check("portfel: świeca w serii", [r["candle"] for r in pf["rows"][::3]], [3, 4, 5])
    check("portfel: stawki od stanu na początku kaskady", [round(r["stake"], 6) for r in pf["rows"][::3]], [500.0, 1000.0, 2000.0])
    check("portfel: koniec", pf["end"], 10000 + 3 * (500 * -0.03 + 1000 * -0.01 + 2000 * 0.04))
    check("portfel: obsunięcie", pf["mdd"], 3 * (15 + 10) / 10000 * 100)
    pf = portfolio(st["OD-004"], data3, ["A", "B", "C"])  # krok 3: 3 × 4500 > stan portfela → przycięcie
    e3 = 10000 - 3 * (500 * 0.03 + 1500 * 0.01)
    check("portfel: przycięcie dnia", (pf["scaled"], round(pf["rows"][-1]["scale"], 9)), (1, round(e3 / 13500, 9)))
    check("portfel: koniec po przycięciu", pf["end"], e3 * 1.04)
    check("portfel: suma stawek ≤ stan", round(sum(r["stake"] for r in pf["rows"][-3:]), 6), round(e3, 6))
    pf = portfolio(st["OD-003"], {"A": daily(_frame(seq6)), "B": dd}, ["A", "B"])
    check("portfel: upadłość", (pf["ruined"], pf["end"] <= 1000), (True, True))
    portfolios = [(st["OD-003"], portfolio(st["OD-003"], data3, ["A", "B", "C"]))]

    # Plik Excel z jednym instrumentem.
    out = Path(tempfile.mkdtemp()) / "t.xlsx"
    strats = strategies()
    results = {s.id: {"AAPL": simulate(s, dd)} for s in strats}
    info = {"AAPL": {"first": dd.dates[0], "last": dd.dates[-1], "days": len(dd.dates), "valid": 67, "start": dd.dates[55]}}
    write_xlsx(out, strats, results, [("AAPL", "main")], info, meta_rows(strats, "selftest", dd.dates[0], dd.dates[-1]),
               portfolios)
    from openpyxl import load_workbook
    wbk = load_workbook(out)
    check("Excel: arkusze", wbk.sheetnames, ["ONE-DAY-STRATEGY", "PORTFEL OD-003", "OPIS"])
    pw = wbk["PORTFEL OD-003"]
    check("Excel: portfel wiersze", pw.max_row, 6 + 9)
    check("Excel: portfel ostatni stan", pw.cell(15, 10).value, round(10000 + 3 * (500 * -0.03 + 1000 * -0.01 + 2000 * 0.04), 2))
    ws = wbk["ONE-DAY-STRATEGY"]
    check("Excel: wiersze i kolumny", (ws.max_row, ws.max_column), (102, 3 + len(SUMMARY) + 4))
    check("Excel: OD-003 wynik", ws.cell(5, 3 + len(SUMMARY) + 3).value, round(results["OD-003"]["AAPL"]["ret"], 4))

    print("\nWYNIK:", "OK" if not fails else f"BŁĘDY ({len(fails)}): " + ", ".join(fails))
    sys.exit(0 if not fails else 1)


def main() -> None:
    ap = argparse.ArgumentParser(description="IA 4 — test One-Day (instrukcja, sekcja 6c)")
    ap.add_argument("--no-sync", action="store_true", help="bez synchronizacji z Firestore")
    ap.add_argument("--out", type=Path, help="ścieżka pliku (domyślnie ~/Desktop/IA4/ONE-DAY-STRATEGY.xlsx)")
    ap.add_argument("--portfolio", nargs="+", default=list(PORTFOLIO_STRATS), metavar="OD-NNN",
                    help="strategie do arkuszy PORTFEL (domyślnie OD-080)")
    ap.add_argument("--capital", type=float, default=PORTFOLIO_CAPITAL, help="kapitał portfela (domyślnie 10 000 $)")
    ap.add_argument("--selftest", action="store_true", help="test na danych syntetycznych")
    a = ap.parse_args()
    if a.selftest:
        selftest()
    else:
        run(no_sync=a.no_sync, out=a.out, pf_ids=tuple(x.upper() for x in a.portfolio), pf_capital=a.capital)


if __name__ == "__main__":
    main()
