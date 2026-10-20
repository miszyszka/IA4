"""
IA 4 — FX-ODKRYWCA: automatyczne poszukiwanie strategii FX na pierwszym roku i sprawdzian 100 najlepszych w skarbcu
(drugi rok, niewidziany przez poszukiwanie) (instrukcja, sekcja 7b). Wersja projektu: 1.54 (2026-10-20) — musi zgadzać
się z IA4_INSTRUKCJA.md

Zadanie człowieka (20.10.2026): zamiast zgadywać pojedyncze pomysły — algorytm, który przelicza miliony sygnałów
i wyjść, wybiera serię 100 najlepszych strategii WYŁĄCZNIE na pierwszym roku (największy wzrost, najwięcej wystąpień,
najmniejsze obsunięcie, największa konsekwencja), a potem przelicza je raz na drugim roku (skarbiec) i pokazuje całość
w Excelu, także jako portfel wielu strategii grających równolegle (1000 $, dźwignia 1:30).

Jak liczy:
  * CECHY — ok. 90 zmiennych z każdej świecy 1h (ruch 1–120 świec, kształt świecy, położenie w zakresie 12–240 świec,
    odległość i nachylenie średnich, RSI / stochastyk / CCI / Bollinger / MACD, zmienność i jej reżim, ruch dnia,
    wczorajszy zakres, wybicie zakresu Azji, siła walut z 10 par, BTC jako nastrój rynku, świece od ostatniego impulsu)
    oraz cechy kategorii (godzina, dzień tygodnia, początek / koniec miesiąca, NR4 / NR7, świeca wewnętrzna / objęcia).
  * ATOMY — warunek „cecha ≤ / ≥ próg”, progi = kwantyle 5–95% z PIERWSZEGO ROKU (wszystkie pary razem).
  * WEJŚCIE — 1 do 4 atomów naraz (różne cechy), kierunek long albo short, po open następnej świecy.
  * WYJŚCIA — ok. 270: SL / TP w ATR (0,5–3 albo bez) × limit 1–64 świec, stop kroczący 1–3 ATR, cel = EMA20,
    zamknięcie po sesji NY. Rozliczenie pesymistyczne (zasada 5), spread raz na transakcję (EUR/USD 1,1 pipsa).
  * ETAPY — 1: pojedyncze atomy; 2: wszystkie pary atomów; 3 i 4: najlepsze kombinacje + jeszcze jeden atom
    (szerokość wiązki dobrana do budżetu czasu `--godziny`). Liczy na wielu procesach (`--procesy`).
  * PULA — reguły z roku 1: ≥ 30 transakcji, suma > 0, PF ≥ 1,15, t ≥ 2, ≥ 6 z 10 okresów na plusie.
  * OCENA — średnia z 4 miejsc w rankingu puli (percentyle): wzrost (suma % netto), wystąpienia (transakcje),
    obsunięcie (mniejsze lepiej), konsekwencja (½ okresy na plusie + ½ R² krzywej). Seria 100: od najlepszej oceny,
    bez powtórzeń tego samego wejścia i z najwyżej 50% wspólnych transakcji z już wybraną strategią.
  * SKARBIEC — rok 2 liczony dopiero po zamrożeniu listy 100 (zapis listy i każdego otwarcia w data/fxodkrywca/).
    Dla porównania w skarbcu: 100 losowych reguł z puli i 100 najlepszych wg samego t.

    python -m ia4.fxodkrywca                  # ok. 3 h (--godziny), wszystkie rdzenie → ~/Desktop/IA4/FX-ODKRYWCA.xlsx
    python -m ia4.fxodkrywca --godziny 6      # dłużej = szersze etapy 3–4
    python -m ia4.fxodkrywca --kierunek short # tylko shorty
    python -m ia4.fxodkrywca --selftest       # „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import sys
import time
from pathlib import Path

import numba as nb
import numpy as np
import pandas as pd

from . import __version__, config
from . import fxshort as fs
from .fxshort import PAIRS, PIP

GRANICA = '2025-10-15'          # początek skarbca (drugi rok danych)
WIDZIANE_DO = '2026-10-17'      # do tej daty Claude oglądał dane przy wcześniejszych badaniach (1.49–1.50)
NPER = 10                        # okresów w każdym roku
MIN_N = 30
WARM = 600
CACHE = config.CACHE_DIR / 'fxodkrywca'
CUR = ['USD', 'EUR', 'JPY', 'GBP', 'AUD', 'CAD', 'CHF', 'NZD']

# ---------------------------------------------------------------- cechy


def _sma(x, n):
    return pd.Series(x).rolling(n, min_periods=n).mean().values


def _ema(x, n):
    return pd.Series(x).ewm(span=n, adjust=False).mean().values


def _std(x, n):
    return pd.Series(x).rolling(n, min_periods=n).std().values


def _rmax(x, n):
    return pd.Series(x).rolling(n, min_periods=n).max().values


def _rmin(x, n):
    return pd.Series(x).rolling(n, min_periods=n).min().values


def _lag(x, k):
    return np.r_[np.full(k, np.nan), x[:-k]]


def _rsi(c, n):
    d = np.diff(c, prepend=c[0]); up = np.maximum(d, 0); dn = np.maximum(-d, 0)
    au = pd.Series(up).ewm(alpha=1 / n, adjust=False).mean().values
    ad = pd.Series(dn).ewm(alpha=1 / n, adjust=False).mean().values
    return 100 - 100 / (1 + au / np.maximum(ad, 1e-12))


@nb.njit(cache=True)
def _since(flag, cap):
    out = np.empty(flag.size); last = -100000
    for i in range(flag.size):
        if flag[i]:
            last = i
        out[i] = min(i - last, cap)
    return out


@nb.njit(cache=True)
def _mad(x, n):
    out = np.full(x.size, np.nan)
    for i in range(n - 1, x.size):
        m = 0.0
        for j in range(i - n + 1, i + 1):
            m += x[j]
        m /= n
        s = 0.0
        for j in range(i - n + 1, i + 1):
            s += abs(x[j] - m)
        out[i] = s / n
    return out


@nb.njit(cache=True)
def _streak(g, r):
    out = np.zeros(g.size)
    for i in range(g.size):
        if g[i]:
            out[i] = out[i - 1] + 1 if i > 0 and out[i - 1] > 0 else 1
        elif r[i]:
            out[i] = out[i - 1] - 1 if i > 0 and out[i - 1] < 0 else -1
    return out


class Rynek:
    """10 par FX (+ BTC jako nastrój) na wspólnej siatce czasu, cechy wszystkich świec obu lat."""

    def __init__(self, frames: dict, spread: dict, btc: pd.DataFrame | None = None, granica=GRANICA):
        common = None
        for p in PAIRS:
            tt = set(frames[p].time.values.tolist())
            common = tt if common is None else common & tt
        t = np.array(sorted(common), np.int64)
        self.t = t; self.n = t.size; self.spread = spread
        self.b = int(np.searchsorted(t, int(pd.Timestamp(granica, tz='UTC').timestamp())))
        self.swieze = int(np.searchsorted(t, int(pd.Timestamp(WIDZIANE_DO, tz='UTC').timestamp())))
        D = {p: frames[p].set_index('time').loc[t] for p in PAIRS}
        self.o = np.array([D[p].o.values for p in PAIRS], np.float64)
        self.h = np.array([D[p].h.values for p in PAIRS], np.float64)
        self.l = np.array([D[p].l.values for p in PAIRS], np.float64)
        self.c = np.array([D[p].c.values for p in PAIRS], np.float64)
        filled = np.array([D[p].filled.values for p in PAIRS], bool)
        self.F, self.OPIS, self.KAT = {}, {}, {}
        import warnings
        with warnings.catch_warnings(), np.errstate(all='ignore'):
            warnings.simplefilter('ignore', RuntimeWarning)
            self._cechy(btc)
        nt = np.r_[t[1:], t[-1] + 10 ** 9]
        self.valid = (np.isfinite(self.atr) & np.isfinite(self.F['zm_100_500']) & ((nt - t) == 3600)[None, :] & ~filled)
        self.valid[:, :WARM] = False
        # okresy: 10 w roku 1 i 10 w skarbcu
        T0, TB, T1 = int(t[min(WARM, self.n - 1)]), int(t[self.b]) if self.b < self.n else int(t[-1]), int(t[-1]) + 3600
        self.per = np.where(np.arange(self.n) < self.b,
                            np.clip(((t - T0) / max(TB - T0, 1) * NPER).astype(np.int64), 0, NPER - 1),
                            NPER + np.minimum(((t - TB) / max(T1 - TB, 1) * NPER).astype(np.int64), NPER - 1))
        self.periods = ([(pd.Timestamp(T0 + (TB - T0) * k / NPER, unit='s'), pd.Timestamp(T0 + (TB - T0) * (k + 1) / NPER, unit='s')) for k in range(NPER)] +
                        [(pd.Timestamp(TB + (T1 - TB) * k / NPER, unit='s'), pd.Timestamp(TB + (T1 - TB) * (k + 1) / NPER, unit='s')) for k in range(NPER)])

    def add(self, name, arr, opis, kat=False):
        self.F[name] = np.asarray(arr, np.float64); self.OPIS[name] = opis; self.KAT[name] = kat

    def _cechy(self, btc):
        o, h, l, c = self.o, self.h, self.l, self.c
        npair, n = c.shape
        pc = np.concatenate([c[:, :1], c[:, :-1]], axis=1)
        tr = np.maximum(h - l, np.maximum(np.abs(h - pc), np.abs(l - pc)))
        atr = np.array([_sma(tr[k], 100) for k in range(npair)]); self.atr = atr
        atr24 = np.array([_sma(tr[k], 24) for k in range(npair)]); atr500 = np.array([_sma(tr[k], 500) for k in range(npair)])
        A = atr
        R = {}
        for k in [1, 2, 3, 4, 6, 8, 12, 24, 48, 96, 120]:
            R[k] = (c - np.array([_lag(c[q], k) for q in range(npair)])) / A
            self.add(f'ruch_{k}', R[k], f'zmiana ceny w {k} św. (w ATR)')
        self.add('korpus', (c - o) / A, 'korpus świecy (close − open, w ATR)')
        self.add('knot_gora', (h - np.maximum(o, c)) / A, 'górny knot (w ATR)')
        self.add('knot_dol', (np.minimum(o, c) - l) / A, 'dolny knot (w ATR)')
        self.add('zakres', (h - l) / A, 'zakres świecy (w ATR)')
        self.add('poz_w_swiecy', (c - l) / np.maximum(h - l, 1e-12), 'położenie close w świecy (0 = dołek, 1 = szczyt)')
        self.add('luka', (o - pc) / A, 'luka na otwarciu (w ATR)')
        self.add('korpus_poprz', np.array([_lag(((c - o) / A)[q], 1) for q in range(npair)]), 'korpus poprzedniej świecy (w ATR)')
        self.add('seria', np.array([_streak(c[q] > o[q], c[q] < o[q]) for q in range(npair)]), 'seria świec (+ zielone, − czerwone)')
        for N in [12, 24, 48, 120, 240]:
            hh = np.array([_rmax(h[q], N) for q in range(npair)]); ll = np.array([_rmin(l[q], N) for q in range(npair)])
            self.add(f'od_szczytu_{N}', (c - hh) / A, f'odległość od szczytu {N} św. (w ATR)')
            self.add(f'od_dolka_{N}', (c - ll) / A, f'odległość od dołka {N} św. (w ATR)')
            self.add(f'poz_zakres_{N}', (c - ll) / np.maximum(hh - ll, 1e-12), f'położenie w zakresie {N} św. (0–1)')
            self.add(f'szer_zakres_{N}', (hh - ll) / A, f'szerokość zakresu {N} św. (w ATR, mała = konsolidacja)')
        E = {}
        for m in [10, 20, 50, 100, 200, 500]:
            E[m] = np.array([_ema(c[q], m) for q in range(npair)])
            self.add(f'od_ema_{m}', (c - E[m]) / A, f'odległość od EMA{m} (w ATR)')
        for m in [20, 50, 200, 500]:
            self.add(f'nachyl_ema_{m}', (E[m] - np.array([_lag(E[m][q], 24) for q in range(npair)])) / A, f'nachylenie EMA{m} w 24 św. (w ATR)')
        self.ema20 = E[20]
        self.add('ema20_50', (E[20] - E[50]) / A, 'EMA20 − EMA50 (w ATR)')
        self.add('ema50_200', (E[50] - E[200]) / A, 'EMA50 − EMA200 (w ATR)')
        self.add('ema200_500', (E[200] - E[500]) / A, 'EMA200 − EMA500 (w ATR)')
        for m in [2, 3, 7, 14, 28]:
            self.add(f'rsi_{m}', np.array([_rsi(c[q], m) for q in range(npair)]), f'RSI({m})')
        hh14 = np.array([_rmax(h[q], 14) for q in range(npair)]); ll14 = np.array([_rmin(l[q], 14) for q in range(npair)])
        self.add('stoch_14', (c - ll14) / np.maximum(hh14 - ll14, 1e-12) * 100, 'stochastyk %K(14)')
        tp = (h + l + c) / 3
        sma_tp = np.array([_sma(tp[q], 20) for q in range(npair)])
        mad = np.array([_mad(tp[q], 20) for q in range(npair)])
        self.add('cci_20', (tp - sma_tp) / np.maximum(0.015 * mad, 1e-12), 'CCI(20)')
        for m in [20, 50]:
            sm = np.array([_sma(c[q], m) for q in range(npair)]); sd = np.array([_std(c[q], m) for q in range(npair)])
            self.add(f'bb_z_{m}', (c - sm) / np.maximum(sd, 1e-12), f'Bollinger: odchylenie od SMA{m} w σ')
            if m == 20:
                self.add('bb_szer_20', 4 * sd / A, 'szerokość Bollingera(20) w ATR (mała = ściśnięcie)')
        macd = np.array([_ema(c[q], 12) - _ema(c[q], 26) for q in range(npair)])
        sig = np.array([_ema(macd[q], 9) for q in range(npair)])
        hist = (macd - sig) / A
        self.add('macd_hist', hist, 'MACD histogram (w ATR)')
        self.add('macd_hist_zm', hist - np.array([_lag(hist[q], 3) for q in range(npair)]), 'zmiana histogramu MACD w 3 św.')
        self.add('zm_24_100', atr24 / A, 'zmienność 24 św. / 100 św. (> 1 = rośnie)')
        self.add('zm_100_500', A / atr500, 'zmienność 100 św. / 500 św.')
        self.add('atr_proc', A / c * 100, 'ATR(100) w % ceny (poziom zmienności)')
        # dzień UTC, wczoraj, Azja
        day = self.t // 86400; hour = (self.t % 86400) // 3600
        dz_ruch = np.full((npair, n), np.nan); dz_zakres = np.full((npair, n), np.nan)
        wcz_h = np.full((npair, n), np.nan); wcz_l = np.full((npair, n), np.nan); wcz_ruch = np.full((npair, n), np.nan)
        az_h = np.full((npair, n), np.nan); az_l = np.full((npair, n), np.nan)
        for q in range(npair):
            df = pd.DataFrame({'d': day, 'hr': hour, 'o': o[q], 'h': h[q], 'l': l[q], 'c': c[q]})
            g = df.groupby('d')
            dopen = g.o.transform('first'); dh = g.h.cummax(); dl = g.l.cummin()
            dz_ruch[q] = (c[q] - dopen.values) / A[q]; dz_zakres[q] = (dh.values - dl.values) / A[q]
            daily = g.agg(o=('o', 'first'), h=('h', 'max'), l=('l', 'min'), c=('c', 'last'))
            prev = daily.shift(1)
            wcz_h[q] = (c[q] - df.d.map(prev.h).values) / A[q]; wcz_l[q] = (c[q] - df.d.map(prev.l).values) / A[q]
            wcz_ruch[q] = (df.d.map(prev.c).values - df.d.map(prev.o).values) / A[q]
            asia = df[df.hr <= 6].groupby('d').agg(h=('h', 'max'), l=('l', 'min'))
            m7 = (df.hr >= 7).values
            az_h[q] = np.where(m7, (c[q] - df.d.map(asia.h).values) / A[q], np.nan)
            az_l[q] = np.where(m7, (c[q] - df.d.map(asia.l).values) / A[q], np.nan)
        self.add('dzien_ruch', dz_ruch, 'ruch od otwarcia dnia UTC (w ATR)')
        self.add('dzien_zakres', dz_zakres, 'zakres dnia UTC do tej świecy (w ATR)')
        self.add('od_wczoraj_szczytu', wcz_h, 'odległość od wczorajszego szczytu (w ATR)')
        self.add('od_wczoraj_dolka', wcz_l, 'odległość od wczorajszego dołka (w ATR)')
        self.add('wczoraj_ruch', wcz_ruch, 'wczorajsza zmiana dnia (w ATR)')
        self.add('od_azji_szczytu', az_h, 'po 7 UTC: odległość od szczytu sesji azjatyckiej 0–6 UTC (w ATR)')
        self.add('od_azji_dolka', az_l, 'po 7 UTC: odległość od dołka sesji azjatyckiej (w ATR)')
        # siła walut
        for k in [4, 24]:
            strength = {}
            for cur in CUR:
                parts = [R[k][q] if PAIRS[q][:3] == cur else -R[k][q] for q in range(npair) if cur in (PAIRS[q][:3], PAIRS[q][3:])]
                strength[cur] = np.nanmean(np.array(parts), axis=0)
            sb = np.array([strength[p[:3]] for p in PAIRS]); sq = np.array([strength[p[3:]] for p in PAIRS])
            self.add(f'sila_bazowej_{k}', sb, f'siła waluty bazowej w {k} św. (średnio wobec pozostałych, w ATR)')
            self.add(f'sila_kwotowanej_{k}', sq, f'siła waluty kwotowanej w {k} św.')
            self.add(f'rezydualny_{k}', R[k] - (sb - sq), f'ruch pary w {k} św. ponad różnicę sił walut (ruch własny pary)')
        self.add('rozproszenie_12', np.repeat(np.nanstd(R[12], axis=0)[None, :], npair, axis=0), 'rozrzut ruchów 12 św. między parami (aktywność rynku)')
        if btc is not None and len(btc):
            b = btc.set_index('time').c.reindex(self.t).ffill().values
            bt = pd.Series(b)
            batr = (bt.diff().abs().rolling(100).mean() / bt * 100).values
            for k in [4, 24]:
                v = ((b - _lag(b, k)) / _lag(b, k) * 100) / batr
                self.add(f'btc_ruch_{k}', np.repeat(v[None, :], npair, axis=0), f'BTC: zmiana w {k} św. (w ATR BTC) — nastrój rynku')
        up = R[1] >= 2; dn = R[1] <= -2
        self.add('swiec_od_impulsu_gora', np.array([_since(up[q], 240) for q in range(npair)]), 'świec od ostatniej świecy wzrostowej ≥ 2 ATR')
        self.add('swiec_od_impulsu_dol', np.array([_since(dn[q], 240) for q in range(npair)]), 'świec od ostatniej świecy spadkowej ≥ 2 ATR')
        # kategorie
        H = np.repeat(hour[None, :], npair, axis=0)
        wd = np.repeat(((day + 3) % 7)[None, :], npair, axis=0)          # 0 = poniedziałek
        dom = np.repeat(pd.to_datetime(self.t, unit='s').day.values[None, :], npair, axis=0)
        self.add('godzina', H, 'godzina świecy (UTC)', kat=True)
        self.add('dzien_tyg', wd, 'dzień tygodnia (0 = pon)', kat=True)
        self.add('dzien_mies', dom, 'dzień miesiąca', kat=True)
        rng = h - l
        nr4 = np.array([rng[q] <= _rmin(rng[q], 4) for q in range(npair)]); nr7 = np.array([rng[q] <= _rmin(rng[q], 7) for q in range(npair)])
        ph = np.concatenate([h[:, :1], h[:, :-1]], axis=1); plw = np.concatenate([l[:, :1], l[:, :-1]], axis=1)
        self.add('nr4', nr4, 'NR4: najwęższa świeca z 4', kat=True)
        self.add('nr7', nr7, 'NR7: najwęższa świeca z 7', kat=True)
        self.add('wewnetrzna', (h < ph) & (l > plw), 'świeca wewnętrzna (w zakresie poprzedniej)', kat=True)
        self.add('objecia', (h > ph) & (l < plw), 'świeca objęcia (szersza z obu stron)', kat=True)


# ---------------------------------------------------------------- atomy

KWANTYLE = [('≤', 0.05), ('≤', 0.10), ('≤', 0.20), ('≤', 0.33), ('≥', 0.67), ('≥', 0.80), ('≥', 0.90), ('≥', 0.95)]
GODZINY = [(0, 2), (3, 5), (6, 7), (8, 10), (11, 13), (14, 16), (17, 19), (20, 23), (7, 16), (12, 16), (0, 6)]
DNI = ['poniedziałek', 'wtorek', 'środa', 'czwartek', 'piątek']


def _fmt(v):
    a = abs(v)
    s = f'{v:.0f}' if a >= 100 else (f'{v:.1f}' if a >= 10 else (f'{v:.2f}' if a >= 0.1 else f'{v:.3f}'))
    return s.replace('.', ',')


def atomy(R: Rynek):
    """Lista atomów (cecha, opis) i maski [atom, para, świeca] na obu latach; progi tylko z roku 1."""
    lab, feat, masks = [], [], []
    fidx = {}
    tr = np.zeros(R.n, bool); tr[WARM:R.b] = True
    for name, x in R.F.items():
        fi = fidx.setdefault(name, len(fidx))
        if R.KAT[name]:
            if name == 'godzina':
                for a, b in GODZINY:
                    lab.append(f'godzina świecy {a}–{b} UTC'); feat.append(fi); masks.append((x >= a) & (x <= b))
            elif name == 'dzien_tyg':
                for d in range(5):
                    lab.append(f'dzień: {DNI[d]}'); feat.append(fi); masks.append(x == d)
            elif name == 'dzien_mies':
                lab.append('początek miesiąca (dni 1–3)'); feat.append(fi); masks.append(x <= 3)
                lab.append('koniec miesiąca (dni 27–31)'); feat.append(fi); masks.append(x >= 27)
            else:
                lab.append(R.OPIS[name]); feat.append(fi); masks.append(x > 0.5)
            continue
        v = x[:, tr & np.ones(R.n, bool)][R.valid[:, tr]]
        v = v[np.isfinite(v)]
        if v.size < 1000:
            continue
        seen = set()
        for op, q in KWANTYLE:
            th = float(np.quantile(v, q))
            if (op, round(th, 10)) in seen:
                continue
            seen.add((op, round(th, 10)))
            with np.errstate(invalid='ignore'):
                m = (x <= th) if op == '≤' else (x >= th)
            m &= np.isfinite(x)
            dol = q <= 0.5
            lab.append(f"{R.OPIS[name]} {op} {_fmt(th)} ({'dolne' if dol else 'górne'} {round((q if dol else 1 - q) * 100):.0f}%)")
            feat.append(fi); masks.append(m)
    return lab, np.array(feat, np.int32), np.array(masks, bool), list(fidx)


# ---------------------------------------------------------------- wyjścia

def wyjscia():
    E = []
    for sl in [0.5, 1.0, 1.5, 2.0, 3.0, 0.0]:
        for tp in [0.5, 1.0, 1.5, 2.0, 3.0, 0.0]:
            for H in [1, 2, 4, 8, 16, 32, 64]:
                E.append((0, sl, tp, H))
    for k in [1.0, 2.0, 3.0]:
        for H in [16, 64]:
            E.append((1, k, 0.0, H))
    for sl in [1.0, 2.0, 0.0]:
        for H in [16, 64]:
            E.append((2, sl, 0.0, H))
    for sl in [1.0, 2.0, 0.0]:
        E.append((3, sl, 0.0, 24))
    return E


EXITS = wyjscia()
NE = len(EXITS)
MA_SL = np.array([(k == 1) or (a > 0) for k, a, b, H in EXITS])
SERIE = ['seria 100', 'seria 100 z SL']
EX = np.array([[e[0], e[1], e[2], e[3]] for e in EXITS], np.float64)


def _n(x):
    return f'{x:g}'.replace('.', ',')


def exit_name(e: int) -> str:
    k, a, b, H = EXITS[e]; H = int(H)
    sl = 'bez SL' if a == 0 else f'SL {_n(a)} ATR'
    if k == 0:
        tp = 'bez TP' if b == 0 else f'TP {_n(b)} ATR'
        return f'{sl} · {tp} · {H} św.'
    if k == 1:
        return f'stop kroczący {_n(a)} ATR · {H} św.'
    if k == 2:
        return f'{sl} · cel EMA20 · {H} św.'
    return f'{sl} · zamknięcie po świecy 20 UTC (najwyżej 24 św.)'


@nb.njit(cache=True)
def wyniki(o, h, l, c, atr, ema, hour, valid, sp, d, ex):
    """Wynik % netto (kierunek d: +1 long, −1 short) i liczba świec dla każdej świecy sygnału i każdego wyjścia.
    Wejście po open świecy i+1. Pesymistycznie: w świecy najpierw ruch przeciw pozycji (SL), luka przez SL → open."""
    n = o.size; ne = ex.shape[0]
    P = np.full((n, ne), np.nan, np.float32); X = np.zeros((n, ne), np.int16)
    for i in range(n - 1):
        if not valid[i]:
            continue
        e0 = i + 1; en = o[e0]; a = atr[i]
        for e in range(ne):
            kind = int(ex[e, 0]); p1 = ex[e, 1]; p2 = ex[e, 2]; H = int(ex[e, 3])
            if e0 + H - 1 >= n:
                continue
            # poziomy w „cenie kierunkowej”: dla shortu odwracamy znak, żeby jedna logika obsłużyła oba kierunki
            if kind == 1:
                stop = en - d * p1 * a
            elif p1 > 0:
                stop = en - d * p1 * a
            else:
                stop = np.nan
            tgt = np.nan
            if kind == 0 and p2 > 0:
                tgt = en + d * p2 * a
            elif kind == 2:
                tgt = ema[i]
                if d * (tgt - en) <= 0:
                    continue                      # cel EMA20 już minięty przy wejściu
            best = en; ex_ = np.nan; nb_ = 0
            for j in range(e0, e0 + H):
                nb_ = j - e0 + 1
                adv = l[j] if d > 0 else h[j]     # najgorsza cena świecy dla pozycji
                fav = h[j] if d > 0 else l[j]
                if j > e0:
                    if stop == stop and d * (o[j] - stop) <= 0:
                        ex_ = o[j]; break
                    if tgt == tgt and d * (o[j] - tgt) >= 0:
                        ex_ = tgt; break
                if stop == stop and d * (adv - stop) <= 0:
                    ex_ = stop; break
                if tgt == tgt and d * (fav - tgt) >= 0:
                    ex_ = tgt; break
                if kind == 1:
                    if d * (fav - best) > 0:
                        best = fav
                        ns = best - d * p1 * a
                        if d * (ns - stop) > 0:
                            stop = ns
                if kind == 2:
                    tgt = ema[j]
                if kind == 3 and hour[j] == 20:
                    ex_ = c[j]; break
                if j == e0 + H - 1:
                    ex_ = c[j]
            P[i, e] = d * (ex_ - en) / en * 100.0 - sp / en * 100.0
            X[i, e] = nb_
    return P, X


# ---------------------------------------------------------------- ocena reguły

POLA = ['n', 'suma', 'srednia', 't', 'trafne', 'PF', 'obsuniecie', 'dodatnie_okresy', 'R2', 'swiece']


@nb.njit(cache=True)
def ocen(pp, ii, P, X, per, nper, min_n):
    """Transakcje reguły (jedna pozycja na parę) dla każdego wyjścia → POLA. n < min_n → tylko n."""
    ne = P.shape[2]; npair = P.shape[0]
    out = np.zeros((ne, 10), np.float32)
    last = np.empty(npair, np.int64); ps = np.zeros(nper)
    m = pp.size
    for e in range(ne):
        for q in range(npair):
            last[q] = -1
        n = 0; s = 0.0; s2 = 0.0; w = 0; gp = 0.0; gl = 0.0; cum = 0.0; pk = 0.0; dd = 0.0; hold = 0.0
        sx = 0.0; sy = 0.0; sxy = 0.0; sxx = 0.0; syy = 0.0
        ps[:] = 0.0
        for k in range(m):
            q = pp[k]; i = ii[k]
            if i <= last[q]:
                continue
            v = P[q, i, e]
            if v != v:
                continue
            x = X[q, i, e]
            last[q] = i + x
            n += 1; s += v; s2 += v * v; hold += x
            if v > 0:
                w += 1; gp += v
            else:
                gl -= v
            cum += v
            if cum > pk: pk = cum
            if pk - cum > dd: dd = pk - cum
            ps[per[i]] += v
            sx += n; sy += cum; sxy += n * cum; sxx += n * n; syy += cum * cum
        out[e, 0] = n
        if n < min_n or s <= 0:
            out[e, 1] = s
            continue
        mu = s / n; var = s2 / n - mu * mu
        out[e, 1] = s; out[e, 2] = mu
        out[e, 3] = mu / math.sqrt(var) * math.sqrt(n) if var > 0 else 0.0
        out[e, 4] = w / n; out[e, 5] = gp / gl if gl > 0 else 99.0; out[e, 6] = dd
        pos = 0
        for qq in range(nper):
            if ps[qq] > 0: pos += 1
        out[e, 7] = pos
        cov = sxy - sx * sy / n; vx = sxx - sx * sx / n; vy = syy - sy * sy / n
        out[e, 8] = cov * cov / (vx * vy) if vx > 0 and vy > 0 else 0.0
        out[e, 9] = hold / n
    return out


# ---------------------------------------------------------------- przeszukanie (rok 1, wiele procesów)

FILTR = dict(min_n=MIN_N, pf=1.15, t=2.0, okresy=6)
_G = {}


def _init(cache_dir, min_n):
    d = Path(cache_dir)
    _G['A'] = np.load(d / 'atomy.npy', mmap_mode='r')
    _G['feat'] = np.load(d / 'cechy_atomow.npy')
    _G['per'] = np.load(d / 'okresy.npy')
    _G['P'] = {1: np.load(d / 'P_long.npy', mmap_mode='r'), -1: np.load(d / 'P_short.npy', mmap_mode='r')}
    _G['X'] = {1: np.load(d / 'X_long.npy', mmap_mode='r'), -1: np.load(d / 'X_short.npy', mmap_mode='r')}
    _G['min_n'] = min_n
    _G['cdir'] = cache_dir


def _eval_entry(atoms, mask, dirs, rows, beam, stats):
    ii, pp = np.nonzero(mask.T)
    if ii.size < _G['min_n']:
        return
    ii = ii.astype(np.int64); pp = pp.astype(np.int64)
    for d in dirs:
        out = ocen(pp, ii, _G['P'][d], _G['X'][d], _G['per'], NPER, _G['min_n'])
        stats[0] += 1; stats[1] += NE
        n = out[:, 0]; s = out[:, 1]; t = out[:, 3]
        ok = (n >= FILTR['min_n']) & (s > 0) & (out[:, 5] >= FILTR['pf']) & (t >= FILTR['t']) & (out[:, 7] >= FILTR['okresy'])
        stats[2] += int(ok.sum())
        if ok.any():
            idx = np.nonzero(ok)[0]
            idx = idx[np.argsort(-t[idx])][:3]
            for e in idx:
                rows.append((*atoms, *([-1] * (4 - len(atoms))), d, int(e), *out[e].tolist()))
        valid_t = np.where(n >= 3 * _G['min_n'], t, -9)
        e = int(np.argmax(valid_t))
        if valid_t[e] >= 1.0:
            beam.append((tuple(atoms), d, float(valid_t[e]), float(n[e])))


def _task(task):
    """Zadanie: ('2', a, kierunki) — wszystkie pary (a, b>a); ('k', baza, kierunek, rangi) — rozszerzenie o jeden atom."""
    A = _G['A']; feat = _G['feat']; rows = []; beam = []; stats = [0, 0, 0]
    if task[0] == '1':
        for a in task[1]:
            _eval_entry([a], A[a], task[2], rows, beam, stats)
    elif task[0] == '2':
        a, dirs = task[1], task[2]
        Aa = np.asarray(A[a])
        for b in range(a + 1, A.shape[0]):
            if feat[b] == feat[a]:
                continue
            _eval_entry([a, b], Aa & A[b], dirs, rows, beam, stats)
    else:
        depth, k0, k1 = task[1], task[2], task[3]
        if _G.get('depth') != depth:
            import pickle
            _G['wiazka'] = pickle.load(open(Path(_G['cdir']) / f'wiazka_{depth}.pkl', 'rb'))
            _G['rank'] = {(x[0], x[1]): k for k, x in enumerate(_G['wiazka'])}
            _G['depth'] = depth
        rank = _G['rank']
        for my in range(k0, k1):
            base, d = _G['wiazka'][my][0], _G['wiazka'][my][1]
            m0 = np.asarray(A[base[0]])
            for x in base[1:]:
                m0 = m0 & A[x]
            fset = {int(feat[x]) for x in base}
            for c in range(A.shape[0]):
                if int(feat[c]) in fset:
                    continue
                full = tuple(sorted(base + (c,)))
                dup = False                       # ten sam zestaw wygeneruje lepiej oceniona baza z wiązki — pomiń
                for drop in full:
                    sub = tuple(x for x in full if x != drop)
                    if sub != base and rank.get((sub, d), 10 ** 9) < my:
                        dup = True; break
                if dup:
                    continue
                _eval_entry(list(full), m0 & A[c], [d], rows, beam, stats)
    if task[0] == 'k' and task[1] >= 4:
        beam = []                                     # ostatni etap — wiązka niepotrzebna
    return np.asarray(rows, np.float32).reshape(-1, len(ROW_COLS)), beam, stats    # float32: pula zajmuje mało pamięci


ROW_COLS = ['a1', 'a2', 'a3', 'a4', 'kier', 'e'] + POLA


def przeszukaj(R: Rynek, lab, feat, M, godziny=3.0, procesy=None, kierunki=(1, -1), log=print):
    """Etapy 1–4 na roku 1. Zwraca pulę reguł (DataFrame) i statystyki etapów."""
    import multiprocessing as mp
    t_start = time.time()
    b = R.b
    CACHE.mkdir(parents=True, exist_ok=True)
    cdir = CACHE / 'rok1'; cdir.mkdir(exist_ok=True)
    hour = ((R.t % 86400) // 3600).astype(np.int64)
    np.save(cdir / 'atomy.npy', M[:, :, :b] & R.valid[None, :, :b])
    np.save(cdir / 'cechy_atomow.npy', feat)
    np.save(cdir / 'okresy.npy', R.per[:b].astype(np.int64))
    for d, nm in [(1, 'long'), (-1, 'short')]:
        Pd = np.empty((len(PAIRS), b, NE), np.float32); Xd = np.empty((len(PAIRS), b, NE), np.int16)
        for q, p in enumerate(PAIRS):
            Pd[q], Xd[q] = wyniki(R.o[q, :b], R.h[q, :b], R.l[q, :b], R.c[q, :b], R.atr[q, :b], R.ema20[q, :b], hour[:b],
                                  R.valid[q, :b], R.spread[p] * PIP[p], float(d), EX)
        np.save(cdir / f'P_{nm}.npy', Pd); np.save(cdir / f'X_{nm}.npy', Xd)
        del Pd, Xd
    procesy = procesy or max(1, (os.cpu_count() or 2) - 1)
    ctx = mp.get_context('spawn')
    parts, ST = [], []
    nA = M.shape[0]
    with ctx.Pool(procesy, initializer=_init, initargs=(str(cdir), FILTR['min_n'])) as pool:
        def run_stage(name, tasks):
            t0 = time.time(); st = [0, 0, 0]; bm = []
            for k, (r_, b_, s_) in enumerate(pool.imap_unordered(_task, tasks, chunksize=1)):
                if len(r_):
                    parts.append(r_)
                bm.extend(b_)
                if len(bm) > 2_000_000:                   # pamięć: zostaje milion najlepszych kandydatów do wiązki
                    bm = sorted(bm, key=lambda x: -x[2])[:1_000_000]
                st = [st[0] + s_[0], st[1] + s_[1], st[2] + s_[2]]
                if log and (k % max(1, len(tasks) // 10) == 0):
                    log(f'    {name}: {k + 1}/{len(tasks)} zadań, {st[0]:,} reguł wejścia, {time.time() - t0:.0f} s')
            dt = time.time() - t0
            ST.append(dict(etap=name, reguly_wejscia=st[0], reguly=st[1], w_puli=st[2], sekundy=dt))
            log(f'  {name}: {st[0]:,} reguł wejścia × {NE} wyjść = {st[1]:,} reguł, w puli {st[2]:,}, {dt / 60:.1f} min')
            return bm, (st[0] / dt if dt > 0 else 1.0)
        log(f'  procesy: {procesy}, atomy: {nA}, wyjścia: {NE}, kierunki: {len(kierunki)}, budżet: {godziny:g} h')
        bm1, _ = run_stage('etap 1 — jeden warunek', [('1', list(range(k, min(k + 20, nA))), list(kierunki)) for k in range(0, nA, 20)])
        bm2, rate2 = run_stage('etap 2 — dwa warunki', [('2', a, list(kierunki)) for a in range(nA)])
        prev = bm2
        for depth in (3, 4):
            left = godziny * 3600 - (time.time() - t_start) - 900          # 15 min na skarbiec i Excel
            share = 0.6 if depth == 3 else 1.0
            B = int(max(0, left * share) * rate2 / nA)
            if B < 20 or not prev:
                log(f'  etap {depth}: pominięty (brak czasu w budżecie)'); break
            prev = sorted({(x[0], x[1]): x for x in prev}.values(), key=lambda x: -x[2])[:B]
            import pickle
            pickle.dump(prev, open(cdir / f'wiazka_{depth}.pkl', 'wb'))
            tasks = [('k', depth, k, min(k + 10, len(prev))) for k in range(0, len(prev), 10)]
            log(f'  etap {depth}: wiązka {len(prev):,} najlepszych reguł × {nA} atomów')
            bmd, r_ = run_stage(f'etap {depth} — {depth} warunki', tasks)
            rate2 = r_ if r_ > 0 else rate2
            prev = bmd
    pula = pd.DataFrame(np.concatenate(parts) if parts else np.zeros((0, len(ROW_COLS)), np.float32), columns=ROW_COLS)
    for c_ in ['a1', 'a2', 'a3', 'a4', 'kier', 'e']:
        pula[c_] = pula[c_].astype(np.int64)
    if len(pula):                                       # kolejność niezależna od procesów — wynik powtarzalny
        pula = pula.drop_duplicates(subset=['a1', 'a2', 'a3', 'a4', 'kier', 'e']).sort_values(['a1', 'a2', 'a3', 'a4', 'kier', 'e']).reset_index(drop=True)
    return pula, pd.DataFrame(ST)


# ---------------------------------------------------------------- wybór 100 i skarbiec

def ocena(pula: pd.DataFrame) -> pd.DataFrame:
    """Średnia z 4 percentyli w puli: wzrost, wystąpienia, obsunięcie (mniejsze lepiej), konsekwencja."""
    p = pula.copy()
    p['konsekwencja'] = 0.5 * p.dodatnie_okresy / NPER + 0.5 * p.R2
    p['m_wzrost'] = p.suma.rank(pct=True) * 100
    p['m_wystapienia'] = p.n.rank(pct=True) * 100
    p['m_obsuniecie'] = (-p.obsuniecie).rank(pct=True) * 100
    p['m_konsekwencja'] = p.konsekwencja.rank(pct=True) * 100
    p['ocena'] = p[['m_wzrost', 'm_wystapienia', 'm_obsuniecie', 'm_konsekwencja']].mean(axis=1)
    return p.sort_values('ocena', ascending=False, kind='stable').reset_index(drop=True)


def atoms_of(r):
    return tuple(int(x) for x in (r.a1, r.a2, r.a3, r.a4) if x >= 0)


class Ksiega:
    """Wyniki transakcji: rok 1 (macierze ucięte na granicy — bez podglądania) i całość (tylko dla skarbca)."""

    def __init__(self, R: Rynek, M, rok2=False):
        self.R = R; self.M = M
        hour = ((R.t % 86400) // 3600).astype(np.int64)
        n = R.b if not rok2 else R.n
        self.n = n
        self.P, self.X = {}, {}
        for d in (1, -1):
            Pd = np.empty((len(PAIRS), n, NE), np.float32); Xd = np.empty((len(PAIRS), n, NE), np.int16)
            for q, p in enumerate(PAIRS):
                Pd[q], Xd[q] = wyniki(R.o[q, :n], R.h[q, :n], R.l[q, :n], R.c[q, :n], R.atr[q, :n], R.ema20[q, :n], hour[:n],
                                      R.valid[q, :n], R.spread[p] * PIP[p], float(d), EX)
            self.P[d] = Pd; self.X[d] = Xd

    def trades(self, atoms, d, e, od=0, do=None):
        do = self.n if do is None else do
        m = self.R.valid[:, :self.n].copy()
        for a in atoms:
            m &= self.M[a][:, :self.n]
        m[:, :od] = False; m[:, do:] = False
        ii, pp = np.nonzero(m.T)
        last = np.full(len(PAIRS), -1); out = []
        for i, q in zip(ii, pp):
            if i <= last[q]:
                continue
            v = self.P[d][q, i, e]
            if v != v:
                continue
            x = int(self.X[d][q, i, e]); last[q] = i + x
            out.append((int(q), int(i), x, float(v)))
        return out


def staty(tr, per, okresy):
    """POLA dla listy transakcji (q, i, świece, wynik) i zakresu okresów."""
    if not tr:
        return dict(n=0, suma=0.0, srednia=0.0, t=0.0, trafne=0.0, PF=0.0, obsuniecie=0.0, dodatnie_okresy=0, R2=0.0, swiece=0.0)
    w = np.array([x[3] for x in tr]); n = w.size; cum = np.cumsum(w)
    ps = np.zeros(len(okresy))
    for x in tr:
        ps[okresy.index(per[x[1]])] += x[3]
    sd = w.std()
    k = np.arange(1, n + 1)
    r2 = float(np.corrcoef(k, cum)[0, 1] ** 2) if n > 2 and cum.std() > 0 else 0.0
    return dict(n=n, suma=float(w.sum()), srednia=float(w.mean()), t=float(w.mean() / sd * np.sqrt(n)) if sd > 0 else 0.0,
                trafne=float((w > 0).mean()), PF=float(w[w > 0].sum() / max(-w[w < 0].sum(), 1e-9)),
                obsuniecie=float(np.max(np.maximum.accumulate(np.r_[0, cum])[1:] - cum)), dodatnie_okresy=int((ps > 0).sum()), R2=r2,
                swiece=float(np.mean([x[2] for x in tr])), okresy=ps)


def wybierz(pula: pd.DataFrame, K1: Ksiega, ile=100, max_wspolne=0.5, log=print):
    """Seria `ile` strategii: od najlepszej oceny, bez powtórzeń wejścia, ≤ 50% wspólnych transakcji (rok 1)."""
    wyb, zbiory, wejscia = [], [], set()
    for k, r in pula.iterrows():
        at = atoms_of(r); key = (at, int(r.kier))
        if key in wejscia:
            continue
        tr = K1.trades(at, int(r.kier), int(r.e))
        S = {(q, i) for q, i, _, _ in tr}
        if any(len(S & Z) / max(1, len(S | Z)) > max_wspolne for Z in zbiory):
            continue
        wyb.append(k); zbiory.append(S); wejscia.add(key)
        if len(wyb) >= ile:
            break
    return pula.loc[wyb].reset_index(drop=True)


def zamroz(listy: dict, lab, R: Rynek):
    """Zapis list przed otwarciem skarbca + dziennik otwarć (zasada 4: każde otwarcie liczone i zapisane)."""
    CACHE.mkdir(parents=True, exist_ok=True)
    rec = {name: [dict(atomy=[lab[a] for a in atoms_of(r)], kier='long' if r.kier > 0 else 'short', wyjscie=exit_name(int(r.e))) for _, r in df.iterrows()]
           for name, df in listy.items()}
    blob = json.dumps(rec, ensure_ascii=False, sort_keys=True)
    h = hashlib.sha256(blob.encode()).hexdigest()[:16]
    (CACHE / f'lista_{h}.json').write_text(blob, encoding='utf-8')
    logf = CACHE / 'skarbiec.jsonl'
    prev = logf.read_text(encoding='utf-8').strip().splitlines() if logf.exists() else []
    ent = dict(czas=pd.Timestamp.now(tz='UTC').isoformat(), lista=h, otwarcie=len(prev) + 1,
               granica=str(pd.Timestamp(int(R.t[R.b]), unit='s')), dane_do=str(pd.Timestamp(int(R.t[-1]), unit='s')), wersja=__version__)
    with open(logf, 'a', encoding='utf-8') as f:
        f.write(json.dumps(ent, ensure_ascii=False) + '\n')
    return ent, [json.loads(x) for x in prev] + [ent]


# ---------------------------------------------------------------- portfel wielu strategii

def portfel(T: pd.DataFrame, R: Rynek, f=0.05, lev=30.0, start=1000.0, swap_pa=0.0, stop_out=0.5, lot=1000, od=0):
    """Rachunek CFD z wieloma strategiami naraz (long i short): każda strategia — najwyżej jedna pozycja na parę,
    depozyt pozycji = f × saldo, wartość = depozyt × dźwignia, lot 0,01, brak wolnego depozytu → mniej / pominięta.
    Co godzinę wycena po najgorszej cenie świecy, poziom depozytu < 50% → stop out wszystkich pozycji."""
    BASE = fs.BASE_USD
    ent = {}
    for k, r in enumerate(T.itertuples()):
        ent.setdefault(int(r.i) + 1, []).append(k)
    rows = list(T.itertuples())
    cash = start; open_ = {}; eq = np.full(R.n, np.nan); events = []; done = []; skipped = 0
    for j in range(od, R.n):
        for k in ent.get(j, []):
            r = rows[k]; key = (r.sid, r.q)
            if key in open_:
                continue
            used = sum(x['margin'] for x in open_.values())
            margin = min(f * cash, cash - used)
            p = PAIRS[r.q]
            bu = 1.0 if BASE[p] is None else float(R.o[PAIRS.index(BASE[p]), j])
            units = int(margin * lev / bu // lot) * lot if margin > 0 else 0
            if units < lot:
                skipped += 1; continue
            en = float(R.o[r.q, j]); notional = units * bu
            open_[key] = dict(q=r.q, d=r.d, j0=j, j1=j + int(r.swiece) - 1, notional=notional, margin=notional / lev, en=en, pnl=float(r.wynik),
                              cap=float(r.cap), unreal=0.0, sid=r.sid)
        if open_:
            worst = 0.0; used = 0.0
            for x in open_.values():
                sp = R.spread[PAIRS[x['q']]] * PIP[PAIRS[x['q']]]
                adv_px = R.l[x['q'], j] if x['d'] > 0 else R.h[x['q'], j]
                adv = min((x['d'] * (x['en'] - adv_px) + sp) / x['en'], x['cap'])
                worst -= x['notional'] * adv; used += x['margin']
            if cash + worst < stop_out * used:
                tot = 0.0
                for x in open_.values():
                    sp = R.spread[PAIRS[x['q']]] * PIP[PAIRS[x['q']]]
                    adv_px = R.l[x['q'], j] if x['d'] > 0 else R.h[x['q'], j]
                    pl = -x['notional'] * min((x['d'] * (x['en'] - adv_px) + sp) / x['en'], x['cap'])
                    cash += pl; tot += pl
                events.append(dict(czas=pd.Timestamp(int(R.t[j]), unit='s'), kwota=tot)); open_.clear()
                if cash <= 0:
                    cash = 0.0; eq[j:] = 0.0; break
        for key in list(open_):
            x = open_[key]
            if x['j1'] == j:
                pl = x['notional'] * x['pnl'] / 100.0
                if swap_pa:
                    pl -= x['notional'] * swap_pa / 100.0 / 365.0 * fs.rollovers(pd.Timestamp(int(R.t[x['j0']]), unit='s'), pd.Timestamp(int(R.t[j]) + 3600, unit='s'))
                cash += pl; done.append((x['sid'], pl)); del open_[key]
        for x in open_.values():
            sp = R.spread[PAIRS[x['q']]] * PIP[PAIRS[x['q']]]
            x['unreal'] = x['notional'] * (x['d'] * (R.c[x['q'], j] - x['en']) / x['en'] - sp / x['en'])
        eq[j] = cash + sum(x['unreal'] for x in open_.values())
    curve = pd.Series(eq, index=pd.to_datetime(R.t, unit='s')).iloc[od:].ffill().fillna(start)
    return curve, dict(transakcje=len(done), pominiete=skipped, stop_out=len(events))


# ---------------------------------------------------------------- całość

def _cap(e, atr, en, sp):
    k, a, b, H = EXITS[e]
    if k == 1 or a > 0:
        return (a * atr + sp) / en
    return 9e9


def compute(R: Rynek, godziny=3.0, procesy=None, kierunki=(1, -1), f=0.05, lev=30.0, start=1000.0, log=print, szybko=False) -> dict:
    t0 = time.time()
    O = dict(f=f, lev=lev, start=start, godziny=godziny, kierunki=kierunki, od=pd.Timestamp(int(R.t[0]), unit='s'),
             granica=pd.Timestamp(int(R.t[R.b]), unit='s'), do=pd.Timestamp(int(R.t[-1]), unit='s'), swiec=R.n, periods=R.periods, spread=R.spread)
    lab, feat, M, fnames = atomy(R)
    if szybko:                                             # test: co 6. atom
        keep = np.arange(0, len(lab), 6)
        lab = [lab[k] for k in keep]; feat = feat[keep]; M = M[keep]
    O['atomy'] = lab; O['cechy'] = fnames
    log(f'Atomy: {len(lab)} z {len(fnames)} cech; wyjścia: {NE}')
    pula, ST = przeszukaj(R, lab, feat, M, godziny, procesy, kierunki, log)
    O['etapy'] = ST; O['pula_n'] = len(pula)
    if len(pula) == 0:
        raise RuntimeError('Pula pusta — żadna reguła nie spełniła warunków na roku 1.')
    pula = ocena(pula)
    log(f'Pula: {len(pula):,} reguł (rok 1). Wybór serii 100…')
    K1 = Ksiega(R, M, rok2=False)
    top = wybierz(pula, K1, 100)
    log('  seria 100 z SL…')
    top_sl = wybierz(pula[MA_SL[pula.e.values]].reset_index(drop=True), K1, 100)
    u = pula.loc[~pula[['a1', 'a2', 'a3', 'a4', 'kier']].duplicated()]
    losowe = u.sample(min(100, len(u)), random_state=7).reset_index(drop=True)
    top_t = u.sort_values('t', ascending=False).head(100).reset_index(drop=True)
    depth = (u[['a1', 'a2', 'a3', 'a4']] >= 0).sum(axis=1).values
    G = [('seria 100', top), ('seria 100 z SL', top_sl), ('100 losowych z puli', losowe), ('100 najlepszych wg t', top_t)]
    for dp in (1, 2, 3, 4):
        ud = u[depth == dp]
        if len(ud):
            G.append((f'50 losowych z puli: {dp} warunk{"i" if dp > 1 else ""}' if dp > 1 else '50 losowych z puli: 1 warunek', ud.sample(min(50, len(ud)), random_state=11 + dp).reset_index(drop=True)))
    # --- bramka skarbca: wszystkie listy zamrożone przed policzeniem roku 2
    ent, dziennik = zamroz(dict(G), lab, R)
    O['skarbiec'] = dziennik
    log(f'Skarbiec otwarty (otwarcie nr {ent["otwarcie"]}, listy {ent["lista"]}) — liczę rok 2…')
    K2 = Ksiega(R, M, rok2=True)
    per = R.per
    o1 = list(range(NPER)); o2 = list(range(NPER, 2 * NPER))
    grupy = {}
    for name, Gd in G:
        rows = []
        for k, r in Gd.iterrows():
            at = atoms_of(r); d = int(r.kier); e = int(r.e)
            tr1 = K1.trades(at, d, e); tr2 = K2.trades(at, d, e, od=R.b)
            s1 = staty(tr1, per, o1); s2 = staty(tr2, per, o2)
            fresh = [x for x in tr2 if x[1] >= R.swieze]
            rows.append(dict(nr=k + 1, atomy=at, warunki=' ORAZ '.join(lab[a] for a in at), kier='long' if d > 0 else 'short', e=e,
                             wyjscie=exit_name(e), ocena=float(r.ocena),
                             m_wzrost=float(r.m_wzrost), m_wystapienia=float(r.m_wystapienia), m_obsuniecie=float(r.m_obsuniecie),
                             m_konsekwencja=float(r.m_konsekwencja),
                             **{f'r1_{k_}': v for k_, v in s1.items() if k_ != 'okresy'}, **{f'r2_{k_}': v for k_, v in s2.items() if k_ != 'okresy'},
                             okresy=np.r_[s1.get('okresy', np.zeros(NPER)), s2.get('okresy', np.zeros(NPER))],
                             swieze_n=len(fresh), swieze_suma=float(sum(x[3] for x in fresh)), tr1=tr1, tr2=tr2))
        grupy[name] = rows
    O['grupy'] = grupy
    # transakcje skarbca obu serii (do arkusza) i portfele
    vt = []; it = []
    for seria in SERIE:
        for r in grupy[seria]:
            d = 1 if r['kier'] == 'long' else -1
            for q, i, x, v in r['tr2']:
                p = PAIRS[q]; en = float(R.o[q, i + 1]); sp = R.spread[p] * PIP[p]
                vt.append(dict(seria=seria, sid=r['nr'], para=p, q=q, i=i, d=d, kier=r['kier'], sygnal_t=pd.Timestamp(int(R.t[i]), unit='s'),
                               wyjscie_t=pd.Timestamp(int(R.t[i + x]) + 3600, unit='s'), swiece=x, wynik=v, pipsy=v / 100 * en / PIP[p],
                               okres=int(per[i]) + 1, swieze=i >= R.swieze, cap=_cap(r['e'], float(R.atr[q, i]), en, sp)))
            for q, i, x, v in r['tr1']:
                p = PAIRS[q]; en = float(R.o[q, i + 1])
                it.append(dict(seria=seria, sid=r['nr'], q=q, i=i, d=d, swiece=x, wynik=v, cap=_cap(r['e'], float(R.atr[q, i]), en, R.spread[p] * PIP[p])))
    VT = pd.DataFrame(vt).sort_values(['seria', 'i', 'sid']).reset_index(drop=True) if vt else pd.DataFrame()
    IT = pd.DataFrame(it).sort_values(['seria', 'i', 'sid']).reset_index(drop=True)
    O['vt'] = VT
    log('Portfele wielu strategii naraz…')
    pv = []; curves = {}; curves1 = {}
    for seria in SERIE:
        for Kn in [10, 25, 50, 100]:
            for fv in sorted({0.02, 0.05, 0.10, f}):
                for sw in ([0.0, 3.0] if fv == f else [0.0]):
                    for rok, TT, od in [('skarbiec (rok 2)', VT, R.b), ('rok 1 (w próbie)', IT, WARM)]:
                        sub = TT[(TT.seria == seria) & (TT.sid <= Kn)] if len(TT) else TT
                        if len(sub) == 0:
                            continue
                        cv, info = portfel(sub, R, f=fv, lev=lev, start=start, swap_pa=sw, od=od)
                        if rok.startswith('rok 1'):
                            cv = cv[cv.index < pd.Timestamp(int(R.t[R.b]), unit='s')]
                        st = fs.curve_stats(cv, start)
                        pv.append(dict(seria=seria, strategii=Kn, f=fv, swap=sw, rok=rok, **st, **info))
                        if fv == f and sw == 0.0:
                            (curves if rok.startswith('skarb') else curves1)[(seria, Kn)] = fs.daily(cv)
    O['portfele'] = pd.DataFrame(pv); O['curves'] = curves; O['curves1'] = curves1
    # cechy w serii 100
    use = {}
    for r in grupy['seria 100'] + grupy['seria 100 z SL']:
        for a in r['atomy']:
            nm = fnames[feat[a]]
            use.setdefault(nm, []).append(r['r2_suma'])
    O['cechy_tab'] = pd.DataFrame([dict(cecha=nm, opis=R.OPIS[nm], atomow=int((feat == k).sum()), w_serii=len(use.get(nm, [])),  # obie serie razem
                                        skarbiec_srednio=float(np.mean(use[nm])) if nm in use else np.nan)
                                   for k, nm in enumerate(fnames)]).sort_values(['w_serii', 'cecha'], ascending=[False, True])
    O['czas_min'] = (time.time() - t0) / 60
    return O


# ---------------------------------------------------------------- Excel

def write_xlsx(O: dict, path: Path, wnioski=None) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference, ScatterChart, Series
    from openpyxl.chart.axis import DateAxis
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    FN = 'Arial'
    fN = Font(name=FN, size=10); fB = Font(name=FN, size=10, bold=True); fT = Font(name=FN, size=14, bold=True)
    fH = Font(name=FN, size=10, bold=True, color='FFFFFF'); fS = Font(name=FN, size=9, italic=True, color='555555')
    HDR = PatternFill('solid', fgColor='1F3A5F'); HDR2 = PatternFill('solid', fgColor='7A4B00'); GRY = PatternFill('solid', fgColor='EEF2F7')
    GOOD = PatternFill('solid', fgColor='E2F0D9'); BAD = PatternFill('solid', fgColor='FBE5D6')
    thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')
    PCT = '0.0%'; PCT3 = '0.000%'; USD = '#,##0.00 [$$-409];[Red]-#,##0.00 [$$-409]'; USD0 = '#,##0 [$$-409]'; DT = 'yyyy-mm-dd hh:mm'; DD = 'yyyy-mm-dd'

    def pl(x, d=1):
        return f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')

    wb = Workbook()

    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name; ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def hdr(ws, r, c, names, widths=None, fill=HDR):
        for k, nm in enumerate(names):
            x = ws.cell(r, c + k, nm); x.font = fH; x.fill = fill if not isinstance(fill, list) else fill[k]; x.border = BOX
            x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for k, w in enumerate(widths or []):
            if w:
                ws.column_dimensions[CL(c + k)].width = w
        ws.row_dimensions[r].height = 48

    def put(ws, r, c, v, font=fN, fmt=None, fill=None, al=None):
        if isinstance(v, (np.floating,)):
            v = float(v)
        if isinstance(v, (np.integer,)):
            v = int(v)
        if isinstance(v, float) and not np.isfinite(v):
            v = None
        x = ws.cell(r, c, v); x.font = font; x.border = BOX
        if fmt: x.number_format = fmt
        if fill: x.fill = fill
        if al: x.alignment = al
        return x

    def note(ws, r, c, text, span=10, h=None):
        x = ws.cell(r, c, text); x.font = fS; x.alignment = WR
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        if h: ws.row_dimensions[r].height = h

    def lchart(title, h, w):
        ch = LineChart(); ch.title = title; ch.height = h; ch.width = w
        ch.y_axis.crossAx = 500; ch.x_axis = DateAxis(crossAx=100); ch.x_axis.number_format = 'yyyy-mm'; ch.x_axis.majorTimeUnit = 'months'
        ch.x_axis.delete = False; ch.y_axis.delete = False; ch.display_blanks = 'span'
        return ch

    py = lambda ts: pd.Timestamp(ts).to_pydatetime()
    S = O['grupy']['seria 100']
    wsP = sheet('PODSUMOWANIE', first=True)

    # TRANSAKCJE SKARBCA
    ws = sheet('TRANSAKCJE SKARBCA')
    ws['A1'] = 'TRANSAKCJE SKARBCA — wszystkie transakcje obu serii w roku 2 (dane, których poszukiwanie nie widziało)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Wynik % netto po spreadzie, bez dźwigni, rozliczenie pesymistyczne. „Świeże” = świece po 16.10.2026 — nikt (także Claude) nie oglądał ich przy żadnym badaniu.', 12, 30)
    VT = O['vt']; R0 = 5; R1 = R0 + max(len(VT), 1) - 1
    hdr(ws, 4, 1, ['Seria', 'Strategia nr', 'Para', 'Kierunek', 'Sygnał — świeca (UTC)', 'Wyjście (UTC)', 'Świec w pozycji', 'Wynik % (netto)', 'Pipsy (netto)', 'Okres (11–20)', 'Świeże'],
        [15, 9, 8, 8, 16, 16, 8, 11, 9, 8, 7])
    for k, r in enumerate(VT.itertuples()):
        row = R0 + k
        for c, (v, fm) in enumerate([(r.seria, None), (int(r.sid), '0'), (r.para, None), (r.kier, None), (py(r.sygnal_t), DT), (py(r.wyjscie_t), DT), (int(r.swiece), '0'),
                                     (r.wynik / 100, PCT3), (round(r.pipsy, 1), '0.0'), (int(r.okres), '0'), ('tak' if r.swieze else 'nie', None)]):
            put(ws, row, 1 + c, v, fmt=fm)
    ws.freeze_panes = 'A5'; ws.auto_filter.ref = f'A4:K{R1}'
    TR = lambda col: f"'TRANSAKCJE SKARBCA'!${col}${R0}:${col}${R1}"

    # STRATEGIE (obie serie)
    cols1 = ['Nr', 'Ocena (0–100)', 'Warunki wejścia (wszystkie naraz)', 'Kierunek', 'Wyjście',
             'Rok 1: transakcje', 'Rok 1: suma %', 'Rok 1: średnia', 'Rok 1: trafne', 'Rok 1: PF', 'Rok 1: t', 'Rok 1: obsunięcie', 'Rok 1: okresy na plusie / 10', 'Rok 1: R²',
             'Miejsce: wzrost', 'Miejsce: wystąpienia', 'Miejsce: obsunięcie', 'Miejsce: konsekwencja']
    cols2 = ['Skarbiec: transakcje', 'Skarbiec: suma %', 'Skarbiec: średnia', 'Skarbiec: trafne', 'Skarbiec: PF', 'Skarbiec: t', 'Skarbiec: obsunięcie', 'Skarbiec: okresy na plusie / 10',
             'Świeże: transakcje', 'Świeże: suma %']
    SUMROW = {}
    for seria, title, opis in [('seria 100', 'STRATEGIE', 'SERIA 100 — strategie wybrane wyłącznie na roku 1 (4 kryteria człowieka) i ich wynik w skarbcu (rok 2)'),
                               ('seria 100 z SL', 'STRATEGIE Z SL', 'SERIA 100 Z SL — to samo, ale tylko wyjścia ze stop lossem (stały SL albo kroczący) — ochrona przed rzadką dużą stratą')]:
        Sx = O['grupy'][seria]
        ws = sheet(title)
        ws['A1'] = opis; ws['A1'].font = fT
        note(ws, 2, 1, 'Ocena = średnia z 4 miejsc w puli roku 1 (percentyl 0–100): wzrost (suma % netto), wystąpienia (transakcje), obsunięcie (mniejsze lepiej), konsekwencja (½ okresy na plusie z 10 + ½ R² krzywej). Kolumny skarbca (brązowe) liczą formuły z arkusza TRANSAKCJE SKARBCA; t, obsunięcie i okresy — Python. Wynik % bez dźwigni, suma wszystkich transakcji.', 20, 42)
        hdr(ws, 4, 1, cols1 + cols2, [5, 8, 70, 7, 30] + [9] * 13 + [9] * 10, fill=[HDR] * len(cols1) + [HDR2] * len(cols2))
        for k, r in enumerate(Sx):
            row = 5 + k
            vals = [(r['nr'], '0'), (r['ocena'], '0.0'), (r['warunki'], None), (r['kier'], None), (r['wyjscie'], None),
                    (r['r1_n'], '0'), (r['r1_suma'] / 100, PCT3), (r['r1_srednia'] / 100, PCT3), (r['r1_trafne'], PCT), (min(r['r1_PF'], 99), '0.00'),
                    (r['r1_t'], '0.0'), (r['r1_obsuniecie'] / 100, PCT3), (r['r1_dodatnie_okresy'], '0'), (r['r1_R2'], '0.00'),
                    (r['m_wzrost'], '0'), (r['m_wystapienia'], '0'), (r['m_obsuniecie'], '0'), (r['m_konsekwencja'], '0')]
            for c, (v, fm) in enumerate(vals):
                put(ws, row, 1 + c, v, fmt=fm, al=WR if c == 2 else None)
            c0 = len(cols1) + 1; A = CL(c0); crit = f'{TR("A")},"{seria}",{TR("B")},A{row}'
            put(ws, row, c0, f'=COUNTIFS({crit})', fmt='0')
            x = put(ws, row, c0 + 1, f'=SUMIFS({TR("H")},{crit})', fmt=PCT3)
            put(ws, row, c0 + 2, f'=IF({A}{row}=0,0,{CL(c0 + 1)}{row}/{A}{row})', fmt=PCT3)
            put(ws, row, c0 + 3, f'=IF({A}{row}=0,0,COUNTIFS({crit},{TR("H")},">0")/{A}{row})', fmt=PCT)
            put(ws, row, c0 + 4, f'=IFERROR(SUMIFS({TR("H")},{crit},{TR("H")},">0")/-SUMIFS({TR("H")},{crit},{TR("H")},"<0"),99)', fmt='0.00')
            put(ws, row, c0 + 5, r['r2_t'], fmt='0.0'); put(ws, row, c0 + 6, r['r2_obsuniecie'] / 100, fmt=PCT3); put(ws, row, c0 + 7, r['r2_dodatnie_okresy'], fmt='0')
            put(ws, row, c0 + 8, f'=COUNTIFS({crit},{TR("K")},"tak")', fmt='0')
            put(ws, row, c0 + 9, f'=SUMIFS({TR("H")},{crit},{TR("K")},"tak")', fmt=PCT3)
            x.fill = GOOD if r['r2_suma'] > 0 else BAD
            ws.row_dimensions[row].height = 15 * max(1, min(4, len(r['warunki']) // 70 + 1))
        rs = 5 + len(Sx); c0 = len(cols1) + 1
        put(ws, rs, 3, 'Razem / średnio', fB, fill=GRY)
        put(ws, rs, 6, f'=SUM(F5:F{rs-1})', fB, '0'); put(ws, rs, 7, f'=AVERAGE(G5:G{rs-1})', fB, PCT3)
        put(ws, rs, c0, f'=SUM({CL(c0)}5:{CL(c0)}{rs-1})', fB, '0'); put(ws, rs, c0 + 1, f'=AVERAGE({CL(c0+1)}5:{CL(c0+1)}{rs-1})', fB, PCT3)
        put(ws, rs + 1, 3, 'Strategii z sumą > 0', fB, fill=GRY)
        put(ws, rs + 1, 7, f'=COUNTIF(G5:G{rs-1},">0")', fB, '0'); put(ws, rs + 1, c0 + 1, f'=COUNTIF({CL(c0+1)}5:{CL(c0+1)}{rs-1},">0")', fB, '0')
        ws.freeze_panes = 'D5'
        SUMROW[seria] = (title, rs, c0)

    # OKRESY — mapa
    ws = sheet('OKRESY')
    ws['A1'] = 'OKRESY — suma wyników % każdej strategii w 20 okresach (10 w roku 1 | 10 w skarbcu)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Zielony — zysk, czerwony — strata. Strategia, która „działa”, ma zielone pola po obu stronach granicy (kolumny brązowe = skarbiec).', 22, 30)
    per = O['periods']; r = 4
    for seria in SERIE:
        Sx = O['grupy'][seria]
        ws.cell(r, 1, seria.upper()).font = fB
        hdr(ws, r + 1, 1, ['Nr', 'Kierunek'] + [f'R1 · {k + 1}\n{per[k][0]:%Y-%m-%d}' for k in range(NPER)] + [f'SKARBIEC · {k + 1}\n{per[NPER + k][0]:%Y-%m-%d}' for k in range(NPER)],
            [5, 7] + [10] * (2 * NPER), fill=[HDR] * (2 + NPER) + [HDR2] * NPER)
        for k, rr in enumerate(Sx):
            row = r + 2 + k; put(ws, row, 1, rr['nr']); put(ws, row, 2, rr['kier'])
            for j in range(2 * NPER):
                put(ws, row, 3 + j, float(rr['okresy'][j]) / 100, fmt='0.00%')
        last = r + 1 + len(Sx)
        ws.conditional_formatting.add(f'C{r + 2}:{CL(2 + 2 * NPER)}{last}', ColorScaleRule(start_type='percentile', start_value=5, start_color='F8696B',
                                                                                             mid_type='num', mid_value=0, mid_color='FFFFFF', end_type='percentile', end_value=95, end_color='63BE7B'))
        put(ws, last + 1, 2, 'na plusie', fB, fill=GRY)
        for j in range(2 * NPER):
            put(ws, last + 1, 3 + j, f'=COUNTIF({CL(3 + j)}{r + 2}:{CL(3 + j)}{last},">0")', fB, '0')
        r = last + 4

    # PORTFEL
    ws = sheet('PORTFEL')
    ws['A1'] = f"PORTFEL — kilka strategii naraz, {pl(O['start'], 0)} $, dźwignia 1:{O['lev']:g}"; ws['A1'].font = fT
    note(ws, 2, 1, 'Strategie 1–K z serii grają równolegle: każda najwyżej jedna pozycja na parę, depozyt pozycji = % salda, brak wolnego depozytu → pozycja mniejsza albo pominięta, wycena co godzinę, stop out < 50%. Long i short tej samej pary z różnych strategii mogą być otwarte naraz (rachunek bez kompensaty pozycji). Koszt nocny 3% rocznie — szacunek. Rok 1 = dane, na których wybierano strategie — wynik zawyżony.', 12, 54)
    hdr(ws, 4, 1, ['Seria', 'Rok', 'Strategii', 'Depozyt na pozycję', 'Koszt nocny', 'Saldo końcowe ($)', 'Zysk', 'Zysk roczny (CAGR)', 'Największe obsunięcie', 'Transakcje', 'Pominięte', 'Stop out'],
        [15, 18, 9, 10, 9, 14, 11, 11, 12, 10, 10, 8])
    for k, v in enumerate(O['portfele'].sort_values(['seria', 'rok', 'strategii', 'f', 'swap'], ascending=[True, False, True, True, True]).itertuples()):
        row = 5 + k
        for c, (val, fm) in enumerate([(v.seria, None), (v.rok, None), (v.strategii, '0'), (v.f, PCT), (v.swap / 100, '0%'), (round(v.koniec, 2), USD), (v.zysk_pct / 100, PCT),
                                       (v.cagr_pct / 100, PCT), (-v.max_obs_pct / 100, PCT), (v.transakcje, '0'), (v.pominiete, '0'), (v.stop_out, '0')]):
            x = put(ws, row, 1 + c, val, fmt=fm)
            if c == 6:
                x.fill = GOOD if v.zysk_pct > 0 else BAD
    r0 = max(6 + len(O['portfele']), 96); col = 15; anchors = iter(['N4', 'N26', 'N48', 'N70'])     # dane krzywych pod wykresami
    for seria in SERIE:
        for title, cur in [('skarbiec (rok 2)', O['curves']), ('rok 1 (w próbie)', O['curves1'])]:
            cs = {Kn: v for (sr, Kn), v in cur.items() if sr == seria}
            if not cs:
                continue
            cv = pd.DataFrame(cs); cv = cv.reindex(columns=sorted(cv.columns))
            hdr(ws, r0, col, ['Dzień'] + [f'{k} strategii' for k in cv.columns], [12] + [11] * len(cv.columns))
            for k, (d, rr) in enumerate(cv.iterrows()):
                put(ws, r0 + 1 + k, col, py(d), fmt=DD)
                for j, cc in enumerate(cv.columns):
                    put(ws, r0 + 1 + k, col + 1 + j, None if pd.isna(rr[cc]) else round(float(rr[cc]), 2), fmt=USD0)
            ch = lchart(f"{seria}, {title}: {pl(O['start'], 0)} $, 1:{O['lev']:g}, depozyt {pl(O['f'] * 100, 0)}%", 10, 22)
            ch.add_data(Reference(ws, min_col=col + 1, max_col=col + len(cv.columns), min_row=r0, max_row=r0 + len(cv)), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=col, min_row=r0 + 1, max_row=r0 + len(cv))); ch.legend.position = 'b'
            for s_, colr in zip(ch.series, ['9DC3E6', '1F3A5F', 'ED7D31', 'C00000']):
                s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 15000
            ws.add_chart(ch, next(anchors))
            col += len(cv.columns) + 2

    # PORÓWNANIE
    ws = sheet('PORÓWNANIE')
    ws['A1'] = 'PORÓWNANIE — czy wybór na roku 1 coś daje w skarbcu?'; ws['A1'].font = fT
    note(ws, 2, 1, 'Ta sama pula reguł z roku 1. Seria 100 — wybór wg 4 kryteriów i różnorodności; 100 losowych — los z puli (punkt odniesienia: tyle daje sama pula); 100 najlepszych wg t — sama siła statystyczna. Jeśli seria nie wypada lepiej niż losowe, ranking z roku 1 nie przenosi się na rok 2.', 12, 42)
    hdr(ws, 4, 1, ['Grupa', 'Rok 1: średnia suma % na strategię', 'Rok 1: średnio transakcji', 'Skarbiec: średnia suma % na strategię', 'Skarbiec: mediana', 'Skarbiec: strategii z sumą > 0',
                   'Skarbiec: średnie t', 'Skarbiec: średnio na transakcję', 'Skarbiec: transakcji razem'], [24, 13, 11, 13, 11, 12, 10, 12, 11])
    for k, (name, rows) in enumerate(O['grupy'].items()):
        row = 5 + k
        s1 = np.array([r['r1_suma'] for r in rows]); s2 = np.array([r['r2_suma'] for r in rows]); n2 = np.array([r['r2_n'] for r in rows])
        for c, (v, fm) in enumerate([(name, None), (s1.mean() / 100, PCT3), (np.mean([r['r1_n'] for r in rows]), '0'), (s2.mean() / 100, PCT3), (np.median(s2) / 100, PCT3),
                                     (int((s2 > 0).sum()), '0'), (np.mean([r['r2_t'] for r in rows]), '0.00'), (s2.sum() / max(n2.sum(), 1) / 100, PCT3), (int(n2.sum()), '0')]):
            put(ws, row, 1 + c, v, fmt=fm, fill=GRY if k == 0 else None)
    # wykres: ocena roku 1 vs skarbiec (dane pod tabelą grup — bez nakładania)
    r0 = 5 + len(O['grupy']) + 2
    ws.cell(r0 - 1, 1, 'Seria 100 — każda strategia: rok 1 obok skarbca (dane wykresu)').font = fB
    hdr(ws, r0, 1, ['Nr', 'Rok 1: suma %', 'Skarbiec: suma %', 'Rok 1: t', 'Skarbiec: t'])
    for k, r in enumerate(S):
        put(ws, r0 + 1 + k, 1, r['nr']); put(ws, r0 + 1 + k, 2, r['r1_suma'] / 100, fmt=PCT3); put(ws, r0 + 1 + k, 3, r['r2_suma'] / 100, fmt=PCT3)
        put(ws, r0 + 1 + k, 4, r['r1_t'], fmt='0.0'); put(ws, r0 + 1 + k, 5, r['r2_t'], fmt='0.0')
    sc = ScatterChart(); sc.title = 'Seria 100: wynik w roku 1 (oś X) vs w skarbcu (oś Y)'; sc.style = 13; sc.height = 10; sc.width = 18
    sc.x_axis.title = 'rok 1: suma %'; sc.y_axis.title = 'skarbiec: suma %'; sc.x_axis.delete = False; sc.y_axis.delete = False
    sc.x_axis.tickLblPos = 'low'; sc.y_axis.tickLblPos = 'low'
    xs = Reference(ws, min_col=2, min_row=r0 + 1, max_row=r0 + len(S)); ys = Reference(ws, min_col=3, min_row=r0 + 1, max_row=r0 + len(S))
    se = Series(ys, xs, title='strategia'); se.marker.symbol = 'circle'; se.marker.size = 6; se.graphicalProperties.line.noFill = True
    se.marker.graphicalProperties.solidFill = '1F3A5F'; se.marker.graphicalProperties.line.solidFill = '1F3A5F'
    sc.series.append(se); sc.legend = None; sc.x_axis.number_format = '0%'; sc.y_axis.number_format = '0%'
    ws.add_chart(sc, f'G{r0}')

    # GRUPY — wszystkie grupy porównawcze, strategia po strategii
    ws = sheet('GRUPY')
    ws['A1'] = 'GRUPY PORÓWNAWCZE — każda strategia z grup losowych i „najlepszych wg t” (zamrożone razem z serią, przed skarbcem)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Losowe grupy wg liczby warunków pokazują, czy bardziej złożone reguły (dopasowane do roku 1) gorzej przenoszą się na rok 2. Kolumna „Skarbiec: suma %” — zielona = zysk.', 10, 30)
    r1 = 4
    for name in [g for g in O['grupy'] if g not in SERIE]:
        ws.cell(r1, 1, name.upper()).font = fB
        hdr(ws, r1 + 1, 1, ['Nr', 'Warunki', 'Kierunek', 'Wyjście', 'Rok 1: n', 'Rok 1: suma %', 'Rok 1: t', 'Skarbiec: n', 'Skarbiec: suma %', 'Skarbiec: t'], [5, 60, 9, 36, 8, 10, 8, 9, 10, 9])
        for k, r in enumerate(O['grupy'][name]):
            row = r1 + 2 + k
            for c, (v, fm) in enumerate([(r['nr'], '0'), (r['warunki'], None), (r['kier'], None), (r['wyjscie'], None), (r['r1_n'], '0'), (r['r1_suma'] / 100, PCT3),
                                         (r['r1_t'], '0.0'), (r['r2_n'], '0'), (r['r2_suma'] / 100, PCT3), (r['r2_t'], '0.0')]):
                x = put(ws, row, 1 + c, v, fmt=fm, al=WR if c == 1 else None)
            x = ws.cell(row, 9); x.fill = GOOD if r['r2_suma'] > 0 else BAD
            ws.row_dimensions[row].height = 13 * max(1, min(5, len(r['warunki']) // 75 + 1))
        r1 += len(O['grupy'][name]) + 4

    # CECHY
    ws = sheet('CECHY')
    ws['A1'] = 'CECHY — katalog zmiennych i ich udział w serii 100'; ws['A1'].font = fT
    hdr(ws, 3, 1, ['Cecha', 'Opis', 'Atomów (progów)', 'W serii 100 (razy)', 'Skarbiec: średnia suma % strategii z tą cechą'], [22, 70, 10, 10, 14])
    for k, r in enumerate(O['cechy_tab'].itertuples()):
        put(ws, 4 + k, 1, r.cecha); put(ws, 4 + k, 2, r.opis); put(ws, 4 + k, 3, r.atomow, fmt='0'); put(ws, 4 + k, 4, r.w_serii, fmt='0')
        x = put(ws, 4 + k, 5, None if pd.isna(r.skarbiec_srednio) else r.skarbiec_srednio / 100, fmt=PCT3)
        if not pd.isna(r.skarbiec_srednio):
            x.fill = GOOD if r.skarbiec_srednio > 0 else BAD

    # PRZESZUKANIE
    ws = sheet('PRZESZUKANIE')
    ws['A1'] = 'PRZESZUKANIE — ile policzono (wyłącznie rok 1)'; ws['A1'].font = fT
    hdr(ws, 3, 1, ['Etap', 'Reguł wejścia (warunki + kierunek)', 'Reguł (× wyjścia)', 'Spełnia warunki puli (wejście × wyjście)', 'Minuty'], [34, 14, 16, 16, 9])
    for k, r in enumerate(O['etapy'].itertuples()):
        put(ws, 4 + k, 1, r.etap); put(ws, 4 + k, 2, r.reguly_wejscia, fmt='#,##0'); put(ws, 4 + k, 3, r.reguly, fmt='#,##0'); put(ws, 4 + k, 4, r.w_puli, fmt='#,##0')
        put(ws, 4 + k, 5, r.sekundy / 60, fmt='0.0')
    rr = 5 + len(O['etapy'])
    put(ws, rr, 1, 'Razem (w puli: najwyżej 3 wyjścia na wejście)', fB, fill=GRY); put(ws, rr, 2, f'=SUM(B4:B{rr-1})', fB, '#,##0'); put(ws, rr, 3, f'=SUM(C4:C{rr-1})', fB, '#,##0')
    put(ws, rr, 4, O['pula_n'], fB, '#,##0'); put(ws, rr, 5, f'=SUM(E4:E{rr-1})', fB, '0.0')
    params = [('Rok 1 (poszukiwanie)', f"{O['od']:%Y-%m-%d} – {O['granica']:%Y-%m-%d}"), ('Skarbiec (rok 2)', f"{O['granica']:%Y-%m-%d} – {O['do']:%Y-%m-%d %H:%M} UTC"),
              ('Pary', ', '.join(PAIRS)), ('Spread (pipsy)', ', '.join(f'{p} {pl(v, 2)}' for p, v in O['spread'].items())),
              ('Cech', str(len(O['cechy']))), ('Atomów (cecha ≤ / ≥ próg z roku 1)', str(len(O['atomy']))), ('Wyjść', str(NE)),
              ('Kierunki', ', '.join('long' if d > 0 else 'short' for d in O['kierunki'])),
              ('Pula: warunki', f"≥ {FILTR['min_n']} transakcji, suma > 0, PF ≥ {pl(FILTR['pf'], 2)}, t ≥ {pl(FILTR['t'], 1)}, ≥ {FILTR['okresy']} z {NPER} okresów na plusie; do puli najwyżej 3 wyjścia na regułę wejścia (najlepsze t)"),
              ('Ocena', 'średnia z 4 percentyli w puli: suma %, liczba transakcji, obsunięcie (mniejsze lepiej), konsekwencja (½ okresy + ½ R²)'),
              ('Różnorodność serii', 'każde wejście raz; najwyżej 50% wspólnych transakcji (rok 1) z już wybraną strategią'),
              ('Budżet czasu', f"{pl(O['godziny'], 1)} h; policzono w {pl(O['czas_min'], 0)} min")]
    for k, (a, b_) in enumerate(params):
        put(ws, rr + 2 + k, 1, a, fB, fill=GRY); put(ws, rr + 2 + k, 2, b_, al=WR); ws.merge_cells(start_row=rr + 2 + k, start_column=2, end_row=rr + 2 + k, end_column=5)
        ws.row_dimensions[rr + 2 + k].height = 30

    # SKARBIEC
    ws = sheet('SKARBIEC')
    ws['A1'] = 'SKARBIEC — dziennik otwarć (zasada 4: każde otwarcie liczone i zapisane)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Każde uruchomienie, które przelicza rok 2, zapisuje się tutaj. Wiele otwarć = skarbiec traci wartość sprawdzianu (wybór „po obejrzeniu”). Plik: ia4-research/data/fxodkrywca/skarbiec.jsonl.', 8, 30)
    hdr(ws, 4, 1, ['Otwarcie', 'Czas (UTC)', 'Lista (skrót)', 'Granica', 'Dane do', 'Wersja', 'Uwaga'], [9, 26, 18, 18, 18, 8, 60])
    for k, e in enumerate(O['skarbiec']):
        for c, key in enumerate(['otwarcie', 'czas', 'lista', 'granica', 'dane_do', 'wersja', 'uwaga']):
            put(ws, 5 + k, 1 + c, e.get(key))

    # PODSUMOWANIE
    ws = wsP
    ws.column_dimensions['A'].width = 46; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 16; ws.column_dimensions['D'].width = 16
    ws.column_dimensions['E'].width = 16; ws.column_dimensions['F'].width = 30
    ws['A1'] = 'FX-ODKRYWCA — 100 strategii wybranych na roku 1, sprawdzonych w skarbcu (rok 2)'; ws['A1'].font = fT
    ws['A2'] = (f"IA4 {__version__} · badanie informacyjne · rok 1: {O['od']:%Y-%m-%d} – {O['granica']:%Y-%m-%d} · skarbiec: {O['granica']:%Y-%m-%d} – {O['do']:%Y-%m-%d} · "
                f"policzono {pd.Timestamp.now():%Y-%m-%d %H:%M}")
    ws['A2'].font = fS
    put(ws, 4, 1, 'JAK POWSTAŁA SERIA', fH, fill=HDR); ws.merge_cells('A4:F4')
    tot = int(O['etapy'].reguly.sum())
    desc = [f"Policzono {pl(tot, 0)} reguł na roku 1: {len(O['atomy'])} warunków z {len(O['cechy'])} cech (1–4 warunki naraz) × kierunek × {NE} wyjść; szczegóły — PRZESZUKANIE, katalog cech — CECHY.",
            f"Do puli weszło {pl(O['pula_n'], 0)} reguł (≥ {FILTR['min_n']} transakcji, zysk, PF ≥ 1,15, t ≥ 2, ≥ 6/10 okresów na plusie). Ocena = średnia z miejsc: wzrost, wystąpienia, obsunięcie, konsekwencja.",
            'Seria 100 = najlepsze oceny, każde wejście raz, ≤ 50% wspólnych transakcji; seria 100 z SL — to samo tylko z wyjściami ze stop lossem. Obie serie i grupy porównawcze (losowe z puli, najlepsze wg t, losowe wg liczby warunków) zamrożono i zapisano PRZED przeliczeniem roku 2 (SKARBIEC).']
    for k, s_ in enumerate(desc):
        x = ws.cell(5 + k, 1, s_); x.font = fN; x.alignment = WR; ws.merge_cells(start_row=5 + k, start_column=1, end_row=5 + k, end_column=6); ws.row_dimensions[5 + k].height = 30
    r = 9
    put(ws, r, 1, 'WYNIK W SKARBCU (rok 2, bez dźwigni; formuły z arkuszy STRATEGIE)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    hdr(ws, r + 1, 1, ['', 'Seria 100 (4 kryteria)', 'Seria 100 z SL'])
    lines = [('Strategii z zyskiem w roku 1', 'G', 1, '0'), ('Strategii z zyskiem w skarbcu', 'V', 1, '0'), ('Średnia suma % strategii — rok 1', 'G', 0, PCT3),
             ('Średnia suma % strategii — skarbiec', 'V', 0, PCT3), ('Transakcje w roku 1', 'F', 0, '#,##0'), ('Transakcje w skarbcu', 'U', 0, '#,##0')]
    for k, (labx, colx, off, fm) in enumerate(lines):
        put(ws, r + 2 + k, 1, labx, fB, fill=GRY)
        for j, seria in enumerate(SERIE):
            title, rs, c0 = SUMROW[seria]
            cc = {'G': 'G', 'F': 'F', 'V': CL(c0 + 1), 'U': CL(c0)}[colx]
            put(ws, r + 2 + k, 2 + j, f"='{title}'!{cc}{rs + off}", fB, fm)
    r += 3 + len(lines)
    put(ws, r, 1, 'PORÓWNANIE GRUP W SKARBCU (czy wybór na roku 1 coś daje?)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    hdr(ws, r + 1, 1, ['Grupa', 'Rok 1: średnia suma %', 'Skarbiec: średnia suma %', 'Skarbiec: z zyskiem', 'Skarbiec: średnio na transakcję', 'Strategii w grupie'])
    for k, name in enumerate(O['grupy']):
        for c in range(5):
            put(ws, r + 2 + k, 1 + c, f"='PORÓWNANIE'!{['A', 'B', 'D', 'F', 'H'][c]}{5 + k}", fmt=[None, PCT3, PCT3, '0', PCT3][c])
        put(ws, r + 2 + k, 6, len(O['grupy'][name]), fmt='0')
    r += 3 + len(O['grupy'])
    put(ws, r, 1, f"PORTFEL W SKARBCU ({pl(O['start'], 0)} $, 1:{O['lev']:g}, depozyt {pl(O['f'] * 100, 0)}% na pozycję)", fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    hdr(ws, r + 1, 1, ['Seria · strategii naraz', 'Skarbiec: saldo końcowe ($)', 'Skarbiec: zysk', 'Skarbiec: największe obsunięcie', 'Skarbiec: transakcje', 'Rok 1 (w próbie): saldo ($)'])
    pv = O['portfele']; k = 0
    for seria in SERIE:
        for Kn in [10, 25, 50, 100]:
            sel = (pv.seria == seria) & (pv.strategii == Kn) & (abs(pv.f - O['f']) < 1e-9) & (pv.swap == 0)
            a_ = pv[sel & pv.rok.str.startswith('skarb')]; b_ = pv[sel & pv.rok.str.startswith('rok 1')]
            if not len(a_):
                continue
            a_ = a_.iloc[0]
            for c, (v, fm) in enumerate([(f'{seria} · {Kn}', None), (round(a_.koniec, 2), USD), (a_.zysk_pct / 100, PCT), (-a_.max_obs_pct / 100, PCT), (a_.transakcje, '0'),
                                         (round(b_.iloc[0].koniec, 2) if len(b_) else None, USD)]):
                x = put(ws, r + 2 + k, 1 + c, v, fmt=fm)
                if c == 2:
                    x.fill = GOOD if a_.zysk_pct > 0 else BAD
            k += 1
    r += 3 + k
    put(ws, r, 1, 'WNIOSKI' if wnioski else 'JAK CZYTAĆ', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    txt = wnioski or ['STRATEGIE — każda z 100 strategii: rok 1 i skarbiec obok siebie.', 'OKRESY — mapa 100 × 20 okresów (granica = początek skarbca).',
                      'PORÓWNANIE — czy seria jest lepsza od losowych reguł z tej samej puli; GRUPY — strategie grup porównawczych.', 'PORTFEL — kilka strategii naraz na dźwigni.']
    for k, s_ in enumerate(txt):
        x = ws.cell(r + 1 + k, 1, s_); x.font = fN; x.alignment = WR
        ws.merge_cells(start_row=r + 1 + k, start_column=1, end_row=r + 1 + k, end_column=6); ws.row_dimensions[r + 1 + k].height = 56 if wnioski else 18
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- uruchomienie

def run(no_sync=False, out=None, godziny=3.0, procesy=None, kierunki=(1, -1), f=0.05, lev=30.0, start=1000.0, frames=None, btc=None,
        log=print, keep=None, szybko=False) -> Path:
    t0 = time.time()
    if frames is None:
        from .fx import sync as fxs
        if not no_sync:
            log('Kopia fx1h — odświeżenie (python -m ia4.fx.sync)')
            fxs.sync(quiet=True, syms=list(PAIRS) + ['BTCUSD'])
        frames = {p: fxs.load(p) for p in PAIRS}
        try:
            btc = fxs.load('BTCUSD')
        except FileNotFoundError:
            btc = None
    R = Rynek(frames, fs.spreads(), btc)
    log(f'Dane: {R.n} świec na parę; rok 1 do {pd.Timestamp(int(R.t[R.b]), unit="s"):%Y-%m-%d} ({R.b} świec), skarbiec {R.n - R.b} świec')
    O = compute(R, godziny=godziny, procesy=procesy, kierunki=kierunki, f=f, lev=lev, start=start, log=log, szybko=szybko)
    if keep is not None:
        keep.update(O)
    path = write_xlsx(O, out or (config.RESEARCH_DIR.parent / 'FX-ODKRYWCA.xlsx'))
    S = O['grupy']['seria 100']
    s2 = np.array([r['r2_suma'] for r in S])
    log(f"\nSeria {len(S)}: w skarbcu z zyskiem {int((s2 > 0).sum())}, średnia suma {s2.mean():+.2f}% na strategię")
    for name, rows in O['grupy'].items():
        x = np.array([r['r2_suma'] for r in rows]); log(f"  {name:22s} skarbiec: średnio {x.mean():+.2f}%, z zyskiem {int((x > 0).sum())}/{len(x)}")
    log(f'Zapisano: {path}  ({(time.time() - t0) / 60:.0f} min)')
    return path


def selftest() -> bool:
    ok = True

    def check(name, cond):
        nonlocal ok
        print(f"  {'✓' if cond else '✗'} {name}")
        ok &= bool(cond)

    n = 12
    o = np.full(n, 1.0); h = o + 0.002; l = o - 0.002; c = o.copy(); atr = np.full(n, 0.01); ema = np.full(n, 1.0)
    hour = np.arange(n, dtype=np.int64); valid = np.zeros(n, bool); valid[0] = True
    h[2] = 1.025; l[2] = 0.985                         # ta sama świeca sięga TP i SL → SL
    ex = np.array([[0, 1.0, 1.0, 5], [0, 1.0, 1.0, 1], [0, 0.0, 0.0, 3]], np.float64)
    P, X = wyniki(o, h, l, c, atr, ema, hour, valid, 0.0001, 1.0, ex)
    check('long: SL i TP w jednej świecy → SL (−1% − spread)', abs(P[0, 0] - (-1.0 - 0.01)) < 1e-4 and X[0, 0] == 2)
    check('long: limit 1 świecy → wyjście po close, tylko spread', abs(P[0, 1] + 0.01) < 1e-4 and X[0, 1] == 1)
    P, X = wyniki(o, h, l, c, atr, ema, hour, valid, 0.0001, -1.0, ex)
    check('short: SL i TP w jednej świecy → SL', abs(P[0, 0] - (-1.0 - 0.01)) < 1e-4)
    o2 = np.full(n, 1.0); o2[3:] = 0.97; h2 = o2 + 0.001; l2 = o2 - 0.001; c2 = o2.copy()
    P, X = wyniki(o2, h2, l2, c2, atr, ema, hour, valid, 0.0001, 1.0, np.array([[0, 1.0, 0.0, 6]], np.float64))
    check('long: luka przez SL → wyjście po open (−3%)', abs(P[0, 0] - (-3.0 - 0.01)) < 1e-4)
    o3 = np.array([1.0, 1.0, 1.004, 1.019, 1.029, 1.03, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0])
    h3 = np.array([1.0, 1.005, 1.02, 1.03, 1.031, 1.031, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0])
    l3 = np.array([1.0, 0.999, 1.004, 1.019, 1.029, 1.015, 1.0, 1.0, 1.0, 1.0, 1.0, 1.0]); c3 = o3.copy()
    P, X = wyniki(o3, h3, l3, c3, atr, ema, hour, valid, 0.0, 1.0, np.array([[1, 1.0, 0.0, 10]], np.float64))
    check('stop kroczący 1 ATR: szczyt 1,031 → stop 1,021 (+2,1%)', abs(P[0, 0] - 2.1) < 1e-6)
    pp = np.array([0, 0, 0], np.int64); ii = np.array([0, 2, 4], np.int64)
    Pm = np.zeros((1, 5, 1), np.float32); Pm[0, [0, 2, 4], 0] = [1.0, 2.0, -0.5]; Xm = np.ones((1, 5, 1), np.int16)
    out = ocen(pp, ii, Pm, Xm, np.zeros(5, np.int64), NPER, 2)
    check('ocena: 3 transakcje, suma 2,5, PF 6', out[0, 0] == 3 and abs(out[0, 1] - 2.5) < 1e-6 and abs(out[0, 5] - 6.0) < 1e-5)
    # cała ścieżka na losowym spacerze (bez przewagi)
    fr = fs._synthetic(n=3200)
    R = Rynek(fr, fs.spreads(), None, granica=str(pd.Timestamp(int(fr['EURUSD'].time.iloc[1700]), unit='s')))
    lab, feat, M, fn = atomy(R)
    check(f'cechy {len(R.F)} i atomy {len(lab)}', len(R.F) > 60 and len(lab) > 400)
    check('progi atomów tylko z roku 1 (maski istnieją w roku 2)', M[:, :, R.b:].any())
    print(f"\nWYNIK: {'OK' if ok else 'BŁĄD'}")
    return ok


def main():
    ap = argparse.ArgumentParser(description='IA 4 — FX-ODKRYWCA: 100 strategii z roku 1, sprawdzian w skarbcu (sekcja 7b)')
    ap.add_argument('--godziny', type=float, default=3.0, help='budżet czasu przeszukania (h), domyślnie 3')
    ap.add_argument('--procesy', type=int, default=None, help='liczba procesów (domyślnie rdzenie − 1)')
    ap.add_argument('--kierunek', choices=['oba', 'long', 'short'], default='oba')
    ap.add_argument('--depozyt', type=float, default=0.05, help='depozyt na pozycję w portfelu (część salda), domyślnie 0.05')
    ap.add_argument('--dzwignia', type=float, default=30.0)
    ap.add_argument('--start', type=float, default=1000.0)
    ap.add_argument('--bez-sync', action='store_true')
    ap.add_argument('--szybko', action='store_true', help='test: co 6. warunek, krótko')
    ap.add_argument('--out', type=Path)
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args()
    if a.selftest:
        sys.exit(0 if selftest() else 1)
    kier = {'oba': (1, -1), 'long': (1,), 'short': (-1,)}[a.kierunek]
    print(f'IA 4 {__version__} — FX-ODKRYWCA (rok 1 → seria 100 → skarbiec)\n')
    run(no_sync=a.bez_sync, out=a.out, godziny=a.godziny, procesy=a.procesy, kierunki=kier, f=a.depozyt, lev=a.dzwignia, start=a.start, szybko=a.szybko)


if __name__ == '__main__':
    main()
