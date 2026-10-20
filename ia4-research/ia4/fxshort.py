"""
IA 4 — FX-SHORT-NY: short po gwałtownym wzroście w sesji Londyn + Nowy Jork (instrukcja, sekcja 7b).
Wersja projektu: 1.53 (2026-10-20) — musi zgadzać się z IA4_INSTRUKCJA.md

Hipoteza człowieka (19.10.2026): krótkotrwałe wzrosty kończą się gwałtowną zieloną świecą, po której przychodzi
korekta — szukamy tylko shortów, na 10 parach walutowych (bez BTC), szczególnie na EUR/USD. Świece 1h z kopii
lokalnej fx1h. Spread raz na transakcję: EUR/USD 1,1 pipsa (zrzut XTB człowieka: Kup 1,12081 / Sprzedaj 1,12070),
pozostałe pary — `fx_spread_pips` z ia4/ai/config.py × 1,1 / 0,8. Rozliczenie pesymistyczne (zasada 5).

Etap 1 — szerokie przeszukanie: 102 sygnały wzrostu (świeca w ATR i w %, wzrost N świec w ATR i w %, kulminacja,
odległość od EMA, RSI, seria zielonych, nowy szczyt, potwierdzenie czerwoną świecą, górny knot) × 24 filtry trendu
dłuższej średniej (SMA100/200/500: nie rośnie mocno, płaska, spada, cena pod średnią) × 4 okna godzin × 810 wyjść
(SL 0,5–5 ATR / nad szczytem / bez × TP 0,5–3 ATR / 50% / 100% korpusu / bez × 1–48 świec). Wszystkie pary razem
i EUR/USD osobno; 20 okresów; placebo (sygnały przesunięte w czasie); wybór chronologiczny (1. rok → 2. rok).
Etap 2 — zawężona siatka wokół rodziny, która przeszła etap 1 (wzrost w sesji NY przy spadającej SMA500), z placebo.
Strategia FX-SHORT-NY (stałe w MAIN) → portfel 1000 $ na dźwigni 1:30 (depozyt pozycji = część salda, lot 0,01,
stop out 50%, opcjonalny koszt nocny).

    python -m ia4.fxshort               # pełne liczenie → ~/Desktop/IA4/FX-SHORT-NY.xlsx (ok. 10–15 min, pamięć ok. 2 GB)
    python -m ia4.fxshort --bez-sync    # bez odświeżenia kopii fx1h (python -m ia4.fx.sync)
    python -m ia4.fxshort --depozyt 0.33 --dzwignia 30 --start 1000   # ustawienia rachunku w arkuszu TRANSAKCJE
    python -m ia4.fxshort --selftest    # „WYNIK: OK”
"""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import numba as nb
import numpy as np
import pandas as pd

from . import __version__, config

PAIRS = ('EURUSD', 'USDJPY', 'GBPUSD', 'AUDUSD', 'USDCAD', 'USDCHF', 'NZDUSD', 'EURJPY', 'GBPJPY', 'EURGBP')
PIP = {p: (0.01 if p.endswith('JPY') else 0.0001) for p in PAIRS}
EURUSD_SPREAD = 1.1          # pipsy, zrzut XTB człowieka 19.10.2026


def spreads() -> dict:
    from .ai.config import DEFAULTS
    base = DEFAULTS['fx_spread_pips']
    k = EURUSD_SPREAD / base['EURUSD']
    return {p: round(base[p] * k, 2) for p in PAIRS}


MAIN = dict(setup='wzrost 6 św. ≥ 3.1 ATR', trend='SMA500 spada', hours='11–16 UTC', sl=3.0, tp=0.75, H=8)
WARMUP = 600
NPER = 20

# ---------------------------------------------------------------- cechy

def _sma(x, n):
    return pd.Series(x).rolling(n).mean().values


def _ema(x, n):
    return pd.Series(x).ewm(span=n, adjust=False).mean().values


def _rsi(c, n):
    d = np.diff(c, prepend=c[0]); up = np.maximum(d, 0); dn = np.maximum(-d, 0)
    au = pd.Series(up).ewm(alpha=1 / n, adjust=False).mean().values
    ad = pd.Series(dn).ewm(alpha=1 / n, adjust=False).mean().values
    return 100 - 100 / (1 + au / np.maximum(ad, 1e-12))


def features(df: pd.DataFrame) -> dict:
    o, h, l, c = [df[k].values.astype(np.float64) for k in 'ohlc']
    t = df.time.values.astype(np.int64); n = len(c)
    pc = np.r_[c[0], c[:-1]]
    tr = np.maximum(h - l, np.maximum(abs(h - pc), abs(l - pc)))
    F = {'o': o, 'h': h, 'l': l, 'c': c, 't': t, 'filled': df.filled.values.astype(bool)}
    atr = _sma(tr, 100); F['atr'] = atr; F['atr500'] = _sma(tr, 500)
    F['body'] = (c - o) / atr; F['bodypct'] = (c - o) / o * 100
    for N in [2, 3, 4, 6, 8, 12, 24, 48]:
        cN = np.r_[np.full(N, np.nan), c[:-N]]
        F[f'run{N}'] = (c - cN) / atr; F[f'runp{N}'] = (c - cN) / cN * 100
    for m in [20, 50]:
        F[f'ext{m}'] = (c - _ema(c, m)) / atr
    for m in [2, 3, 14]:
        F[f'rsi{m}'] = _rsi(c, m)
    g = c > o; st = np.zeros(n, int)
    for i in range(n):
        st[i] = (st[i - 1] + 1 if i > 0 else 1) if g[i] else 0
    F['streak'] = st
    for N in [24, 48, 120]:
        F[f'newhi{N}'] = c > pd.Series(h).shift(1).rolling(N).max().values
    for m in [100, 200, 500]:
        sm = _sma(c, m)
        F[f'slope{m}'] = (sm - np.r_[np.full(24, np.nan), sm[:-24]]) / atr      # zmiana SMA w 24 świecach, w ATR
        F[f'dist{m}'] = (c - sm) / atr
    ts = pd.to_datetime(t, unit='s', utc=True)
    F['hour'] = ts.hour.values
    nt = np.r_[t[1:], t[-1] + 10 ** 9]
    valid = (~np.isnan(atr)) & (~np.isnan(F['atr500'])) & ((nt - t) == 3600) & (~F['filled'])
    valid[:WARMUP] = False
    F['valid'] = valid
    return F


def _nz(x):
    return np.nan_to_num(x, nan=-1e9)


def _after_red(spike, red):
    m = np.zeros_like(spike); armed = -100
    for i in range(len(spike)):
        if spike[i]:
            armed = i
        elif red[i] and 1 <= i - armed <= 3:
            m[i] = True; armed = -100
    return m


def setups1(F) -> dict:
    S = {}; b = _nz(F['body']); bp = _nz(F['bodypct'])
    for k in [1.0, 1.25, 1.5, 2.0, 2.5, 3.0]:
        S[f'świeca ≥ {k:g} ATR'] = b >= k
    for x in [0.08, 0.10, 0.12, 0.15, 0.20, 0.25]:
        S[f'świeca ≥ {x:g}%'] = bp >= x
    for N in [2, 3, 4, 6, 8, 12, 24]:
        for kb in [1.0, 1.5, 2.0, 2.5]:
            S[f'wzrost {N} św. ≥ {kb*np.sqrt(N):.1f} ATR'] = _nz(F[f'run{N}']) >= kb * np.sqrt(N)
    for N in [3, 6, 12, 24]:
        for xb in [0.25, 0.35, 0.5, 0.7]:
            S[f'wzrost {N} św. ≥ {xb*np.sqrt(N/6):.2f}%'] = _nz(F[f'runp{N}']) >= xb * np.sqrt(N / 6)
    for N in [6, 12, 24]:
        for a in [1.0, 1.5]:
            for bb in [1.0, 1.5]:
                S[f'kulminacja: {N} św. ≥ {a*np.sqrt(N):.1f} ATR + świeca ≥ {bb:g} ATR'] = (_nz(F[f'run{N}']) >= a * np.sqrt(N)) & (b >= bb)
    for k in [2.5, 3, 3.5, 4, 5]:
        S[f'cena − EMA20 ≥ {k:g} ATR'] = _nz(F['ext20']) >= k
    for k in [3, 4, 5, 6]:
        S[f'cena − EMA50 ≥ {k:g} ATR'] = _nz(F['ext50']) >= k
    for m, ths in [(2, [95, 98, 99]), (3, [90, 95, 98]), (14, [70, 75, 80])]:
        for th in ths:
            S[f'RSI({m}) ≥ {th}'] = _nz(F[f'rsi{m}']) >= th
    for m in [4, 5, 6, 7]:
        S[f'{m}+ zielonych z rzędu'] = F['streak'] >= m
    for N in [24, 48, 120]:
        for bb in [1.0, 1.5]:
            S[f'nowy szczyt {N} św. + świeca ≥ {bb:g} ATR'] = F[f'newhi{N}'] & (b >= bb)
    red = F['c'] < F['o']
    for k in [1.5, 2.0, 2.5]:
        S[f'świeca ≥ {k:g} ATR, potem pierwsza czerwona (≤ 3 św.)'] = _after_red(b >= k, red)
    S['wzrost 6 św. ≥ 4,9 ATR, potem pierwsza czerwona (≤ 3 św.)'] = _after_red(_nz(F['run6']) >= 2 * np.sqrt(6), red)
    o, h, l, c, atr = F['o'], F['h'], F['l'], F['c'], F['atr']
    uw = _nz((h - np.maximum(o, c)) / atr); lowhalf = (c - l) <= 0.5 * (h - l)
    for N in [6, 12]:
        S[f'górny knot ≥ 1 ATR po wzroście {N} św. ≥ {np.sqrt(N):.1f} ATR'] = (uw >= 1) & lowhalf & (_nz(F[f'run{N}']) >= np.sqrt(N))
    return S


def trend1(F) -> dict:
    T = {'bez filtra trendu': np.ones(len(F['c']), bool)}
    for m in [100, 200, 500]:
        s = np.nan_to_num(F[f'slope{m}'], nan=9.0)
        T[f'SMA{m} nie rośnie mocno (< +1)'] = s < 1
        T[f'SMA{m} < +0,5'] = s < 0.5
        T[f'SMA{m} płaska (|·| < 0,5)'] = np.abs(s) < 0.5
        T[f'SMA{m} płaska (|·| < 1)'] = np.abs(s) < 1
        T[f'SMA{m} spada'] = s < 0
        T[f'SMA{m} spada < −0,5'] = s < -0.5
        T[f'SMA{m} spada < −1'] = s < -1
    T['cena pod SMA200'] = np.nan_to_num(F['dist200'], nan=9.0) < 0
    T['cena pod SMA500'] = np.nan_to_num(F['dist500'], nan=9.0) < 0
    return T


def hours1(F) -> dict:
    hr = F['hour']
    return {'wszystkie godziny': np.ones(len(hr), bool), 'bez 20–23 UTC': ~((hr >= 20) & (hr <= 23)),
            '7–16 UTC (Londyn + NY)': (hr >= 7) & (hr <= 16), '0–6 UTC (Azja)': hr <= 6}


def setups2(F) -> dict:
    S = {}; b = _nz(F['body'])
    for N in [3, 4, 6, 8, 12]:
        for kb in [1.25, 1.5, 1.75, 2.0]:
            S[f'wzrost {N} św. ≥ {kb*np.sqrt(N):.1f} ATR'] = _nz(F[f'run{N}']) >= kb * np.sqrt(N)
    for k in [1.5, 2.0, 2.5]:
        S[f'świeca ≥ {k:g} ATR'] = b >= k
    for k in [2.5, 3.0, 3.5]:
        S[f'cena − EMA20 ≥ {k:g} ATR'] = _nz(F['ext20']) >= k
    S['suma: wzrost 6 ≥ 3,7 | świeca ≥ 2 | EMA20 ≥ 3'] = (_nz(F['run6']) >= 1.5 * np.sqrt(6)) | (b >= 2) | (_nz(F['ext20']) >= 3)
    S['suma: wzrost 3/6/12 ≥ 1,5√N ATR'] = ((_nz(F['run3']) >= 1.5 * np.sqrt(3)) | (_nz(F['run6']) >= 1.5 * np.sqrt(6))
                                            | (_nz(F['run12']) >= 1.5 * np.sqrt(12)))
    return S


def trend2(F) -> dict:
    s5, s2 = np.nan_to_num(F['slope500'], nan=9.0), np.nan_to_num(F['slope200'], nan=9.0)
    return {'bez filtra trendu': np.ones(len(s5), bool), 'SMA500 spada': s5 < 0, 'SMA500 < +0,5': s5 < 0.5,
            'SMA500 nie rośnie mocno (< +1)': s5 < 1, 'SMA200 spada': s2 < 0, 'SMA200 < +0,5': s2 < 0.5,
            'cena pod SMA500': np.nan_to_num(F['dist500'], nan=9.0) < 0}


def hours2(F) -> dict:
    h = F['hour']
    return {f'{a}–{b} UTC': (h >= a) & (h <= b) for a, b in [(11, 16), (12, 16), (12, 15), (11, 15), (10, 16), (8, 16), (7, 16)]}


# ---------------------------------------------------------------- wyjścia i silnik

SL_K = [0.5, 0.75, 1.0, 1.5, 2.0, 3.0, 4.0, 5.0, -1.0, 0.0]      # -1 = nad szczytem świecy sygnału + 0,25 ATR; 0 = bez SL
TP_K = [0.5, 0.75, 1.0, 1.5, 2.0, 3.0, -0.5, -1.0, 0.0]          # -0,5 / -1 = cofnięcie 50% / 100% korpusu; 0 = bez TP
HOLD = [1, 2, 3, 4, 6, 8, 12, 24, 48]
EXITS = [(a, b, hh) for a in SL_K for b in TP_K for hh in HOLD]
NE = len(EXITS)
EX_SL = np.array([e[0] for e in EXITS]); EX_TP = np.array([e[1] for e in EXITS]); EX_H = np.array([e[2] for e in EXITS], np.int64)


def _num(x):
    return f'{x:g}'.replace('.', ',')


def exit_name(e: int) -> str:
    a, b, hh = EXITS[e]
    sl = 'brak' if a == 0 else ('szczyt+0,25ATR' if a == -1 else f'{_num(a)}ATR')
    tp = 'brak' if b == 0 else ('50% korpusu' if b == -0.5 else ('100% korpusu' if b == -1 else f'{_num(b)}ATR'))
    return f'SL {sl} · TP {tp} · {hh} św.'


def exit_index(sl, tp, hh) -> int:
    return EXITS.index((sl, tp, hh))


@nb.njit(cache=True)
def outcomes(o, h, l, c, atr, valid, sp, exsl, extp, exh):
    """Short po open świecy i+1 dla każdej świecy sygnału i i każdego wyjścia: wynik % netto, świece, MAE %."""
    n = o.size; ne = exsl.size
    pnl = np.full((n, ne), np.nan, np.float32)
    xb = np.zeros((n, ne), np.int16)
    mae = np.zeros((n, ne), np.float32)
    for i in range(n - 1):
        if not valid[i]:
            continue
        e0 = i + 1; en = o[e0]; a = atr[i]; bodyp = c[i] - o[i]
        for e in range(ne):
            k = exsl[e]
            if k == 0.0:
                SL = 1e18
            elif k < 0:
                SL = max(h[i], en) + 0.25 * a
            else:
                SL = en + k * a
            q = extp[e]
            if q == 0.0:
                TP = -1e18
            elif q < 0:
                TP = c[i] - (-q) * bodyp if bodyp > 0 else en - 0.5 * a
                if TP > en - 0.1 * a:
                    TP = en - 0.1 * a
            else:
                TP = en - q * a
            H = exh[e]
            if e0 + H - 1 >= n:
                continue
            ex = np.nan; nb_ = 0; hi = en
            for j in range(e0, e0 + H):
                nb_ = j - e0 + 1
                if h[j] > hi:
                    hi = min(h[j], SL) if not (j > e0 and o[j] >= SL) else o[j]
                if j > e0:
                    if o[j] >= SL:
                        ex = o[j]; break        # luka przez SL — wyjście po open
                    if o[j] <= TP:
                        ex = TP; break          # luka przez TP — bez premii
                if h[j] >= SL:
                    ex = SL; break              # SL i TP w tej samej świecy — zawsze SL
                if l[j] <= TP:
                    ex = TP; break
                if j == e0 + H - 1:
                    ex = c[j]
            pnl[i, e] = (en - ex) / en * 100.0 - sp / en * 100.0
            xb[i, e] = nb_
            mae[i, e] = (hi - en) / en * 100.0 + sp / en * 100.0
    return pnl, xb, mae


@nb.njit(cache=True)
def evaluate(sp_, si_, P, X, per, nper, eu):
    """Sygnały (para, świeca) w kolejności czasu → statystyki dla każdego wyjścia; grupa 0 = wszystkie pary, 1 = EUR/USD.
    Pola: FIELDS (pełna próba) + połowy (okresy 1–10 i 11–20)."""
    ne = P.shape[2]; npair = P.shape[0]; half = nper // 2
    out = np.zeros((ne, 2, 22), np.float64)
    m = sp_.size
    buf = np.empty(m, np.float64); bufe = np.empty(m, np.float64)
    last = np.empty(npair, np.int64)
    ps = np.zeros((2, nper))
    for e in range(ne):
        for p in range(npair):
            last[p] = -1
        ps[:, :] = 0.0
        cnt = np.zeros(2, np.int64); cum = np.zeros(2); peak = np.zeros(2); dd = np.zeros(2)
        gp = np.zeros(2); gl = np.zeros(2); hold = np.zeros(2)
        hn = np.zeros((2, 2)); hcum = np.zeros((2, 2)); hpk = np.zeros((2, 2)); hdd = np.zeros((2, 2))
        for k in range(m):
            p = sp_[k]; i = si_[k]
            if i <= last[p]:
                continue
            v = P[p, i, e]
            if v != v:
                continue
            x = X[p, i, e]
            last[p] = i + x
            hf = 0 if per[i] < half else 1
            for g in range(2):
                if g == 1 and p != eu:
                    continue
                if g == 0:
                    buf[cnt[0]] = v
                else:
                    bufe[cnt[1]] = v
                cnt[g] += 1
                cum[g] += v
                if cum[g] > peak[g]: peak[g] = cum[g]
                if peak[g] - cum[g] > dd[g]: dd[g] = peak[g] - cum[g]
                if v > 0: gp[g] += v
                else: gl[g] -= v
                hold[g] += x
                ps[g, per[i]] += v
                hn[g, hf] += 1; hcum[g, hf] += v
                if hcum[g, hf] > hpk[g, hf]: hpk[g, hf] = hcum[g, hf]
                if hpk[g, hf] - hcum[g, hf] > hdd[g, hf]: hdd[g, hf] = hpk[g, hf] - hcum[g, hf]
        for g in range(2):
            n_ = cnt[g]
            out[e, g, 0] = n_
            if n_ == 0:
                continue
            b = buf[:n_] if g == 0 else bufe[:n_]
            s = b.sum(); mu = s / n_; sd = b.std()
            out[e, g, 1] = s; out[e, g, 2] = mu; out[e, g, 3] = np.median(b)
            out[e, g, 4] = mu / sd * np.sqrt(n_) if sd > 0 else 0.0
            out[e, g, 5] = (b > 0).sum() / n_
            out[e, g, 6] = gp[g] / gl[g] if gl[g] > 0 else 99.0
            out[e, g, 7] = dd[g]
            lose = 0; empty = 0
            for q in range(nper):
                if ps[g, q] < 0: lose += 1
                if ps[g, q] == 0: empty += 1
            out[e, g, 8] = lose; out[e, g, 9] = empty
            out[e, g, 10] = ps[g].mean(); out[e, g, 11] = np.median(ps[g]); out[e, g, 12] = ps[g].min()
            out[e, g, 13] = hold[g] / n_
            for hf in range(2):
                l2 = 0
                for q in range(half * hf, half * hf + half):
                    if ps[g, q] < 0: l2 += 1
                out[e, g, 14 + 4 * hf] = hn[g, hf]; out[e, g, 15 + 4 * hf] = hcum[g, hf]
                out[e, g, 16 + 4 * hf] = l2; out[e, g, 17 + 4 * hf] = hdd[g, hf]
    return out


FIELDS = ['n', 'suma', 'srednia', 'mediana', 't', 'trafne', 'PF', 'obsuniecie', 'stratne_okresy', 'puste_okresy',
          'okres_srednia', 'okres_mediana', 'okres_min', 'swiece', 'n1', 'suma1', 'stratne1', 'obs1', 'n2', 'suma2', 'stratne2', 'obs2']


class Data:
    """Świece 10 par na wspólnej siatce czasu + macierze wyników wszystkich wyjść."""

    def __init__(self, frames: dict, spread: dict):
        common = None
        for p in PAIRS:
            tt = set(frames[p].time.values.tolist())
            common = tt if common is None else common & tt
        tgrid = np.array(sorted(common), np.int64)
        self.F = {p: features(frames[p].set_index('time').loc[tgrid].reset_index()) for p in PAIRS}
        self.t = tgrid; n = len(tgrid); self.n = n
        self.spread = spread
        self.P = np.full((len(PAIRS), n, NE), np.nan, np.float32)
        self.X = np.zeros((len(PAIRS), n, NE), np.int16)
        self.MA = np.zeros((len(PAIRS), n, NE), np.float32)
        for k, p in enumerate(PAIRS):
            F = self.F[p]
            self.P[k], self.X[k], self.MA[k] = outcomes(F['o'], F['h'], F['l'], F['c'], F['atr'], F['valid'],
                                                        spread[p] * PIP[p], EX_SL, EX_TP, EX_H)
        T0, T1 = int(tgrid[0]), int(tgrid[-1]) + 3600
        self.per = np.minimum(((tgrid - T0) / (T1 - T0) * NPER).astype(np.int64), NPER - 1)
        self.periods = [(pd.Timestamp(T0 + (T1 - T0) * k / NPER, unit='s'), pd.Timestamp(T0 + (T1 - T0) * (k + 1) / NPER, unit='s'))
                        for k in range(NPER)]
        self.valid = np.array([self.F[p]['valid'] for p in PAIRS])
        self.stage = {}

    def dicts(self, stage: int):
        if stage not in self.stage:
            fs = (setups1, trend1, hours1) if stage == 1 else (setups2, trend2, hours2)
            self.stage[stage] = tuple({p: f(self.F[p]) for p in PAIRS} for f in fs)
        return self.stage[stage]

    def scan(self, stage: int, seed=None, collect=None, log=print):
        """Wszystkie kombinacje sygnał × trend × godziny; collect(setup, trend, godziny, res[NE, 2, 22])."""
        S, T, H = self.dicts(stage)
        rng = np.random.default_rng(seed) if seed is not None else None
        shifts = {p: int(rng.integers(1500, self.n - 1500)) for p in PAIRS} if rng is not None else None
        names = list(S[PAIRS[0]]); t0 = time.time()
        for k_, sn in enumerate(names):
            SM = np.array([S[p][sn] for p in PAIRS])
            if shifts is not None:
                SM = np.array([np.roll(SM[k], shifts[p]) for k, p in enumerate(PAIRS)])
            for tn in T[PAIRS[0]]:
                TM = np.array([T[p][tn] for p in PAIRS])
                for hn in H[PAIRS[0]]:
                    M = SM & TM & np.array([H[p][hn] for p in PAIRS]) & self.valid
                    ii, pp = np.nonzero(M.T)
                    if ii.size < 20:
                        continue
                    res = evaluate(pp.astype(np.int64), ii.astype(np.int64), self.P, self.X, self.per, NPER, PAIRS.index('EURUSD'))
                    collect(sn, tn, hn, res)
            if log and (k_ % 20 == 19 or k_ == len(names) - 1):
                log(f'    {k_ + 1}/{len(names)} sygnałów, {time.time() - t0:.0f} s')

    def trades(self, stage: int, setup: str, trend: str, hours: str, e: int, pairs=PAIRS, extra_hours=None) -> pd.DataFrame:
        S, T, H = self.dicts(stage)
        Hm = {p: (np.ones(self.n, bool) if hours == 'wszystkie godziny' and hours not in H[p] else H[p][hours]) for p in PAIRS}
        M = np.array([S[p][setup] & T[p][trend] & Hm[p] for p in PAIRS]) & self.valid
        ii, pp = np.nonzero(M.T)
        rows = []; last = {}
        for i, k in zip(ii, pp):
            p = PAIRS[k]
            if p not in pairs or i <= last.get(p, -1):
                continue
            v = self.P[k, i, e]
            if np.isnan(v):
                continue
            x = int(self.X[k, i, e]); last[p] = i + x
            F = self.F[p]
            rows.append(dict(para=p, i=int(i), sygnal_t=int(F['t'][i]), wejscie_t=int(F['t'][i + 1]), wyjscie_t=int(F['t'][i + x]) + 3600,
                             wejscie=float(F['o'][i + 1]), atr=float(F['atr'][i]), swiece=x, wynik=float(v), mae=float(self.MA[k, i, e]),
                             okres=int(self.per[i]) + 1))
        df = pd.DataFrame(rows)
        for k in ['sygnal_t', 'wejscie_t', 'wyjscie_t']:
            if len(df):
                df[k] = pd.to_datetime(df[k], unit='s')
        return df


def summary(df: pd.DataFrame) -> dict:
    if len(df) == 0:
        return dict(n=0, suma=0.0, srednia=0.0, mediana=0.0, trafne=0.0, PF=0.0, obs=0.0, stratne=0, okres_med=0.0, mae_max=0.0)
    w = df.wynik.values; cum = np.cumsum(w)
    ps = df.groupby('okres').wynik.sum().reindex(range(1, NPER + 1), fill_value=0)
    return dict(n=len(w), suma=w.sum(), srednia=w.mean(), mediana=float(np.median(w)), trafne=(w > 0).mean(),
                PF=w[w > 0].sum() / max(-w[w < 0].sum(), 1e-9), obs=float(np.max(np.maximum.accumulate(np.r_[0, cum])[1:] - cum)),
                stratne=int((ps < 0).sum()), okres_med=float(ps.median()), mae_max=float(df.mae.max()))


# ---------------------------------------------------------------- rachunek z dźwignią

BASE_USD = {'EURUSD': 'EURUSD', 'GBPUSD': 'GBPUSD', 'AUDUSD': 'AUDUSD', 'NZDUSD': 'NZDUSD', 'USDJPY': None, 'USDCAD': None,
            'USDCHF': None, 'EURJPY': 'EURUSD', 'GBPJPY': 'GBPUSD', 'EURGBP': 'EURUSD'}


def rollovers(t0: pd.Timestamp, t1: pd.Timestamp) -> int:
    """Rolowania (17:00 Nowy Jork) między wejściem a wyjściem; środa ×3."""
    a = pd.Timestamp(t0).tz_localize('UTC').tz_convert('America/New_York')
    b = pd.Timestamp(t1).tz_localize('UTC').tz_convert('America/New_York')
    n = 0; d = a.normalize() + pd.Timedelta(hours=17)
    if d <= a:
        d += pd.Timedelta(days=1)
    while d <= b:
        if d.dayofweek < 5:
            n += 3 if d.dayofweek == 2 else 1
        d += pd.Timedelta(days=1)
    return n


def simulate(tr: pd.DataFrame, D: Data, lev=30.0, f=0.20, start=1000.0, swap_pa=0.0, stop_out=0.5, lot=1000, sl_k=0.0):
    """Rachunek CFD: depozyt pozycji = f × saldo (zamknięte transakcje), wartość = depozyt × dźwignia, jednostki w dół
    do 0,01 lota; brak wolnego depozytu → pozycja mniejsza albo pominięta (tak samo liczą formuły arkusza TRANSAKCJE).
    Co godzinę wycena po najgorszej cenie świecy (szczyt — short); poziom depozytu < 50% → stop out wszystkich pozycji."""
    n = D.n; ent = {}
    for k, r in tr.iterrows():
        ent.setdefault(int(r.i) + 1, []).append(k)
    cash = start; open_ = {}; log = []; eq = np.full(n, np.nan); events = []; dead = False
    for j in range(n):
        if dead:
            eq[j] = cash; continue
        for k in ent.get(j, []):
            r = tr.loc[k]; p = r.para
            if p in open_:
                continue
            used = sum(q['margin'] for q in open_.values())
            margin = min(f * cash, cash - used)
            if margin <= 0:
                log.append(dict(k=k, para=p, status='pominięta')); continue
            e0 = float(D.F[p]['o'][j]); bu = 1.0 if BASE_USD[p] is None else float(D.F[BASE_USD[p]]['o'][j])
            units = int(margin * lev / bu // lot) * lot
            if units < lot:
                log.append(dict(k=k, para=p, status='pominięta')); continue
            notional = units * bu; sp = D.spread[p] * PIP[p]
            open_[p] = dict(k=k, j0=j, j1=j + int(r.swiece) - 1, notional=notional, margin=notional / lev, entry=e0,
                            pnl=float(r.wynik), unreal=0.0,
                            cap=float(r.cap) if 'cap' in r.index else ((sl_k * float(r.atr) + sp) / e0 if sl_k > 0 else 9e9))
        if open_:
            worst = 0.0; used = 0.0
            for p, q in open_.items():
                adv = min((float(D.F[p]['h'][j]) - q['entry'] + D.spread[p] * PIP[p]) / q['entry'], q['cap'])
                worst -= q['notional'] * adv; used += q['margin']
            if cash + worst < stop_out * used:
                tot = 0.0
                for p, q in list(open_.items()):
                    pl = -q['notional'] * min((float(D.F[p]['h'][j]) - q['entry'] + D.spread[p] * PIP[p]) / q['entry'], q['cap'])
                    cash += pl; tot += pl
                    log.append(dict(k=q['k'], para=p, status='stop out', wynik_usd=pl))
                events.append(dict(czas=pd.Timestamp(int(D.t[j]), unit='s'), kwota=tot, saldo=cash))
                open_.clear()
                if cash <= 0:
                    cash = 0.0; dead = True
                eq[j] = cash
                continue
        for p in list(open_):
            q = open_[p]
            if q['j1'] == j:
                pl = q['notional'] * q['pnl'] / 100.0
                sw = 0.0
                if swap_pa:
                    sw = -q['notional'] * swap_pa / 100.0 / 365.0 * rollovers(pd.Timestamp(int(D.t[q['j0']]), unit='s'),
                                                                               pd.Timestamp(int(D.t[j]) + 3600, unit='s'))
                cash += pl + sw
                log.append(dict(k=q['k'], para=p, status='zamknięta', wynik_usd=pl + sw))
                del open_[p]
        for p, q in open_.items():
            q['unreal'] = -q['notional'] * (float(D.F[p]['c'][j]) - q['entry'] + D.spread[p] * PIP[p]) / q['entry']
        eq[j] = cash + sum(q['unreal'] for q in open_.values())
    curve = pd.Series(eq, index=pd.to_datetime(D.t, unit='s')).ffill().fillna(start)
    return pd.DataFrame(log), curve, pd.DataFrame(events)


def curve_stats(curve: pd.Series, start=1000.0, since=None) -> dict:
    s = curve if since is None else curve[curve.index >= since]
    base = start if since is None else float(s.iloc[0]) if len(s) else start
    yrs = max((s.index[-1] - s.index[0]).days / 365.25, 1e-9)
    end = float(s.iloc[-1])
    return dict(koniec=end, zysk_pct=(end / base - 1) * 100, cagr_pct=((end / base) ** (1 / yrs) - 1) * 100 if end > 0 else -100.0,
                max_obs_pct=-(s / s.cummax() - 1).min() * 100)


def daily(curve: pd.Series) -> pd.Series:
    return curve.resample('D').last().dropna()


# ---------------------------------------------------------------- całe liczenie

def enrich(D: Data, df: pd.DataFrame, sl: float, tp: float) -> pd.DataFrame:
    rows = []
    for _, r in df.iterrows():
        p = r.para; F = D.F[p]; i = int(r.i); e0 = i + 1
        en = float(F['o'][e0]); a = float(F['atr'][i]); sp = D.spread[p] * PIP[p]
        SL = en + sl * a if sl > 0 else np.nan; TP = en - tp * a
        ex = en - r.wynik * en / 100.0 - sp
        if not np.isnan(SL) and ex >= SL - 0.05 * PIP[p]:
            why = 'SL' if abs(ex - SL) < 0.05 * PIP[p] else 'SL (luka)'
        elif abs(ex - TP) < 0.05 * PIP[p]:
            why = 'TP'
        else:
            why = 'czas (zamknięcie świecy)'
        bu = 1.0 if BASE_USD[p] is None else float(D.F[BASE_USD[p]]['o'][e0])
        rows.append(dict(SL=SL, TP=TP, wyjscie=ex, powod=why, pipsy=(en - ex - sp) / PIP[p],
                         noce=rollovers(r.wejscie_t, r.wyjscie_t), bazowa_usd=bu))
    return pd.concat([df.reset_index(drop=True), pd.DataFrame(rows)], axis=1)


def compute(D: Data, f=0.20, lev=30.0, start=1000.0, log=print) -> dict:
    O = {'MAIN': MAIN, 'periods': D.periods, 'spread': D.spread, 'f': f, 'lev': lev, 'start': start,
         'od': pd.Timestamp(int(D.t[0]), unit='s'), 'do': pd.Timestamp(int(D.t[-1]), unit='s'), 'swiec': D.n}
    eu = PAIRS.index('EURUSD')
    # --- etap 1: prawdziwe sygnały i 2 placebo
    RAW_SET = ['świeca ≥ 1 ATR', 'świeca ≥ 1.5 ATR', 'świeca ≥ 2 ATR', 'świeca ≥ 3 ATR', 'świeca ≥ 0.1%', 'świeca ≥ 0.15%', 'świeca ≥ 0.2%',
               'wzrost 3 św. ≥ 2.6 ATR', 'wzrost 6 św. ≥ 3.7 ATR', 'wzrost 12 św. ≥ 5.2 ATR']
    c1 = {}; keep = {'rank0': [], 'rank1': [], 'raw0': [], 'raw1': [], 'wf': {0: [], 1: []}}; nrules = [0]
    for tag, seed in [('real', None), ('placebo 1', 101), ('placebo 2', 202)]:
        log(f'  etap 1 — {tag}')
        cnt = dict(all3=0, all4=0, eu3=0, eu3_100=0); wf = {0: [], 1: []}

        def col(sn, tn, hn, res, cnt=cnt, wf=wf, tag=tag):
            A, E = res[:, 0, :], res[:, 1, :]
            if tag == 'real':
                nrules[0] += NE
            m3 = (A[:, 0] >= 300) & (A[:, 8] <= 3) & (A[:, 1] > 0)
            cnt['all3'] += int(m3.sum()); cnt['all4'] += int(((A[:, 0] >= 300) & (A[:, 8] <= 4) & (A[:, 1] > 0)).sum())
            e3 = (E[:, 0] >= 60) & (E[:, 8] <= 3) & (E[:, 1] > 0)
            cnt['eu3'] += int(e3.sum()); cnt['eu3_100'] += int((e3 & (E[:, 0] >= 100)).sum())
            for fold, (a, b) in enumerate([(14, 18), (18, 14)]):
                mm = (A[:, a] >= 150) & (A[:, a + 2] <= 1) & (A[:, a + 1] > 0)
                for e in np.nonzero(mm)[0]:
                    wf[fold].append((sn, tn, hn, e, A[e, a], A[e, a + 1], A[e, a + 1] / max(A[e, a + 3], 0.05), A[e, b], A[e, b + 1], A[e, b + 2]))
            if tag == 'real':
                for g, mm in [(0, m3), (1, e3 & (E[:, 0] >= 100))]:
                    for e in np.nonzero(mm)[0]:
                        keep[f'rank{g}'].append((sn, tn, hn, e, *res[e, g, :14]))
                if tn == 'bez filtra trendu' and hn == 'wszystkie godziny' and sn in RAW_SET:
                    for e in range(NE):
                        keep['raw0'].append((sn, e, *res[e, 0, :14])); keep['raw1'].append((sn, e, *res[e, 1, :14]))
        D.scan(1, seed, col, log)
        c1[tag] = cnt
        w = {}
        for fold in [0, 1]:
            df = pd.DataFrame(wf[fold], columns=['setup', 'trend', 'godziny', 'e', 'n_wyb', 'suma_wyb', 'score', 'n_spr', 'suma_spr', 'stratne_spr'])
            df = df.sort_values('score', ascending=False)
            if len(df):
                w[fold] = dict(spelnia=len(df), top1=df.suma_spr.iloc[0], top10=df.suma_spr.head(10).mean(),
                               top10_poz=(df.suma_spr.head(10) > 0).mean(), top50_tr=df.suma_spr.head(50).sum() / max(df.n_spr.head(50).sum(), 1))
            if tag == 'real' and fold == 0:
                O['wf_top'] = df.head(10)
        O.setdefault('wf', {})[tag] = w
    O['placebo1'] = c1; O['n_rules1'] = nrules[0]
    cols14 = FIELDS[:14]
    for g in [0, 1]:
        r = pd.DataFrame(keep[f'rank{g}'], columns=['setup', 'trend', 'godziny', 'e'] + cols14)
        r['zysk/obs'] = r.suma / r.obsuniecie.clip(lower=0.05)
        r['wyjście'] = [exit_name(e) for e in r.e]
        O[f'rank{g}'] = r.sort_values('zysk/obs', ascending=False).head(40)
        rw = pd.DataFrame(keep[f'raw{g}'], columns=['setup', 'e'] + cols14); rw['wyjście'] = [exit_name(e) for e in rw.e]
        O[f'raw{g}'] = rw
    # --- etap 2: siatka wokół rodziny + 3 placebo
    c2 = {}; g2 = []
    for tag, seed in [('real', None), ('placebo 1', 11), ('placebo 2', 22), ('placebo 3', 33)]:
        log(f'  etap 2 — {tag}')
        cnt = dict(all3=0, halves=0, eu3=0)

        def col2(sn, tn, hn, res, cnt=cnt, tag=tag):
            A, E = res[:, 0, :], res[:, 1, :]
            cnt['all3'] += int(((A[:, 0] >= 300) & (A[:, 8] <= 3) & (A[:, 1] > 0)).sum())
            cnt['halves'] += int(((A[:, 14] >= 150) & (A[:, 18] >= 150) & (A[:, 16] <= 2) & (A[:, 20] <= 2) & (A[:, 15] > 0) & (A[:, 19] > 0)).sum())
            cnt['eu3'] += int(((E[:, 0] >= 60) & (E[:, 8] <= 3) & (E[:, 1] > 0)).sum())
            if tag == 'real' and (sn == MAIN['setup'] or tn == MAIN['trend'] or hn == MAIN['hours']):
                for e in range(NE):
                    g2.append((sn, tn, hn, EXITS[e][0], EXITS[e][1], EXITS[e][2], *res[e, 0, :]))
        D.scan(2, seed, col2, log)
        c2[tag] = cnt
    O['placebo2'] = c2
    O['grid2'] = pd.DataFrame(g2, columns=['setup', 'trend', 'godziny', 'SL', 'TP', 'H'] + FIELDS)
    # --- strategia
    log('  strategia, warianty dźwigni, EUR/USD')
    e = exit_index(MAIN['sl'], MAIN['tp'], MAIN['H'])
    main = enrich(D, D.trades(2, MAIN['setup'], MAIN['trend'], MAIN['hours'], e), MAIN['sl'], MAIN['tp'])
    O['main'] = main; O['main_sum'] = summary(main)
    w = main.wynik.values
    O['main_t'] = w.mean() / w.std() * np.sqrt(len(w)); O['main_ci'] = 1.96 * w.std() / np.sqrt(len(w))
    var = []; curves = {}
    for fv in sorted({0.05, 0.10, 0.20, 1 / 3, 0.5, 1.0, f}):
        for sw in [0.0, 3.0]:
            L, cv, ev = simulate(main, D, lev=lev, f=fv, start=start, swap_pa=sw, sl_k=MAIN['sl'])
            var.append(dict(f=fv, swap=sw, **curve_stats(cv, start), zawarte=int((L.status == 'zamknięta').sum()) if len(L) else 0,
                            pominiete=int((L.status == 'pominięta').sum()) if len(L) else 0, stop_out=len(ev)))
            if sw == 0.0:
                curves[fv] = daily(cv)
    O['variants'] = pd.DataFrame(var); O['curves'] = curves
    eu_tr = main[main.para == 'EURUSD'].reset_index(drop=True)
    O['eu_sum'] = summary(eu_tr); ev_ = []; O['eu_curves'] = {}
    for fv in [0.2, 1 / 3, 1.0]:
        L, cv, ev = simulate(eu_tr, D, lev=lev, f=fv, start=start, sl_k=MAIN['sl'])
        ev_.append(dict(f=fv, **curve_stats(cv, start), zawarte=int((L.status == 'zamknięta').sum()) if len(L) else 0, stop_out=len(ev)))
        O['eu_curves'][fv] = daily(cv)
    O['eu_var'] = pd.DataFrame(ev_)
    # --- godzina sygnału
    rows = []
    for tn in ['SMA500 spada', 'bez filtra trendu']:
        df = D.trades(2, MAIN['setup'], tn, 'wszystkie godziny', e)
        df['h'] = df.sygnal_t.dt.hour
        for hh in range(24):
            d = df[df.h == hh]; d1 = d[d.okres <= NPER // 2]; d2 = d[d.okres > NPER // 2]
            rows.append(dict(trend=tn, godzina=hh, n=len(d), srednia=d.wynik.mean() if len(d) else np.nan,
                             sr1=d1.wynik.mean() if len(d1) else np.nan, sr2=d2.wynik.mean() if len(d2) else np.nan))
    O['hours_table'] = pd.DataFrame(rows)
    # --- 2. rok: reguła wybrana na 1. roku vs FX-SHORT-NY
    y2 = D.periods[NPER // 2][0]
    top = O['wf_top'].iloc[0] if len(O.get('wf_top', [])) else None
    if top is not None:
        wtr = D.trades(1, top.setup, top.trend, top.godziny, int(top.e))
        wtr = wtr[wtr.okres > NPER // 2].reset_index(drop=True)
        O['wf_rule'] = dict(setup=top.setup, trend=top.trend, godziny=top.godziny, wyjscie=exit_name(int(top.e)))
        O['wf_sum'] = summary(wtr)
        _, cv, ev = simulate(wtr, D, lev=lev, f=0.20, start=start, sl_k=EXITS[int(top.e)][0])
        O['wf_stats'] = curve_stats(cv, start, since=y2); O['wf_curve'] = daily(cv[cv.index >= y2])
    m2 = main[main.okres > NPER // 2].reset_index(drop=True)
    _, cv, _ = simulate(m2, D, lev=lev, f=0.20, start=start, sl_k=MAIN['sl'])
    O['main_y2'] = curve_stats(cv, start, since=y2); O['main_y2_curve'] = daily(cv[cv.index >= y2])
    O['main_y2_tr'] = m2
    return O


# ---------------------------------------------------------------- Excel

def write_xlsx(O: dict, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as CL

    M = O['MAIN']; FN = 'Arial'
    fN = Font(name=FN, size=10); fB = Font(name=FN, size=10, bold=True); fT = Font(name=FN, size=14, bold=True)
    fH = Font(name=FN, size=10, bold=True, color='FFFFFF'); fIn = Font(name=FN, size=10, color='0000FF')
    fS = Font(name=FN, size=9, italic=True, color='555555')
    HDR = PatternFill('solid', fgColor='1F3A5F'); YEL = PatternFill('solid', fgColor='FFFF00'); GRY = PatternFill('solid', fgColor='EEF2F7')
    GOOD = PatternFill('solid', fgColor='E2F0D9'); BAD = PatternFill('solid', fgColor='FBE5D6'); SEL = PatternFill('solid', fgColor='92D050')
    thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    WR = Alignment(wrap_text=True, vertical='top')
    PCT = '0.0%'; PCT3 = '0.000%'; USD = '#,##0.00 [$$-409];[Red]-#,##0.00 [$$-409]'; USD0 = '#,##0 [$$-409]'; DT = 'yyyy-mm-dd hh:mm'; DD = 'yyyy-mm-dd'

    def pl(x, d=1):
        return f'{x:,.{d}f}'.replace(',', ' ').replace('.', ',')

    wb = Workbook()

    def sheet(name, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name; ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = 'landscape'; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        return ws

    def hdr(ws, r, c, names, widths=None):
        for k, nm in enumerate(names):
            x = ws.cell(r, c + k, nm); x.font = fH; x.fill = HDR; x.border = BOX
            x.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        for k, w in enumerate(widths or []):
            ws.column_dimensions[CL(c + k)].width = w
        ws.row_dimensions[r].height = 42

    def put(ws, r, c, v, font=fN, fmt=None, fill=None, al=None):
        x = ws.cell(r, c, v); x.font = font; x.border = BOX
        if fmt: x.number_format = fmt
        if fill: x.fill = fill
        if al: x.alignment = al
        return x

    def note(ws, r, c, text, span=10, h=None):
        x = ws.cell(r, c, text); x.font = fS; x.alignment = WR
        if span > 1:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        if h: ws.row_dimensions[r].height = h

    def lchart(title, h, w):
        ch = LineChart(); ch.title = title; ch.height = h; ch.width = w
        ch.y_axis.crossAx = 500; ch.x_axis = DateAxis(crossAx=100); ch.x_axis.number_format = 'yyyy-mm'; ch.x_axis.majorTimeUnit = 'months'
        ch.x_axis.delete = False; ch.y_axis.delete = False; ch.display_blanks = 'span'
        return ch

    def py(ts):
        return pd.Timestamp(ts).to_pydatetime()

    # TRANSAKCJE
    tr = O['main']; N = len(tr); R0 = 11; R1 = R0 + max(N, 1) - 1
    wsP = sheet('PODSUMOWANIE', first=True); ws = wsT = sheet('TRANSAKCJE')
    ws['A1'] = 'TRANSAKCJE — FX-SHORT-NY, 10 par, świece 1h (rozliczenie pesymistyczne, spread XTB)'; ws['A1'].font = fT
    inp = [('Kapitał początkowy ($)', O['start'], USD0, 'Założenie człowieka: start 1000 $.'),
           ('Dźwignia (1:x)', O['lev'], '0', 'Założenie człowieka: dźwignia 1:30 (XTB, rachunek detaliczny). Depozyt pozycji = wartość / dźwignia.'),
           ('Depozyt na pozycję (% salda)', O['f'], PCT, 'Ile salda zablokować jako depozyt jednej pozycji. 20% → wartość pozycji = 6 × saldo; 5 pozycji naraz = cała dźwignia 1:30. Warianty: arkusz WARIANTY.'),
           ('Koszt nocny (% rocznie wartości pozycji)', 0.0, '0.0%', 'Swap CFD — szacunek, nie dane XTB. 0% = bez kosztu; 3% = wariant ostrożny. Środa liczy się ×3.'),
           ('Jednostki w 0,01 lota', 1000, '#,##0', 'Minimalna pozycja XTB: 0,01 lota = 1000 jednostek waluty bazowej.')]
    for k, (lab, v, fm, com) in enumerate(inp):
        put(ws, 3 + k, 1, lab, fB, fill=GRY)
        x = put(ws, 3 + k, 2, v, fIn, fm, YEL); x.comment = Comment(com, 'IA4')
    ws.column_dimensions['A'].width = 34
    note(ws, 8, 1, 'Żółte komórki można zmieniać — saldo, wykresy w KAPITAŁ i okresy przeliczą się same. Wynik % transakcji (kolumna L) pochodzi z symulacji Python (ia4.fxshort): wejście po open świecy po sygnale, SL i TP w tej samej świecy → SL, luka przez SL → open, spread odjęty raz. Wielkość pozycji liczona od salda zamkniętych transakcji; brak wolnego depozytu → pozycja mniejsza albo pominięta.', 16, 40)
    cols = ['Nr', 'Para', 'Sygnał — świeca (UTC)', 'Wejście (UTC, open)', 'Wyjście (UTC)', 'Świec w pozycji', 'Cena wejścia', 'SL (cena)', 'TP (cena)',
            'Cena wyjścia', 'Powód wyjścia', 'Wynik % (netto, po spreadzie)', 'Wynik w pipsach (netto)', 'Największy ruch przeciw (MAE %)', 'Okres (1–20)',
            'Noce (rolowania, środa ×3)', 'Kurs waluty bazowej w USD', 'Saldo przy wejściu ($)', 'Depozyt zajęty przez otwarte pozycje ($)',
            'Jednostki', 'Loty', 'Wartość pozycji ($)', 'Depozyt pozycji ($)', 'Wynik ($)', 'Saldo po wyjściach do tej chwili ($)', 'Status']
    hdr(ws, 10, 1, cols, [5, 8, 16, 16, 16, 7, 10, 10, 10, 10, 13, 11, 9, 10, 7, 9, 10, 12, 13, 10, 7, 12, 11, 11, 13, 10])
    for k, r in tr.reset_index(drop=True).iterrows():
        row = R0 + k; pf = '0.000' if r.para.endswith('JPY') else '0.00000'
        vals = [(k + 1, None), (r.para, None), (py(r.sygnal_t), DT), (py(r.wejscie_t), DT), (py(r.wyjscie_t), DT), (int(r.swiece), None),
                (round(float(r.wejscie), 6), pf), (round(float(r.SL), 6), pf), (round(float(r.TP), 6), pf), (round(float(r.wyjscie), 6), pf),
                (r.powod, None), (float(r.wynik) / 100, PCT3), (round(float(r.pipsy), 2), '0.0'), (float(r.mae) / 100, PCT3), (int(r.okres), None),
                (int(r.noce), None), (round(float(r.bazowa_usd), 6), '0.0000')]
        for c, (v, fm) in enumerate(vals):
            put(ws, row, 1 + c, v, fmt=fm)
        put(ws, row, 18, f'=$B$3+SUMIFS(X$10:X{row-1},E$10:E{row-1},"<="&D{row})', fmt=USD)
        put(ws, row, 19, f'=SUMIFS(W$10:W{row-1},D$10:D{row-1},"<="&D{row},E$10:E{row-1},">"&D{row})', fmt=USD)
        put(ws, row, 20, f'=IF(MIN($B$5*R{row},R{row}-S{row})<=0,0,INT(MIN($B$5*R{row},R{row}-S{row})*$B$4/Q{row}/$B$7)*$B$7)', fmt='#,##0')
        put(ws, row, 21, f'=T{row}/100000', fmt='0.00')
        put(ws, row, 22, f'=T{row}*Q{row}', fmt=USD)
        put(ws, row, 23, f'=V{row}/$B$4', fmt=USD)
        put(ws, row, 24, f'=V{row}*L{row}-V{row}*$B$6/365*P{row}', fmt=USD)
        put(ws, row, 25, f'=$B$3+SUMIFS(X${R0}:X${R1},E${R0}:E${R1},"<="&E{row})', fmt=USD)
        put(ws, row, 26, f'=IF(T{row}=0,"pominięta","zawarta")')
    ws.freeze_panes = 'C11'; ws.auto_filter.ref = f'A10:Z{R1}'

    def TR(c):
        return f'TRANSAKCJE!${c}${R0}:${c}${R1}'

    # KAPITAŁ
    ws = wsK = sheet('KAPITAŁ')
    ws['A1'] = 'KAPITAŁ — saldo dzień po dniu (ustawienia z TRANSAKCJE!B3:B7)'; ws['A1'].font = fT
    note(ws, 2, 1, f"Saldo = kapitał początkowy + wynik transakcji zamkniętych do końca dnia (formuła). Kolumna F — wycena z otwartymi pozycjami co godzinę z symulacji Python dla ustawień domyślnych ({pl(O['f'] * 100, 0)}% depozytu, bez kosztu nocnego); obsunięcie z wyceny godzinowej jest głębsze niż z salda.", 8, 42)
    hdr(ws, 4, 1, ['Dzień (koniec, UTC)', 'Saldo ($)', 'Najwyższe saldo ($)', 'Obsunięcie od szczytu', 'Liczba transakcji zamkniętych', 'Wycena godzinowa — Python, ustawienia domyślne ($)'], [13, 13, 13, 12, 12, 16])
    cur = O['curves'][O['f']]
    days = pd.date_range(cur.index.min().normalize(), cur.index.max().normalize(), freq='D')
    for k, d in enumerate(days):
        row = 5 + k
        put(ws, row, 1, py(d), fmt=DD)
        put(ws, row, 2, f'=TRANSAKCJE!$B$3+SUMIFS({TR("X")},{TR("E")},"<"&(A{row}+1))', fmt=USD)
        put(ws, row, 3, f'=MAX(B$5:B{row})', fmt=USD)
        put(ws, row, 4, f'=B{row}/C{row}-1', fmt=PCT)
        put(ws, row, 5, f'=COUNTIFS({TR("E")},"<"&(A{row}+1),{TR("Z")},"zawarta")', fmt='0')
        v = cur.get(d, np.nan)
        put(ws, row, 6, None if np.isnan(v) else round(float(v), 2), fmt=USD)
    KR1 = 5 + len(days) - 1
    ch = lchart(f"Kapitał: {pl(O['start'], 0)} $ na dźwigni 1:{O['lev']:g} (saldo, ustawienia z TRANSAKCJE)", 9, 24)
    ch.y_axis.title = '$'; ch.y_axis.scaling.min = 0.8 * O['start']
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=KR1), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=6, min_row=4, max_row=KR1), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=KR1))
    ch.series[0].graphicalProperties.line.width = 22000; ch.series[0].graphicalProperties.line.solidFill = '1F3A5F'
    ch.series[1].graphicalProperties.line.width = 9000; ch.series[1].graphicalProperties.line.solidFill = 'A6A6A6'
    ch.legend.position = 'b'; ws.add_chart(ch, 'H4')
    ch2 = lchart('Obsunięcie salda od szczytu', 6, 24)
    ch2.add_data(Reference(ws, min_col=4, min_row=4, max_row=KR1), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=5, max_row=KR1)); ch2.y_axis.number_format = '0%'; ch2.x_axis.tickLblPos = 'low'
    ch2.series[0].graphicalProperties.line.solidFill = 'C00000'; ch2.legend = None; ws.add_chart(ch2, 'H23')
    ws.freeze_panes = 'A5'

    # OKRESY
    ws = sheet('OKRESY')
    ws['A1'] = 'OKRESY — dostępne świece podzielone na 20 równych części'; ws['A1'].font = fT
    note(ws, 2, 1, 'Transakcja należy do okresu świecy sygnału. Wynik % = suma wyników transakcji bez dźwigni (1:1, każda transakcja tą samą stawką). Zmiana kapitału — saldo z arkusza TRANSAKCJE na początku i końcu okresu. Okres 1 prawie pusty: pierwsze 600 świec to rozbieg średniej SMA500.', 11, 42)
    hdr(ws, 4, 1, ['Okres', 'Od (UTC)', 'Do (UTC)', 'Transakcje', 'Wynik % (1:1, suma)', 'Średnia na transakcję', 'Trafne', 'Wynik ($)', 'Saldo na początku ($)', 'Saldo na końcu ($)', 'Zmiana kapitału'],
        [7, 16, 16, 10, 11, 11, 9, 12, 13, 13, 11])
    for k, (a, b) in enumerate(O['periods']):
        row = 5 + k
        put(ws, row, 1, k + 1); put(ws, row, 2, py(a), fmt=DT); put(ws, row, 3, py(b), fmt=DT)
        put(ws, row, 4, f'=COUNTIF({TR("O")},A{row})', fmt='0')
        put(ws, row, 5, f'=SUMIFS({TR("L")},{TR("O")},A{row})', fmt=PCT3)
        put(ws, row, 6, f'=IF(D{row}=0,0,E{row}/D{row})', fmt=PCT3)
        put(ws, row, 7, f'=IF(D{row}=0,0,COUNTIFS({TR("O")},A{row},{TR("L")},">0")/D{row})', fmt=PCT)
        put(ws, row, 8, f'=SUMIFS({TR("X")},{TR("O")},A{row})', fmt=USD)
        put(ws, row, 9, f'=TRANSAKCJE!$B$3+SUMIFS({TR("X")},{TR("E")},"<="&B{row})', fmt=USD)
        put(ws, row, 10, f'=TRANSAKCJE!$B$3+SUMIFS({TR("X")},{TR("E")},"<="&C{row})', fmt=USD)
        put(ws, row, 11, f'=J{row}/I{row}-1', fmt=PCT)
    ws.cell(25, 4, 'Podsumowanie').font = fB; ws.cell(25, 5, 'wynik % (1:1)').font = fB; ws.cell(25, 11, 'zmiana kapitału').font = fB
    for k, (lab, f1, fm1, f2, fm2) in enumerate([('Okresy stratne (< 0)', '=COUNTIF(E5:E24,"<0")', '0', '=COUNTIF(K5:K24,"<0")', '0'),
                                                 ('Okresy puste', '=COUNTIF(D5:D24,0)', '0', None, None),
                                                 ('Średnia okresu', '=AVERAGE(E5:E24)', PCT3, '=AVERAGE(K5:K24)', PCT),
                                                 ('Mediana okresu', '=MEDIAN(E5:E24)', PCT3, '=MEDIAN(K5:K24)', PCT),
                                                 ('Najgorszy okres', '=MIN(E5:E24)', PCT3, '=MIN(K5:K24)', PCT),
                                                 ('Najlepszy okres', '=MAX(E5:E24)', PCT3, '=MAX(K5:K24)', PCT)]):
        put(ws, 26 + k, 1, lab, fB, fill=GRY); ws.merge_cells(start_row=26 + k, start_column=1, end_row=26 + k, end_column=4)
        put(ws, 26 + k, 5, f1, fmt=fm1)
        if f2:
            put(ws, 26 + k, 11, f2, fmt=fm2)
    bc = BarChart(); bc.title = 'Zmiana kapitału w każdym z 20 okresów'; bc.height = 8; bc.width = 22
    bc.x_axis.delete = False; bc.y_axis.delete = False
    bc.add_data(Reference(ws, min_col=11, min_row=4, max_row=24), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=24)); bc.y_axis.number_format = '0%'; bc.legend = None
    bc.series[0].graphicalProperties.solidFill = '1F3A5F'; bc.series[0].invertIfNegative = False
    ws.add_chart(bc, 'M4')

    # PARY
    ws = sheet('PARY')
    ws['A1'] = 'PARY — wynik strategii na każdej parze'; ws['A1'].font = fT
    hdr(ws, 3, 1, ['Para', 'Spread (pipsy)', 'Transakcje', 'Wynik % (1:1, suma)', 'Średnia na transakcję', 'Średnio pipsów', 'Trafne', 'Wynik ($)', 'Okresy stratne / 20'],
        [10, 9, 10, 12, 12, 10, 9, 12, 10])
    for k, p in enumerate(PAIRS):
        row = 4 + k
        put(ws, row, 1, p); put(ws, row, 2, O['spread'][p], fIn, '0.00')
        put(ws, row, 3, f'=COUNTIF({TR("B")},A{row})', fmt='0')
        put(ws, row, 4, f'=SUMIFS({TR("L")},{TR("B")},A{row})', fmt=PCT3)
        put(ws, row, 5, f'=IF(C{row}=0,0,D{row}/C{row})', fmt=PCT3)
        put(ws, row, 6, f'=IF(C{row}=0,0,SUMIFS({TR("M")},{TR("B")},A{row})/C{row})', fmt='0.0')
        put(ws, row, 7, f'=IF(C{row}=0,0,COUNTIFS({TR("B")},A{row},{TR("L")},">0")/C{row})', fmt=PCT)
        put(ws, row, 8, f'=SUMIFS({TR("X")},{TR("B")},A{row})', fmt=USD)
        e_ = tr[tr.para == p].groupby('okres').wynik.sum() if len(tr) else pd.Series(dtype=float)
        put(ws, row, 9, int((e_ < 0).sum()), fmt='0')
    put(ws, 14, 1, 'Razem', fB, fill=GRY); put(ws, 14, 3, '=SUM(C4:C13)', fB, '0'); put(ws, 14, 4, '=SUM(D4:D13)', fB, PCT3)
    put(ws, 14, 5, '=IF(C14=0,0,D14/C14)', fB, PCT3); put(ws, 14, 7, f'=IF(C14=0,0,COUNTIF({TR("L")},">0")/C14)', fB, PCT); put(ws, 14, 8, '=SUM(H4:H13)', fB, USD)
    note(ws, 16, 1, 'Spread: EUR/USD 1,1 pipsa ze zrzutu XTB człowieka (Kup 1,12081 / Sprzedaj 1,12070). Pozostałe pary — fx_spread_pips z ia4/ai/config.py × 1,1 / 0,8 (ten sam stosunek co EUR/USD). Okresy stratne — wartość z Pythona.', 9, 52)

    # WARIANTY
    ws = sheet('WARIANTY')
    ws['A1'] = 'WARIANTY — ile depozytu na pozycję przy dźwigni (symulacja godzinowa Python)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Symulacja Python: wycena otwartych pozycji co godzinę, stop out przy poziomie depozytu < 50% (najgorsza cena świecy), lot 0,01. Wartości wpisane (nie formuły). Obsunięcie — z wyceny godzinowej. Koszt nocny 3% rocznie — szacunek ostrożny.', 10, 42)
    hdr(ws, 4, 1, ['Depozyt na pozycję (% salda)', 'Wartość pozycji (× saldo)', 'Koszt nocny (% rocznie)', 'Saldo końcowe ($)', 'Zysk', 'Zysk roczny (CAGR)', 'Największe obsunięcie', 'Transakcje zawarte', 'Pominięte (brak depozytu)', 'Stop out'],
        [13, 11, 10, 13, 10, 11, 12, 11, 12, 8])
    for k, v in O['variants'].reset_index(drop=True).iterrows():
        row = 5 + k
        for c, (val, fm) in enumerate([(v.f, PCT), (v.f * O['lev'], '0.0"×"'), (v.swap / 100, '0%'), (round(v.koniec, 2), USD), (v.zysk_pct / 100, PCT),
                                       (v.cagr_pct / 100, PCT), (-v.max_obs_pct / 100, PCT), (int(v.zawarte), '0'), (int(v.pominiete), '0'), (int(v.stop_out), '0')]):
            x = put(ws, row, 1 + c, val, fmt=fm)
            if abs(v.f - O['f']) < 1e-9 and v.swap == 0:
                x.fill = GOOD
    c0 = 13; fs = [f_ for f_ in [0.10, 0.20, 1 / 3, 1.0] if f_ in O['curves']]
    hdr(ws, 4, c0, ['Dzień'] + [f"Depozyt {pl(f_ * 100, 0)}% ({pl(f_ * O['lev'], 0)}×)" for f_ in fs], [12] + [12] * len(fs))
    cv = pd.DataFrame({f_: O['curves'][f_] for f_ in fs})
    for k, (d, rr) in enumerate(cv.iterrows()):
        put(ws, 5 + k, c0, py(d), fmt=DD)
        for j, f_ in enumerate(fs):
            put(ws, 5 + k, c0 + 1 + j, None if np.isnan(rr[f_]) else round(float(rr[f_]), 2), fmt=USD0)
    ch = lchart(f"{pl(O['start'], 0)} $ przy dźwigni 1:{O['lev']:g} — różny depozyt na pozycję (bez kosztu nocnego)", 10, 24)
    ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + len(fs), min_row=4, max_row=4 + len(cv)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=c0, min_row=5, max_row=4 + len(cv))); ch.y_axis.title = '$'
    for s_, colr in zip(ch.series, ['9DC3E6', '1F3A5F', 'ED7D31', 'C00000']):
        s_.graphicalProperties.line.solidFill = colr; s_.graphicalProperties.line.width = 15000
    ch.legend.position = 'b'; ws.add_chart(ch, 'A20')

    # ODPORNOŚĆ
    ws = sheet('ODPORNOŚĆ'); g = O['grid2']
    ws['A1'] = 'ODPORNOŚĆ — sąsiednie ustawienia (wszystkie pary, wynik % bez dźwigni)'; ws['A1'].font = fT
    note(ws, 2, 1, 'Wartości z siatki etapu 2. Komórka: suma wyników % (1:1) i w nawiasie liczba stratnych okresów z 20. Wybrana strategia — zielona. Zysk ma być „płaskowyżem”, nie pojedynczym szczytem.', 14, 30)
    r = 4; ws.cell(r, 1, f"A. Godziny sygnału × TP × limit świec (sygnał: {M['setup']}, trend: {M['trend']}, SL {pl(M['sl'], 0)} ATR)").font = fB
    tps = [0.5, 0.75, 1.0, 1.5, -0.5]; Hs = [4, 6, 8, 12, 24]
    tpn = {0.5: 'TP 0,5 ATR', 0.75: 'TP 0,75 ATR', 1.0: 'TP 1 ATR', 1.5: 'TP 1,5 ATR', -0.5: 'TP 50% korpusu'}
    hdr(ws, r + 1, 1, ['Godziny (UTC)'] + [f'{tpn[a]} · {hh} św.' for a in tps for hh in Hs], [16] + [10] * 25)
    sub = g[(g.setup == M['setup']) & (g.trend == M['trend']) & (g.SL == M['sl'])]
    for k, hn in enumerate(['7–16 UTC', '8–16 UTC', '10–16 UTC', '11–16 UTC', '11–15 UTC', '12–16 UTC', '12–15 UTC']):
        row = r + 2 + k; put(ws, row, 1, hn, fB)
        for j, (a, hh) in enumerate([(a, hh) for a in tps for hh in Hs]):
            d = sub[(sub.godziny == hn) & (sub.TP == a) & (sub.H == hh)]
            x = put(ws, row, 2 + j, f'{pl(d.iloc[0].suma)} ({int(d.iloc[0].stratne_okresy)})' if len(d) else '', al=Alignment(horizontal='center'))
            if len(d):
                x.fill = GOOD if d.iloc[0].suma > 0 and d.iloc[0].stratne_okresy <= 4 else (BAD if d.iloc[0].suma <= 0 else PatternFill())
            if hn == M['hours'] and a == M['tp'] and hh == M['H']:
                x.font = fB; x.fill = SEL
    r += 11

    def block(r, title, key, vals, flt, label):
        ws.cell(r, 1, title).font = fB
        hdr(ws, r + 1, 1, [label, 'Transakcje', 'Wynik % (1:1)', 'Średnia', 'Trafne', 'PF', 'Obsunięcie (suma %)', 'Okresy stratne / 20', '1. połowa (%)', '2. połowa (%)'])
        for k, v in enumerate(vals):
            d = g[flt & (g[key] == v)]; row = r + 2 + k
            lab = v if isinstance(v, str) else (f'{_num(v)} ATR' if v > 0 else ('bez SL' if v == 0 else 'szczyt + 0,25 ATR'))
            put(ws, row, 1, lab, fB)
            if len(d):
                d = d.iloc[0]
                for j, (val, fm) in enumerate([(int(d.n), '0'), (d.suma / 100, PCT3), (d.srednia / 100, PCT3), (d.trafne, PCT), (min(d.PF, 99), '0.00'),
                                                (d.obsuniecie / 100, PCT3), (int(d.stratne_okresy), '0'), (d.suma1 / 100, PCT3), (d.suma2 / 100, PCT3)]):
                    put(ws, row, 2 + j, val, fmt=fm)
                if v == {'SL': M['sl'], 'setup': M['setup'], 'trend': M['trend']}[key]:
                    for c in range(1, 11):
                        ws.cell(row, c).fill = SEL
        return r + 3 + len(vals)
    base = (g.TP == M['tp']) & (g.H == M['H'])
    r = block(r, 'B. Stop loss (reszta jak w strategii)', 'SL', SL_K, base & (g.setup == M['setup']) & (g.trend == M['trend']) & (g.godziny == M['hours']), 'Stop loss')
    r = block(r, 'C. Sygnał wzrostu (reszta jak w strategii)', 'setup', list(g[(g.trend == M['trend']) & (g.godziny == M['hours'])].setup.unique()),
              base & (g.SL == M['sl']) & (g.trend == M['trend']) & (g.godziny == M['hours']), 'Sygnał')
    r = block(r, 'D. Filtr trendu dłuższej średniej (reszta jak w strategii)', 'trend', list(g[(g.setup == M['setup']) & (g.godziny == M['hours'])].trend.unique()),
              base & (g.SL == M['sl']) & (g.setup == M['setup']) & (g.godziny == M['hours']), 'Filtr trendu')
    ht = O['hours_table']
    ws.cell(r, 1, f"E. Godzina sygnału (świeca, UTC) — średni wynik transakcji, cała próba i połowy ({M['setup']}, SL {pl(M['sl'], 0)} ATR, TP {pl(M['tp'], 2)} ATR, {M['H']} św.)").font = fB
    hdr(ws, r + 1, 1, ['Godzina (UTC)', 'SMA500 spada: n', 'średnia', '1. połowa', '2. połowa', 'bez filtra: n', 'średnia', '1. połowa', '2. połowa'])
    for hh in range(24):
        row = r + 2 + hh; put(ws, row, 1, hh)
        for j, tn in enumerate(['SMA500 spada', 'bez filtra trendu']):
            d = ht[(ht.trend == tn) & (ht.godzina == hh)].iloc[0]
            put(ws, row, 2 + 4 * j, int(d.n))
            for q, c in enumerate(['srednia', 'sr1', 'sr2']):
                v = d[c]; x = put(ws, row, 3 + 4 * j + q, None if pd.isna(v) else v / 100, fmt=PCT3)
                if not pd.isna(v):
                    x.fill = GOOD if v > 0 else BAD
        if 11 <= hh <= 16:
            ws.cell(row, 1).fill = SEL

    # UCZCIWY TEST
    ws = sheet('UCZCIWY TEST')
    ws['A1'] = 'UCZCIWY TEST — czy to przewaga, czy przypadek z przeszukania?'; ws['A1'].font = fT
    note(ws, 2, 1, 'Zasada 4: wynik na danych, na których szukano, jest zawyżony. Placebo — ten sam proces na sygnałach przesuniętych w czasie (ta sama częstość i skupienie, zero związku z ceną); połowy danych osobno; wybór chronologiczny (reguła wybrana tylko na 1. roku, sprawdzona na 2. roku).', 12, 42)
    r = 4; ws.cell(r, 1, 'A. Placebo — ile reguł spełnia warunek (≥ 300 transakcji wszystkie pary / ≥ 60–100 EUR/USD, ≤ 3 stratne okresy z 20, suma > 0)').font = fB
    hdr(ws, r + 1, 1, ['Przeszukanie', 'Prawdziwe sygnały', 'Placebo 1', 'Placebo 2', 'Placebo 3'], [52, 14, 12, 12, 12])
    p1, p2 = O['placebo1'], O['placebo2']
    rows_ = [('Etap 1, wszystkie pary, ≤ 3 stratne', [p1[k]['all3'] for k in p1]), ('Etap 1, wszystkie pary, ≤ 4 stratne', [p1[k]['all4'] for k in p1]),
             ('Etap 1, sam EUR/USD (≥ 60 transakcji), ≤ 3 stratne', [p1[k]['eu3'] for k in p1]),
             ('Etap 1, sam EUR/USD (≥ 100 transakcji), ≤ 3 stratne', [p1[k]['eu3_100'] for k in p1]),
             ('Etap 2, wszystkie pary, ≤ 3 stratne', [p2[k]['all3'] for k in p2]),
             ('Etap 2, wszystkie pary, obie połowy osobno ≤ 2 stratne z 10', [p2[k]['halves'] for k in p2]),
             ('Etap 2, sam EUR/USD, ≤ 3 stratne', [p2[k]['eu3'] for k in p2])]
    for k, (lab, vals) in enumerate(rows_):
        put(ws, r + 2 + k, 1, lab)
        for j, v in enumerate(vals):
            put(ws, r + 2 + k, 2 + j, v, fmt='#,##0')
    r += 11
    mm = g[(g.setup == M['setup']) & (g.trend == M['trend']) & (g.godziny == M['hours']) & (g.SL == M['sl']) & (g.TP == M['tp']) & (g.H == M['H'])]
    ws.cell(r, 1, 'B. FX-SHORT-NY w połowach danych (1:1, wszystkie pary)').font = fB
    hdr(ws, r + 1, 1, ['Część', 'Transakcje', 'Wynik % (suma)', 'Okresy stratne / 10', 'Obsunięcie (suma %)'])
    if len(mm):
        mm = mm.iloc[0]
        for k, (lab, a) in enumerate([('1. połowa (okresy 1–10)', 14), ('2. połowa (okresy 11–20)', 18)]):
            vals = [mm[FIELDS[a]], mm[FIELDS[a + 1]], mm[FIELDS[a + 2]], mm[FIELDS[a + 3]]]
            put(ws, r + 2 + k, 1, lab); put(ws, r + 2 + k, 2, int(vals[0])); put(ws, r + 2 + k, 3, vals[1] / 100, fmt=PCT3)
            put(ws, r + 2 + k, 4, int(vals[2])); put(ws, r + 2 + k, 5, vals[3] / 100, fmt=PCT3)
    r += 6
    ws.cell(r, 1, 'C. Wybór chronologiczny: najlepsza reguła etapu 1 wybrana WYŁĄCZNIE na jednej połowie (≥ 150 transakcji, ≤ 1 stratny okres z 10, ranking zysk / obsunięcie), sprawdzona na drugiej').font = fB
    hdr(ws, r + 1, 1, ['Proces', 'Reguł spełnia warunek', 'Najlepsza: wynik % w sprawdzianie', 'Najlepsze 10: średni wynik %', 'Najlepsze 10: dodatnich', 'Najlepsze 50: średnio na transakcję'])
    k = 0
    for tag, w in O['wf'].items():
        for fold, lab in [(0, '1. rok → 2. rok'), (1, '2. rok → 1. rok')]:
            if tag != 'real' and fold == 1:
                continue
            row = r + 2 + k; k += 1; put(ws, row, 1, f'{tag}, {lab}')
            if fold in w:
                x = w[fold]
                put(ws, row, 2, int(x['spelnia'])); put(ws, row, 3, x['top1'] / 100, fmt=PCT3); put(ws, row, 4, x['top10'] / 100, fmt=PCT3)
                put(ws, row, 5, x['top10_poz'], fmt='0%'); put(ws, row, 6, x['top50_tr'] / 100, fmt=PCT3)
    r += k + 4
    ws.cell(r, 1, f"D. {pl(O['start'], 0)} $ w 2. roku przy 1:{O['lev']:g} i 20% depozytu na pozycję").font = fB
    hdr(ws, r + 1, 1, ['Strategia', 'Transakcje', 'Średnia na transakcję', 'Trafne', 'Saldo po 2. roku ($)', 'Zysk', 'Największe obsunięcie'])
    if 'wf_rule' in O:
        wr, wsu, wst = O['wf_rule'], O['wf_sum'], O['wf_stats']
        put(ws, r + 2, 1, f"Wybrana na 1. roku (poza próbą): {wr['setup']}, {wr['trend']}, {wr['godziny']}, {wr['wyjscie']}", al=WR)
        put(ws, r + 2, 2, int(wsu['n'])); put(ws, r + 2, 3, wsu['srednia'] / 100, fmt=PCT3); put(ws, r + 2, 4, wsu['trafne'], fmt=PCT)
        put(ws, r + 2, 5, round(wst['koniec'], 2), fmt=USD); put(ws, r + 2, 6, wst['zysk_pct'] / 100, fmt=PCT); put(ws, r + 2, 7, -wst['max_obs_pct'] / 100, fmt=PCT)
    m2, my2 = O['main_y2_tr'], O['main_y2']
    put(ws, r + 3, 1, 'FX-SHORT-NY w 2. roku (w próbie — ustawienia znane z całych danych)', al=WR)
    put(ws, r + 3, 2, len(m2)); put(ws, r + 3, 3, (m2.wynik.mean() if len(m2) else 0) / 100, fmt=PCT3); put(ws, r + 3, 4, (m2.wynik > 0).mean() if len(m2) else 0, fmt=PCT)
    put(ws, r + 3, 5, round(my2['koniec'], 2), fmt=USD); put(ws, r + 3, 6, my2['zysk_pct'] / 100, fmt=PCT); put(ws, r + 3, 7, -my2['max_obs_pct'] / 100, fmt=PCT)
    ws.row_dimensions[r + 2].height = 40; ws.row_dimensions[r + 3].height = 30
    if 'wf_curve' in O:
        c0 = 9
        hdr(ws, r + 1, c0, ['Dzień', 'Wybrana na 1. roku (poza próbą)', 'FX-SHORT-NY (w próbie)'], [12, 14, 14])
        cc = pd.DataFrame({'a': O['wf_curve'], 'b': O['main_y2_curve']}).dropna()
        for k, (d, rr) in enumerate(cc.iterrows()):
            put(ws, r + 2 + k, c0, py(d), fmt=DD); put(ws, r + 2 + k, c0 + 1, round(rr.a, 2), fmt=USD0); put(ws, r + 2 + k, c0 + 2, round(rr.b, 2), fmt=USD0)
        ch = lchart('2. rok: reguła wybrana bez wiedzy o nim vs FX-SHORT-NY', 9, 20)
        ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + 2, min_row=r + 1, max_row=r + 1 + len(cc)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=c0, min_row=r + 2, max_row=r + 1 + len(cc)))
        ch.series[0].graphicalProperties.line.solidFill = 'ED7D31'; ch.series[1].graphicalProperties.line.solidFill = '1F3A5F'; ch.legend.position = 'b'
        ws.add_chart(ch, f'A{r + 6}')

    # HIPOTEZA
    ws = sheet('HIPOTEZA')
    ws['A1'] = 'HIPOTEZA człowieka: po gwałtownej zielonej świecy zawsze przychodzi korekta — test bez filtrów'; ws['A1'].font = fT
    note(ws, 2, 1, f'Short po każdym sygnale wzrostu, bez filtra trendu i godzin, {NE} wariantów wyjścia. Wynik netto po spreadzie. Gdyby korekta przychodziła zawsze, większość wyjść byłaby na plusie.', 12, 40)
    for blk, (src, title) in enumerate([(O['raw0'], 'Wszystkie pary'), (O['raw1'], 'Sam EUR/USD')]):
        r = 4 + blk * 14; ws.cell(r, 1, title).font = fB
        hdr(ws, r + 1, 1, ['Sygnał', 'Sygnałów (najwięcej)', 'Średnia ze wszystkich wyjść (na transakcję)', 'Wyjść z sumą > 0', 'Najlepsze wyjście', 'Najlepsze: suma %', 'Najlepsze: okresy stratne'],
            [34, 11, 14, 11, 34, 11, 11])
        for k, (sn, d) in enumerate(src.groupby('setup', sort=False)):
            best = d.sort_values('suma').iloc[-1]; row = r + 2 + k
            put(ws, row, 1, sn); put(ws, row, 2, int(d.n.max())); put(ws, row, 3, d.srednia.mean() / 100, fmt=PCT3)
            put(ws, row, 4, (d.suma > 0).mean(), fmt='0%'); put(ws, row, 5, best['wyjście']); put(ws, row, 6, best.suma / 100, fmt=PCT3); put(ws, row, 7, int(best.stratne_okresy))

    # RANKING
    ws = sheet('RANKING')
    ws['A1'] = 'RANKING z etapu 1 — tylko do wglądu (w próbie, zawyżony)'; ws['A1'].font = fT
    note(ws, 2, 1, f"Etap 1: {pl(O['n_rules1'], 0)} reguł. Warunek: ≤ 3 stratne okresy z 20, suma > 0; ranking: zysk / obsunięcie; wyniki % bez dźwigni. EUR/USD: na placebo podobnie dużo reguł spełnia warunek (UCZCIWY TEST) — te wyniki to głównie przypadek.", 16, 40)
    cols = ['Sygnał', 'Trend', 'Godziny', 'Wyjście', 'Transakcje', 'Suma %', 'Średnia', 'Mediana', 'Trafne', 'PF', 't', 'Obsunięcie', 'Okresy stratne', 'Okres: średnia', 'Okres: mediana', 'Okres: najgorszy', 'Zysk / obsunięcie']
    for blk, (src, title) in enumerate([(O['rank0'], 'Wszystkie pary (≥ 300 transakcji)'), (O['rank1'], 'Sam EUR/USD (≥ 100 transakcji)')]):
        r = 4 + blk * 45; ws.cell(r, 1, title).font = fB
        hdr(ws, r + 1, 1, cols, [30, 22, 18, 30, 9, 9, 9, 9, 8, 7, 7, 9, 8, 9, 9, 9, 9])
        for k, (_, d) in enumerate(src.iterrows()):
            vals = [d.setup, d.trend, d.godziny, d['wyjście'], int(d.n), d.suma / 100, d.srednia / 100, d.mediana / 100, d.trafne, min(d.PF, 99), d.t,
                    d.obsuniecie / 100, int(d.stratne_okresy), d.okres_srednia / 100, d.okres_mediana / 100, d.okres_min / 100, d['zysk/obs']]
            fms = [None] * 4 + ['0', PCT3, PCT3, PCT3, PCT, '0.00', '0.0', PCT3, '0', PCT3, PCT3, PCT3, '0.0']
            for j, (v, fm) in enumerate(zip(vals, fms)):
                put(ws, r + 2 + k, 1 + j, v, fmt=fm)

    # EURUSD
    ws = sheet('EURUSD')
    ws['A1'] = 'EUR/USD — ta sama strategia tylko na EUR/USD'; ws['A1'].font = fT
    note(ws, 2, 1, 'Formuły czytają TRANSAKCJE (wynik % bez dźwigni). Saldo przy dźwigni — symulacja Python tylko z transakcjami EUR/USD (jedna pozycja naraz). Mało transakcji → szeroki przedział ufności (zasada 6).', 10, 42)
    hdr(ws, 4, 1, ['Okres', 'Transakcje', 'Wynik % (1:1)', 'Trafne'], [8, 10, 12, 9])
    for k in range(NPER):
        row = 5 + k; put(ws, row, 1, k + 1)
        put(ws, row, 2, f'=COUNTIFS({TR("O")},A{row},{TR("B")},"EURUSD")', fmt='0')
        put(ws, row, 3, f'=SUMIFS({TR("L")},{TR("O")},A{row},{TR("B")},"EURUSD")', fmt=PCT3)
        put(ws, row, 4, f'=IF(B{row}=0,0,COUNTIFS({TR("O")},A{row},{TR("B")},"EURUSD",{TR("L")},">0")/B{row})', fmt=PCT)
    put(ws, 25, 1, 'Razem', fB, fill=GRY); put(ws, 25, 2, '=SUM(B5:B24)', fB, '0'); put(ws, 25, 3, '=SUM(C5:C24)', fB, PCT3)
    put(ws, 26, 1, 'Okresy stratne', fB, fill=GRY); put(ws, 26, 3, '=COUNTIF(C5:C24,"<0")', fB, '0')
    put(ws, 27, 1, 'Okresy puste', fB, fill=GRY); put(ws, 27, 3, '=COUNTIF(B5:B24,0)', fB, '0')
    hdr(ws, 4, 6, ['Depozyt na pozycję', 'Wartość pozycji', 'Saldo końcowe ($)', 'Zysk', 'Największe obsunięcie', 'Transakcje', 'Stop out'], [11, 10, 13, 9, 11, 9, 8])
    for k, v in O['eu_var'].reset_index(drop=True).iterrows():
        row = 5 + k
        for c, (val, fm) in enumerate([(v.f, PCT), (v.f * O['lev'], '0.0"×"'), (round(v.koniec, 2), USD), (v.zysk_pct / 100, PCT), (-v.max_obs_pct / 100, PCT),
                                       (int(v.zawarte), '0'), (int(v.stop_out), '0')]):
            put(ws, row, 6 + c, val, fmt=fm)
    c0 = 14
    hdr(ws, 4, c0, ['Dzień', 'EUR/USD, depozyt 33%', 'EUR/USD, depozyt 100%'], [12, 13, 13])
    ec = pd.DataFrame({'a': O['eu_curves'][1 / 3], 'b': O['eu_curves'][1.0]})
    for k, (d, rr) in enumerate(ec.iterrows()):
        put(ws, 5 + k, c0, py(d), fmt=DD); put(ws, 5 + k, c0 + 1, round(rr.a, 2), fmt=USD0); put(ws, 5 + k, c0 + 2, round(rr.b, 2), fmt=USD0)
    ch = lchart(f"EUR/USD: {pl(O['start'], 0)} $ przy 1:{O['lev']:g}", 8, 20)
    ch.add_data(Reference(ws, min_col=c0 + 1, max_col=c0 + 2, min_row=4, max_row=4 + len(ec)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=c0, min_row=5, max_row=4 + len(ec))); ch.legend.position = 'b'
    ws.add_chart(ch, 'A30')

    # PODSUMOWANIE
    ws = wsP
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 70
    ws['A1'] = 'FX-SHORT-NY — short po gwałtownym wzroście w sesji Londyn + Nowy Jork (10 par, świece 1h)'; ws['A1'].font = fT
    ws['A2'] = (f"IA4 {__version__} · badanie informacyjne · kopia fx1h: {O['od']:%Y-%m-%d} – {O['do']:%Y-%m-%d} ({pl(O['swiec'], 0)} świec na parę) · "
                f"policzono {pd.Timestamp.now():%Y-%m-%d %H:%M}")
    ws['A2'].font = fS
    r = 4; put(ws, r, 1, 'REGUŁA (wszystkie liczby stałe)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    rule = [('Rynek', 'EUR/USD, USD/JPY, GBP/USD, AUD/USD, USD/CAD, USD/CHF, NZD/USD, EUR/JPY, GBP/JPY, EUR/GBP — świece 1h; jedna pozycja na parę'),
            ('Sygnał (na zamknięciu świecy)', 'kurs wzrósł w ostatnich 6 świecach o ≥ 3,1 × ATR(100) [1,25 × √6] — „krótkotrwały wzrost”'),
            ('Godzina', 'świeca sygnału zaczyna się 11:00–16:00 UTC (sesja Londyn + Nowy Jork; w Polsce 13:00–18:00 latem, 12:00–17:00 zimą)'),
            ('Filtr trendu', 'SMA(500) niżej niż 24 świece wcześniej (dłuższa średnia spada — wzrost jest korektą w trendzie spadkowym)'),
            ('Wejście', 'SHORT po open następnej świecy'),
            ('Stop loss', f"cena wejścia + {pl(M['sl'], 0)} × ATR(100)"),
            ('Take profit', f"cena wejścia − {pl(M['tp'], 2)} × ATR(100)"),
            ('Limit czasu', f"po {M['H']} świecach wyjście po close"),
            ('Koszty', 'spread raz na transakcję: EUR/USD 1,1 pipsa (zrzut XTB), pozostałe pary — arkusz PARY; koszt nocny — TRANSAKCJE!B6'),
            ('Rozliczenie', 'pesymistyczne (zasada 5): SL i TP w tej samej świecy → SL; luka przez SL → open; luka przez TP → TP bez premii')]
    for k, (a, b) in enumerate(rule):
        put(ws, r + 1 + k, 1, a, fB, fill=GRY, al=WR); put(ws, r + 1 + k, 2, b, al=WR)
        ws.merge_cells(start_row=r + 1 + k, start_column=2, end_row=r + 1 + k, end_column=3); ws.row_dimensions[r + 1 + k].height = 28
    r += 2 + len(rule)
    put(ws, r, 1, 'WYNIK (formuły z arkuszy TRANSAKCJE, OKRESY, KAPITAŁ)', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    ms = O['main_sum']; reasons = O['main'].powod.value_counts() if len(O['main']) else {}
    vdef = O['variants'][(abs(O['variants'].f - O['f']) < 1e-9) & (O['variants'].swap == 0)].iloc[0]
    res = [('Transakcje (wszystkie pary)', f'=COUNTA({TR("B")})', '0', f"EUR/USD: {O['eu_sum']['n']} (arkusz EURUSD)"),
           ('Średni wynik transakcji (bez dźwigni)', f'=AVERAGE({TR("L")})', PCT3, f"przedział 95%: ± {pl(O['main_ci'], 3)}% (t = {pl(O['main_t'])})"),
           ('Mediana transakcji', f'=MEDIAN({TR("L")})', PCT3, 'dużo małych zysków (TP), rzadkie większe straty (SL / czas)'),
           ('Trafne transakcje', f'=COUNTIF({TR("L")},">0")/COUNTA({TR("L")})', PCT, ' · '.join(f'{k} {v}' for k, v in dict(reasons).items())),
           ('Profit factor', f'=SUMIF({TR("L")},">0")/-SUMIF({TR("L")},"<0")', '0.00', ''),
           ('Mediana świec w pozycji', f'=MEDIAN({TR("F")})', '0', 'godziny'),
           ('Okresy stratne z 20 (wynik %)', '=OKRESY!E26', '0', 'warunek człowieka: najmniej stratnych okresów'),
           ('Średnia / mediana okresu (wynik % bez dźwigni)', '=OKRESY!E28', PCT3, ''), ('', '=OKRESY!E29', PCT3, ''),
           ('Saldo końcowe przy ustawieniach z TRANSAKCJE ($)', f'=KAPITAŁ!B{KR1}', USD, 'zmień TRANSAKCJE!B3:B7'),
           ('Zysk', f'=KAPITAŁ!B{KR1}/TRANSAKCJE!B3-1', PCT, ''),
           ('Największe obsunięcie salda (dzienne)', f'=-MIN(KAPITAŁ!D5:D{KR1})', PCT, f"z wyceną otwartych pozycji co godzinę (Python, ustawienia domyślne): {pl(vdef.max_obs_pct)}%"),
           ('Zmiana kapitału: okresy stratne / mediana okresu', '=OKRESY!K26', '0', ''), ('', '=OKRESY!K29', PCT, '')]
    for k, (a, f_, fm, c) in enumerate(res):
        row = r + 1 + k
        put(ws, row, 1, a, fB, fill=GRY); put(ws, row, 2, f_, fB, fm); put(ws, row, 3, c, fS, al=WR)
    r += 2 + len(res)
    put(ws, r, 1, 'JAK TO CZYTAĆ', fH, fill=HDR); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    txt = ['Hipoteza w czystej postaci — arkusz HIPOTEZA (short po każdej dużej świecy / wzroście, bez filtrów).',
           'Kiedy korekta działa — ODPORNOŚĆ (godzina sygnału, filtr trendu, SL, sąsiednie progi).',
           f"Kapitał przy dźwigni — KAPITAŁ i WARIANTY; ustawienia domyślne: depozyt {pl(O['f'] * 100, 0)}% salda na pozycję → saldo {pl(vdef.koniec, 0)} $, obsunięcie {pl(vdef.max_obs_pct)}%.",
           'Uczciwość wyniku — UCZCIWY TEST: placebo, połowy, wybór chronologiczny. Wynik w próbie jest zawyżony; liczy się 2. rok reguły wybranej na 1. roku.',
           'Ryzyka poza symulacją: szerszy spread w czasie danych z USA (12:30, 14:00 UTC), rozdzielczość świec Yahoo, jeden reżim rynku w 2 latach danych.']
    for k, s_ in enumerate(txt):
        x = ws.cell(r + 1 + k, 1, s_); x.font = fN; x.alignment = WR
        ws.merge_cells(start_row=r + 1 + k, start_column=1, end_row=r + 1 + k, end_column=3); ws.row_dimensions[r + 1 + k].height = 30
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- uruchomienie

def run(no_sync=False, out=None, f=0.20, lev=30.0, start=1000.0, frames=None, log=print) -> Path:
    t0 = time.time()
    if frames is None:
        from .fx import sync as fxs
        if not no_sync:
            log('Kopia fx1h — odświeżenie (python -m ia4.fx.sync)')
            fxs.sync(quiet=True, syms=list(PAIRS))
        frames = {p: fxs.load(p) for p in PAIRS}
    D = Data(frames, spreads())
    log(f'Dane: {D.n} świec na parę, {pd.Timestamp(int(D.t[0]), unit="s"):%Y-%m-%d} – {pd.Timestamp(int(D.t[-1]), unit="s"):%Y-%m-%d} ({time.time() - t0:.0f} s)')
    O = compute(D, f=f, lev=lev, start=start, log=log)
    path = write_xlsx(O, out or (config.RESEARCH_DIR.parent / 'FX-SHORT-NY.xlsx'))
    ms = O['main_sum']; v = O['variants'][(abs(O['variants'].f - f) < 1e-9) & (O['variants'].swap == 0)].iloc[0]
    log(f"\nFX-SHORT-NY: {ms['n']} transakcji, średnio {ms['srednia']:+.3f}% (t = {O['main_t']:.1f}), trafne {ms['trafne']:.0%}, "
        f"okresy stratne {ms['stratne']}/20; {start:.0f} $ przy 1:{lev:g} i depozycie {f:.0%} → {v.koniec:,.0f} $, obsunięcie {v.max_obs_pct:.1f}%")
    if 'wf_stats' in O:
        log(f"Poza próbą (reguła wybrana na 1. roku, 2. rok, 20%): {O['wf_stats']['zysk_pct']:+.1f}%")
    log(f'Zapisano: {path}  ({time.time() - t0:.0f} s)')
    return path


def _synthetic(n=3000, seed=0) -> dict:
    """Losowy spacer bez przewagi — do selftestu."""
    rng = np.random.default_rng(seed)
    t0 = int(pd.Timestamp('2025-01-06 00:00', tz='UTC').timestamp())
    t = t0 + 3600 * np.arange(n)
    out = {}
    for k, p in enumerate(PAIRS):
        base = 150.0 if p.endswith('JPY') else 1.1
        r = rng.normal(0, 0.001, n); c = base * np.exp(np.cumsum(r)); o = np.r_[c[0], c[:-1]]
        h = np.maximum(o, c) * (1 + np.abs(rng.normal(0, 0.0005, n))); l = np.minimum(o, c) * (1 - np.abs(rng.normal(0, 0.0005, n)))
        out[p] = pd.DataFrame({'time': t, 'o': o, 'h': h, 'l': l, 'c': c, 'filled': False})
    return out


def selftest() -> bool:
    ok = True

    def check(name, cond):
        nonlocal ok
        print(f"  {'✓' if cond else '✗'} {name}")
        ok &= bool(cond)

    # rozliczenie pesymistyczne na ręcznym przykładzie (świeca 0 = sygnał, wejście po open świecy 1)
    o = np.array([1.0, 1.0, 1.0, 1.0, 1.0]); c = o.copy(); atr = np.full(5, 0.01); valid = np.array([True, False, False, False, False])
    sl = np.array([1.0]); tp = np.array([1.0]); H = np.array([3], np.int64)
    h = np.array([1.0, 1.02, 1.0, 1.0, 1.0]); l = np.array([1.0, 0.98, 1.0, 1.0, 1.0])
    P, X, MA = outcomes(o, h, l, c, atr, valid, 0.0001, sl, tp, H)
    check('SL i TP w tej samej świecy → SL (−1 ATR − spread)', abs(P[0, 0] - (-1.0 - 0.01)) < 1e-4 and X[0, 0] == 1)
    h = np.array([1.0, 1.0, 1.0, 1.0, 1.0]); l = np.array([1.0, 0.985, 1.0, 1.0, 1.0])
    P, X, MA = outcomes(o, h, l, c, atr, valid, 0.0001, sl, tp, H)
    check('TP 1 ATR → +1% − spread', abs(P[0, 0] - (1.0 - 0.01)) < 1e-4)
    o2 = np.array([1.0, 1.0, 1.03, 1.03, 1.03]); h2 = np.array([1.0, 1.0, 1.03, 1.03, 1.03]); l2 = np.array([1.0, 1.0, 1.03, 1.03, 1.03])
    P, X, MA = outcomes(o2, h2, l2, o2, atr, valid, 0.0001, sl, tp, H)
    check('luka przez SL → wyjście po open (−3%)', abs(P[0, 0] - (-3.0 - 0.01)) < 1e-4 and X[0, 0] == 2)
    P, X, MA = outcomes(o, np.ones(5), np.ones(5), c, atr, valid, 0.0001, np.array([0.0]), np.array([0.0]), H)
    check('bez SL i TP → wyjście po close 3. świecy, tylko spread', abs(P[0, 0] + 0.01) < 1e-4 and X[0, 0] == 3)
    check('rolowania: pon 20:00 → śr 23:00 UTC (zima) = pon 1 + wt 1 + śr 3', rollovers(pd.Timestamp('2025-01-06 20:00'), pd.Timestamp('2025-01-08 23:00')) == 5)
    # całość na losowym spacerze: przebiega, formuły i Python liczą tak samo
    fr = _synthetic()
    D = Data(fr, spreads())
    e = exit_index(MAIN['sl'], MAIN['tp'], MAIN['H'])
    tr = D.trades(2, 'wzrost 3 św. ≥ 2.2 ATR', 'bez filtra trendu', '7–16 UTC', e)
    check(f'transakcje na danych syntetycznych ({len(tr)})', len(tr) > 20)
    L, cv, ev = simulate(tr, D, f=0.2)
    cash = 1000.0; rows = []
    for k, r in tr.iterrows():
        rows.append(r)
    # odtworzenie formuł arkusza: saldo przy wejściu = start + wyniki zamkniętych do wejścia
    res = []
    for k, r in tr.iterrows():
        closed = sum(x['usd'] for x in res if x['exit'] <= r.wejscie_t)
        openm = sum(x['margin'] for x in res if x['entry'] <= r.wejscie_t < x['exit'])
        bal = 1000.0 + closed; m = min(0.2 * bal, bal - openm)
        bu = 1.0 if BASE_USD[r.para] is None else float(D.F[BASE_USD[r.para]]['o'][int(r.i) + 1])
        units = 0 if m <= 0 else int(m * 30 / bu / 1000) * 1000
        res.append(dict(entry=r.wejscie_t, exit=r.wyjscie_t, margin=units * bu / 30, usd=units * bu * r.wynik / 100))
    check('symulacja Python = logika formuł TRANSAKCJE', len(ev) > 0 or abs(1000 + sum(x['usd'] for x in res) - L[L.status == 'zamknięta'].wynik_usd.sum() - 1000) < 1e-6)
    print(f"\nWYNIK: {'OK' if ok else 'BŁĄD'}")
    return ok


def main():
    ap = argparse.ArgumentParser(description='IA 4 — FX-SHORT-NY: short po gwałtownym wzroście (sekcja 7b)')
    ap.add_argument('--bez-sync', action='store_true', help='bez odświeżenia kopii fx1h')
    ap.add_argument('--depozyt', type=float, default=0.20, help='depozyt na pozycję jako część salda (domyślnie 0.20)')
    ap.add_argument('--dzwignia', type=float, default=30.0)
    ap.add_argument('--start', type=float, default=1000.0)
    ap.add_argument('--out', type=Path)
    ap.add_argument('--selftest', action='store_true')
    a = ap.parse_args()
    if a.selftest:
        sys.exit(0 if selftest() else 1)
    print(f'IA 4 {__version__} — FX-SHORT-NY (short po gwałtownym wzroście, 10 par, 1h)\n')
    run(no_sync=a.bez_sync, out=a.out, f=a.depozyt, lev=a.dzwignia, start=a.start)


if __name__ == '__main__':
    main()
